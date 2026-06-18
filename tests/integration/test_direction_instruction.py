"""Engine tests for the dir: scan-direction-switch instruction.

Uses a small synthetic 3×2 grid where LR and TD scan orders diverge clearly:

        A      B
   1    a1     b1
   2    a2     b2
   3    a3     b3

  LR order: a1, b1, a2, b2, a3, b3
  TD order: a1, a2, a3, b1, b2, b3
"""
import os
import tempfile

import openpyxl
import pytest

from grepxcel import Engine, Logger, VerbosityLevel


def _make_grid() -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Grid'
    ws['A1'], ws['B1'] = 'a1', 'b1'
    ws['A2'], ws['B2'] = 'a2', 'b2'
    ws['A3'], ws['B3'] = 'a3', 'b3'
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _make_pattern(rows: list) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _run(pattern_rows: list) -> dict:
    data = _make_grid()
    pat = _make_pattern(pattern_rows)
    lg = Logger(level=VerbosityLevel.QUIET)
    try:
        return Engine().process(pattern_file=pat, data_file=data, logger=lg)
    finally:
        os.unlink(data)
        os.unlink(pat)


_FIELDS = [
    ['var:', 'v1', 'string', '.*'],
    ['var:', 'v2', 'string', '.*'],
    ['var:', 'v3', 'string', '.*'],
    ['var:', 'v4', 'string', '.*'],
    ['var:', 'v5', 'string', '.*'],
    ['var:', 'v6', 'string', '.*'],
]


def test_td_switch_to_lr_midway():
    """Read first cell in TD, switch to LR, read the rest.

    TD first read = a1 (consumed). Switch to LR: scan order a1(skip), b1, a2, ...
    so v2 must be b1 (LR behaviour), not a2 (pure-TD behaviour).
    """
    result = _run([
        ['config:', 'read.direction', 'TD'],
        *_FIELDS,
        ['START:'],
        ['cell:next', 'v1'],   # a1
        ['dir:LR'],
        ['cell:next', 'v2'],   # b1  (LR), would be a2 in pure TD
        ['cell:next', 'v3'],   # a2
        ['END:'],
    ])
    assert result['v1'] == 'a1'
    assert result['v2'] == 'b1'
    assert result['v3'] == 'a2'


def test_lr_switch_to_td_midway():
    """Read first cell in LR (default), switch to TD, read the rest.

    LR first read = a1 (consumed). Switch to TD: scan order a1(skip), a2, a3, ...
    so v2 must be a2 (TD behaviour), not b1 (pure-LR behaviour).
    """
    result = _run([
        *_FIELDS,
        ['START:'],
        ['cell:next', 'v1'],   # a1
        ['dir:TD'],
        ['cell:next', 'v2'],   # a2  (TD), would be b1 in pure LR
        ['cell:next', 'v3'],   # a3
        ['END:'],
    ])
    assert result['v1'] == 'a1'
    assert result['v2'] == 'a2'
    assert result['v3'] == 'a3'


def test_dir_does_not_consume_a_cell():
    """dir: only changes scan direction; it must not read or consume any cell."""
    result = _run([
        *_FIELDS,
        ['START:'],
        ['dir:TD'],            # no read
        ['cell:next', 'v1'],   # a1 — first TD cell, proving dir consumed nothing
        ['END:'],
    ])
    assert result['v1'] == 'a1'


def test_dir_skips_already_consumed_cells():
    """After a direction switch, consumed cells are skipped in the new order."""
    result = _run([
        ['config:', 'read.direction', 'TD'],
        *_FIELDS,
        ['START:'],
        ['cell:next', 'v1'],   # a1
        ['cell:next', 'v2'],   # a2
        ['dir:LR'],            # switch; a1,a2 already consumed
        ['cell:next', 'v3'],   # LR: a1(skip),b1,a2(skip),b2 → b1
        ['cell:next', 'v4'],   # b2
        ['END:'],
    ])
    assert result['v1'] == 'a1'
    assert result['v2'] == 'a2'
    assert result['v3'] == 'b1'
    assert result['v4'] == 'b2'


def test_multiple_direction_switches():
    """Direction can be switched more than once in a sequence."""
    result = _run([
        ['config:', 'read.direction', 'LR'],
        *_FIELDS,
        ['START:'],
        ['cell:next', 'v1'],   # LR: a1
        ['dir:TD'],
        ['cell:next', 'v2'],   # TD: a2
        ['dir:LR'],
        ['cell:next', 'v3'],   # LR: a1(c),b1 → b1
        ['END:'],
    ])
    assert result['v1'] == 'a1'
    assert result['v2'] == 'a2'
    assert result['v3'] == 'b1'
