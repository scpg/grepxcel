"""Web wizard request guard: CSRF and DNS-rebinding protection.

The wizard binds to 127.0.0.1, but any page open in the user's browser can
still reach it. Before this guard, a cross-site
``<form method=POST enctype=text/plain>`` could call ``/api/save`` (and every
other JSON endpoint) because handlers parse bodies with ``request.json()``,
which ignores Content-Type — so a "simple" request needing no CORS preflight
was accepted, and wrote files into the data file's folder. Any Host header was
also accepted, which lets a DNS-rebinding page read the session.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

try:
    from fastapi.testclient import TestClient
    from grepxcel.wizard_api import create_app
    _API_OK = True
except ImportError:
    _API_OK = False

pytestmark = pytest.mark.skipif(
    not _API_OK,
    reason='fastapi / uvicorn not installed (pip install "grepxcel[web]")',
)

LOCAL = 'http://localhost:8765'
EVIL = 'http://evil.example'
FIXTURE_DIR = Path(__file__).parent.parent / 'fixtures/01_simple_invoice'
PATTERN_01 = FIXTURE_DIR / '01_simple_invoice_pattern-manual.xlsx'


def _client(tmp_path: Path, base_url: str = LOCAL) -> 'TestClient':
    xlsx_path = tmp_path / 'data.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Invoice No:'
    ws['B1'] = 'AB123456'
    wb.save(xlsx_path)
    return TestClient(create_app(str(xlsx_path)), base_url=base_url)


def _save_body(name: str) -> str:
    return json.dumps({'filename': name, 'confirm_overwrite': True})


class TestCrossSiteFormPost:
    """The exact request a malicious page can send without a CORS preflight."""

    def test_text_plain_cross_origin_save_rejected(self, tmp_path):
        client = _client(tmp_path)
        r = client.post('/api/save', content=_save_body('pwned.txt'),
                        headers={'Content-Type': 'text/plain', 'Origin': EVIL})
        assert r.status_code == 403
        assert not (tmp_path / 'pwned.txt').exists()

    def test_text_plain_without_origin_rejected(self, tmp_path):
        """Content-Type alone is refused, even if a browser omitted Origin."""
        client = _client(tmp_path)
        r = client.post('/api/save', content=_save_body('pwned.txt'),
                        headers={'Content-Type': 'text/plain'})
        assert r.status_code == 415
        assert not (tmp_path / 'pwned.txt').exists()

    def test_urlencoded_form_rejected(self, tmp_path):
        client = _client(tmp_path)
        r = client.post('/api/classify', content='a=b',
                        headers={'Content-Type': 'application/x-www-form-urlencoded'})
        assert r.status_code == 415

    def test_multipart_outside_upload_endpoint_rejected(self, tmp_path):
        client = _client(tmp_path)
        r = client.post('/api/save', files={'file': ('x.txt', b'x')})
        assert r.status_code == 415

    @pytest.mark.parametrize('origin', [EVIL, 'https://localhost:8765', 'null',
                                        'http://localhost.evil.example'])
    def test_json_from_foreign_origin_rejected(self, tmp_path, origin):
        client = _client(tmp_path)
        r = client.post('/api/save', content=_save_body('out.csv'),
                        headers={'Content-Type': 'application/json', 'Origin': origin})
        assert r.status_code == 403
        assert not (tmp_path / 'out.csv').exists()

    def test_multipart_upload_from_foreign_origin_rejected(self, tmp_path):
        client = _client(tmp_path)
        r = client.post('/api/load-pattern', headers={'Origin': EVIL},
                        files={'file': ('p.xlsx', PATTERN_01.read_bytes())})
        assert r.status_code == 403


class TestDnsRebinding:
    @pytest.mark.parametrize('host', ['attacker.example:8765', 'attacker.example',
                                      'localhost.attacker.example', ''])
    def test_foreign_host_rejected_on_reads(self, tmp_path, host):
        client = _client(tmp_path)
        r = client.get('/api/state', headers={'Host': host})
        assert r.status_code == 403

    @pytest.mark.parametrize('base_url', ['http://localhost:8765', 'http://127.0.0.1:8765',
                                          'http://localhost'])
    def test_loopback_hosts_accepted(self, tmp_path, base_url):
        client = _client(tmp_path, base_url=base_url)
        assert client.get('/api/state').status_code == 200
        assert client.get('/').status_code == 200

    @pytest.mark.parametrize('host', ['[::1]:8765', '[::1]'])
    def test_ipv6_loopback_host_accepted(self, host):
        """Asserted against the predicate rather than through TestClient.

        Driving this with `base_url='http://[::1]:8765'` fails inside Starlette's
        TestClient — it splits the netloc on ':' and calls int() on the result,
        so a bracketed IPv6 literal raises
        `ValueError: invalid literal for int() with base 10: ':1]:8765'`
        before the request ever reaches the guard. That is the client's
        limitation, not the wizard's, and it is version-dependent: the case
        passed in CI and failed locally on starlette 1.6.0 / httpx2 2.12.0.

        A test whose result depends on a test client's URL parser is testing the
        wrong thing. The guard reads the Host header, so the header is what gets
        asserted.

        Note a bare `::1` is deliberately *not* accepted: RFC 7230 requires an
        IPv6 literal in a Host header to be bracketed, so an unbracketed one is
        malformed rather than loopback.
        """
        from grepxcel.wizard_api import _is_loopback_host

        assert _is_loopback_host(host) is True


class TestWizardOwnRequestsStillWork:
    """Requests shaped like the ones static/wizard.js sends must pass."""

    @pytest.mark.parametrize('origin', [LOCAL, 'http://127.0.0.1:8765', None])
    def test_same_origin_json_save(self, tmp_path, origin):
        client = _client(tmp_path)
        headers = {'Origin': origin} if origin else {}
        r = client.post('/api/save', json={'filename': 'out.csv'}, headers=headers)
        assert r.status_code == 200, r.text
        assert (tmp_path / 'out.csv').exists()

    def test_bodyless_post(self, tmp_path):
        """wizard.js calls POST /api/undo with no body and no Content-Type."""
        client = _client(tmp_path)
        r = client.post('/api/undo', headers={'Origin': LOCAL})
        assert r.status_code not in (403, 415)

    def test_multipart_pattern_upload(self, tmp_path):
        # No skip when python-multipart is missing: the [web] extra declares it,
        # so a missing package is exactly the regression this should catch.
        client = _client(tmp_path)
        r = client.post('/api/load-pattern', headers={'Origin': LOCAL},
                        files={'file': ('p.xlsx', PATTERN_01.read_bytes())})
        assert r.status_code == 200, r.text

    def test_upload_without_multipart_explains_fix(self, tmp_path, monkeypatch):
        """An env set up before [web] declared python-multipart gets an actionable
        message (shown by wizard.js as a toast) instead of Starlette's bare 500."""
        import grepxcel.wizard_api as wizard_api
        monkeypatch.setattr(wizard_api, '_MULTIPART_OK', False)
        client = _client(tmp_path)
        r = client.post('/api/load-pattern', headers={'Origin': LOCAL},
                        files={'file': ('p.xlsx', PATTERN_01.read_bytes())})
        assert r.status_code == 500
        assert 'pip install -U "grepxcel[web]"' in r.json()['detail']

    def test_static_assets_served(self, tmp_path):
        client = _client(tmp_path)
        assert client.get('/static/wizard.js').status_code == 200
