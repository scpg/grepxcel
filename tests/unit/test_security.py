"""
Direct tests for the security module's load-bearing guarantees:
  - XXE protection (defusedxml) is live and fail-closed
  - the ReDoS detector is wired to the single re._parser import path
  - ZIP-bomb guards catch both absolute-size and expansion-ratio attacks
  - blocked extensions (.xlsm / .xlsb / .xls) are rejected before any file I/O
  - the engine never evaluates formulas (data_only=True)
"""

import zipfile

import openpyxl
import pytest

from grepxcel.engine import Engine
from grepxcel.logger import Logger, VerbosityLevel
from grepxcel.security import (
    SecurityError,
    assert_xxe_protection,
    check_regex_safety,
    validate_file,
    _check_zip_safety,
)


# ─── XXE / defusedxml guard ───────────────────────────────────────────────────

def test_xxe_protection_is_enabled():
    """openpyxl must be using defusedxml in this environment."""
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML is True
    # assert_xxe_protection() must not raise when protection is on
    assert_xxe_protection() is None


def test_xxe_guard_fails_closed_when_disabled(monkeypatch):
    """If defusedxml is somehow disabled, the guard must raise SecurityError."""
    import openpyxl.xml
    monkeypatch.setattr(openpyxl.xml, 'DEFUSEDXML', False, raising=True)
    with pytest.raises(SecurityError, match='XXE'):
        assert_xxe_protection()


def test_validate_file_runs_xxe_guard_first(monkeypatch, tmp_path):
    """validate_file must refuse to proceed when XXE protection is off,
    before it even looks at the path."""
    import openpyxl.xml
    monkeypatch.setattr(openpyxl.xml, 'DEFUSEDXML', False, raising=True)
    missing = tmp_path / 'nope.xlsx'
    with pytest.raises(SecurityError, match='XXE'):
        validate_file(str(missing))


# ─── ReDoS detector single import path ────────────────────────────────────────

def test_regex_parser_uses_re_internal():
    """We require Python >=3.11, so the parser must come from re._parser
    (no fallback to the deprecated top-level sre_parse module)."""
    from grepxcel import security
    import re._parser as expected
    assert security._sre_parse is expected


def test_safe_regex_accepted():
    # A normal anchored quantifier is fine.
    check_regex_safety(r'PO-[0-9]{4}', 'po_number')


def test_nested_unbounded_quantifier_rejected():
    with pytest.raises(SecurityError, match='ReDoS|backtracking'):
        check_regex_safety(r'(a+)+', 'evil')


def test_invalid_regex_rejected():
    with pytest.raises(SecurityError, match='Invalid regex'):
        check_regex_safety(r'(unclosed', 'bad')


# ─── ZIP-bomb guard measures REAL decompressed size ──────────────────────────

def _make_zip(path, name='big.xml', payload=b'A', repeat=1):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(name, payload * repeat)


def test_legit_small_zip_passes(tmp_path):
    path = tmp_path / 'ok.xlsx'
    _make_zip(path, payload=b'hello world\n', repeat=10)
    _check_zip_safety(str(path), max_uncompressed_mb=50)   # must not raise


def test_real_decompressed_size_over_cap_rejected(tmp_path):
    """A highly compressible member whose real expansion exceeds the cap is
    rejected even though the file on disk is tiny."""
    path = tmp_path / 'big.xlsx'
    _make_zip(path, payload=b'A', repeat=8 * 1024 * 1024)  # 8 MB → compresses tiny
    with pytest.raises(SecurityError, match='ZIP|uncompressed|bomb'):
        _check_zip_safety(str(path), max_uncompressed_mb=1)


def test_falsified_metadata_cannot_bypass_guard(tmp_path, monkeypatch):
    """The guard must count ACTUAL decompressed bytes, not the central-directory
    file_size — so an entry that lies about its size (file_size=0) is still
    caught by the real expansion."""
    path = tmp_path / 'liar.xlsx'
    _make_zip(path, payload=b'A', repeat=8 * 1024 * 1024)  # real 8 MB

    real_infolist = zipfile.ZipFile.infolist

    def lying_infolist(self):
        infos = real_infolist(self)
        for info in infos:
            info.file_size = 0          # forge the declared uncompressed size
        return infos

    monkeypatch.setattr(zipfile.ZipFile, 'infolist', lying_infolist)
    with pytest.raises(SecurityError, match='ZIP|uncompressed|bomb'):
        _check_zip_safety(str(path), max_uncompressed_mb=1)


# ─── ZIP-bomb: expansion-ratio ceiling ───────────────────────────────────────

def test_expansion_ratio_ceiling_rejected(tmp_path):
    """A ZIP that passes the absolute-size cap but has an abnormally high
    expansion ratio (>50×) must still be caught as a likely ZIP bomb.

    200 KB of identical bytes compresses to ~250 bytes on disk (DEFLATE achieves
    ~800× on a run of the same byte), which is well under the 50 MB absolute cap
    but far over the 50× ratio ceiling.
    """
    path = tmp_path / 'ratio_bomb.xlsx'
    _make_zip(path, payload=b'A', repeat=200 * 1024)
    with pytest.raises(SecurityError, match='ratio|bomb'):
        _check_zip_safety(str(path), max_uncompressed_mb=50)


# ─── Blocked extensions ───────────────────────────────────────────────────────

@pytest.mark.parametrize('ext', ['.xlsm', '.xlsb', '.xls'])
def test_blocked_extension_rejected_by_validate_file(tmp_path, ext):
    """Macro-enabled and legacy Excel extensions (.xlsm, .xlsb, .xls) must be
    rejected by validate_file() before any attempt to open or decompress the
    file — matching SECURITY.md claim 5."""
    path = tmp_path / f'file{ext}'
    path.write_bytes(b'')
    with pytest.raises(SecurityError, match='not accepted|blocked'):
        validate_file(str(path))


# ─── Engine: data_only=True ───────────────────────────────────────────────────

def _make_xlsx(rows_or_cells, path, is_data=False):
    """Write a minimal .xlsx for testing.  Pass a list of rows for a pattern
    file or a dict of {coord: value} for a data file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    if is_data:
        for coord, val in rows_or_cells.items():
            ws[coord] = val
    else:
        for row in rows_or_cells:
            ws.append(row)
    wb.save(str(path))
    return str(path)


def test_engine_passes_data_only_true_to_openpyxl(tmp_path, monkeypatch):
    """Engine.process() must pass data_only=True when loading the data file so
    that formulas are never evaluated — matching SECURITY.md claim 'formulas in
    data sheets are not executed'.  Pattern files are loaded separately and are
    not subject to this requirement."""
    import grepxcel.engine as _engine

    calls = []
    real_load = _engine.openpyxl.load_workbook

    def recording_load(filename, **kwargs):
        calls.append({'filename': str(filename), 'data_only': kwargs.get('data_only')})
        return real_load(filename, **kwargs)

    monkeypatch.setattr(_engine.openpyxl, 'load_workbook', recording_load)

    pat = _make_xlsx(
        [['lbl:', 'inv', 'string', 'INV-001'],
         ['START:'], ['cell:next', 'num'], ['END:']],
        tmp_path / 'pattern.xlsx',
    )
    dat = _make_xlsx({'A1': 'INV-001', 'B1': '42'}, tmp_path / 'data.xlsx', is_data=True)
    dat_str = str(tmp_path / 'data.xlsx')

    Engine().process(pat, dat, logger=Logger(level=VerbosityLevel.QUIET))

    data_calls = [c for c in calls if c['filename'] == dat_str]
    assert data_calls, 'load_workbook was never called for the data file'
    assert all(c['data_only'] is True for c in data_calls), (
        f'Expected data_only=True for every data-file load; got {data_calls}'
    )
