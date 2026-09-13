"""Unit tests for the grepxcel web wizard FastAPI backend (wizard_api.py).

Uses FastAPI's built-in TestClient (sync wrapper around httpx) so no real
server or browser is needed.  Tests are skipped automatically when FastAPI /
uvicorn are not installed (same pattern as test_wizard_tui.py for Textual).
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
import pytest

# ── Optional import guard ──────────────────────────────────────────────────

try:
    from fastapi.testclient import TestClient
    from grepxcel.wizard_api import create_app, _STATE
    _API_OK = True
except ImportError:
    _API_OK = False

_skip_no_api = pytest.mark.skipif(
    not _API_OK,
    reason='fastapi / uvicorn not installed (pip install "grepxcel[web]")',
)

# ── Fixture helpers ────────────────────────────────────────────────────────

FIXTURE_01 = Path(__file__).parent.parent / 'fixtures/01_simple_invoice/01_simple_invoice_data.xlsx'
# manual pattern is used read-only; it has labels that actually match the data (preload works)
PATTERN_01 = Path(__file__).parent.parent / 'fixtures/01_simple_invoice/01_simple_invoice_pattern-manual.xlsx'


def _make_client(tmp_path: Path, cells: dict | None = None) -> 'TestClient':
    """Create a TestClient backed by a minimal xlsx workbook."""
    if cells is None:
        cells = {
            (1, 1): 'Invoice No:',
            (1, 2): 'AB123456',
            (1, 3): 'Date:',
            (1, 4): '2026-01-15',
            (2, 1): 'Amount:',
            (2, 2): 1234.56,
        }
    xlsx_path = str(tmp_path / 'test.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    for (r, c), v in cells.items():
        ws.cell(row=r, column=c, value=v)
    wb.save(xlsx_path)

    app = create_app(xlsx_path)
    return TestClient(app)


def _make_real_client() -> 'TestClient':
    """Create a TestClient backed by fixture 01 (real file)."""
    assert FIXTURE_01.exists(), f'Fixture not found: {FIXTURE_01}'
    app = create_app(str(FIXTURE_01))
    return TestClient(app)


def _make_real_client_with_pattern() -> 'TestClient':
    """Create a TestClient backed by fixture 01 data + pattern (for extract tests)."""
    assert FIXTURE_01.exists(), f'Data fixture not found: {FIXTURE_01}'
    assert PATTERN_01.exists(), f'Pattern fixture not found: {PATTERN_01}'
    app = create_app(str(FIXTURE_01), pattern_path=str(PATTERN_01))
    return TestClient(app)


# ── Tests: page + basic state ──────────────────────────────────────────────

@_skip_no_api
class TestIndexPage:
    def test_returns_html(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/')
        assert r.status_code == 200
        assert 'text/html' in r.headers['content-type']

    def test_html_contains_grid_element(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/')
        assert 'grid-table' in r.text

    def test_html_contains_sheet_name(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/')
        assert 'Sheet1' in r.text


@_skip_no_api
class TestSheetEndpoint:
    def test_returns_rows(self, tmp_path):
        client = _make_client(tmp_path)
        d = client.get('/api/sheet').json()
        assert 'rows' in d
        assert len(d['rows']) > 0

    def test_col_letters_present(self, tmp_path):
        client = _make_client(tmp_path)
        d = client.get('/api/sheet').json()
        assert 'col_letters' in d
        assert d['col_letters'][0] == 'A'

    def test_cells_have_ref_value_choice(self, tmp_path):
        client = _make_client(tmp_path)
        d = client.get('/api/sheet').json()
        cell = d['rows'][0][0]
        assert 'ref' in cell
        assert 'value' in cell
        assert 'choice' in cell

    def test_cell_has_name_field(self, tmp_path):
        """Cells must carry a 'name' field for Phase B provenance support."""
        client = _make_client(tmp_path)
        d = client.get('/api/sheet').json()
        cell = d['rows'][0][0]
        assert 'name' in cell  # may be empty string if unclassified

    def test_choice_empty_before_classification(self, tmp_path):
        client = _make_client(tmp_path)
        d = client.get('/api/sheet').json()
        cell = d['rows'][0][0]
        assert cell['choice'] == ''

    def test_merged_cells_have_colspan_rowspan(self, tmp_path):
        # Build a workbook with a 2-column merge in A1:B1
        xlsx_path = str(tmp_path / 'merged.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Merged Header'
        ws.merge_cells('A1:B1')
        wb.save(xlsx_path)
        app = create_app(xlsx_path)
        client = TestClient(app)
        d = client.get('/api/sheet').json()
        a1 = d['rows'][0][0]
        assert a1['ref'] == 'A1'
        assert a1['colspan'] == 2
        assert a1['rowspan'] == 1
        assert a1['merged'] is True

    def test_ghost_cells_have_skip_true(self, tmp_path):
        # B1 is the ghost cell in a A1:B1 merge — skip=True, absent from grid
        xlsx_path = str(tmp_path / 'merged2.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Header'
        ws.merge_cells('A1:B1')
        wb.save(xlsx_path)
        app = create_app(xlsx_path)
        client = TestClient(app)
        d = client.get('/api/sheet').json()
        b1 = d['rows'][0][1]
        assert b1['ref'] == 'B1'
        assert b1['skip'] is True


@_skip_no_api
class TestClassifyEndpoint:
    def test_classify_label(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'inv_label', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        assert r.status_code == 200
        assert r.json()['ok'] is True

    def test_classify_value(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'B1', 'action': 'V',
            'fields': {'name': 'invoice_no', 'type': 'string',
                       'match_mode': '(default)', 'match': r'.*', 'notes': ''},
        })
        assert r.status_code == 200
        assert r.json()['ok'] is True

    def test_classify_constant(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'C1', 'action': 'C',
            'fields': {'name': 'date_header'},
        })
        assert r.status_code == 200
        assert r.json()['ok'] is True

    def test_classify_ignore(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'A1', 'action': 'I', 'fields': {},
        })
        assert r.status_code == 200

    def test_classify_table(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'A1', 'action': 'T',
            'fields': {'name': 'items_table', 'notes': ''},
        })
        assert r.status_code == 200

    def test_invalid_action_rejected(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'A1', 'action': 'X', 'fields': {},
        })
        assert r.status_code == 400

    def test_clear_action(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'label', 'type': 'string',
                       'match_mode': '(default)', 'match': 'x', 'notes': ''},
        })
        r = client.post('/api/classify', json={'ref': 'A1', 'action': 'CLEAR', 'fields': {}})
        assert r.status_code == 200
        assert r.json()['action'] == 'CLEAR'
        # Cell should now be unclassified
        cell = client.get('/api/cell/A1').json()
        assert cell['choice'] == ''

    def test_ghost_cell_classification_rejected(self, tmp_path):
        # Build workbook with merge A1:B1; classifying B1 (ghost) must fail
        xlsx_path = str(tmp_path / 'merge_test.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Title'
        ws.merge_cells('A1:B1')
        wb.save(xlsx_path)
        app = create_app(xlsx_path)
        client = TestClient(app)
        r = client.post('/api/classify', json={
            'ref': 'B1', 'action': 'L',
            'fields': {'name': 'title', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Title', 'notes': ''},
        })
        assert r.status_code == 400
        assert 'merged' in r.json()['detail'].lower()


@_skip_no_api
class TestPreviewEndpoint:
    def test_returns_csv(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/api/preview')
        assert r.status_code == 200
        assert 'text/plain' in r.headers['content-type']

    def test_csv_has_config_header(self, tmp_path):
        """CSV output must start with config rows (e.g. config:,read.direction,LR)."""
        client = _make_client(tmp_path)
        text = client.get('/api/preview').text
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        # First row is a config row: 'config:', 'read.direction', 'LR'
        assert rows[0][0] == 'config:', f'unexpected first row: {rows[0]}'

    def test_classified_cell_appears_in_preview(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'inv_label', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        text = client.get('/api/preview').text
        assert 'inv_label' in text


@_skip_no_api
class TestUndoEndpoint:
    def test_undo_removes_last_classification(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'x', 'notes': ''},
        })
        client.post('/api/undo')
        cell = client.get('/api/cell/A1').json()
        assert cell['choice'] == ''

    def test_undo_on_empty_stack_is_noop(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/undo')
        assert r.status_code == 200  # must not crash

    def test_undo_multiple_steps(self, tmp_path):
        client = _make_client(tmp_path)
        # Classify A1 → L, then B1 → V
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'x', 'notes': ''},
        })
        client.post('/api/classify', json={
            'ref': 'B1', 'action': 'V',
            'fields': {'name': 'val', 'type': 'string',
                       'match_mode': '(default)', 'match': r'.*', 'notes': ''},
        })
        # Undo B1 → B1 unclassified, A1 still L
        client.post('/api/undo')
        assert client.get('/api/cell/B1').json()['choice'] == ''
        assert client.get('/api/cell/A1').json()['choice'] == 'L'
        # Undo A1 → both unclassified
        client.post('/api/undo')
        assert client.get('/api/cell/A1').json()['choice'] == ''


@_skip_no_api
class TestConfigEndpoint:
    def test_save_config_direction(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.post('/api/config', json={'direction': 'TD'})
        assert r.status_code == 200
        state = client.get('/api/state').json()
        assert state['config']['direction'] == 'TD'

    def test_save_config_ignore_case(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/config', json={'ignore_case_labels': False,
                                         'ignore_case_values': False})
        state = client.get('/api/state').json()
        assert state['config']['ignore_case_labels'] is False
        assert state['config']['ignore_case_values'] is False

    def test_save_config_currency(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/config', json={'currency_sign': '$'})
        state = client.get('/api/state').json()
        assert state['config']['currency_sign'] == '$'


@_skip_no_api
class TestCellEndpoint:
    def test_cell_value(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/api/cell/A1')
        assert r.status_code == 200
        d = r.json()
        assert d['ref'] == 'A1'
        assert 'Invoice No:' in d['value']

    def test_cell_classified_choice_returned(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        d = client.get('/api/cell/A1').json()
        assert d['choice'] == 'L'
        assert d['name'] == 'lbl'

    def test_invalid_ref_returns_404(self, tmp_path):
        client = _make_client(tmp_path)
        r = client.get('/api/cell/ZZZZZ9999999')
        assert r.status_code == 404


@_skip_no_api
class TestStatsEndpoint:
    def test_stats_via_state(self, tmp_path):
        client = _make_client(tmp_path)
        state = client.get('/api/state').json()
        stats = state['stats']
        assert 'total' in stats
        assert 'classified' in stats
        assert stats['classified'] == 0  # nothing classified yet

    def test_classified_count_increments(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'x', 'notes': ''},
        })
        stats = client.get('/api/state').json()['stats']
        assert stats['classified'] == 1
        assert stats['L'] == 1

    def test_classified_never_exceeds_total(self, tmp_path):
        """Classified percentage must never exceed 100%."""
        client = _make_client(tmp_path)
        # Classify several cells
        for ref in ['A1', 'B1', 'C1', 'D1', 'A2', 'B2']:
            client.post('/api/classify', json={
                'ref': ref, 'action': 'V',
                'fields': {'name': f'f_{ref.lower()}', 'type': 'string',
                           'match_mode': '(default)', 'match': r'.*', 'notes': ''},
            })
        stats = client.get('/api/state').json()['stats']
        assert stats['classified'] <= stats['total'], (
            f"classified ({stats['classified']}) > total ({stats['total']}) — stats > 100%"
        )

    def test_T_HEAD_T_DATA_counted_in_stats(self, tmp_path):
        """T-HEAD and T-DATA preloaded choices must appear in stats (not as unclassified)."""
        client = _make_client(tmp_path)
        # Manually inject T-HEAD / T-DATA into the state (simulates preload)
        from grepxcel.wizard_api import _STATE
        _STATE['choices']['A1'] = {'choice': 'T-HEAD', 'name': 'col_a'}
        _STATE['choices']['A2'] = {'choice': 'T-DATA', 'name': 'items'}
        stats = client.get('/api/state').json()['stats']
        assert stats['T_HEAD'] == 1
        assert stats['T_DATA'] == 1
        # Their refs count as classified, so classified ≥ 2
        assert stats['classified'] >= 2


@_skip_no_api
class TestSaveEndpoints:
    def test_save_csv_creates_file(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        r = client.post('/api/save', json={})
        assert r.status_code == 200
        assert r.json()['ok'] is True
        saved_to = r.json()['saved_to']
        assert Path(saved_to).exists()

    def test_save_xlsx_creates_file(self, tmp_path):
        client = _make_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'lbl', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        r = client.post('/api/save-xlsx', json={})
        assert r.status_code == 200
        assert r.json()['ok'] is True
        saved_to = r.json()['saved_to']
        assert Path(saved_to).exists()
        # Verify it's a valid xlsx (not just a renamed CSV)
        wb = openpyxl.load_workbook(saved_to)
        assert wb is not None


@_skip_no_api
class TestExtractEndpoint:
    """Tests for the POST /api/extract endpoint (Phase A)."""

    def test_no_pattern_returns_ok_false(self, tmp_path):
        """Without a pattern, extraction must return ok=False gracefully."""
        client = _make_client(tmp_path)  # no pattern_path
        r = client.post('/api/extract')
        assert r.status_code == 200
        data = r.json()
        assert data['ok'] is False
        assert 'error' in data
        assert 'pattern' in data['error'].lower()

    def test_with_real_pattern_returns_result(self):
        """With a real pattern, extraction must return ok=True and a non-empty result."""
        if not FIXTURE_01.exists() or not PATTERN_01.exists():
            pytest.skip('fixture 01 or pattern not found')
        client = _make_real_client_with_pattern()
        r = client.post('/api/extract')
        assert r.status_code == 200
        data = r.json()
        assert data['ok'] is True, f"Extraction failed: {data.get('error')}"
        assert isinstance(data['result'], dict)
        assert len(data['result']) > 0

    def test_extract_result_no_datetimes(self):
        """Result must be JSON-safe (no raw datetime objects)."""
        if not FIXTURE_01.exists() or not PATTERN_01.exists():
            pytest.skip('fixture 01 or pattern not found')
        client = _make_real_client_with_pattern()
        r = client.post('/api/extract')
        assert r.status_code == 200
        import json
        # If the response body parsed without error, datetimes are serialised
        data = json.loads(r.content)
        assert data['ok'] is True

    def test_extract_result_has_provenance(self):
        """When pattern is preloaded, provenance must map field names to cell refs."""
        if not FIXTURE_01.exists() or not PATTERN_01.exists():
            pytest.skip('fixture 01 or pattern not found')
        client = _make_real_client_with_pattern()
        r = client.post('/api/extract')
        data = r.json()
        assert data['ok'] is True
        prov = data.get('provenance', {})
        # Provenance should be non-empty since choices were preloaded from pattern
        assert isinstance(prov, dict)
        assert len(prov) > 0, 'expected provenance from preloaded pattern choices'

    def test_extract_provenance_values_are_cell_refs(self):
        """Each provenance entry must be a list of valid cell refs like ['A1', 'B3']."""
        if not FIXTURE_01.exists() or not PATTERN_01.exists():
            pytest.skip('fixture 01 or pattern not found')
        client = _make_real_client_with_pattern()
        data = client.post('/api/extract').json()
        assert data['ok'] is True
        import re
        ref_re = re.compile(r'^[A-Z]+\d+$')
        for name, refs in data['provenance'].items():
            assert isinstance(refs, list), f'provenance[{name!r}] should be a list'
            for ref in refs:
                assert ref_re.match(ref), f'provenance[{name!r}] contains invalid ref: {ref!r}'

    def test_extract_classified_never_exceeds_total_after_extract(self):
        """Stats total must remain ≥ classified even after extraction is run."""
        if not FIXTURE_01.exists() or not PATTERN_01.exists():
            pytest.skip('fixture 01 or pattern not found')
        client = _make_real_client_with_pattern()
        client.post('/api/extract')
        stats = client.get('/api/state').json()['stats']
        assert stats['classified'] <= stats['total'], (
            f"After extract: classified ({stats['classified']}) > total ({stats['total']})"
        )


@_skip_no_api
class TestRealFixture:
    def test_fixture01_sheet_loads(self):
        if not FIXTURE_01.exists():
            pytest.skip('fixture 01 not found')
        client = _make_real_client()
        d = client.get('/api/sheet').json()
        assert d['max_row'] >= 1
        assert d['max_col'] >= 1
        # Fixture 01 has Invoice No: in A1
        a1 = d['rows'][0][0]
        assert a1['value'] == 'Invoice No:'

    def test_fixture01_classify_and_preview(self):
        if not FIXTURE_01.exists():
            pytest.skip('fixture 01 not found')
        client = _make_real_client()
        client.post('/api/classify', json={
            'ref': 'A1', 'action': 'L',
            'fields': {'name': 'invoice_no_label', 'type': 'string',
                       'match_mode': '(default)', 'match': 'Invoice No:', 'notes': ''},
        })
        client.post('/api/classify', json={
            'ref': 'B1', 'action': 'V',
            'fields': {'name': 'invoice_no', 'type': 'string',
                       'match_mode': '(default)', 'match': r'[A-Z]{2}\d+',
                       'modifiers': 'not-null', 'notes': ''},
        })
        text = client.get('/api/preview').text
        assert 'invoice_no_label' in text
        assert 'invoice_no' in text
        assert 'not-null' in text


# ── Table Editor tests ─────────────────────────────────────────────────────────

FIXTURE_03_DATA = Path(__file__).parent.parent / 'fixtures/03_purchase_order/03_purchase_order_data.xlsx'
FIXTURE_03_PATTERN = Path(__file__).parent.parent / 'fixtures/03_purchase_order/03_purchase_order_pattern-from-local.csv'


def _make_table_client(tmp_path: Path) -> 'TestClient':
    """Workbook with a KV row + table that mirrors fixture 3 layout.

       Row 1: PO Number: | PO-2026 | Date: | 2026-05-01
       Row 3: Item | Description | Qty | Unit Price | Total  (header)
       Row 4: Laptop | Dell XPS 15 | 2 | 1200 | 2400
       Row 5: X | Monitor Stand | 4 | 45 | 180
       Row 6: Dock | USB-C Hub | 6 | 75 | 450
       Row 7: Grand Total | | | | 3030
    """
    xlsx_path = str(tmp_path / 'po.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['A1'] = 'PO Number:'; ws['B1'] = 'PO-2026'
    ws['C1'] = 'Date:';      ws['D1'] = '2026-05-01'
    ws['A3'] = 'Item';       ws['B3'] = 'Description'
    ws['C3'] = 'Qty';        ws['D3'] = 'Unit Price'; ws['E3'] = 'Total'
    ws['A4'] = 'Laptop';     ws['B4'] = 'Dell XPS 15'
    ws['C4'] = 2;            ws['D4'] = 1200;         ws['E4'] = 2400
    ws['A5'] = 'X';          ws['B5'] = 'Monitor Stand'
    ws['C5'] = 4;            ws['D5'] = 45;           ws['E5'] = 180
    ws['A6'] = 'Dock';       ws['B6'] = 'USB-C Hub'
    ws['C6'] = 6;            ws['D6'] = 75;           ws['E6'] = 450
    ws['A7'] = 'Grand Total';ws['E7'] = 3030
    wb.save(xlsx_path)
    app = create_app(xlsx_path)
    return TestClient(app)


_TABLE_ROW_CONFIGS = [
    # HEADER row: role L — each col matches actual spreadsheet column header text
    {'sheet_row': 3, 'row_type': 'header', 'row_n': 1, 'cols': [
        {'role': 'L', 'name': 'item',        'lmatch': 'Item'},
        {'role': 'L', 'name': 'description', 'lmatch': 'Description'},
        {'role': 'L', 'name': 'qty',         'lmatch': 'Qty'},
        {'role': 'L', 'name': 'unit_price',  'lmatch': 'Unit Price'},
        {'role': 'L', 'name': 'total',       'lmatch': 'Total'},
    ]},
    {'sheet_row': 4, 'row_type': 'data', 'cols': [
        {'role': 'V', 'name': 'item',        'ftype': 'string',   'match': '.*'},
        {'role': 'V', 'name': 'description', 'ftype': 'string',   'match': '.*'},
        {'role': 'V', 'name': 'qty',         'ftype': 'integer',  'match': r'\d+'},
        {'role': 'V', 'name': 'unit_price',  'ftype': 'currency', 'match': '.*'},
        {'role': 'V', 'name': 'total',       'ftype': 'currency', 'match': '.*'},
    ]},
    {'sheet_row': 5, 'row_type': 'data_inherited'},
    {'sheet_row': 6, 'row_type': 'data_inherited'},
    {'sheet_row': 7, 'row_type': 'skip', 'cols': [
        {'condition': 'label', 'lmatch': 'Grand Total'},
        {'condition': 'IGNORE'},
        {'condition': 'IGNORE'},
        {'condition': 'IGNORE'},
        {'condition': 'IGNORE'},
    ]},
]


@_skip_no_api
class TestTableRangeEndpoint:
    def test_returns_rows_and_cells(self, tmp_path):
        client = _make_table_client(tmp_path)
        r = client.get('/api/table-range/A3/E7')
        assert r.status_code == 200
        d = r.json()
        assert d['start_ref'] == 'A3'
        assert d['end_ref']   == 'E7'
        assert d['row_count'] == 5
        assert d['col_count'] == 5
        assert len(d['rows']) == 5

    def test_first_row_contains_header_values(self, tmp_path):
        client = _make_table_client(tmp_path)
        d = client.get('/api/table-range/A3/E7').json()
        first_row_values = [c['raw'] for c in d['rows'][0]['cells']]
        assert 'Item' in first_row_values
        assert 'Description' in first_row_values
        assert 'Qty' in first_row_values

    def test_cells_have_ref_and_choice(self, tmp_path):
        client = _make_table_client(tmp_path)
        d = client.get('/api/table-range/A3/E7').json()
        cell = d['rows'][0]['cells'][0]
        assert 'ref' in cell
        assert 'choice' in cell
        assert cell['ref'] == 'A3'

    def test_choice_reflects_existing_classification(self, tmp_path):
        client = _make_table_client(tmp_path)
        # Classify A3 as T anchor first
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        d = client.get('/api/table-range/A3/E7').json()
        a3_cell = d['rows'][0]['cells'][0]
        assert a3_cell['choice'] == 'T'

    def test_invalid_refs_return_400(self, tmp_path):
        client = _make_table_client(tmp_path)
        r = client.get('/api/table-range/ZZZZZ/QQQQQ')
        assert r.status_code == 400


@_skip_no_api
class TestClassifyTableFull:
    def test_anchor_cell_stored_as_T(self, tmp_path):
        client = _make_table_client(tmp_path)
        r = client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        assert r.status_code == 200
        cell = client.get('/api/cell/A3').json()
        assert cell['choice'] == 'T'

    def test_header_cells_marked_T_HEAD(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        for col_ref in ('B3', 'C3', 'D3', 'E3'):
            cell = client.get(f'/api/cell/{col_ref}').json()
            assert cell['choice'] == 'T-HEAD', f'{col_ref} should be T-HEAD'
            assert cell['anchor'] == 'A3'

    def test_data_cells_marked_T_DATA(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        for data_ref in ('A4', 'B4', 'C4', 'A5', 'A6'):
            cell = client.get(f'/api/cell/{data_ref}').json()
            assert cell['choice'] == 'T-DATA', f'{data_ref} should be T-DATA'
            assert cell['anchor'] == 'A3'

    def test_csv_contains_table_block(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        csv_text = client.get('/api/preview').text
        assert 'table:*' in csv_text
        assert 'HEADER:1' in csv_text
        assert 'DATA:*' in csv_text

    def test_csv_contains_column_names(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        csv_text = client.get('/api/preview').text
        assert 'line_items.item' in csv_text
        assert 'line_items.description' in csv_text
        assert 'line_items.qty' in csv_text

    def test_csv_contains_skip_if_for_grand_total(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        csv_text = client.get('/api/preview').text
        assert 'SKIP_IF' in csv_text
        assert 'Grand Total' in csv_text

    def test_reapply_clears_stale_cells(self, tmp_path):
        """Re-classifying a table with a different range clears old T-HEAD/T-DATA."""
        client = _make_table_client(tmp_path)
        # First classify with full range
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        # Re-classify with only 2 columns (smaller range)
        small_configs = [
            {'sheet_row': 3, 'row_type': 'header', 'row_n': 1, 'cols': [
                {'role': 'V', 'name': 'item', 'ftype': 'string', 'match': '.*'},
                {'role': 'V', 'name': 'description', 'ftype': 'string', 'match': '.*'},
            ]},
            {'sheet_row': 4, 'row_type': 'data', 'cols': [
                {'role': 'V', 'name': 'item', 'ftype': 'string', 'match': '.*'},
                {'role': 'V', 'name': 'description', 'ftype': 'string', 'match': '.*'},
            ]},
        ]
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'B6',
                       'row_configs': small_configs},
        })
        # C3 was T-HEAD before; after re-apply with smaller range it should be gone
        cell_c3 = client.get('/api/cell/C3').json()
        assert cell_c3['choice'] != 'T-HEAD', 'Stale T-HEAD at C3 should be cleared'


@_skip_no_api
class TestTableContext:
    def test_anchor_returns_context(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        ctx = client.get('/api/table-context/A3').json()
        assert ctx['anchor'] == 'A3'
        assert ctx['name']   == 'line_items'
        assert ctx['mult']   == '*'
        assert ctx['end_ref'] == 'E7'

    def test_thead_returns_anchor_context(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        ctx = client.get('/api/table-context/B3').json()
        assert ctx['anchor'] == 'A3'
        assert ctx['name']   == 'line_items'

    def test_tdata_returns_anchor_context(self, tmp_path):
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        ctx = client.get('/api/table-context/C5').json()
        assert ctx['anchor'] == 'A3'

    def test_web_row_configs_preserved_for_edit(self, tmp_path):
        """Round-trip: classify → context → _web_row_configs present."""
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        ctx = client.get('/api/table-context/A3').json()
        wrc = ctx.get('_web_row_configs', [])
        assert len(wrc) == len(_TABLE_ROW_CONFIGS)
        assert wrc[0]['row_type'] == 'header'
        assert wrc[4]['row_type'] == 'skip'

    def test_non_table_cell_returns_404(self, tmp_path):
        client = _make_table_client(tmp_path)
        r = client.get('/api/table-context/A1')
        assert r.status_code == 404


@_skip_no_api
class TestTableWithFooterRow:
    """Table that has an explicit FOOTER:1 row instead of a Skip row."""

    def test_csv_contains_footer(self, tmp_path):
        footer_configs = [
            {'sheet_row': 3, 'row_type': 'header', 'row_n': 1, 'cols': [
                {'role': 'L', 'name': 'item',  'lmatch': 'Item'},
                {'role': 'L', 'name': 'total', 'lmatch': 'Total'},
            ]},
            {'sheet_row': 4, 'row_type': 'data', 'cols': [
                {'role': 'V', 'name': 'item',  'ftype': 'string', 'match': '.*'},
                {'role': 'V', 'name': 'total', 'ftype': 'currency', 'match': '.*'},
            ]},
            {'sheet_row': 5, 'row_type': 'footer', 'row_n': 1, 'cols': [
                {'role': 'L', 'name': 'grand_total_lbl', 'ftype': 'string', 'lmatch': 'Grand Total'},
                {'role': 'V', 'name': 'grand_total',     'ftype': 'currency', 'match': '.*'},
            ]},
        ]
        cells = {
            (3, 1): 'Item',  (3, 2): 'Total',
            (4, 1): 'Laptop', (4, 2): 2400,
            (5, 1): 'Grand Total', (5, 2): 3030,
        }
        xlsx_path = str(tmp_path / 'footer_test.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        for (r, c), v in cells.items():
            ws.cell(row=r, column=c, value=v)
        wb.save(xlsx_path)

        from grepxcel.wizard_api import create_app
        app = create_app(xlsx_path)
        client = TestClient(app)

        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'sales', 'mult': '*', 'end_ref': 'B5',
                       'row_configs': footer_configs},
        })
        csv_text = client.get('/api/preview').text
        assert 'FOOTER:1' in csv_text
        assert 'Grand Total' in csv_text


@_skip_no_api
class TestTableScanEndpoint:
    def test_detects_columns_from_header_row(self, tmp_path):
        client = _make_table_client(tmp_path)
        d = client.get('/api/table-scan/A3').json()
        assert d['anchor'] == 'A3'
        col_headers = [c['header'] for c in d['columns']]
        assert 'Item' in col_headers
        assert 'Description' in col_headers
        assert 'Qty' in col_headers

    def test_stops_at_empty_column(self, tmp_path):
        client = _make_table_client(tmp_path)
        # Row 3 has A3:E3 filled, F3 is empty → stops at 5 columns
        d = client.get('/api/table-scan/A3').json()
        assert d['anchor'] == 'A3'
        assert len(d['columns']) == 5


@_skip_no_api
class TestTableIntegrationFixture3:
    """End-to-end: classify the fixture-3 layout via web wizard → generate CSV
    → parse with PatternParser → extract → verify field values."""

    def test_csv_parses_without_error(self, tmp_path):
        from grepxcel.pattern_parser import PatternParser
        client = _make_table_client(tmp_path)
        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E7',
                       'row_configs': _TABLE_ROW_CONFIGS},
        })
        csv_text = client.get('/api/preview').text
        csv_path = str(tmp_path / 'pattern.csv')
        Path(csv_path).write_text(csv_text, encoding='utf-8')
        cfg, defs, seq = PatternParser().parse(csv_path)
        assert cfg is not None
        assert any('line_items' in (d.name or '') for d in defs.values())

    def test_extraction_finds_table_rows(self, tmp_path):
        """Full round-trip: classify → CSV → extract → table rows present."""
        import grepxcel as gx
        from grepxcel.pattern_parser import PatternParser

        # Write the workbook
        xlsx_path = str(tmp_path / 'po.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A3'] = 'Item';     ws['B3'] = 'Description'
        ws['C3'] = 'Qty';      ws['D3'] = 'Unit Price'; ws['E3'] = 'Total'
        ws['A4'] = 'Laptop';   ws['B4'] = 'Dell XPS 15'
        ws['C4'] = 2;          ws['D4'] = 1200;         ws['E4'] = 2400
        ws['A5'] = 'Dock';     ws['B5'] = 'USB-C Hub'
        ws['C5'] = 6;          ws['D5'] = 75;           ws['E5'] = 450
        ws['A6'] = 'Grand Total'; ws['E6'] = 2850
        wb.save(xlsx_path)

        from grepxcel.wizard_api import create_app
        app = create_app(xlsx_path)
        client = TestClient(app)

        client.post('/api/classify', json={
            'ref': 'A3', 'action': 'T',
            'fields': {'name': 'line_items', 'mult': '*', 'end_ref': 'E6',
                       'row_configs': [
                           {'sheet_row': 3, 'row_type': 'header', 'row_n': 1, 'cols': [
                               {'role': 'L', 'name': 'item',        'lmatch': 'Item'},
                               {'role': 'L', 'name': 'description', 'lmatch': 'Description'},
                               {'role': 'L', 'name': 'qty',         'lmatch': 'Qty'},
                               {'role': 'L', 'name': 'unit_price',  'lmatch': 'Unit Price'},
                               {'role': 'L', 'name': 'total',       'lmatch': 'Total'},
                           ]},
                           {'sheet_row': 4, 'row_type': 'data', 'cols': [
                               {'role': 'V', 'name': 'item',        'ftype': 'string',   'match': '.*'},
                               {'role': 'V', 'name': 'description', 'ftype': 'string',   'match': '.*'},
                               {'role': 'V', 'name': 'qty',         'ftype': 'integer',  'match': '.*'},
                               {'role': 'V', 'name': 'unit_price',  'ftype': 'currency', 'match': '.*'},
                               {'role': 'V', 'name': 'total',       'ftype': 'currency', 'match': '.*'},
                           ]},
                           {'sheet_row': 5, 'row_type': 'data_inherited'},
                           {'sheet_row': 6, 'row_type': 'skip', 'cols': [
                               {'condition': 'label', 'lmatch': 'Grand Total'},
                               {'condition': 'IGNORE'}, {'condition': 'IGNORE'},
                               {'condition': 'IGNORE'}, {'condition': 'IGNORE'},
                           ]},
                       ]},
        })

        csv_text = client.get('/api/preview').text
        csv_path = str(tmp_path / 'pattern.csv')
        Path(csv_path).write_text(csv_text, encoding='utf-8')

        result = gx.extract(csv_path, xlsx_path, output_format='nested')
        # Table is keyed by the common var prefix 'line_items'
        # (no global lbl: defs for HEADER cols → key derived from DATA var prefix)
        tables = result.get('line_items', [])
        assert isinstance(tables, list), 'Expected list of table rows'
        assert len(tables) >= 1, 'Should have at least one table group'
        # Data rows are nested under the 'data' key within each group
        first_group = tables[0]
        data_rows = first_group.get('data', [])
        assert len(data_rows) >= 1, 'Should have at least one data row'
        # Grand Total row should be skipped by SKIP_IF (engine label match)
        items = [r.get('item') for r in data_rows]
        assert 'Grand Total' not in items, 'Grand Total should be skipped by SKIP_IF'
