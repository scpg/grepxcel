"""Unit tests for the xlsx Audit sheet (grepxcel.xlsx_writer._add_audit_sheet)."""

import os
import tempfile

import pytest
import openpyxl

from grepxcel.xlsx_writer import nested_to_xlsx
from grepxcel.logger import Logger, VerbosityLevel, Severity, Category, LogRecord


# ── Helpers ───────────────────────────────────────────────────────────────────

def _make_logger(records: list[dict]) -> Logger:
    """Create a Logger with pre-populated _records."""
    logger = Logger(level=VerbosityLevel.QUIET)
    for r in records:
        rec = LogRecord(
            severity=r.get('severity', Severity.INFO),
            category=r.get('category', Category.EXTRACTION),
            message=r.get('message', ''),
            location=r.get('location', ''),
            field=r.get('field', ''),
            hint=r.get('hint', ''),
        )
        logger._records.append(rec)
    return logger


def _open_audit(tmp_path, result, logger):
    """Write xlsx and return the Audit sheet, or None if not present."""
    path = tmp_path / 'out.xlsx'
    nested_to_xlsx(result, str(path), logger=logger)
    wb = openpyxl.load_workbook(str(path))
    return wb['Audit'] if 'Audit' in wb.sheetnames else None


SIMPLE_RESULT = {
    'inv_number': 'INV-001',
    'inv_total': '110',
}


# ── Audit sheet presence ──────────────────────────────────────────────────────

class TestAuditSheetPresence:
    def test_no_logger_no_audit_sheet(self, tmp_path):
        path = tmp_path / 'out.xlsx'
        nested_to_xlsx(SIMPLE_RESULT, str(path), logger=None)
        wb = openpyxl.load_workbook(str(path))
        assert 'Audit' not in wb.sheetnames

    def test_logger_produces_audit_sheet(self, tmp_path):
        logger = _make_logger([
            {'field': 'inv_number', 'location': 'Sheet1!B3',
             'message': "inv_number = 'INV-001'"},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        assert ws is not None

    def test_extraction_sheet_still_present(self, tmp_path):
        logger = Logger(level=VerbosityLevel.QUIET)
        path = tmp_path / 'out.xlsx'
        nested_to_xlsx(SIMPLE_RESULT, str(path), logger=logger)
        wb = openpyxl.load_workbook(str(path))
        assert 'Extraction' in wb.sheetnames
        assert 'Audit' in wb.sheetnames


# ── Header row ────────────────────────────────────────────────────────────────

class TestAuditHeader:
    def test_header_columns(self, tmp_path):
        logger = Logger(level=VerbosityLevel.QUIET)
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
        assert headers == ['Field', 'Source Cell', 'Value', 'Status', 'Details']


# ── Extracted fields appear in audit ─────────────────────────────────────────

class TestAuditFieldRows:
    def test_clean_field_appears(self, tmp_path):
        logger = _make_logger([
            {'field': 'inv_number', 'location': 'Sheet1!B3',
             'category': 'EXTRACTION', 'severity': 'INFO',
             'message': "inv_number = 'INV-001'"},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        # Find the row with 'inv_number' in column 1
        field_row = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == 'inv_number':
                field_row = row
                break
        assert field_row is not None, 'inv_number row not found'
        assert field_row[1] == 'Sheet1!B3'     # source cell
        assert field_row[3] == 'Clean'          # status

    def test_warning_field_status(self, tmp_path):
        logger = _make_logger([
            {'field': 'inv_total', 'location': 'Sheet1!B5',
             'category': 'EXTRACTION', 'severity': 'INFO',
             'message': "inv_total = '110'"},
            {'field': 'inv_total', 'location': 'Sheet1!B5',
             'category': 'VALIDATION', 'severity': 'WARNING',
             'message': 'value does not match expected pattern'},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        field_row = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == 'inv_total':
                field_row = row
                break
        assert field_row is not None
        assert field_row[3] == 'Warning'

    def test_missing_field_status(self, tmp_path):
        logger = _make_logger([
            {'field': 'inv_number', 'location': '',
             'category': 'VALIDATION', 'severity': 'WARNING',
             'message': 'field is missing or empty'},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        field_row = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == 'inv_number':
                field_row = row
                break
        assert field_row is not None
        assert field_row[3] == 'Missing'

    def test_empty_logger_empty_audit_body(self, tmp_path):
        """No records → no data rows (just header + maybe legend)."""
        logger = Logger(level=VerbosityLevel.QUIET)
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        # Check no field data rows (row 2 should be empty or contain legend)
        row2_vals = [ws.cell(row=2, column=c).value for c in range(1, 6)]
        # All None means no data row was written
        assert all(v is None for v in row2_vals)


# ── Assert failures section ───────────────────────────────────────────────────

class TestAuditAssertSection:
    def test_assert_failure_appears(self, tmp_path):
        logger = _make_logger([
            {'field': '',
             'category': 'VALIDATION', 'severity': 'WARNING',
             'message': 'Assertion failed: total == net + vat',
             'hint': "Expression: 'total == net + vat' evaluated to False"},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        # Find any row with 'Assertion failed' in the Details column (col 5)
        found = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[4] and 'Assertion failed' in str(row[4]):
                found = True
                break
        assert found, 'Assert failure row not found in Audit sheet'


# ── Legend ────────────────────────────────────────────────────────────────────

class TestAuditLegend:
    def test_legend_present(self, tmp_path):
        logger = _make_logger([
            {'field': 'f', 'category': 'EXTRACTION', 'severity': 'INFO',
             'message': "f = 'val'", 'location': 'A1'},
        ])
        ws = _open_audit(tmp_path, SIMPLE_RESULT, logger)
        # Find 'Legend' somewhere in column A
        legend_found = any(
            ws.cell(row=r, column=1).value == 'Legend'
            for r in range(1, ws.max_row + 1)
        )
        assert legend_found


# ── Formula injection prevention ──────────────────────────────────────────────

class TestAuditSecurity:
    def test_formula_in_field_value_neutralized(self, tmp_path):
        evil_result = {'amount': '=CMD|"/c calc"!A0'}
        logger = _make_logger([
            {'field': 'amount', 'location': 'Sheet1!C2',
             'category': 'EXTRACTION', 'severity': 'INFO',
             'message': "amount = '=CMD|\"/c calc\"!A0'"},
        ])
        ws = _open_audit(tmp_path, evil_result, logger)
        # Find amount row; value cell (col 3) must not start with '='
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == 'amount':
                val = str(row[2]) if row[2] else ''
                assert not val.startswith('='), f'Formula not neutralized: {val}'
                break
