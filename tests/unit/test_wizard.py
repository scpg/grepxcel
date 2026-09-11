"""Unit tests for the grepxcel wizard command."""

import csv
import datetime
import io
import json
import os

import openpyxl
import pytest

from grepxcel.wizard import (
    WizardState,
    _propose_type,
    _save_state_json,
    _slugify,
    _col_label,
    _write_pattern,
    run_wizard,
)


# ── _propose_type ─────────────────────────────────────────────────────────────

class TestProposeType:
    def test_none_is_skip(self):
        assert _propose_type(None) == 'skip'

    def test_empty_string_is_skip(self):
        assert _propose_type('') == 'skip'

    def test_blank_string_is_skip(self):
        assert _propose_type('   ') == 'skip'

    # ── Colon suffix → label (scalar label convention) ──────────────────────

    def test_colon_suffix_is_label(self):
        assert _propose_type('Invoice No:') == 'label'

    def test_colon_suffix_multiword_is_label(self):
        assert _propose_type('Account Holder:') == 'label'

    # ── Single-word pure alpha → label (column header default) ─────────────

    def test_single_word_alpha_is_label(self):
        assert _propose_type('Description') == 'label'

    def test_single_word_alpha_short_is_label(self):
        assert _propose_type('SKU') == 'label'

    # ── Alpha+digit code → var:string ────────────────────────────────────────

    def test_alpha_digit_code_is_var_string(self):
        assert _propose_type('AB123456') == 'var:string'

    def test_alpha_digit_mixed_code_is_var_string(self):
        assert _propose_type('ELC001') == 'var:string'

    # ── Email → var:string ────────────────────────────────────────────────────

    def test_email_is_var_string(self):
        assert _propose_type('alice@wonderland.example') == 'var:string'

    # ── Multi-word pure alpha → var:string (name, description, title) ────────

    def test_multi_word_pure_alpha_is_var_string(self):
        assert _propose_type('Alice Wonderland') == 'var:string'

    def test_multi_word_title_no_colon_is_var_string(self):
        assert _propose_type('Invoice Number') == 'var:string'

    def test_long_description_is_var_string(self):
        assert _propose_type('This is a very long description that exceeds forty chars') == 'var:string'

    # ── Multi-word with digit → var:string (period text, range) ─────────────

    def test_multi_word_with_digit_is_var_string(self):
        assert _propose_type('Q1 2026') == 'var:string'

    def test_multi_word_month_year_is_var_string(self):
        assert _propose_type('January 2026') == 'var:string'

    # ── Left-neighbour context: single-word value after a colon-label ────────

    def test_single_word_after_colon_label_neighbour_is_var_string(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Department:'
        ws['B1'] = 'Engineering'
        assert _propose_type('Engineering', ws=ws, row=1, col=2) == 'var:string'

    def test_single_word_without_colon_neighbour_stays_label(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Product'   # no colon → neighbour heuristic does NOT fire
        ws['B1'] = 'SKU'
        assert _propose_type('SKU', ws=ws, row=1, col=2) == 'label'

    # ── Numeric types ────────────────────────────────────────────────────────

    def test_integer_is_var_integer(self):
        assert _propose_type(42) == 'var:integer'

    def test_whole_float_is_var_integer(self):
        assert _propose_type(3.0) == 'var:integer'

    def test_fractional_float_is_var_currency(self):
        assert _propose_type(3.14) == 'var:currency'

    def test_date_is_var_date(self):
        assert _propose_type(datetime.date(2026, 1, 1)) == 'var:date'

    def test_datetime_is_var_datetime(self):
        assert _propose_type(datetime.datetime(2026, 1, 1, 12, 0)) == 'var:datetime'


# ── _slugify ──────────────────────────────────────────────────────────────────

class TestSlugify:
    def test_lower_and_replace_spaces(self):
        assert _slugify('Invoice Number') == 'invoice_number'

    def test_strips_special_chars(self):
        assert _slugify('Total (€)') == 'total'

    def test_already_slug(self):
        assert _slugify('amount_net') == 'amount_net'


# ── _col_label ────────────────────────────────────────────────────────────────

class TestColLabel:
    def test_col_1_is_A(self):
        assert _col_label(1) == 'A'

    def test_col_26_is_Z(self):
        assert _col_label(26) == 'Z'

    def test_col_27_is_AA(self):
        assert _col_label(27) == 'AA'

    def test_col_28_is_AB(self):
        assert _col_label(28) == 'AB'


# ── _write_pattern ────────────────────────────────────────────────────────────

class TestWritePattern:
    def _read_csv(self, path):
        with open(path, newline='', encoding='utf-8') as f:
            return list(csv.reader(f))

    def test_config_rows_written(self, tmp_path):
        state = WizardState(direction='LR', currency_sign='€')
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['config:', 'read.direction', 'LR'] in rows
        assert ['config:', 'currency.sign', '€'] in rows

    def test_ignore_case_written_when_true(self, tmp_path):
        state = WizardState(direction='TD', ignore_case=True, currency_sign='$')
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['config:', 'ignore.case', 'yes'] in rows

    def test_ignore_case_absent_when_false(self, tmp_path):
        state = WizardState(ignore_case=False)
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert not any(r[:2] == ['config:', 'ignore.case'] for r in rows)

    def test_lbl_def_written(self, tmp_path):
        state = WizardState()
        state.lbl_defs.append(('inv_number_lbl', 'string', 'Invoice Number'))
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['lbl:', 'inv_number_lbl', 'string', 'Invoice Number'] in rows

    def test_var_def_written(self, tmp_path):
        state = WizardState()
        state.var_defs.append(('inv.number', 'string', '[A-Z]{2}[0-9]+'))
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['var:', 'inv.number', 'string', '[A-Z]{2}[0-9]+'] in rows

    def test_start_end_always_present(self, tmp_path):
        state = WizardState()
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        first_col = [r[0] for r in rows]
        assert 'START:' in first_col
        assert 'END:' in first_col

    def test_scalar_cell_instruction(self, tmp_path):
        state = WizardState()
        state.body_rows.append(['cell:1', 'inv.number'])
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['cell:1', 'inv.number'] in rows

    def test_table_block_written(self, tmp_path):
        state = WizardState()
        state.body_rows.append(['table:*'])
        state.body_rows.append(['', 'HEADER:1', 'item.name', 'item.qty'])
        state.body_rows.append(['', 'DATA:*', 'item.name', 'item.qty'])
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['table:*'] in rows
        assert ['', 'HEADER:1', 'item.name', 'item.qty'] in rows
        assert ['', 'DATA:*', 'item.name', 'item.qty'] in rows

    def test_footer_row_written(self, tmp_path):
        state = WizardState()
        state.body_rows.append(['table:1'])
        state.body_rows.append(['', 'HEADER:1', 'label', 'total'])
        state.body_rows.append(['', 'DATA:*', 'label', 'total'])
        state.body_rows.append(['', 'FOOTER:1', 'footer.label', 'footer.total'])
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        rows = self._read_csv(out)
        assert ['', 'FOOTER:1', 'footer.label', 'footer.total'] in rows

    def test_output_is_valid_utf8_csv(self, tmp_path):
        state = WizardState(currency_sign='€')
        state.lbl_defs.append(('lbl', 'string', 'Értéke'))
        out = str(tmp_path / 'p.csv')
        _write_pattern(state, out)
        content = open(out, encoding='utf-8').read()
        assert 'Értéke' in content

    # B1: xlsx output must be a real openpyxl workbook, not a CSV-inside-xlsx
    def test_xlsx_output_is_real_workbook(self, tmp_path):
        state = WizardState(direction='LR', currency_sign='$')
        state.lbl_defs.append(('inv_lbl', 'string', 'Invoice:'))
        state.var_defs.append(('inv.number', 'string', '[A-Z]+[0-9]+'))
        state.body_rows.append(['cell:1', 'inv_lbl'])
        state.body_rows.append(['cell:1', 'inv.number'])
        out = str(tmp_path / 'pattern.xlsx')
        _write_pattern(state, out)

        # Must be a valid xlsx (openpyxl can open it)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        values = [[c.value for c in row] for row in ws.iter_rows()]
        wb.close()

        first_col = [row[0] for row in values if row]
        assert 'config:' in first_col
        assert 'lbl:' in first_col
        assert 'var:' in first_col
        assert 'START:' in first_col
        assert 'END:' in first_col

    def test_xlsx_contains_lbl_row(self, tmp_path):
        state = WizardState()
        state.lbl_defs.append(('total_lbl', 'string', 'Total:'))
        out = str(tmp_path / 'p.xlsx')
        _write_pattern(state, out)
        wb = openpyxl.load_workbook(out)
        rows = [[c.value for c in r] for r in wb.active.iter_rows()]
        wb.close()
        assert any(r[0] == 'lbl:' and r[1] == 'total_lbl' and r[3] == 'Total:'
                   for r in rows)

    def test_xlsx_contains_var_row(self, tmp_path):
        state = WizardState()
        state.var_defs.append(('amount', 'currency', r'\d+\.\d{2}'))
        out = str(tmp_path / 'p.xlsx')
        _write_pattern(state, out)
        wb = openpyxl.load_workbook(out)
        rows = [[c.value for c in r] for r in wb.active.iter_rows()]
        wb.close()
        assert any(r[0] == 'var:' and r[1] == 'amount' and r[2] == 'currency'
                   for r in rows)


# ── WizardState defaults ──────────────────────────────────────────────────────

class TestWizardState:
    def test_defaults(self):
        s = WizardState()
        assert s.direction == 'LR'
        assert s.ignore_case is False
        assert s.currency_sign == '€'
        assert s.lbl_defs == []
        assert s.var_defs == []

    # ── to_dict / from_dict (JSON round-trip) ────────────────────────────────

    def test_to_dict_is_json_serialisable(self):
        s = WizardState(direction='TD', currency_sign='$', ignore_case=True)
        s.lbl_defs.append(('inv_lbl', 'string', 'Invoice:'))
        s.var_defs.append(('inv.number', 'integer', r'\d+'))
        s.body_rows.append(['cell:1', 'inv.number'])
        d = s.to_dict()
        # Must round-trip through JSON without error
        text = json.dumps(d)
        assert '"direction"' in text
        assert '"inv_lbl"' in text
        assert '"inv.number"' in text

    def test_from_dict_restores_all_fields(self):
        original = WizardState(
            direction='TD',
            ignore_case=True,
            currency_sign='$',
            sheet_name='Sheet2',
        )
        original.lbl_defs.append(('dept_lbl', 'string', 'Department:', ''))
        original.var_defs.append(('dept', 'string', r'[A-Z]+', ''))
        original.body_rows.append(['cell:1', 'dept'])

        restored = WizardState.from_dict(original.to_dict())
        assert restored.direction == 'TD'
        assert restored.ignore_case is True
        assert restored.currency_sign == '$'
        assert restored.sheet_name == 'Sheet2'
        assert restored.lbl_defs == [('dept_lbl', 'string', 'Department:', '')]
        assert restored.var_defs == [('dept', 'string', r'[A-Z]+', '')]
        assert restored.body_rows == [['cell:1', 'dept']]

    def test_from_dict_uses_defaults_for_missing_keys(self):
        restored = WizardState.from_dict({})
        assert restored.direction == 'LR'
        assert restored.ignore_case is False
        assert restored.currency_sign == '€'
        assert restored.sheet_name is None
        assert restored.lbl_defs == []
        assert restored.var_defs == []
        assert restored.body_rows == []

    def test_round_trip_via_json_string(self):
        s = WizardState(direction='LR', currency_sign='£')
        s.lbl_defs.append(('total_lbl', 'currency', 'Total:', ''))
        s.body_rows.append(['table:*'])
        restored = WizardState.from_dict(json.loads(json.dumps(s.to_dict())))
        assert restored.direction == 'LR'
        assert restored.currency_sign == '£'
        assert restored.lbl_defs == [('total_lbl', 'currency', 'Total:', '')]
        assert restored.body_rows == [['table:*']]


# ── _save_state_json ──────────────────────────────────────────────────────────

class TestSaveStateJson:
    def test_writes_valid_json(self, tmp_path):
        state = WizardState(direction='TD', currency_sign='$')
        state.lbl_defs.append(('lbl', 'string', 'Invoice:'))
        path = str(tmp_path / 'state.json')
        _save_state_json(state, path)
        with open(path, encoding='utf-8') as fh:
            d = json.load(fh)
        assert d['direction'] == 'TD'
        assert d['currency_sign'] == '$'
        assert d['lbl_defs'] == [['lbl', 'string', 'Invoice:']]

    def test_round_trip_preserves_body_rows(self, tmp_path):
        state = WizardState()
        state.body_rows.append(['table:*'])
        state.body_rows.append(['', 'HEADER:1', 'col_a'])
        state.body_rows.append(['', 'DATA:*', 'col_a'])
        path = str(tmp_path / 'state.json')
        _save_state_json(state, path)
        with open(path, encoding='utf-8') as fh:
            restored = WizardState.from_dict(json.load(fh))
        assert restored.body_rows == state.body_rows


# ── run_wizard --load-state / --save-state ────────────────────────────────────

FIXTURE_DIR = os.path.join(os.path.dirname(__file__), '..', 'fixtures')


def _make_state_json(tmp_path, name='invoice') -> str:
    """Write a minimal WizardState JSON file and return its path."""
    state = WizardState(direction='LR', currency_sign='€')
    state.lbl_defs.append(('inv_lbl', 'string', 'Invoice No:', ''))
    state.var_defs.append(('inv.number', 'string', r'[A-Z]+\d+', ''))
    state.body_rows.append(['cell:1', 'inv_lbl'])
    state.body_rows.append(['cell:1', 'inv.number'])
    path = str(tmp_path / f'{name}.json')
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(state.to_dict(), fh, indent=2)
    return path


class TestRunWizardLoadState:
    """run_wizard(load_state=...) converts JSON → pattern without an Excel file."""

    def test_writes_csv_pattern(self, tmp_path):
        state_path = _make_state_json(tmp_path)
        out = str(tmp_path / 'out.csv')
        rc = run_wizard(data_file=None, load_state=state_path, output=out)
        assert rc == 0
        assert os.path.exists(out)
        with open(out, newline='', encoding='utf-8') as fh:
            rows = list(csv.reader(fh))
        first_col = [r[0] for r in rows]
        assert 'lbl:' in first_col
        assert 'var:' in first_col
        assert 'START:' in first_col
        assert 'END:' in first_col

    def test_lbl_and_var_rows_match_state(self, tmp_path):
        state_path = _make_state_json(tmp_path)
        out = str(tmp_path / 'out.csv')
        run_wizard(data_file=None, load_state=state_path, output=out)
        with open(out, newline='', encoding='utf-8') as fh:
            rows = list(csv.reader(fh))
        assert ['lbl:', 'inv_lbl', 'string', 'Invoice No:'] in rows
        assert any(r[0] == 'var:' and r[1] == 'inv.number' for r in rows)

    def test_writes_xlsx_pattern_when_output_is_xlsx(self, tmp_path):
        state_path = _make_state_json(tmp_path)
        out = str(tmp_path / 'out.xlsx')
        rc = run_wizard(data_file=None, load_state=state_path, output=out)
        assert rc == 0
        wb = openpyxl.load_workbook(out)
        rows = [[c.value for c in r] for r in wb.active.iter_rows()]
        wb.close()
        first_col = [r[0] for r in rows if r]
        assert 'lbl:' in first_col
        assert 'var:' in first_col

    def test_default_output_name_derived_from_state_file(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        state_path = _make_state_json(tmp_path, name='mystate')
        rc = run_wizard(data_file=None, load_state=state_path)
        assert rc == 0
        assert (tmp_path / 'pattern-mystate.csv').exists()

    def test_error_on_missing_state_file(self, tmp_path):
        rc = run_wizard(data_file=None, load_state=str(tmp_path / 'no_such.json'))
        assert rc == 1

    def test_error_on_invalid_json(self, tmp_path):
        bad = str(tmp_path / 'bad.json')
        with open(bad, 'w') as fh:
            fh.write('not json {{{')
        rc = run_wizard(data_file=None, load_state=bad)
        assert rc == 1


class TestRunWizardSaveState:
    """run_wizard(save_state=...) writes state JSON after the pattern is saved."""

    def test_save_state_written_alongside_pattern(self, tmp_path):
        state_path = _make_state_json(tmp_path)
        out_pattern = str(tmp_path / 'out.csv')
        out_state   = str(tmp_path / 'saved.json')
        rc = run_wizard(data_file=None, load_state=state_path,
                        output=out_pattern, save_state=out_state)
        assert rc == 0
        assert os.path.exists(out_state)
        with open(out_state, encoding='utf-8') as fh:
            d = json.load(fh)
        assert 'lbl_defs' in d
        assert 'var_defs' in d

    def test_save_state_is_valid_round_trip(self, tmp_path):
        state_path = _make_state_json(tmp_path)
        out_pattern = str(tmp_path / 'out.csv')
        out_state   = str(tmp_path / 'saved.json')
        run_wizard(data_file=None, load_state=state_path,
                   output=out_pattern, save_state=out_state)
        with open(out_state, encoding='utf-8') as fh:
            restored = WizardState.from_dict(json.load(fh))
        assert restored.lbl_defs == [('inv_lbl', 'string', 'Invoice No:', '')]
        assert any(t[0] == 'inv.number' for t in restored.var_defs)
