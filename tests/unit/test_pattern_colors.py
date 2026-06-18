"""Tests for grepxcel pattern_colors — colorize pattern xlsx files."""

import openpyxl
from openpyxl.styles import PatternFill
import pytest

from grepxcel.pattern_colors import colorize_pattern_file, ROW_FILLS


# ── helpers ───────────────────────────────────────────────────────────────────

def _make_pattern(tmp_path, rows, name='pattern.xlsx'):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    path = str(tmp_path / name)
    wb.save(path)
    return path


def _fill_color(ws, row_num):
    """Return the fgColor hex of row's first cell fill, or None."""
    fill = ws.cell(row=row_num, column=1).fill
    if fill.patternType == 'solid' and fill.fgColor and fill.fgColor.rgb:
        return fill.fgColor.rgb
    return None


def _any_fill_in_row(ws, row_num, max_col=6):
    """Return True if any cell in the row has a solid fill."""
    for c in range(1, max_col + 1):
        fill = ws.cell(row=row_num, column=c).fill
        if fill.patternType == 'solid' and fill.fgColor and fill.fgColor.rgb:
            rgb = fill.fgColor.rgb
            if rgb != '00000000':
                return True
    return False


# ── core behavior ─────────────────────────────────────────────────────────────

class TestColorizePatternFile:
    def test_config_row_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['config:', 'read.direction', 'LR'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_lbl_row_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['lbl:', 'inv_label', 'string', 'Invoice No:'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_var_row_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['var:', 'inv.number', 'string', r'INV-\d+'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_start_end_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['START:'],
            ['cell:next', 'inv_label'],
            ['END:'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)
        assert _any_fill_in_row(wb.active, 3)

    def test_comment_row_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['# This is a comment'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_table_keyword_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['table:*'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_dir_row_colored(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['dir:TD'],
            ['dir:LR'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)
        assert _any_fill_in_row(wb.active, 2)

    def test_uppercase_cell_keyword_colored(self, tmp_path):
        """A CELL: row (uppercase keyword) must be colored like cell:."""
        path = _make_pattern(tmp_path, [
            ['CELL:J59', 'totals.x'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        assert _any_fill_in_row(wb.active, 1)

    def test_template_row_colored(self, tmp_path):
        """HEADER/DATA/FOOTER/SKIP_IF rows (column A blank, keyword in B)."""
        path = _make_pattern(tmp_path, [
            [None, 'HEADER:1', 'col_a', 'col_b'],
            [None, 'DATA:*', 'line.a', 'line.b'],
            [None, 'SKIP_IF', 'EMPTY', 'IGNORE'],
            [None, 'FOOTER:1', 'total_label', 'inv.total'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        for row in range(1, 5):
            assert _any_fill_in_row(wb.active, row), f'row {row} should be colored'

    def test_different_types_get_different_colors(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['lbl:', 'x', 'string', '.*'],
            ['var:', 'y', 'string', '.*'],
            ['config:', 'ignore.case', 'yes'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        colors = [_fill_color(wb.active, r) for r in range(1, 4)]
        assert len(set(colors)) >= 3, f'Expected at least 3 distinct colors, got {colors}'


class TestContentPreserved:
    def test_content_unchanged(self, tmp_path):
        rows = [
            ['config:', 'read.direction', 'LR'],
            ['lbl:', 'inv_label', 'string', 'Invoice No:'],
            ['var:', 'inv.number', 'string', r'INV-\d+'],
            ['START:'],
            ['cell:next', 'inv_label'],
            ['cell:next', 'inv.number'],
            ['END:'],
        ]
        path = _make_pattern(tmp_path, rows)
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                assert ws.cell(row=r_idx, column=c_idx).value == val


class TestColumnWidths:
    def test_columns_auto_widened(self, tmp_path):
        path = _make_pattern(tmp_path, [
            ['var:', 'very_long_field_name_here', 'string', r'some-pattern'],
        ])
        colorize_pattern_file(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        col_b_width = ws.column_dimensions['B'].width
        assert col_b_width is not None and col_b_width > 8


class TestReturnValue:
    def test_returns_path(self, tmp_path):
        path = _make_pattern(tmp_path, [['var:', 'x', 'string', '.*']])
        result = colorize_pattern_file(path)
        assert result == path
