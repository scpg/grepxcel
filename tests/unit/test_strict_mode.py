"""Tests for --strict mode (exit non-zero on missing fields)."""

import json
import os
import shutil

import openpyxl
import pytest

from grepxcel.cli import main as cli_main
from tests.conftest import find_data_file, find_pattern_xlsx

FIXTURES = os.path.join(os.path.dirname(__file__), '..', 'fixtures')


def _fixture(*parts):
    return os.path.join(FIXTURES, *parts)


def _make_mismatched_data(tmp_path):
    """Data file where the 'customer' cell has wrong format (triggers warning)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Invoice Number'
    ws['B1'] = 'AB123456'
    ws['C1'] = 'Customer Email'
    ws['D1'] = 'not-an-email'
    wb.save(tmp_path / 'data.xlsx')
    return str(tmp_path / 'data.xlsx')


def _make_pattern_expecting_email(tmp_path):
    """Pattern that expects an email — will warn on the data above."""
    dst = tmp_path / 'pattern.csv'
    dst.write_text(
        'var:,inv.number,string,.+\n'
        'var:,client.email,string,.+@.+\n'
        'START:,,,\n'
        'cell:1,IGNORE,,\n'
        'cell:1,inv.number,,\n'
        'cell:1,IGNORE,,\n'
        'cell:1,client.email,,\n'
        'END:,,,\n',
        encoding='utf-8',
    )
    return str(dst)


class TestStrictMode:
    """--strict causes exit 2 when extraction has issues."""

    def test_strict_passes_when_all_fields_populated(self, tmp_path):
        pattern = find_pattern_xlsx(_fixture('01_simple_invoice'))
        data = find_data_file(_fixture('01_simple_invoice'))
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit) as exc_info:
            cli_main([
                'extract', '-p', pattern, data,
                '--strict', '-o', out_dir,
            ])
        assert exc_info.value.code == 0

    def test_strict_fails_on_validation_warning(self, tmp_path):
        pattern = _make_pattern_expecting_email(tmp_path)
        data = _make_mismatched_data(tmp_path)
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit) as exc_info:
            cli_main([
                'extract', '-p', pattern, data,
                '--strict', '-o', out_dir,
            ])
        assert exc_info.value.code == 2

    def test_strict_reports_field_names_on_stderr(self, tmp_path, capsys):
        pattern = _make_pattern_expecting_email(tmp_path)
        data = _make_mismatched_data(tmp_path)
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit):
            cli_main([
                'extract', '-p', pattern, data,
                '--strict', '-o', out_dir,
            ])
        stderr = capsys.readouterr().err
        assert 'client.email' in stderr
        assert '--strict' in stderr

    def test_without_strict_exits_one_not_two(self, tmp_path):
        """Without --strict, a validation warning exits 1 (not 2)."""
        pattern = _make_pattern_expecting_email(tmp_path)
        data = _make_mismatched_data(tmp_path)
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit) as exc_info:
            cli_main([
                'extract', '-p', pattern, data,
                '-o', out_dir,
            ])
        assert exc_info.value.code == 1

    def test_strict_still_writes_output(self, tmp_path):
        """Even in strict mode, the JSON output is written before failing."""
        pattern = _make_pattern_expecting_email(tmp_path)
        data = _make_mismatched_data(tmp_path)
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit):
            cli_main([
                'extract', '-p', pattern, data,
                '--strict', '-o', out_dir,
            ])
        jsons = [f for f in os.listdir(out_dir) if f.endswith('.json')]
        assert len(jsons) == 1
        with open(os.path.join(out_dir, jsons[0])) as f:
            result = json.load(f)
        assert isinstance(result, dict)

    def test_strict_with_batch_fails_if_any_file_has_issues(self, tmp_path):
        """In batch mode, --strict fails if any single file has issues."""
        pattern = _make_pattern_expecting_email(tmp_path)
        data = _make_mismatched_data(tmp_path)
        data2 = tmp_path / 'data2.xlsx'
        shutil.copy2(data, data2)
        out_dir = str(tmp_path / 'output')
        with pytest.raises(SystemExit) as exc_info:
            cli_main([
                'extract', '-p', pattern, data, str(data2),
                '--strict', '-o', out_dir,
            ])
        assert exc_info.value.code == 2
