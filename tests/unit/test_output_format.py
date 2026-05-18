"""
Unit tests for the output-format helpers in engine.engine:
_set_nested, _field_local, _field_group, _table_group,
_build_row_obj, _build_nested_output, _expand_merged_cells.
"""

import openpyxl
import pytest

from engine.engine import (
    _set_nested,
    _field_local,
    _field_group,
    _table_group,
    _build_row_obj,
    _build_nested_output,
    _expand_merged_cells,
)
from engine.models import FieldDef


def _fd(name: str, role: str = 'var') -> FieldDef:
    return FieldDef(name=name, type='string', regex='.*', role=role)


def _defs(*pairs) -> dict:
    """pairs: (name, role) or just name (defaults to var)."""
    result = {}
    for item in pairs:
        if isinstance(item, tuple):
            name, role = item
        else:
            name, role = item, 'var'
        result[name] = _fd(name, role)
    return result


# ── _set_nested ───────────────────────────────────────────────────────────────

class TestSetNested:
    def test_single_segment(self):
        d = {}
        _set_nested(d, 'key', 'value')
        assert d == {'key': 'value'}

    def test_two_segments(self):
        d = {}
        _set_nested(d, 'a.b', 42)
        assert d == {'a': {'b': 42}}

    def test_three_segments(self):
        d = {}
        _set_nested(d, 'a.b.c', True)
        assert d == {'a': {'b': {'c': True}}}

    def test_merges_into_existing_group(self):
        d = {'po': {'number': 'PO-001'}}
        _set_nested(d, 'po.date', '2026-01-01')
        assert d == {'po': {'number': 'PO-001', 'date': '2026-01-01'}}

    def test_overwrites_existing_leaf(self):
        d = {}
        _set_nested(d, 'x.y', 1)
        _set_nested(d, 'x.y', 2)
        assert d['x']['y'] == 2

    def test_none_value_stored(self):
        d = {}
        _set_nested(d, 'a.b', None)
        assert d == {'a': {'b': None}}


# ── _field_local ──────────────────────────────────────────────────────────────

class TestFieldLocal:
    def test_with_one_dot(self):
        assert _field_local('po.number') == 'number'

    def test_without_dot(self):
        assert _field_local('number') == 'number'

    def test_two_dots_returns_tail(self):
        assert _field_local('a.b.c') == 'b.c'


# ── _field_group ──────────────────────────────────────────────────────────────

class TestFieldGroup:
    def test_with_dot(self):
        assert _field_group('po.number') == 'po'

    def test_without_dot(self):
        assert _field_group('number') == 'number'

    def test_two_dots_returns_head(self):
        assert _field_group('a.b.c') == 'a'


# ── _table_group ──────────────────────────────────────────────────────────────

class TestTableGroup:
    def test_infers_group_from_data(self):
        defs = _defs('line.item', 'line.qty')
        entry = {'data': [{'line.item': 'Laptop', 'line.qty': 2}]}
        assert _table_group(entry, defs) == 'line'

    def test_lbl_fields_ignored(self):
        defs = {
            'line.item': _fd('line.item', 'var'),
            'col_header': _fd('col_header', 'lbl'),
        }
        entry = {'data': [{'line.item': 'x', 'col_header': 'Item'}]}
        assert _table_group(entry, defs) == 'line'

    def test_fields_without_dot_ignored(self):
        defs = _defs('name')
        entry = {'data': [{'name': 'x'}], 'table_index': 0}
        assert _table_group(entry, defs) == 'table_0'

    def test_fallback_uses_table_index(self):
        entry = {'data': [], 'table_index': 3}
        assert _table_group(entry, {}) == 'table_3'

    def test_majority_group_wins(self):
        defs = _defs('line.item', 'line.qty', 'other.x')
        rows = [{'line.item': 'a', 'line.qty': 1, 'other.x': 'z'}] * 3
        entry = {'data': rows}
        assert _table_group(entry, defs) == 'line'

    def test_empty_data_falls_back(self):
        entry = {'data': [], 'table_index': 7}
        assert _table_group(entry, _defs('line.x')) == 'table_7'


# ── _build_row_obj ────────────────────────────────────────────────────────────

class TestBuildRowObj:
    def test_var_field_nested(self):
        defs = _defs('line.qty')
        obj = _build_row_obj({'line.qty': 5}, defs)
        assert obj == {'qty': 5}

    def test_lbl_field_excluded(self):
        defs = {
            'col_header': _fd('col_header', 'lbl'),
            'line.item': _fd('line.item', 'var'),
        }
        obj = _build_row_obj({'col_header': 'Item', 'line.item': 'Laptop'}, defs)
        assert obj == {'item': 'Laptop'}
        assert 'col_header' not in obj

    def test_unknown_field_included_as_is(self):
        # Fields not in defs pass through using local name
        obj = _build_row_obj({'mystery': 'value'}, {})
        assert obj == {'mystery': 'value'}

    def test_empty_row_returns_empty_dict(self):
        assert _build_row_obj({}, {}) == {}


# ── _build_nested_output ──────────────────────────────────────────────────────

class TestBuildNestedOutput:
    def test_var_cell_nested(self):
        defs = _defs('po.number', 'po.date')
        raw = {'cells': {'po.number': 'PO-001', 'po.date': '2026-01-01'}, 'tables': []}
        out = _build_nested_output(raw, defs)
        assert out == {'po': {'number': 'PO-001', 'date': '2026-01-01'}}

    def test_lbl_cell_excluded(self):
        defs = {
            'po_label': _fd('po_label', 'lbl'),
            'po.number': _fd('po.number', 'var'),
        }
        raw = {'cells': {'po_label': 'PO Number:', 'po.number': 'PO-001'}, 'tables': []}
        out = _build_nested_output(raw, defs)
        assert 'po_label' not in out
        assert out == {'po': {'number': 'PO-001'}}

    def test_table_produces_array(self):
        defs = _defs('line.item', 'line.qty')
        raw = {
            'cells': {},
            'tables': [{
                'headers': [],
                'data': [{'line.item': 'Laptop', 'line.qty': 2}],
                'footers': [],
            }],
        }
        out = _build_nested_output(raw, defs)
        assert isinstance(out['line'], list)
        assert out['line'][0]['data'][0] == {'item': 'Laptop', 'qty': 2}

    def test_table_header_lbl_stripped(self):
        defs = {
            'col_item': _fd('col_item', 'lbl'),
            'line.item': _fd('line.item', 'var'),
        }
        raw = {
            'cells': {},
            'tables': [{
                'headers': [{'col_item': 'Item'}],
                'data': [{'line.item': 'Laptop'}],
                'footers': [],
            }],
        }
        out = _build_nested_output(raw, defs)
        # all-lbl header → no 'header' key in instance
        assert 'header' not in out['line'][0]

    def test_footer_in_instance(self):
        defs = _defs(('footer.label', 'var'), ('footer.value', 'var'), ('line.item', 'var'))
        raw = {
            'cells': {},
            'tables': [{
                'headers': [],
                'data': [{'line.item': 'x'}],
                'footers': [{'footer.label': 'Grand Total', 'footer.value': 3030.0}],
            }],
        }
        out = _build_nested_output(raw, defs)
        assert out['line'][0]['footer'] == {'label': 'Grand Total', 'value': 3030.0}

    def test_multi_instance_array(self):
        defs = _defs('item.name')
        raw = {
            'cells': {},
            'tables': [
                {'headers': [], 'data': [{'item.name': 'Laptop'}], 'footers': []},
                {'headers': [], 'data': [{'item.name': 'Pen'}],    'footers': []},
                {'headers': [], 'data': [{'item.name': 'Desk'}],   'footers': []},
            ],
        }
        out = _build_nested_output(raw, defs)
        assert len(out['item']) == 3
        assert [inst['data'][0]['name'] for inst in out['item']] == ['Laptop', 'Pen', 'Desk']

    def test_cells_and_tables_combined(self):
        defs = _defs('vendor.name', 'line.item')
        raw = {
            'cells': {'vendor.name': 'Acme'},
            'tables': [{'headers': [], 'data': [{'line.item': 'Pen'}], 'footers': []}],
        }
        out = _build_nested_output(raw, defs)
        assert out['vendor'] == {'name': 'Acme'}
        assert out['line'][0]['data'][0] == {'item': 'Pen'}

    def test_empty_raw_returns_empty_dict(self):
        assert _build_nested_output({'cells': {}, 'tables': []}, {}) == {}

    def test_fully_empty_data_rows_dropped(self):
        # data row contains only an lbl: field; after filtering it becomes {}
        # → dropped → instance has no 'data' key.
        # No var: fields with dots in data → _table_group falls back to 'table_0'.
        defs = _defs(('col_h', 'lbl'))
        raw = {
            'cells': {},
            'tables': [{
                'headers': [],
                'data': [{'col_h': 'Header'}],
                'footers': [],
                'table_index': 0,
            }],
        }
        out = _build_nested_output(raw, defs)
        assert 'data' not in out['table_0'][0]


# ── _expand_merged_cells ──────────────────────────────────────────────────────

class TestExpandMergedCells:
    def test_merged_range_filled(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Merged Value'
        ws.merge_cells('A1:C1')
        # Before expansion, B1 and C1 are None
        assert ws['B1'].value is None
        _expand_merged_cells(ws)
        assert ws['B1'].value == 'Merged Value'
        assert ws['C1'].value == 'Merged Value'

    def test_non_merged_cells_unchanged(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Hello'
        ws['B1'] = 'World'
        _expand_merged_cells(ws)
        assert ws['A1'].value == 'Hello'
        assert ws['B1'].value == 'World'

    def test_multiple_merged_ranges(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Alpha'
        ws.merge_cells('A1:B1')
        ws['C1'] = 'Beta'
        ws.merge_cells('C1:D1')
        _expand_merged_cells(ws)
        assert ws['B1'].value == 'Alpha'
        assert ws['D1'].value == 'Beta'

    def test_none_value_merged_range(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        # Merge without setting a value — top-left is None
        ws.merge_cells('A1:B1')
        _expand_merged_cells(ws)
        assert ws['B1'].value is None
