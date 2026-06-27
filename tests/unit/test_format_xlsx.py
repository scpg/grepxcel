"""Unit tests for --format xlsx colored Excel report output."""

import io
import os
import sys

import openpyxl
import pytest

from grepxcel.cli import main
from grepxcel.xlsx_writer import nested_to_xlsx
from tests.conftest import find_pattern_xlsx

_INVOICE_FIX = os.path.join(os.path.dirname(__file__), '..', 'fixtures', '01_simple_invoice')
_INVOICE_PAT = find_pattern_xlsx(_INVOICE_FIX)
_INVOICE_DATA = os.path.join(_INVOICE_FIX, 'data.xlsx')

_CATALOG_FIX = os.path.join(os.path.dirname(__file__), '..', 'fixtures', '02_product_catalog')
_CATALOG_PAT = find_pattern_xlsx(_CATALOG_FIX)
_CATALOG_DATA = os.path.join(_CATALOG_FIX, 'data.xlsx')

_EXPENSE_FIX = os.path.join(os.path.dirname(__file__), '..', 'fixtures', '05_expense_report')
_EXPENSE_PAT = find_pattern_xlsx(_EXPENSE_FIX)
_EXPENSE_DATA = os.path.join(_EXPENSE_FIX, 'data.xlsx')


def _run_cli(*args):
    """Run CLI and capture stdout + return code."""
    old_stdout = sys.stdout
    sys.stdout = buf = io.StringIO()
    try:
        rc = main(list(args))
    except SystemExit as e:
        rc = e.code
    finally:
        sys.stdout = old_stdout
    return rc, buf.getvalue()


# ── nested_to_xlsx unit tests ───────────────────────────────────────────────

class TestNestedToXlsx:
    def test_scalar_only(self, tmp_path):
        data = {'inv': {'number': 'AB123', 'date': '2026-01-01'}}
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        assert out.exists()
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        # Row 1: var: | 1 | inv.number | AB123
        assert ws.cell(row=1, column=1).value == 'var:'
        assert ws.cell(row=1, column=2).value == 1
        assert ws.cell(row=1, column=3).value == 'inv.number'
        assert ws.cell(row=1, column=4).value == 'AB123'
        # Row 2: var: | 2 | inv.date | 2026-01-01
        assert ws.cell(row=2, column=1).value == 'var:'
        assert ws.cell(row=2, column=3).value == 'inv.date'

    def test_table_only(self, tmp_path):
        data = {
            'items': [
                {'data': [
                    {'name': 'A', 'qty': 1},
                    {'name': 'B', 'qty': 2},
                ]},
            ],
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        # Row 1: table: | 1 | HEADER | items.name | items.qty
        assert ws.cell(row=1, column=1).value == 'table:'
        assert ws.cell(row=1, column=2).value == 1
        assert ws.cell(row=1, column=3).value == 'HEADER'
        assert ws.cell(row=1, column=4).value == 'items.name'
        assert ws.cell(row=1, column=5).value == 'items.qty'
        # Row 2: None | None | DATA | A | 1
        assert ws.cell(row=2, column=3).value == 'DATA'
        assert ws.cell(row=2, column=4).value == 'A'
        # Row 3: None | None | DATA | B | 2
        assert ws.cell(row=3, column=4).value == 'B'

    def test_mixed_scalar_and_table(self, tmp_path):
        data = {
            'vendor': {'name': 'Acme'},
            'items': [
                {'data': [
                    {'item': 'Pen', 'price': 2},
                ]},
            ],
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        # Row 1: var: | 1 | vendor.name | Acme (padded to total_width)
        assert ws.cell(row=1, column=1).value == 'var:'
        assert ws.cell(row=1, column=3).value == 'vendor.name'
        assert ws.cell(row=1, column=4).value == 'Acme'
        # Row 2: blank separator between scalars and tables
        assert ws.cell(row=2, column=1).value is None
        # Row 3: table HEADER row
        assert ws.cell(row=3, column=1).value == 'table:'
        assert ws.cell(row=3, column=3).value == 'HEADER'

    def test_source_excluded(self, tmp_path):
        data = {
            'items': [
                {
                    '_source': {'sheet': 'Sheet1', 'ref': 'A1:D5'},
                    'data': [{'name': 'A'}],
                },
            ],
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        all_values = []
        for row in ws.iter_rows(values_only=True):
            all_values.extend(str(v) for v in row if v is not None)
        assert not any('Sheet1' in v for v in all_values)

    def test_footer_shown(self, tmp_path):
        data = {
            'items': [
                {
                    'data': [{'x': 1}],
                    'footer': {'label': 'Total', 'amount': 100},
                },
            ],
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        all_values = []
        for row in ws.iter_rows(values_only=True):
            all_values.extend(v for v in row if v is not None)
        # FOOTER marker present, label skipped, amount value present
        assert 'FOOTER' in all_values
        assert 100 in all_values
        assert 'Total' not in all_values  # label is excluded

    def test_coloring_applied(self, tmp_path):
        data = {'inv': {'number': 'AB123'}}
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        header_fill = ws.cell(row=1, column=1).fill
        assert header_fill.fgColor is not None
        assert header_fill.fgColor.rgb != '00000000'

    def test_meta_excluded(self, tmp_path):
        data = {
            'inv': {'number': 'AB123'},
            '_meta': {'run_id': 'test', 'stats': {}},
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        all_values = []
        for row in ws.iter_rows(values_only=True):
            all_values.extend(str(v) for v in row if v is not None)
        assert not any('_meta' in v for v in all_values)
        assert not any('run_id' in v for v in all_values)

    def test_empty_result(self, tmp_path):
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx({}, str(out))
        assert out.exists()
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        assert ws.cell(row=1, column=1).value is None

    def test_multi_instance_table(self, tmp_path):
        data = {
            'items': [
                {'data': [{'x': 1}]},
                {'data': [{'x': 2}, {'x': 3}]},
            ],
        }
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        values = []
        for row in ws.iter_rows(values_only=True):
            values.extend(v for v in row if v is not None)
        assert values.count(1) >= 1
        assert values.count(2) >= 1
        assert values.count(3) >= 1


# ── CLI integration ─────────────────────────────────────────────────────────

class TestCliXlsx:
    def test_requires_output_dir(self):
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                         '--format', 'xlsx')
        assert rc == 2

    def test_writes_xlsx_file(self, tmp_path):
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                         '--format', 'xlsx', '-o', str(tmp_path))
        assert rc == 0
        xlsx_files = list(tmp_path.glob('*.xlsx'))
        assert len(xlsx_files) == 1

    def test_xlsx_stem_matches_source_filename(self, tmp_path):
        """Output xlsx filename uses the source data file stem (data.xlsx → data.xlsx)."""
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                         '--format', 'xlsx', '-o', str(tmp_path))
        assert rc == 0
        xlsx_files = list(tmp_path.glob('*.xlsx'))
        assert len(xlsx_files) == 1
        assert xlsx_files[0].name == 'data.xlsx'

    def test_output_is_valid_workbook(self, tmp_path):
        _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                 '--format', 'xlsx', '-o', str(tmp_path))
        xlsx_files = list(tmp_path.glob('*.xlsx'))
        wb = openpyxl.load_workbook(str(xlsx_files[0]))
        ws = wb.active
        assert ws.cell(row=1, column=1).value is not None

    def test_table_extraction_xlsx(self, tmp_path):
        rc, _ = _run_cli('extract', '-p', _CATALOG_PAT, _CATALOG_DATA,
                         '--format', 'xlsx', '-o', str(tmp_path))
        assert rc == 0
        xlsx_files = list(tmp_path.glob('*.xlsx'))
        wb = openpyxl.load_workbook(str(xlsx_files[0]))
        ws = wb.active
        values = []
        for row in ws.iter_rows(values_only=True):
            values.extend(v for v in row if v is not None)
        assert any(v == 'Laptop' or v == 'Mouse' for v in values)

    def test_mixed_extraction_xlsx(self, tmp_path):
        rc, _ = _run_cli('extract', '-p', _EXPENSE_PAT, _EXPENSE_DATA,
                         '--format', 'xlsx', '-o', str(tmp_path))
        assert rc == 0

    def test_refuses_overwrite_source(self, tmp_path):
        """Output must not overwrite the source data file."""
        src_dir = os.path.dirname(_INVOICE_DATA)
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                         '--format', 'xlsx', '-o', src_dir)
        assert rc == 2

    def test_refuses_all_sheets(self, tmp_path):
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT, _INVOICE_DATA,
                         '--format', 'xlsx', '--all-sheets', '-o', str(tmp_path))
        assert rc == 2

    def test_output_mirrors_subdir_structure(self, tmp_path):
        """When extracting from a subdirectory, output preserves structure."""
        sub = tmp_path / 'input' / 'invoices'
        sub.mkdir(parents=True)
        import shutil
        shutil.copy(_INVOICE_DATA, sub / 'data.xlsx')
        out = tmp_path / 'output'
        rc, _ = _run_cli('extract', '-p', _INVOICE_PAT,
                         str(sub / 'data.xlsx'),
                         '--format', 'xlsx', '-o', str(out))
        assert rc == 0
        xlsx_files = list(out.rglob('*.xlsx'))
        assert len(xlsx_files) == 1
