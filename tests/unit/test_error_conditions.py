"""
Tests that verify the engine and parser correctly trap error conditions.

Each test either:
  - asserts logger.has_errors() → engine correctly detected a fatal problem
  - raises PatternError from PatternParser.parse() → parser-level validation
  - asserts specific warning counts → non-fatal validation issues logged

These tests ensure that bad input does not produce silent wrong results.
"""

import openpyxl
import pytest

from grepxcel.engine import Engine
from grepxcel.logger import Logger, VerbosityLevel
from grepxcel.pattern_parser import PatternError, PatternParser
from grepxcel.security import SecurityError


# ── helpers ───────────────────────────────────────────────────────────────────

def _write_pattern(rows: list, tmp_path, name='pattern.xlsx') -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    path = str(tmp_path / name)
    wb.save(path)
    return path


def _write_data(cells: dict, tmp_path, name='data.xlsx') -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    for coord, val in cells.items():
        ws[coord] = val
    path = str(tmp_path / name)
    wb.save(path)
    return path


def _engine_run(pattern_rows, data_cells, tmp_path) -> Logger:
    pat = _write_pattern(pattern_rows, tmp_path)
    dat = _write_data(data_cells, tmp_path)
    lg = Logger(level=VerbosityLevel.QUIET)
    Engine().process(pat, dat, logger=lg)
    return lg


# ═════════════════════════════════════════════════════════════════════════════
# PARSER-LEVEL ERRORS (PatternError raised before engine runs)
# ═════════════════════════════════════════════════════════════════════════════

class TestParserOrderingErrors:
    """Parser must reject absolute-ref sequences that are impossible to satisfy."""

    def test_lr_b2_then_a1_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:B2', 'a'], ['cell:A1', 'b'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='unreachable'):
            PatternParser().parse(path)

    def test_lr_same_cell_twice_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:C3', 'a'], ['cell:C3', 'b'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='unreachable'):
            PatternParser().parse(path)

    def test_lr_later_row_accepted(self, tmp_path):
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:A1', 'a'], ['cell:A2', 'b'], ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 2

    def test_lr_same_row_later_col_accepted(self, tmp_path):
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:A1', 'a'], ['cell:C1', 'b'], ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 2

    def test_lr_same_row_earlier_col_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:C1', 'a'], ['cell:A1', 'b'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError):
            PatternParser().parse(path)

    def test_td_b1_then_a2_rejected(self, tmp_path):
        # In TD mode: column primary. B1 (col 2, row 1) comes before A2 (col 1, row 2)?
        # No — TD scans col 1 fully first (A1, A2, ...), then col 2 (B1, B2, ...).
        # So A2 (col 1, row 2) comes BEFORE B1 (col 2, row 1).
        # Pattern: B1 then A2 → B1 position > A2 position in TD → reject
        path = _write_pattern([
            ['config:', 'read.direction', 'TD'],
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:B1', 'a'], ['cell:A2', 'b'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='unreachable'):
            PatternParser().parse(path)

    def test_td_a2_then_b1_accepted(self, tmp_path):
        # TD: A2 comes before B1 (column A fully scanned before column B)
        path = _write_pattern([
            ['config:', 'read.direction', 'TD'],
            ['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
            ['START:'], ['cell:A2', 'a'], ['cell:B1', 'b'], ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].target == 'A2'
        assert seq[1].target == 'B1'

    def test_cell_next_between_absolutes_no_ordering_check(self, tmp_path):
        # cell:next doesn't participate in ordering validation
        path = _write_pattern([
            ['var:', 'a', 'string', '.*'],
            ['var:', 'b', 'string', '.*'],
            ['var:', 'c', 'string', '.*'],
            ['START:'],
            ['cell:A1', 'a'],
            ['cell:next', 'b'],   # no coordinate → skipped in ordering check
            ['cell:C1',  'c'],    # C1 > A1 in LR → OK even though next might be between
            ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 3


class TestParserInvalidSyntax:
    """Parser must reject malformed cell: instructions."""

    def test_cell_foo_rejected(self, tmp_path):
        path = _write_pattern([
            ['START:'], ['cell:foo', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown cell instruction'):
            PatternParser().parse(path)

    def test_cell_0_rejected(self, tmp_path):
        # Row 0 does not exist in Excel
        path = _write_pattern([
            ['START:'], ['cell:0', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown cell instruction'):
            PatternParser().parse(path)

    def test_cell_A0_rejected(self, tmp_path):
        # Row 0 — the regex requires first digit to be [1-9]
        path = _write_pattern([
            ['START:'], ['cell:A0', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown cell instruction'):
            PatternParser().parse(path)

    def test_cell_star_rejected(self, tmp_path):
        path = _write_pattern([
            ['START:'], ['cell:*', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown cell instruction'):
            PatternParser().parse(path)

    def test_cell_2_rejected(self, tmp_path):
        # Only '1' and 'next' are valid sequential multiplicities
        path = _write_pattern([
            ['START:'], ['cell:2', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown cell instruction'):
            PatternParser().parse(path)

    def test_cell_1_accepted(self, tmp_path):
        path = _write_pattern([['START:'], ['cell:1', 'IGNORE'], ['END:']], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].multiplicity == '1'

    def test_cell_next_accepted(self, tmp_path):
        path = _write_pattern([['START:'], ['cell:next', 'IGNORE'], ['END:']], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].multiplicity == 'next'

    def test_cell_A1_accepted(self, tmp_path):
        path = _write_pattern([['START:'], ['cell:A1', 'IGNORE'], ['END:']], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].multiplicity == 'abs'
        assert seq[0].target == 'A1'


class TestParserSecurityErrors:
    """Security validation must reject unsafe pattern content."""

    def test_formula_in_pattern_rejected(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = '=1+2'
        path = str(tmp_path / 'bad.xlsx')
        wb.save(path)
        with pytest.raises(SecurityError, match='Formulas'):
            PatternParser().parse(path)

    def test_numeric_value_in_pattern_rejected(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 99
        path = str(tmp_path / 'bad.xlsx')
        wb.save(path)
        with pytest.raises(SecurityError, match='plain text'):
            PatternParser().parse(path)

    def test_boolean_value_in_pattern_rejected(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = False
        path = str(tmp_path / 'bad.xlsx')
        wb.save(path)
        with pytest.raises(SecurityError, match='plain text'):
            PatternParser().parse(path)


# ═════════════════════════════════════════════════════════════════════════════
# ENGINE-LEVEL FATAL ERRORS (logged as errors, processing stops)
# ═════════════════════════════════════════════════════════════════════════════

class TestEngineFatalErrors:
    """Engine must log a fatal error for unrecoverable conditions."""

    def test_lbl_field_empty_at_absolute_ref(self, tmp_path):
        lg = _engine_run(
            [['lbl:', 'title', 'string', 'Expected Title'],
             ['START:'], ['cell:B2', 'title'], ['END:']],
            {'A1': 'something else'},  # B2 is empty
            tmp_path,
        )
        assert lg.has_errors()

    def test_sequential_next_on_empty_sheet(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.v', 'string', '.*'],
             ['START:'], ['cell:next', 'x.v'], ['END:']],
            {},  # empty sheet
            tmp_path,
        )
        assert lg.has_errors()

    def test_undefined_field_in_start_sequence(self, tmp_path):
        # Field 'ghost' not defined with var:/lbl:
        lg = _engine_run(
            [['var:', 'x.v', 'string', '.*'],
             ['START:'], ['cell:A1', 'ghost'], ['END:']],
            {'A1': 'val'},
            tmp_path,
        )
        assert lg.has_errors()

    def test_cursor_past_absolute_target(self, tmp_path):
        lg = _engine_run(
            [['var:', 'a.v', 'string', '.*'],
             ['var:', 'b.v', 'string', '.*'],
             ['START:'],
             ['cell:next', 'a.v'],   # consumes A1, cursor past index 0
             ['cell:A1',   'b.v'],   # A1 already past → fatal
             ['END:']],
            {'A1': 'val', 'B1': 'other'},
            tmp_path,
        )
        assert lg.has_errors()

    def test_pattern_syntax_error_propagates_to_engine(self, tmp_path):
        # PatternError from parser becomes engine fatal
        lg = _engine_run(
            [['START:'], ['cell:notvalid', 'IGNORE'], ['END:']],
            {'A1': 'val'},
            tmp_path,
        )
        assert lg.has_errors()

    def test_out_of_order_refs_propagate_to_engine(self, tmp_path):
        lg = _engine_run(
            [['var:', 'a', 'string', '.*'], ['var:', 'b', 'string', '.*'],
             ['START:'], ['cell:C3', 'a'], ['cell:A1', 'b'], ['END:']],
            {'A1': 'x', 'C3': 'y'},
            tmp_path,
        )
        assert lg.has_errors()

    def test_pattern_with_formula_propagates_to_engine(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = '=SUM(1,2)'
        pat_path = str(tmp_path / 'pattern.xlsx')
        wb.save(pat_path)
        dat_path = _write_data({'A1': 'data'}, tmp_path)
        lg = Logger(level=VerbosityLevel.QUIET)
        Engine().process(pat_path, dat_path, logger=lg)
        assert lg.has_errors()


# ═════════════════════════════════════════════════════════════════════════════
# NON-FATAL WARNINGS (engine continues, but issues are logged)
# ═════════════════════════════════════════════════════════════════════════════

class TestEngineWarnings:
    """Engine should warn (not error) for validation failures in DATA rows
    and for regex/type mismatches at cell: instructions."""

    def test_regex_mismatch_at_absolute_ref_is_warning(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.code', 'string', r'[A-Z]{3}\d{3}'],
             ['START:'], ['cell:A1', 'x.code'], ['END:']],
            {'A1': 'not-a-code'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1

    def test_regex_mismatch_at_sequential_ref_is_warning(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.code', 'string', r'[A-Z]+'],
             ['START:'], ['cell:next', 'x.code'], ['END:']],
            {'A1': '12345'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1

    def test_null_value_no_warning(self, tmp_path):
        # Empty var: at absolute ref → null; NO warning (nothing to validate)
        lg = _engine_run(
            [['var:', 'x.v', 'string', r'REQUIRED'],
             ['START:'], ['cell:Z99', 'x.v'], ['END:']],
            {'A1': 'elsewhere'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 0

    def test_lbl_regex_mismatch_at_abs_ref_is_warning(self, tmp_path):
        # lbl: with wrong value at absolute ref → warning only, not fatal
        # (fatal only if the cell is EMPTY)
        lg = _engine_run(
            [['lbl:', 'hdr', 'string', 'Expected Header'],
             ['START:'], ['cell:A1', 'hdr'], ['END:']],
            {'A1': 'Wrong Header'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1

    def test_multiple_warnings_accumulated(self, tmp_path):
        lg = _engine_run(
            [['var:', 'a.x', 'integer', r'\d+'],
             ['var:', 'b.y', 'integer', r'\d+'],
             ['var:', 'c.z', 'integer', r'\d+'],
             ['START:'],
             ['cell:A1', 'a.x'],
             ['cell:B1', 'b.y'],
             ['cell:C1', 'c.z'],
             ['END:']],
            {'A1': 'bad', 'B1': 'also bad', 'C1': 'still bad'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 3

    def test_one_valid_one_invalid_mixed(self, tmp_path):
        lg = _engine_run(
            [['var:', 'a.x', 'integer', r'\d+'],
             ['var:', 'b.y', 'integer', r'\d+'],
             ['START:'],
             ['cell:A1', 'a.x'],
             ['cell:B1', 'b.y'],
             ['END:']],
            {'A1': 42, 'B1': 'not-int'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1


# ═════════════════════════════════════════════════════════════════════════════
# PARTIAL RESULTS AFTER FATAL ERRORS
# ═════════════════════════════════════════════════════════════════════════════

class TestPartialResultsAfterFatal:
    """After a fatal error, the engine returns whatever it extracted before
    the error. Verify the partial result is sane."""

    def test_cells_before_fatal_are_in_partial_result(self, tmp_path):
        pat = _write_pattern([
            ['var:', 'good.val', 'string', '.*'],
            ['lbl:', 'missing_lbl', 'string', 'Must Be Here'],
            ['START:'],
            ['cell:A1', 'good.val'],
            ['cell:B1', 'missing_lbl'],  # B1 is empty → fatal
            ['END:'],
        ], tmp_path)
        dat = _write_data({'A1': 'extracted before error'}, tmp_path)
        lg = Logger(level=VerbosityLevel.QUIET)
        result = Engine().process(pat, dat, logger=lg)
        assert lg.has_errors()
        # 'good.val' was extracted before the fatal
        assert result.get('good', {}).get('val') == 'extracted before error'

    def test_partial_result_has_no_crashed_fields(self, tmp_path):
        pat = _write_pattern([
            ['var:', 'a.v', 'string', '.*'],
            ['var:', 'b.v', 'string', '.*'],
            ['START:'],
            ['cell:A1', 'a.v'],
            ['cell:next', 'b.v'],  # sheet has only A1 → exhausted
            ['END:'],
        ], tmp_path)
        dat = _write_data({'A1': 'only cell'}, tmp_path)
        lg = Logger(level=VerbosityLevel.QUIET)
        result = Engine().process(pat, dat, logger=lg)
        assert lg.has_errors()
        # a.v extracted before exhaustion
        assert result.get('a', {}).get('v') == 'only cell'
        # b.v was never extracted
        assert result.get('b') is None


# ═════════════════════════════════════════════════════════════════════════════
# CORRECT BEHAVIOR PRESERVED (regression guard)
# ═════════════════════════════════════════════════════════════════════════════

class TestRegressionGuards:
    """Ensure that adding cell:next/cell:A1 did not break existing patterns
    that use cell:1 syntax."""

    def test_cell_1_still_works_sequentially(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.v', 'string', '.*'],
             ['START:'], ['cell:1', 'x.v'], ['END:']],
            {'A1': 'backward compat'},
            tmp_path,
        )
        assert not lg.has_errors()

    def test_cell_1_and_absolute_mixed(self, tmp_path):
        lg = _engine_run(
            [['var:', 'a.v', 'string', '.*'],
             ['var:', 'b.v', 'string', '.*'],
             ['START:'],
             ['cell:1',  'a.v'],
             ['cell:C1', 'b.v'],
             ['END:']],
            {'A1': 'seq', 'B1': 'skip', 'C1': 'abs'},
            tmp_path,
        )
        assert not lg.has_errors()

    def test_def_alias_still_works(self, tmp_path):
        # def: is the old alias for var: — must still parse and run
        lg = _engine_run(
            [['def:', 'x.v', 'string', '.*'],
             ['START:'], ['cell:A1', 'x.v'], ['END:']],
            {'A1': 'old alias'},
            tmp_path,
        )
        assert not lg.has_errors()

    def test_doc_rows_ignored(self, tmp_path):
        lg = _engine_run(
            [['doc:', 'This is a comment — ignored'],
             ['var:', 'x.v', 'string', '.*'],
             ['doc:', 'Another comment'],
             ['START:'],
             ['cell:A1', 'x.v'],
             ['END:']],
            {'A1': 'value'},
            tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 0


# ═════════════════════════════════════════════════════════════════════════════
# HARDENED PATTERN SANITY VALIDATION
#   A pattern file is ALWAYS sanity-checked. Malformed patterns must be trapped
#   (PatternError at parse, or a fatal engine error) — never silently tolerated.
# ═════════════════════════════════════════════════════════════════════════════

class TestUnknownFieldType:
    """A var:/lbl: type outside the engine's known set is rejected at parse."""

    def test_typo_type_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'currncy', '.*'],     # typo for 'currency'
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown field type'):
            PatternParser().parse(path)

    @pytest.mark.parametrize('type_name', [
        'string', 'text', 'integer', 'number', 'float', 'decimal',
        'currency', 'percentage', 'boolean', 'bool',
        'date', 'datetime', 'timestamp',
    ])
    def test_all_valid_types_accepted(self, tmp_path, type_name):
        path = _write_pattern([
            ['var:', 'x.v', type_name, '.*'],
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        _, defs, _ = PatternParser().parse(path)
        assert defs['x.v'].type == type_name


class TestMissingFieldName:
    """var:/lbl: rows must name a field."""

    def test_var_without_name_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', None, 'string', '.*'],
            ['START:'], ['cell:A1', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='no field name'):
            PatternParser().parse(path)

    def test_lbl_without_name_rejected(self, tmp_path):
        path = _write_pattern([
            ['lbl:', None, 'string', 'X'],
            ['START:'], ['cell:A1', 'IGNORE'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='no field name'):
            PatternParser().parse(path)


class TestEmptyExtractionSequence:
    """A pattern that defines no extraction steps is a fatal engine error
    (covers: missing START:, empty file, END: before START:, empty block)."""

    def test_missing_start_marker_is_fatal(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.v', 'string', '.*'], ['cell:A1', 'x.v'], ['END:']],
            {'A1': 'val'}, tmp_path,
        )
        assert lg.has_errors()

    def test_empty_pattern_is_fatal(self, tmp_path):
        lg = _engine_run([], {'A1': 'val'}, tmp_path)
        assert lg.has_errors()

    def test_end_before_start_is_fatal(self, tmp_path):
        lg = _engine_run(
            [['var:', 'x.v', 'string', '.*'],
             ['END:'], ['START:'], ['cell:A1', 'x.v']],
            {'A1': 'val'}, tmp_path,
        )
        assert lg.has_errors()

    def test_empty_start_block_is_fatal(self, tmp_path):
        lg = _engine_run([['START:'], ['END:']], {'A1': 'val'}, tmp_path)
        assert lg.has_errors()


class TestTableStructureRules:
    """DATA is required; HEADER must precede DATA; FOOTER must follow DATA."""

    def test_empty_table_block_rejected(self, tmp_path):
        # Previously crashed with IndexError — must be a clean PatternError now.
        path = _write_pattern([
            ['START:'], ['table:*'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='no DATA row'):
            PatternParser().parse(path)

    def test_table_without_data_rejected(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'H'],
            ['START:'], ['table:*'],
            [None, 'HEADER:1', 'h'],            # HEADER but no DATA
            ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='no DATA row'):
            PatternParser().parse(path)

    def test_header_after_data_rejected(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'H'], ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'DATA:*', 'd'],
            [None, 'HEADER:1', 'h'],            # HEADER after DATA
            ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='HEADER row after a DATA row'):
            PatternParser().parse(path)

    def test_footer_before_data_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'], ['var:', 'f', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'FOOTER:1', 'f'],            # FOOTER before DATA
            [None, 'DATA:*', 'd'],
            ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='FOOTER row before a DATA row'):
            PatternParser().parse(path)

    def test_data_only_table_accepted(self, tmp_path):
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'DATA:*', 'd'],
            ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 1

    def test_header_data_footer_in_order_accepted(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'H'],
            ['var:', 'd', 'string', '.*'], ['var:', 'f', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'HEADER:1', 'h'],
            [None, 'DATA:*', 'd'],
            [None, 'FOOTER:1', 'f'],
            ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 1


class TestMultiplicityValidation:
    """table:/DATA:/HEADER:/FOOTER: multiplicities and row keywords are validated."""

    def test_bad_table_multiplicity_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:foo'],
            [None, 'DATA:*', 'd'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Invalid table multiplicity'):
            PatternParser().parse(path)

    def test_bad_data_multiplicity_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'DATA:xyz', 'd'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Invalid DATA multiplicity'):
            PatternParser().parse(path)

    def test_bad_header_multiplicity_rejected(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'H'], ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'HEADER:foo', 'h'],
            [None, 'DATA:*', 'd'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Invalid HEADER multiplicity'):
            PatternParser().parse(path)

    def test_unknown_row_type_rejected(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'H'], ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'HEDER:1', 'h'],             # typo for HEADER
            [None, 'DATA:*', 'd'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='Unknown table row type'):
            PatternParser().parse(path)

    def test_table_numeric_multiplicity_accepted(self, tmp_path):
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:2'],
            [None, 'DATA:*', 'd'], ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].multiplicity == '2'


class TestRegexSanity:
    """Pattern regexes are checked for compilability and ReDoS safety."""

    def test_malformed_regex_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '[A-'],   # unterminated character class
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        with pytest.raises(SecurityError, match='Invalid regex'):
            PatternParser().parse(path)

    def test_redos_pattern_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', r'(a+)+$'],   # classic nested-quantifier ReDoS
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        with pytest.raises(SecurityError, match='nested unbounded quantifiers'):
            PatternParser().parse(path)


class TestNewFieldTypesEndToEnd:
    """number/float/decimal, text, and boolean validate correctly end-to-end."""

    def test_number_accepts_float(self, tmp_path):
        lg = _engine_run(
            [['var:', 'm.v', 'number', r'.*'],
             ['START:'], ['cell:A1', 'm.v'], ['END:']],
            {'A1': 3.5}, tmp_path,
        )
        assert not lg.has_errors()
        assert lg.issues() == []

    def test_number_rejects_text(self, tmp_path):
        lg = _engine_run(
            [['var:', 'm.v', 'number', r'.*'],
             ['START:'], ['cell:A1', 'm.v'], ['END:']],
            {'A1': 'not a number'}, tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1            # validation warning, not fatal

    def test_boolean_accepts_bool(self, tmp_path):
        lg = _engine_run(
            [['var:', 'b.v', 'boolean', r'.*'],
             ['START:'], ['cell:A1', 'b.v'], ['END:']],
            {'A1': True}, tmp_path,
        )
        assert not lg.has_errors()
        assert lg.issues() == []

    def test_boolean_rejects_number(self, tmp_path):
        lg = _engine_run(
            [['var:', 'b.v', 'boolean', r'.*'],
             ['START:'], ['cell:A1', 'b.v'], ['END:']],
            {'A1': 5}, tmp_path,
        )
        assert not lg.has_errors()
        assert len(lg.issues()) == 1

    def test_text_alias_behaves_like_string(self, tmp_path):
        lg = _engine_run(
            [['var:', 't.v', 'text', r'[A-Z]+'],
             ['START:'], ['cell:A1', 't.v'], ['END:']],
            {'A1': 'ABC'}, tmp_path,
        )
        assert not lg.has_errors()
        assert lg.issues() == []


class TestSheetDimensionLimits:
    """The data sheet must fit within the configured row/column limits; an
    oversized sheet fails cleanly and the limit can be raised."""

    _PAT = [['var:', 'x.v', 'string', '.*'],
            ['START:'], ['cell:A1', 'x.v'], ['END:']]

    def _run(self, data_cells, tmp_path, max_rows, max_cols):
        pat = _write_pattern(self._PAT, tmp_path)
        dat = _write_data(data_cells, tmp_path)
        lg = Logger(level=VerbosityLevel.QUIET)
        Engine().process(pat, dat, logger=lg, max_rows=max_rows, max_cols=max_cols)
        return lg

    def test_too_many_rows_is_fatal(self, tmp_path):
        lg = self._run({'A1': 'v', 'A6': 'x'}, tmp_path, max_rows=5, max_cols=100)
        assert lg.has_errors()

    def test_too_many_columns_is_fatal(self, tmp_path):
        lg = self._run({'A1': 'v', 'E1': 'x'}, tmp_path, max_rows=100, max_cols=3)
        assert lg.has_errors()

    def test_within_limits_ok(self, tmp_path):
        lg = self._run({'A1': 'v'}, tmp_path, max_rows=100, max_cols=100)
        assert not lg.has_errors()

    def test_raised_limit_allows_larger_sheet(self, tmp_path):
        # The same sheet that failed at max_rows=5 passes once the limit is raised.
        lg = self._run({'A1': 'v', 'A6': 'x'}, tmp_path, max_rows=10, max_cols=100)
        assert not lg.has_errors()


class TestPatternComments:
    """'#' trailing comments: allowed on config/var/lbl/cell/START rows, fail fast
    on stray non-'#' content, and NOT processed in table rows ('#' is literal there)."""

    def test_comment_on_var_row_keeps_regex(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '.*', '# the value field'],   # col E comment
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        _, defs, _ = PatternParser().parse(path)
        assert defs['x.v'].regex == '.*'   # the comment did not touch the regex (col D)

    def test_comment_on_lbl_row(self, tmp_path):
        path = _write_pattern([
            ['lbl:', 'h', 'string', 'Title', '# header label'],
            ['START:'], ['cell:A1', 'h'], ['END:'],
        ], tmp_path)
        _, defs, _ = PatternParser().parse(path)
        assert defs['h'].role == 'lbl'

    def test_comment_on_config_row(self, tmp_path):
        path = _write_pattern([
            ['config:', 'currency.sign', '$', '# US dollars'],     # col D comment
            ['var:', 'x.v', 'string', '.*'],
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        cfg, _, _ = PatternParser().parse(path)
        assert cfg.currency_sign == '$'

    def test_comment_on_cell_row(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '.*'],
            ['START:'], ['cell:A1', 'x.v', '# grab the value'], ['END:'],   # col C comment
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert seq[0].field == 'x.v'

    def test_comment_on_start_row(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '.*'],
            ['START:', '# begin extraction'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        assert len(seq) == 1

    def test_comment_spans_remaining_cells(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '.*', '# a comment', 'continuing here'],
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        _, defs, _ = PatternParser().parse(path)   # 'continuing here' is part of the comment
        assert 'x.v' in defs

    def test_stray_content_without_hash_rejected(self, tmp_path):
        path = _write_pattern([
            ['var:', 'x.v', 'string', '.*', 'oops wrong column'],   # no '#'
            ['START:'], ['cell:A1', 'x.v'], ['END:'],
        ], tmp_path)
        with pytest.raises(PatternError, match='only contain a comment'):
            PatternParser().parse(path)

    def test_hash_regex_in_column_d_is_not_a_comment(self, tmp_path):
        # A regex (col D) that legitimately starts with '#' is a regex, not a comment.
        path = _write_pattern([
            ['var:', 'x.code', 'string', r'#\d+'],
            ['START:'], ['cell:A1', 'x.code'], ['END:'],
        ], tmp_path)
        _, defs, _ = PatternParser().parse(path)
        assert defs['x.code'].regex == r'#\d+'

    def test_hash_is_literal_in_table_rows(self, tmp_path):
        # Tables get no comment processing: a '#'-leading column cell stays literal.
        path = _write_pattern([
            ['var:', 'd', 'string', '.*'],
            ['START:'], ['table:*'],
            [None, 'DATA:*', 'd', '#notacomment'],
            ['END:'],
        ], tmp_path)
        _, _, seq = PatternParser().parse(path)
        cols = [c.field for c in seq[0].rows[0].columns]
        assert '#notacomment' in cols   # kept as a literal column, not stripped as a comment
