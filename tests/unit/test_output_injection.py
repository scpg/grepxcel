"""Security + helper tests for output writers (formula/CSV injection, flatten)."""

import csv
import io

import openpyxl
import pytest

from grepxcel.utils import neutralize_formula, flatten_nested
from grepxcel.csv_writer import nested_to_csv
from grepxcel.xlsx_writer import nested_to_xlsx


# ── neutralize_formula ───────────────────────────────────────────────────────

class TestNeutralizeFormula:
    @pytest.mark.parametrize('payload', [
        '=1+2',
        '=cmd|"/c calc"!A1',
        '+1+2',
        '-1+2',
        '@SUM(A1)',
        '\t=1',
        '\r=1',
        '=HYPERLINK("http://evil","click")',
    ])
    def test_dangerous_leads_are_prefixed(self, payload):
        out = neutralize_formula(payload)
        assert out.startswith("'")
        assert out[1:] == payload

    @pytest.mark.parametrize('safe', [
        'hello',
        'AB123',
        '2026-01-01',
        'a=b',          # = not in lead position
        '',
        'Acme Ltd.',
    ])
    def test_safe_strings_unchanged(self, safe):
        assert neutralize_formula(safe) == safe

    def test_non_strings_pass_through(self):
        assert neutralize_formula(5) == 5
        assert neutralize_formula(3.14) == 3.14
        assert neutralize_formula(None) is None
        assert neutralize_formula(True) is True


# ── flatten_nested ───────────────────────────────────────────────────────────

class TestFlattenNested:
    def test_flat(self):
        assert flatten_nested({'a': 1, 'b': 2}) == [('a', 1), ('b', 2)]

    def test_nested(self):
        assert flatten_nested({'x': {'y': 1, 'z': 2}}) == [('x.y', 1), ('x.z', 2)]

    def test_prefix(self):
        assert flatten_nested({'y': 1}, 'x') == [('x.y', 1)]

    def test_order_preserved(self):
        out = flatten_nested({'b': 1, 'a': {'d': 2, 'c': 3}})
        assert out == [('b', 1), ('a.d', 2), ('a.c', 3)]


# ── injection neutralized in csv output ──────────────────────────────────────

class TestCsvInjection:
    def test_formula_value_neutralized(self):
        data = {'risk': {'payload': '=1+2'}}
        out, _ = nested_to_csv(data)
        rows = list(csv.DictReader(io.StringIO(out)))
        assert rows[0]['risk.payload'] == "'=1+2"

    def test_table_cell_neutralized(self):
        data = {'items': [{'data': [{'name': '=cmd|calc'}]}]}
        out, _ = nested_to_csv(data)
        rows = list(csv.DictReader(io.StringIO(out)))
        assert rows[0]['items.name'].startswith("'=")


# ── injection neutralized in xlsx output ─────────────────────────────────────

class TestXlsxInjection:
    def test_scalar_not_written_as_formula(self, tmp_path):
        data = {'risk': {'payload': '=cmd|"/c calc"!A1'}}
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        # New format: row 1 = (var:, 1, risk.payload, <value>) — value is in col 4
        cell = ws.cell(row=1, column=4)
        assert cell.data_type != 'f'

    def test_table_cell_not_formula(self, tmp_path):
        data = {'items': [{'data': [{'name': '=1+1'}]}]}
        out = tmp_path / 'out.xlsx'
        nested_to_xlsx(data, str(out))
        wb = openpyxl.load_workbook(str(out))
        ws = wb.active
        formula_cells = [
            c for row in ws.iter_rows() for c in row if c.data_type == 'f'
        ]
        assert not formula_cells
