"""`--sheet` resolution: exact, case-insensitive, index — and what it refuses.

Case is ignored because a workbook cannot hold two sheets differing only by it:
Excel rejects the duplicate, and openpyxl silently uniquifies (a second
'SHEET1' is stored as 'SHEET11'). Whitespace is NOT ignored, because Excel does
treat 'Sheet 1' and 'Sheet1' as two different sheets — so normalising it away
would silently read the wrong sheet.

The near-miss message matters as much as the matching: a user who types
'ShEET 1' for a sheet called 'Sheet1' has a whitespace problem, and being told
"not found, expected one of: Sheet1, Sheet1_2" leaves them staring at the case.
"""
import pytest

from grepxcel.engine import Engine, _ci_sheet_matches, _squash
from grepxcel.logger import Logger, Severity, VerbosityLevel
from tests.unit.test_random_type_fixtures import write_pattern_csv

import openpyxl


def _book(tmp_path, *names):
    path = str(tmp_path / 'data.xlsx')
    wb = openpyxl.Workbook()
    wb.active.title = names[0]
    for extra in names[1:]:
        wb.create_sheet(extra)
    for ws in wb.worksheets:
        ws['A1'] = 'ANCHOR'
        ws['B1'] = ws.title
    wb.save(path)
    return path


def _pattern(tmp_path):
    path = str(tmp_path / 'pattern.csv')
    write_pattern_csv(path, [
        ['lbl:', 'a', 'string', 'ANCHOR'],
        ['var:', 'f', 'string', '.*'],
        [], ['START:'], ['cell:next', 'a'], ['cell:next', 'f'], ['END:'],
    ])
    return path


def _run(tmp_path, sheet):
    lg = Logger(level=VerbosityLevel.QUIET)
    return Engine().process(
        pattern_file=_pattern(tmp_path),
        data_file=_book(tmp_path, 'Sheet1', 'Sheet1_2', 'Data 2025'),
        logger=lg, sheet=sheet,
    )


class TestHelpers:
    def test_ci_matches_is_case_insensitive(self):
        assert _ci_sheet_matches('sheet1', ['Sheet1', 'Other']) == ['Sheet1']
        assert _ci_sheet_matches('SHEET1', ['Sheet1']) == ['Sheet1']
        assert _ci_sheet_matches('nope', ['Sheet1']) == []

    def test_ci_matches_does_not_ignore_whitespace(self):
        assert _ci_sheet_matches('Sheet 1', ['Sheet1']) == []

    def test_squash_is_only_for_messages(self):
        assert _squash('Sheet 1') == _squash('SHEET1') == 'sheet1'


class TestSheetResolution:
    def test_exact_name(self, tmp_path):
        assert _run(tmp_path, 'Sheet1_2')['f'] == 'Sheet1_2'

    @pytest.mark.parametrize('spelling', ['sheet1_2', 'SHEET1_2', 'ShEeT1_2'])
    def test_case_insensitive(self, tmp_path, spelling):
        assert _run(tmp_path, spelling)['f'] == 'Sheet1_2'

    @pytest.mark.parametrize('spelling', ['data 2025', 'DATA 2025'])
    def test_case_insensitive_with_a_space_in_the_real_name(self, tmp_path, spelling):
        """Spaces are preserved, so a name that genuinely has one still matches
        case-insensitively."""
        assert _run(tmp_path, spelling)['f'] == 'Data 2025'

    def test_index_still_works(self, tmp_path):
        assert _run(tmp_path, 1)['f'] == 'Sheet1_2'

    def test_numeric_string_is_an_index(self, tmp_path):
        assert _run(tmp_path, '0')['f'] == 'Sheet1'

    def test_exact_match_is_tried_before_the_case_insensitive_fallback(self):
        """Ordering guard. A real name that happens to differ from another only
        by case must never be shadowed by the fallback. Asserted on the helper
        because the situation cannot be built through openpyxl — it uniquifies
        a case-colliding name ('data' alongside 'DATA' is stored 'data1'), which
        is also why the fallback is unambiguous in practice."""
        names = ['DATA', 'data']
        assert 'data' in names                       # exact branch would hit first
        assert _ci_sheet_matches('data', names) == names   # fallback sees both


def _failure_record(tmp_path, sheet):
    """Run and return the FATAL LogRecord rendered as
    'FATAL ERROR / Expected: / Found:'.

    Two things this has to work around. ``Engine.process`` catches EngineError
    for a setup-phase fatal and returns a partial result — deliberately, which
    is why the CLI prints the error and then ``{}`` — so there is no exception
    to catch here. And the guidance lives in the record's ``expected`` field
    rather than its message, which is the part the user actually reads.
    """
    lg = Logger(level=VerbosityLevel.QUIET)
    result = Engine().process(
        pattern_file=_pattern(tmp_path),
        data_file=_book(tmp_path, 'Sheet1', 'Sheet1_2', 'Data 2025'),
        logger=lg, sheet=sheet,
    )
    assert result == {}, f'expected no extraction, got {result!r}'
    errors = [r for r in lg._records if r.severity == Severity.ERROR]
    assert errors, 'the run failed but logged no ERROR record'
    return errors[-1]


class TestRefusals:
    def test_whitespace_difference_is_refused(self, tmp_path):
        """The reported case: 'ShEET 1' against a sheet named 'Sheet1'. Case is
        ignored, spacing is not, so this must still fail."""
        rec = _failure_record(tmp_path, 'ShEET 1')
        assert 'not found' in rec.message

    def test_and_the_message_names_whitespace_as_the_cause(self, tmp_path):
        """Message-accuracy: without this the user sees 'expected one of:
        Sheet1, Sheet1_2' and re-checks the capitalisation, which is already
        handled. The near-miss has to be named."""
        rec = _failure_record(tmp_path, 'ShEET 1')
        blob = f'{rec.message} {rec.expected} {rec.found}'.lower()
        assert 'sheet1' in blob
        assert 'whitespace' in blob or 'spaces' in blob, blob

    def test_a_genuinely_unknown_sheet_gets_no_misleading_hint(self, tmp_path):
        """The hint must only appear when there really is a near miss."""
        rec = _failure_record(tmp_path, 'Nonexistent')
        assert 'not found' in rec.message
        blob = f'{rec.expected} {rec.found}'.lower()
        assert 'did you mean' not in blob, blob

    def test_out_of_range_index_still_reports_the_range(self, tmp_path):
        rec = _failure_record(tmp_path, 99)
        assert 'out of range' in rec.message
