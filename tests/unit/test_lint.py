"""Tests for grepxcel lint — Excel file inspection before extraction."""

import io
import os
import zipfile
import zlib
import openpyxl
from openpyxl.styles import Font, PatternFill
import pytest

from grepxcel.lint import lint_file, OK, WARN, FAIL, INFO


# ── helpers ───────────────────────────────────────────────────────────────────

def _make_xlsx(tmp_path, cells=None, name='data.xlsx', sheets=None):
    """Create a simple .xlsx with the given cells dict or multiple sheets."""
    wb = openpyxl.Workbook()
    if sheets:
        for i, (sname, scells) in enumerate(sheets.items()):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = sname
            for coord, val in scells.items():
                ws[coord] = val
    else:
        ws = wb.active
        for coord, val in (cells or {'A1': 'hello'}).items():
            ws[coord] = val
    path = str(tmp_path / name)
    wb.save(path)
    return path


# ── basic file checks ─────────────────────────────────────────────────────────

class TestFileAccess:
    def test_nonexistent_file_fails(self, tmp_path):
        results = lint_file(str(tmp_path / 'nope.xlsx'))
        statuses = [r[0] for r in results]
        assert FAIL in statuses
        assert any('not found' in r[2].lower() or 'does not exist' in r[2].lower()
                    for r in results)

    def test_not_a_file_fails(self, tmp_path):
        results = lint_file(str(tmp_path))
        statuses = [r[0] for r in results]
        assert FAIL in statuses

    def test_wrong_extension_fails(self, tmp_path):
        path = str(tmp_path / 'data.txt')
        with open(path, 'w') as f:
            f.write('hello')
        results = lint_file(path)
        statuses = [r[0] for r in results]
        assert FAIL in statuses
        assert any('.xlsx' in r[2] for r in results)

    def test_xlsm_rejected(self, tmp_path):
        path = str(tmp_path / 'data.xlsm')
        with open(path, 'w') as f:
            f.write('fake')
        results = lint_file(path)
        assert any(r[0] == FAIL and 'macro' in r[2].lower() for r in results)

    def test_valid_xlsx_passes_file_check(self, tmp_path):
        path = _make_xlsx(tmp_path)
        results = lint_file(path)
        assert any(r[0] == OK and 'file' in r[1].lower() for r in results)


# ── encryption / IRM / OLE ────────────────────────────────────────────────────

class TestEncryption:
    def test_non_zip_magic_detected(self, tmp_path):
        """An encrypted/IRM-protected .xlsx is OLE, not ZIP — magic bytes differ."""
        path = str(tmp_path / 'encrypted.xlsx')
        # OLE compound document magic: D0 CF 11 E0
        with open(path, 'wb') as f:
            f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 512)
        results = lint_file(path)
        assert any(r[0] == FAIL and ('encrypt' in r[2].lower() or
                                      'protect' in r[2].lower() or
                                      'ole' in r[2].lower())
                   for r in results)

    def test_corrupt_zip_detected(self, tmp_path):
        path = str(tmp_path / 'corrupt.xlsx')
        with open(path, 'wb') as f:
            f.write(b'PK\x03\x04' + b'\xff' * 100)
        results = lint_file(path)
        assert any(r[0] == FAIL for r in results)


# ── ZIP bomb detection ────────────────────────────────────────────────────────

def _make_zip_bomb_ratio(path: str, uncompressed_mb: float = 3.0) -> None:
    """Write a ZIP file whose single member expands to uncompressed_mb of zeros.

    Zeros compress at ~1000:1 with DEFLATE, so a ~3 KB file expands to ~3 MB —
    enough to exceed the 50× ratio guard without allocating significant RAM.
    """
    payload = b'\x00' * int(uncompressed_mb * 1024 * 1024)
    compressed = zlib.compress(payload, level=9)[2:-4]  # strip zlib header/trailer
    crc = zlib.crc32(payload) & 0xFFFFFFFF
    usize = len(payload)
    csize = len(compressed)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        info = zipfile.ZipInfo('bomb.xml')
        info.compress_type = zipfile.ZIP_DEFLATED
        zf.writestr(info, payload)

    with open(path, 'wb') as f:
        f.write(buf.getvalue())


class TestZipBomb:
    def test_high_ratio_zip_fails(self, tmp_path):
        """A ZIP whose content expands >50× is rejected as a likely ZIP bomb."""
        path = str(tmp_path / 'bomb.xlsx')
        _make_zip_bomb_ratio(path, uncompressed_mb=3.0)
        results = lint_file(path)
        assert any(r[0] == FAIL and 'integrity' in r[1].lower() for r in results)
        assert any('ratio' in r[2].lower() or 'bomb' in r[2].lower()
                   or 'zip' in r[2].lower()
                   for r in results if r[0] == FAIL)

    def test_normal_xlsx_passes_zip_guard(self, tmp_path):
        """A legitimate xlsx has a low expansion ratio and must not be flagged."""
        path = _make_xlsx(tmp_path)
        results = lint_file(path)
        assert not any(r[0] == FAIL and 'bomb' in r[2].lower() for r in results)


# ── sheet dimensions / extent ─────────────────────────────────────────────────

class TestSheetDimensions:
    def test_normal_sheet_reports_dimensions(self, tmp_path):
        path = _make_xlsx(tmp_path, {'A1': 'x', 'C5': 'y'})
        results = lint_file(path)
        assert any(r[0] == OK and 'dimension' in r[1].lower() for r in results)

    def test_inflated_extent_warns(self, tmp_path):
        """Styled-but-empty cells inflating declared dimensions → warning."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'data'
        ws['B2'] = 'more'
        ws.cell(row=500, column=1).font = Font(bold=True)
        ws.cell(row=1, column=200).fill = PatternFill('solid', fgColor='FF0000')
        path = str(tmp_path / 'inflated.xlsx')
        wb.save(path)
        results = lint_file(path)
        assert any(r[0] == WARN and 'inflat' in r[2].lower() for r in results)

    def test_empty_sheet_warns(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.title = 'Empty'
        path = str(tmp_path / 'empty.xlsx')
        wb.save(path)
        results = lint_file(path)
        assert any(r[0] == WARN and 'empty' in r[2].lower() for r in results)


# ── merged cells ──────────────────────────────────────────────────────────────

class TestMergedCells:
    def test_merged_cells_warned(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'merged'
        ws.merge_cells('A1:C1')
        path = str(tmp_path / 'merged.xlsx')
        wb.save(path)
        results = lint_file(path)
        assert any(r[0] == WARN and 'merge' in r[2].lower() for r in results)

    def test_no_merged_cells_ok(self, tmp_path):
        path = _make_xlsx(tmp_path)
        results = lint_file(path)
        merged_results = [r for r in results if 'merge' in r[1].lower()]
        assert all(r[0] == OK for r in merged_results)


# ── formulas ──────────────────────────────────────────────────────────────────

class TestFormulas:
    def test_formula_cells_warned(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 10
        ws['A2'] = '=SUM(A1:A1)'
        path = str(tmp_path / 'formulas.xlsx')
        wb.save(path)
        results = lint_file(path)
        assert any(r[0] == WARN and 'formula' in r[2].lower() for r in results)


# ── multi-sheet ───────────────────────────────────────────────────────────────

class TestMultiSheet:
    def test_multi_sheet_reported(self, tmp_path):
        path = _make_xlsx(tmp_path, sheets={
            'Jan': {'A1': 'x'},
            'Feb': {'A1': 'y'},
            'Mar': {'A1': 'z'},
        })
        results = lint_file(path)
        assert any('3' in r[2] and 'sheet' in r[2].lower() for r in results)


# ── known limitations / advisory ──────────────────────────────────────────────

class TestAdvisoryNotes:
    def test_advisory_notes_always_present(self, tmp_path):
        path = _make_xlsx(tmp_path)
        results = lint_file(path)
        info_results = [r for r in results if r[0] == INFO]
        assert len(info_results) >= 1


# ── run_lint output ───────────────────────────────────────────────────────────

class TestRunLint:
    def test_returns_zero_when_no_fail(self, tmp_path):
        from grepxcel.lint import run_lint
        import io
        path = _make_xlsx(tmp_path)
        code = run_lint([path], out=io.StringIO())
        assert code == 0

    def test_returns_one_on_fail(self, tmp_path):
        from grepxcel.lint import run_lint
        import io
        code = run_lint([str(tmp_path / 'nope.xlsx')], out=io.StringIO())
        assert code == 1

    def test_multiple_files(self, tmp_path):
        from grepxcel.lint import run_lint
        import io
        p1 = _make_xlsx(tmp_path, name='a.xlsx')
        p2 = _make_xlsx(tmp_path, name='b.xlsx')
        code = run_lint([p1, p2], out=io.StringIO())
        assert code == 0
