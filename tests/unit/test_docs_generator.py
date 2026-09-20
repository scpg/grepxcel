"""Tests for the `docs` reference generator (grepxcel.docs_generator)."""
import zipfile

import openpyxl

from grepxcel.docs_generator import DocsGenerator


def test_output_is_byte_deterministic(tmp_path):
    """Two generations must be byte-identical so the docs-reference CI check
    (a byte diff) is meaningful — openpyxl otherwise stamps now()."""
    a = tmp_path / 'a'
    b = tmp_path / 'b'
    DocsGenerator().write(str(a))
    DocsGenerator().write(str(b))
    assert (a / 'pattern-reference.xlsx').read_bytes() == (b / 'pattern-reference.xlsx').read_bytes()
    assert (a / 'grepxcel-guide.docx').read_bytes() == (b / 'grepxcel-guide.docx').read_bytes()


def test_output_is_a_valid_workbook(tmp_path):
    DocsGenerator().write(str(tmp_path))
    wb = openpyxl.load_workbook(tmp_path / 'pattern-reference.xlsx')
    assert wb.sheetnames[0] == 'guide'
    assert 'pattern-reference' in wb.sheetnames


def test_source_date_epoch_is_honored(tmp_path, monkeypatch):
    monkeypatch.setenv('SOURCE_DATE_EPOCH', '1577836800')  # 2020-01-01T00:00:00Z
    DocsGenerator().write(str(tmp_path))
    wb = openpyxl.load_workbook(tmp_path / 'pattern-reference.xlsx')
    assert wb.properties.created.year == 2020


def test_lists_the_new_field_types(tmp_path):
    DocsGenerator().write(str(tmp_path))
    wb = openpyxl.load_workbook(tmp_path / 'pattern-reference.xlsx')
    text = ' '.join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value
    )
    for t in ('text', 'number', 'boolean'):
        assert t in text


def test_docx_is_generated(tmp_path):
    DocsGenerator().write(str(tmp_path))
    docx = tmp_path / 'grepxcel-guide.docx'
    assert docx.exists()
    with zipfile.ZipFile(docx) as zf:
        names = zf.namelist()
    assert 'word/document.xml' in names
    assert 'word/styles.xml' in names
    assert '[Content_Types].xml' in names


def test_write_returns_both_paths(tmp_path):
    paths = DocsGenerator().write(str(tmp_path))
    assert len(paths) == 2
    assert any('pattern-reference.xlsx' in p for p in paths)
    assert any('grepxcel-guide.docx' in p for p in paths)
