"""Case generator — builds one (pattern.csv, data.xlsx, manifest.json) triple.

``build_case(params, out_dir)`` realises one point of the §11 matrix: a type, a
shape, a storage case, a number format, a read direction and a seed. It writes
the two input files, returns the :class:`Manifest` that every oracle then checks
them against, and drops ``manifest.json`` beside them for hand inspection.

## Why instructions are emitted in scan order (§9)

``cell:next`` resolves to the **next non-empty cell** and then validates it — it
does *not* search forward for a matching label (``engine._process_cell`` →
``scanner.advance_to_next()``). So a pattern whose instruction order disagrees
with the sheet's scan order does not merely mis-anchor, it fails outright
("sheet is exhausted"). Every element is therefore placed at random and then
sorted by ``scan_index`` for the case's direction before any pattern row is
written. Same reason absolute refs are emitted sorted: the parser rejects an
out-of-order ``cell:<ref>`` at parse time, and it checks the order in the
declared direction (``pattern_parser._coord_after`` — TD compares column-first).

## Where the error cells go, and why it has to be the grid's tail

§4 puts ``error`` cells in scope for the ``lint`` and ``profile`` oracles while
keeping them out of reach of any ``var:`` field. Because ``cell:next`` consumes
the next non-empty cell whatever it contains, an error cell anywhere inside the
scan path would be swallowed by a scalar chain and break it. They are therefore
written to the **last line of the grid in scan order** — row ``GRID`` for LR,
column ``GRID`` for TD — while every extraction block is confined to
``1..GRID-2`` by giving the placer a reduced size. That leaves line ``GRID-1``
entirely empty, which satisfies the one-cell margin against every block, and
guarantees every extraction cell has a lower scan index than every error cell,
so the instruction list is always exhausted before the scanner could reach one.
"""
from __future__ import annotations

import hashlib
import os
import random
from dataclasses import dataclass

import openpyxl

from tests.oracle import catalog
from tests.oracle.manifest import CellSpec, Manifest, TableSpec, normalise
from tests.oracle.placement import GRID, OraclePlacer, sort_blocks, value_pos_for

# Imported, never copied (§0, §7).
from tests.unit.test_random_type_fixtures import write_data_xlsx, write_pattern_csv

SCALAR_SHAPES = ('cells1', 'cells2', 'cellsN')
TABLE_SHAPES = ('table1', 'table2', 'tableN')
SHAPES = SCALAR_SHAPES + TABLE_SHAPES
VIAS = ('next', 'abs')
HEADER_MODES = ('same_header', 'distinct_header')

#: Extraction blocks live in 1..PLACE_SIZE; line PLACE_SIZE+1 is the empty
#: margin; line GRID (= PLACE_SIZE+2) carries the error cells. See the docstring.
PLACE_SIZE = GRID - 2

#: Column D for every generated var:/lbl: field (§5 fixes it at '.*').
REGEX = '.*'


def _slug(text) -> str:
    """Filesystem/pytest-id-safe rendering of a number format.

    A readable part plus a short digest of the raw string. The digest is not
    decoration: sanitising punctuation alone is lossy enough to collide —
    ``$#,##0.00``, ``€#,##0.00`` and ``£#,##0.00`` all reduce to the same
    token, as do ``0.00`` and ``#,##0.00`` — which silently merged distinct
    matrix cases into duplicate case_ids (and pytest ids).
    """
    if text is None:
        return 'nofmt'
    raw = str(text)
    readable = ''.join(ch if (ch.isalnum() or ch in '-_') else '_' for ch in raw)
    readable = readable.strip('_') or 'fmt'
    digest = hashlib.blake2s(raw.encode('utf-8'), digest_size=2).hexdigest()
    return f'{readable}{digest}'


@dataclass(frozen=True)
class CaseParams:
    """One point of the §11 matrix."""
    type_name: str
    shape: str
    storage_key: str
    number_format: str | None
    direction: str
    seed: int
    via: str | None = None            # scalar shapes only
    header_mode: str | None = None    # table shapes only

    @property
    def storage_case(self) -> catalog.StorageCase:
        for case in catalog.STORAGE_MATRIX[self.type_name]:
            if case.key == self.storage_key:
                return case
        raise KeyError(f'{self.type_name} has no storage case {self.storage_key!r}')

    @property
    def case_id(self) -> str:
        """§8's '{type}-{shape}-{storage}-{fmt}-{direction}-seed{n}', plus the
        via / header_mode axis so every generated case has a unique id."""
        extra = self.via or self.header_mode or ''
        extra = f'-{extra}' if extra else ''
        return (f'{self.type_name}-{self.shape}{extra}-{self.storage_key}'
                f'-{_slug(self.number_format)}-{self.direction}-seed{self.seed}')


def _count_for(shape: str, rng: random.Random) -> int:
    if shape in ('cells1', 'table1'):
        return 1
    if shape in ('cells2', 'table2'):
        return 2
    return rng.randint(3, 10)   # cellsN / tableN (§9)


def _error_cell_positions(direction: str, count: int) -> list[tuple[int, int]]:
    """Positions on the grid's last scan line, spaced so the halo holds."""
    positions = []
    for i in range(count):
        offset = 1 + 3 * i
        if offset > GRID:
            break
        positions.append((GRID, offset) if direction == 'LR' else (offset, GRID))
    return positions


def _ref(pos: tuple[int, int]) -> str:
    return openpyxl.utils.get_column_letter(pos[1]) + str(pos[0])


def build_case(params: CaseParams, out_dir: str) -> Manifest:
    """Generate one case's pattern, data file and manifest into *out_dir*."""
    os.makedirs(out_dir, exist_ok=True)
    case = params.storage_case
    fmt = params.number_format
    direction = params.direction

    # §7: the seed selects pool entries; §9: the seed drives placement and counts.
    rng = random.Random(params.seed)
    m = _count_for(params.shape, rng)
    placer = OraclePlacer(params.seed, size=PLACE_SIZE)

    pattern_path = os.path.join(out_dir, 'pattern.csv')
    data_path = os.path.join(out_dir, 'data.xlsx')
    manifest = Manifest(
        case_id=params.case_id, direction=direction,
        pattern_path=pattern_path, data_path=data_path,
    )

    cells: dict[tuple[int, int], tuple] = {}
    defs_rows: list[list[str]] = [['config:', 'read.direction', direction]]
    seq_rows: list[list[str]] = []

    if params.shape in SCALAR_SHAPES:
        _build_scalars(params, case, fmt, m, rng, placer, manifest,
                       cells, defs_rows, seq_rows)
    else:
        _build_tables(params, case, fmt, m, rng, placer, manifest,
                      cells, defs_rows, seq_rows)

    # ── §4 error cells: profile/lint only, on the grid's last scan line ──
    err_storage, err_semantic, err_flags = catalog.ERROR_CELL_PROFILE
    for i, pos in enumerate(_error_cell_positions(direction, catalog.ERROR_CELLS_PER_CASE)):
        code = catalog.ERROR_CODES[(params.seed + i) % len(catalog.ERROR_CODES)]
        cells[pos] = (code, None)
        manifest.cells.append(CellSpec(
            ref=_ref(pos), row=pos[0], col=pos[1], field=None, role='error',
            written_value=code, number_format=None,
            extract_accepted=None, extract_reason=None,
            profile_storage=err_storage, profile_semantic=err_semantic,
            profile_flags=err_flags,
        ))

    write_pattern_csv(pattern_path, defs_rows + [[], ['START:']] + seq_rows + [['END:']])
    write_data_xlsx(data_path, cells)
    manifest.write(os.path.join(out_dir, 'manifest.json'))
    return manifest


def _build_scalars(params, case, fmt, m, rng, placer, manifest,
                   cells, defs_rows, seq_rows) -> None:
    """cells1 / cells2 / cellsN, for via='next' and via='abs'."""
    direction = params.direction
    with_anchor = params.via == 'next'
    blocks = [placer.place_scalar(direction, with_anchor) for _ in range(m)]
    blocks = sort_blocks(blocks, direction, placer.size)

    for i, block in enumerate(blocks):
        field_name = f'f{i:02d}'
        value = rng.choice(case.values)
        expected = normalise(case.extract_value(value))

        if with_anchor:
            anchor_pos = block.anchor
            value_pos = value_pos_for(anchor_pos, direction)
            anchor_text = f'ANCHOR_{i:02d}'
            anchor_field = f'a{i:02d}'
            cells[anchor_pos] = (anchor_text, None)
            manifest.cells.append(CellSpec(
                ref=_ref(anchor_pos), row=anchor_pos[0], col=anchor_pos[1],
                field=None, role='anchor', written_value=anchor_text,
                number_format=None, extract_accepted=None, extract_reason=None,
                profile_storage='s', profile_semantic='string',
                profile_flags=frozenset(),
            ))
            defs_rows.append(['lbl:', anchor_field, 'string', anchor_text])
            defs_rows.append(['var:', field_name, params.type_name, REGEX])
            seq_rows.append(['cell:next', anchor_field])
            seq_rows.append(['cell:next', field_name])
        else:
            value_pos = block.anchor
            defs_rows.append(['var:', field_name, params.type_name, REGEX])
            seq_rows.append([f'cell:{_ref(value_pos)}', field_name])

        cells[value_pos] = (value, fmt)
        manifest.cells.append(CellSpec(
            ref=_ref(value_pos), row=value_pos[0], col=value_pos[1],
            field=field_name, role='value', written_value=value,
            number_format=fmt,
            extract_accepted=case.accepted, extract_reason=case.reason,
            profile_storage=case.profile_storage,
            profile_semantic=case.profile_semantic,
            profile_flags=case.profile_flags,
        ))
        manifest.expected_extract[field_name] = expected


def _build_tables(params, case, fmt, m, rng, placer, manifest,
                  cells, defs_rows, seq_rows) -> None:
    """table1 / table2 / tableN, for same_header and distinct_header (§9)."""
    direction = params.direction
    # k columns, r data rows — seeded, all columns of the same type (§9).
    k = rng.randint(1, 3)
    r = rng.randint(1, 5)
    same = params.header_mode != 'distinct_header'

    blocks = [placer.place_table(data_rows=r, cols=k) for _ in range(m)]
    blocks = sort_blocks(blocks, direction, placer.size)

    if same:
        # ONE table:* block; every table uses identical header text; all
        # instances land under one group key (§9).
        header_texts = [f'COLHEAD_{j:02d}' for j in range(k)]
        header_fields = [f'h{j:02d}' for j in range(k)]
        col_fields = [f'items.col{j:02d}' for j in range(k)]
        for hf, ht in zip(header_fields, header_texts):
            defs_rows.append(['lbl:', hf, 'string', ht])
        for cf in col_fields:
            defs_rows.append(['var:', cf, params.type_name, REGEX])
        seq_rows.append(['table:*'])
        seq_rows.append(['', 'HEADER:1'] + header_fields)
        seq_rows.append(['', 'DATA:*'] + col_fields)

    for i, block in enumerate(blocks):
        if same:
            group_key = 'items'
            expected_instances = m
        else:
            header_texts = [f'HEAD{i:02d}_{j:02d}' for j in range(k)]
            header_fields = [f'h{i:02d}_{j:02d}' for j in range(k)]
            col_fields = [f't{i:02d}.col{j:02d}' for j in range(k)]
            group_key = f't{i:02d}'
            expected_instances = 1
            for hf, ht in zip(header_fields, header_texts):
                defs_rows.append(['lbl:', hf, 'string', ht])
            for cf in col_fields:
                defs_rows.append(['var:', cf, params.type_name, REGEX])
            # m separate table:1 blocks, emitted in scan order of their headers.
            seq_rows.append(['table:1'])
            seq_rows.append(['', 'HEADER:1'] + header_fields)
            seq_rows.append(['', 'DATA:*'] + col_fields)

        top, left = block.anchor
        for j, text in enumerate(header_texts):
            pos = (top, left + j)
            cells[pos] = (text, None)
            manifest.cells.append(CellSpec(
                ref=_ref(pos), row=pos[0], col=pos[1], field=None, role='header',
                written_value=text, number_format=None,
                extract_accepted=None, extract_reason=None,
                profile_storage='s', profile_semantic='string',
                profile_flags=frozenset(),
            ))

        data_rows: list[list] = []
        for ri in range(r):
            row_values = []
            for j in range(k):
                pos = (top + 1 + ri, left + j)
                value = rng.choice(case.values)
                cells[pos] = (value, fmt)
                manifest.cells.append(CellSpec(
                    ref=_ref(pos), row=pos[0], col=pos[1],
                    field=col_fields[j], role='data', written_value=value,
                    number_format=fmt,
                    extract_accepted=case.accepted, extract_reason=case.reason,
                    profile_storage=case.profile_storage,
                    profile_semantic=case.profile_semantic,
                    profile_flags=case.profile_flags,
                ))
                row_values.append(normalise(case.extract_value(value)))
            data_rows.append(row_values)

        manifest.tables.append(TableSpec(
            group_key=group_key, header_ref=_ref(block.anchor),
            columns=list(col_fields), data_rows=data_rows,
            expected_instances=expected_instances,
        ))

    # expected_extract mirrors the engine's nested shape for the data rows only;
    # _source / header / _meta are asserted structurally by check_extract from
    # the TableSpecs instead (§10.4).
    for spec in manifest.tables:
        rows = [
            {name.partition('.')[2] or name: value
             for name, value in zip(spec.columns, row)}
            for row in spec.data_rows
        ]
        manifest.expected_extract.setdefault(spec.group_key, []).append(rows)
