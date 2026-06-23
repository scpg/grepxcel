import csv
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from .models import Config, FieldDef, TemplateColumn, TemplateRow, CellInstruction, TableInstruction, SeekInstruction, DirectionInstruction, LBL_MATCH_MODES
from .security import check_regex_safety, SecurityError

_MAX_PATTERN_CELL_LEN = 1_000  # max characters in any pattern file cell value

_A1_RE = re.compile(r'^[A-Z]{1,3}[1-9][0-9]*$', re.IGNORECASE)
_BOUNDED_DATA_RE = re.compile(r'^\{(\d+),(\d+)\}$')
_POSINT_RE = re.compile(r'^[1-9][0-9]*$')

# Recognised template-row keywords inside a table: block.
_VALID_TABLE_ROW_TYPES = frozenset({'HEADER', 'DATA', 'FOOTER', 'SPLITTER', 'SKIP_IF'})

# Pattern-format/semantics version ("compatibility generation"). A single
# monotonic integer that bumps ONLY on a backward-incompatible change; additive
# changes never bump it. A pattern may declare `config: | pattern.version | N`;
# absent ⇒ MIN (the original format). The engine understands MIN..CURRENT.
CURRENT_PATTERN_VERSION = 1
MIN_SUPPORTED_PATTERN_VERSION = 1


_TRUTHY = frozenset({'1', 'true', 'yes', 'on', 'y'})
_FALSY  = frozenset({'0', 'false', 'no', 'off', 'n', ''})

# Field types the engine knows how to validate (see utils.validate_type).
# Keep this in lockstep with that function — a name here that it can't handle
# would parse cleanly but make every value fail validation.
_VALID_FIELD_TYPES = frozenset({
    'string', 'text',
    'integer',
    'number', 'float', 'decimal',
    'currency', 'percentage',
    'boolean', 'bool',
    'time', 'duration', 'date', 'datetime', 'timestamp',
})


def _truthy(val) -> bool:
    """Interpret a config cell as a boolean.

    Accepts yes/no, true/false, on/off, 1/0 (case-insensitive). Excel may store
    the cell as a real bool, so handle that too. Unknown text → False.
    """
    if isinstance(val, bool):
        return val
    return str(val).strip().lower() in _TRUTHY


class PatternError(Exception):
    """Raised when a pattern file contains a structural or ordering error."""


def _parse_pattern_version(val) -> int:
    """Coerce a `pattern.version` cell to a supported integer version.

    Excel may store the value as int, float (1.0), or text ('1'). Rejects
    non-integers, versions below the supported floor, and versions newer than
    this engine understands (with an upgrade hint)."""
    try:
        if isinstance(val, bool):
            raise ValueError
        if isinstance(val, float):
            if not val.is_integer():
                raise ValueError
            v = int(val)
        else:
            v = int(str(val).strip())
    except (ValueError, TypeError):
        raise PatternError(
            f"Invalid pattern.version {val!r} — must be a whole number (e.g. 1)."
        )
    if v < MIN_SUPPORTED_PATTERN_VERSION:
        raise PatternError(
            f"pattern.version {v} is below the minimum supported "
            f"({MIN_SUPPORTED_PATTERN_VERSION})."
        )
    if v > CURRENT_PATTERN_VERSION:
        raise PatternError(
            f"pattern.version {v} is newer than this grepxcel understands "
            f"(supports up to {CURRENT_PATTERN_VERSION}). Upgrade grepxcel to use "
            f"this pattern."
        )
    return v


def _coord_after(prev: tuple, new: tuple, direction: str) -> bool:
    """Return True if `new` comes strictly after `prev` in `direction` reading order."""
    pr, pc = prev
    nr, nc = new
    if direction == 'LR':
        return (nr, nc) > (pr, pc)
    return (nc, nr) > (pc, pr)  # TD: column primary


def _coord_str(pos: tuple) -> str:
    return f'{get_column_letter(pos[1])}{pos[0]}'


class PatternParser:
    def parse(self, filepath: str) -> tuple:
        """
        Parse a pattern file (.xlsx or .csv) into the same structures regardless
        of source format. Returns (global_config, defs, start_sequence).

        Both formats are read into a common 2D-grid IR (list[list[str|None]]);
        all semantics below operate on that grid, so the two source formats share
        one parser. See _read_grid for the format dispatch.
        """
        rows = self._read_grid(filepath)

        global_config = Config()
        defs = {}
        start_sequence = []

        i = 0
        in_start = False
        last_abs_pos: tuple | None = None  # (row, col) of last cell:XY seen
        # Scan direction in effect for abs-ref ordering checks; starts at the
        # global config value and is updated by each dir: instruction.
        current_direction: str | None = None

        while i < len(rows):
            row = self._pad(rows[i])
            col_a = row[0]
            # Structural keywords are case-insensitive (cell:/CELL:/Cell: all
            # work). Match on a lowercased view; field names and addresses are
            # read from the original col_a so their case is preserved.
            col_a_l = col_a.lower() if isinstance(col_a, str) else col_a

            if col_a_l == 'end:':
                break

            if col_a_l == 'start:':
                self._check_comment_zone(row, 1, i + 1, 'START:')   # B+ may be a # comment
                in_start = True
                i += 1
                continue

            if not in_start:
                if col_a_l == 'config:':
                    self._apply_global_config(row, global_config)
                    self._check_comment_zone(row, 3, i + 1, 'config:')    # D+ comment
                elif col_a_l in ('def:', 'var:'):
                    fd = self._parse_field(row, role='var', row_num=i + 1)
                    defs[fd.name] = fd
                    self._check_comment_zone(row, 4, i + 1, col_a_l)      # E+ comment
                elif isinstance(col_a_l, str) and col_a_l.startswith('lbl:'):
                    suffix = col_a_l[4:]  # '' | 'literal' | 'glob' | 'regexp'
                    if suffix and suffix not in LBL_MATCH_MODES:
                        raise PatternError(
                            f"Unknown lbl: variant {col_a!r} at pattern row {i + 1}. "
                            f"Use lbl: (global default), lbl:literal, lbl:glob, or lbl:regexp."
                        )
                    fd = self._parse_field(row, role='lbl', row_num=i + 1,
                                           lbl_match_override=suffix if suffix else None)
                    defs[fd.name] = fd
                    self._check_comment_zone(row, 4, i + 1, col_a_l)
                elif col_a_l in ('doc:', 'info:'):
                    pass  # inline documentation — ignored by engine
                elif col_a is not None and str(col_a).strip() != '':
                    raise PatternError(
                        f"Unrecognised row {col_a!r} at pattern row {i + 1} "
                        f"(before START:). Expected config:, var:, lbl:, def:, "
                        f"doc:, info:, or START:."
                    )
                i += 1
                continue

            # Inside START: section
            if col_a_l and col_a_l.startswith('cell:'):
                raw = col_a.split(':', 1)[1]
                field = str(row[1] or 'IGNORE')
                self._check_comment_zone(row, 2, i + 1, 'cell:')      # C+ may be a # comment
                if raw.lower() in ('1', 'next'):
                    start_sequence.append(CellInstruction(multiplicity=raw.lower(), field=field))
                elif _A1_RE.match(raw):
                    ref = raw.upper()
                    new_pos = coordinate_to_tuple(ref)
                    _dir = current_direction or global_config.read_direction
                    if last_abs_pos is not None and not _coord_after(last_abs_pos, new_pos, _dir):
                        raise PatternError(
                            f"cell:{ref} at pattern row {i + 1} is before or equal to the "
                            f"previous absolute reference {_coord_str(last_abs_pos)} in "
                            f"{_dir} reading order — unreachable"
                        )
                    last_abs_pos = new_pos
                    start_sequence.append(CellInstruction(multiplicity='abs', field=field, target=ref))
                else:
                    raise PatternError(
                        f"Unknown cell instruction 'cell:{raw}' at pattern row {i + 1}. "
                        f"Use 'cell:next', 'cell:1', or a cell coordinate like 'cell:B5'."
                    )
                i += 1

            elif col_a_l and col_a_l.startswith('seek:'):
                raw = col_a.split(':', 1)[1]
                if not _A1_RE.match(raw):
                    raise PatternError(
                        f"Invalid seek instruction 'seek:{raw}' at pattern row {i + 1}. "
                        f"Use a cell coordinate like 'seek:G5' (A1-notation)."
                    )
                ref = raw.upper()
                # seek: fully resets the abs-ref ordering constraint so that the
                # cell immediately at the seek target (or any cell after it) is
                # accepted. Backward refs after seek are caught at runtime.
                last_abs_pos = None
                self._check_comment_zone(row, 1, i + 1, 'seek:')    # B+ may be a # comment
                start_sequence.append(SeekInstruction(target=ref))
                i += 1

            elif col_a_l and col_a_l.startswith('dir:'):
                raw = col_a.split(':', 1)[1].strip().upper()
                if raw not in ('LR', 'TD'):
                    raise PatternError(
                        f"Invalid dir instruction 'dir:{col_a.split(':', 1)[1]}' at "
                        f"pattern row {i + 1}. Use 'dir:LR' (left-to-right) or "
                        f"'dir:TD' (top-down)."
                    )
                # Switching direction changes the meaning of "forward", so the
                # abs-ref ordering constraint is reset (like seek:). Subsequent
                # abs refs are checked in the new direction.
                current_direction = raw
                last_abs_pos = None
                self._check_comment_zone(row, 1, i + 1, 'dir:')    # B+ may be a # comment
                start_sequence.append(DirectionInstruction(direction=raw))
                i += 1

            elif col_a_l and col_a_l.startswith('table:'):
                mult = col_a.split(':', 1)[1]
                # table:<mult> must be '*' or a positive instance count.
                if mult != '*' and not _POSINT_RE.match(mult):
                    raise PatternError(
                        f"Invalid table multiplicity 'table:{mult}' at pattern row "
                        f"{i + 1}. Use 'table:*' (all instances) or a positive "
                        f"count like 'table:1'."
                    )
                table_config = Config(
                    read_direction=global_config.read_direction,
                    currency_sign=global_config.currency_sign,
                    empty_aliases=list(global_config.empty_aliases),
                    ignore_case=global_config.ignore_case,
                    lbl_match=global_config.lbl_match,
                )
                template_rows = []
                i += 1

                while i < len(rows):
                    sub = self._pad(rows[i])
                    sub_a = sub[0]
                    sub_b = sub[1]

                    # Non-empty col A means we're back at the top level
                    if sub_a is not None:
                        break

                    # Blank row inside table block
                    if sub_b is None:
                        i += 1
                        continue

                    if isinstance(sub_b, str) and sub_b.lower() == 'config:':
                        key = str(sub[2]).lower() if sub[2] is not None else ''
                        val = sub[3]
                        if key == 'read.direction' and val:
                            direction = str(val).strip().upper()
                            if direction not in ('LR', 'TD'):
                                raise PatternError(
                                    f"Invalid read.direction '{val}' in table config. "
                                    f"Valid values: LR (left-to-right), TD (top-down)")
                            table_config.read_direction = direction
                        elif key == 'ignore.case' and val is not None:
                            table_config.ignore_case = _truthy(val)
                        i += 1
                        continue

                    # HEADER / SPLITTER / DATA / FOOTER / SKIP_IF row
                    sub_b_str = str(sub_b) if sub_b else ''
                    is_skip_if = sub_b_str.upper() == 'SKIP_IF'
                    has_colon  = ':' in sub_b_str

                    if is_skip_if or has_colon:
                        if is_skip_if:
                            row_type = 'SKIP_IF'
                            row_mult = ''
                        else:
                            row_type, row_mult = sub_b_str.rsplit(':', 1)
                            # Row-type keyword is case-insensitive (HEADER/header);
                            # the multiplicity (* / n / {n,m}) keeps its case.
                            row_type = row_type.upper()

                        # Reject typo'd row keywords (e.g. 'HEDER:1') — otherwise
                        # they would be silently dropped by the engine.
                        if row_type not in _VALID_TABLE_ROW_TYPES:
                            raise PatternError(
                                f"Unknown table row type {sub_b_str!r} at pattern "
                                f"row {i + 1}. Valid row types: "
                                f"{', '.join(sorted(_VALID_TABLE_ROW_TYPES))}."
                            )

                        # Validate the multiplicity for each row type.
                        #   DATA            → '*', positive int, or '{n,m}'
                        #   HEADER/FOOTER/SPLITTER → a positive int (e.g. ':1')
                        if row_type == 'DATA':
                            if not (row_mult == '*' or _POSINT_RE.match(row_mult)
                                    or _BOUNDED_DATA_RE.match(row_mult)):
                                raise PatternError(
                                    f"Invalid DATA multiplicity 'DATA:{row_mult}' at "
                                    f"pattern row {i + 1}. Use 'DATA:*', 'DATA:1', or "
                                    f"a bounded 'DATA:{{n,m}}'."
                                )
                        elif row_type in ('HEADER', 'FOOTER', 'SPLITTER'):
                            if not _POSINT_RE.match(row_mult):
                                raise PatternError(
                                    f"Invalid {row_type} multiplicity "
                                    f"'{row_type}:{row_mult}' at pattern row {i + 1}. "
                                    f"Use a positive count like '{row_type}:1'."
                                )

                        cols_raw = list(sub[2:])
                        while cols_raw and cols_raw[-1] is None:
                            cols_raw.pop()
                        columns = [
                            TemplateColumn(field=str(v) if v is not None else 'EMPTY')
                            for v in cols_raw
                        ]

                        min_rows, max_rows = 0, None
                        if row_type == 'DATA':
                            m = _BOUNDED_DATA_RE.match(row_mult)
                            if m:
                                min_rows = int(m.group(1))
                                max_rows = int(m.group(2))
                                if min_rows > max_rows:
                                    raise PatternError(
                                        f'DATA:{{{min_rows},{max_rows}}}: '
                                        f'min ({min_rows}) must be ≤ max ({max_rows})'
                                    )

                        template_rows.append(TemplateRow(
                            row_type=row_type,
                            multiplicity=row_mult,
                            columns=columns,
                            min_rows=min_rows,
                            max_rows=max_rows,
                        ))

                    i += 1

                # ── DATA type mutual-exclusivity validation ────────────────────
                data_rows = [r for r in template_rows if r.row_type == 'DATA']
                if data_rows:
                    has_bounded = any(r.max_rows is not None for r in data_rows)
                    has_star    = any(r.multiplicity == '*' for r in data_rows)
                    has_one     = any(r.multiplicity == '1' for r in data_rows)
                    type_count  = sum([has_bounded, has_star, has_one])
                    if type_count > 1:
                        raise PatternError(
                            'Table block mixes DATA types. '
                            'Use only one of: DATA:1 (repeatable), DATA:*, or DATA:{n,m}.'
                        )
                    if (has_bounded or has_star) and len(data_rows) > 1:
                        kind = 'DATA:{n,m}' if has_bounded else 'DATA:*'
                        raise PatternError(
                            f'Only one {kind} row is allowed per table block.'
                        )
                    # SKIP_IF is meaningful with DATA:* and DATA:{n,m} (rows are
                    # filtered from output while scanning continues). It is not
                    # meaningful with DATA:1 — skipping the only row creates
                    # ambiguous extraction semantics.
                    skip_if_rows = [r for r in template_rows if r.row_type == 'SKIP_IF']
                    if skip_if_rows and has_one:
                        raise PatternError(
                            'SKIP_IF requires DATA:{n,m} or DATA:*. '
                            'SKIP_IF has no effect with DATA:1.'
                        )

                # Structural rules for a table block:
                #   • at least one DATA row is required (HEADER and FOOTER are
                #     optional). A block with no DATA extracts nothing and would
                #     otherwise crash the engine.
                #   • any HEADER row must come before the first DATA row.
                #   • any FOOTER row must come after the last DATA row.
                row_types = [r.row_type for r in template_rows]
                if 'DATA' not in row_types:
                    raise PatternError(
                        f"table:{mult} block has no DATA row — a table must define "
                        f"at least one DATA row (HEADER and FOOTER are optional). "
                        f"Add e.g. a 'DATA:*' row."
                    )
                _data_idxs  = [k for k, t in enumerate(row_types) if t == 'DATA']
                _first_data, _last_data = _data_idxs[0], _data_idxs[-1]
                if any(k > _first_data for k, t in enumerate(row_types) if t == 'HEADER'):
                    raise PatternError(
                        f"table:{mult} block has a HEADER row after a DATA row — "
                        f"HEADER rows must come before DATA."
                    )
                if any(k < _last_data for k, t in enumerate(row_types) if t == 'FOOTER'):
                    raise PatternError(
                        f"table:{mult} block has a FOOTER row before a DATA row — "
                        f"FOOTER rows must come after DATA."
                    )

                start_sequence.append(TableInstruction(
                    multiplicity=mult,
                    config=table_config,
                    rows=template_rows,
                ))

            elif col_a_l in ('doc:', 'info:'):
                i += 1   # inline comment inside START — ignored by engine

            elif col_a is None or str(col_a).strip() == '':
                i += 1   # blank row (incl. table template rows) — skip

            else:
                raise PatternError(
                    f"Unrecognised instruction {col_a!r} at pattern row {i + 1} "
                    f"inside START:. Expected cell:, seek:, dir:, table:, doc:, "
                    f"or a blank row. (Keywords are case-insensitive — check for "
                    f"a typo.)"
                )

        self._check_nesting_conflicts(defs, start_sequence)
        return global_config, defs, start_sequence

    def _check_nesting_conflicts(self, defs, start_sequence) -> None:
        """A field cannot be both a value AND the parent of another field, e.g.
        'week1.day' alongside 'week1.day.timeIn'. The nested output would try to
        assign a key into a scalar value and crash. Detect this here and fail
        fast with a clear message instead of a TypeError deep in output building.

        Only fields that actually reach the output are considered: those
        referenced in the sequence, excluding IGNORE/EMPTY and lbl: anchors
        (lbl: values are stripped from the JSON, so they never nest)."""
        names: set[str] = set()
        for instr in start_sequence:
            fields = []
            if isinstance(instr, CellInstruction):
                fields = [instr.field]
            elif isinstance(instr, TableInstruction):
                fields = [c.field for r in instr.rows for c in r.columns]
            for field in fields:
                if field in ('IGNORE', 'EMPTY'):
                    continue
                fd = defs.get(field)
                if fd is not None and fd.role == 'lbl':
                    continue
                names.add(field)

        ordered = sorted(names)
        for idx, parent in enumerate(ordered):
            for child in ordered[idx + 1:]:
                if child.startswith(parent + '.'):
                    raise PatternError(
                        f"Field {parent!r} is used both as a value and as the "
                        f"parent of {child!r} — a field name cannot be both. "
                        f"Rename one (e.g. {parent!r} -> '{parent}.value')."
                    )

    # ── Grid readers (front-ends over the common 2D-grid IR) ────────────────────

    # ── Grid readers (front-ends over the common 2D-grid IR) ────────────────────

    def _read_grid(self, filepath: str) -> list:
        """
        Read a pattern file into the common grid IR, dispatching on extension.

          .csv  → _read_csv_and_validate  (plain text, comma-separated)
          .xlsx → openpyxl + _read_and_validate  (default)

        Both return list[list[str|None]] with empty cells as None, so the
        downstream semantic parser is identical for either source format.
        """
        _, ext = os.path.splitext(filepath)
        if ext.lower() == '.csv':
            return self._read_csv_and_validate(filepath)
        wb = openpyxl.load_workbook(filepath)  # data_only=False: we want to see formulas
        ws = wb.active
        return self._read_and_validate(ws)

    @staticmethod
    def _detect_delimiter(sample: str) -> str:
        """Pick the CSV delimiter from ``,`` ``;`` or tab.

        csv.Sniffer is tried first (it understands quoting), but it raises on
        short or ragged samples — pattern files are ragged by nature (a lone
        ``START:`` line next to multi-column ``var:`` lines). The fallback then
        chooses whichever candidate appears on the most lines, defaulting to
        ``,`` when none is present (a genuinely single-column file).
        """
        try:
            return csv.Sniffer().sniff(sample, delimiters=',;\t').delimiter
        except csv.Error:
            pass
        counts = {d: sum(d in line for line in sample.splitlines())
                  for d in (',', ';', '\t')}
        best = max(counts, key=counts.get)
        return best if counts[best] else ','

    def _read_csv_and_validate(self, filepath: str) -> list:
        """
        Read a CSV pattern file into the grid IR with the same content rules as
        the xlsx reader.

        - Empty fields → None, so a blank column A (which marks table-template
          rows) behaves identically to a blank xlsx cell.
        - Leading '=' is rejected for parity with the xlsx formula guard.
        - Per-cell length is capped at _MAX_PATTERN_CELL_LEN.

        The delimiter is auto-detected (``,`` ``;`` or tab) so CSVs exported by
        Excel in locales that use ``;`` as the list separator parse correctly;
        it falls back to ``,`` when the sample is ambiguous (e.g. one column).

        Authoring note: a regex value containing the active delimiter (e.g. a
        comma in \\d{1,3} when the file is comma-separated) must be quoted in the
        CSV ("\\d{1,3}") — the csv module unquotes it correctly.
        utf-8-sig transparently strips a BOM written by Excel's "Save as CSV".
        """
        rows: list = []
        try:
            with open(filepath, newline='', encoding='utf-8-sig') as fh:
                sample = fh.read(8192)
                fh.seek(0)
                delimiter = self._detect_delimiter(sample)
                for line_no, raw in enumerate(
                        csv.reader(fh, delimiter=delimiter), start=1):
                    row_values = []
                    for col_idx, val in enumerate(raw):
                        if val == '':
                            row_values.append(None)
                            continue
                        if val.startswith('='):
                            coord = f'{get_column_letter(col_idx + 1)}{line_no}'
                            raise SecurityError(
                                f'Formulas are not allowed in pattern files. '
                                f'Cell {coord} contains: {val!r}  '
                                f'Replace it with a plain text value.'
                            )
                        if len(val) > _MAX_PATTERN_CELL_LEN:
                            coord = f'{get_column_letter(col_idx + 1)}{line_no}'
                            raise SecurityError(
                                f'Pattern file cell {coord} value is too long '
                                f'({len(val)} chars, limit is {_MAX_PATTERN_CELL_LEN}). '
                                f'Pattern values should be short identifiers or regex patterns.'
                            )
                        row_values.append(val)
                    rows.append(row_values)
        except UnicodeDecodeError as exc:
            raise SecurityError(
                f'Pattern CSV {filepath!r} is not valid UTF-8 text: {exc}'
            )
        return rows

    def _read_and_validate(self, ws) -> list:
        """
        Read all rows from the pattern worksheet, enforcing that every cell is
        either empty or a plain string.  Formulas, numbers, dates, and booleans
        are all rejected — the pattern file is a configuration document, not a
        spreadsheet.
        """
        rows = []
        for row in ws.iter_rows():
            row_values = []
            for cell in row:
                val = cell.value
                if val is None:
                    row_values.append(None)
                    continue

                # Formulas are never allowed in pattern files
                if cell.data_type == 'f' or (isinstance(val, str) and val.startswith('=')):
                    raise SecurityError(
                        f'Formulas are not allowed in pattern files. '
                        f'Cell {cell.coordinate} contains: {val!r}  '
                        f'Replace it with a plain text value.'
                    )

                # Only plain strings are accepted
                if not isinstance(val, str):
                    raise SecurityError(
                        f'Pattern file cells must contain plain text only. '
                        f'Cell {cell.coordinate} contains a {type(val).__name__} value: {val!r}  '
                        f'All values in a pattern file must be strings.'
                    )

                # Guard against excessively long values
                if len(val) > _MAX_PATTERN_CELL_LEN:
                    raise SecurityError(
                        f'Pattern file cell {cell.coordinate} value is too long '
                        f'({len(val)} chars, limit is {_MAX_PATTERN_CELL_LEN}). '
                        f'Pattern values should be short identifiers or regex patterns.'
                    )

                row_values.append(val)
            rows.append(row_values)
        return rows

    def _pad(self, row, length=10) -> list:
        """Ensure a row list has at least `length` elements."""
        row = list(row)
        while len(row) < length:
            row.append(None)
        return row

    def _check_comment_zone(self, row, start_idx: int, row_num: int, context: str) -> None:
        """Validate the trailing cells (from start_idx) of a non-table row.

        The trailing cells may be empty, or a comment: a cell whose text starts
        with '#' begins a comment that runs to the end of the row. Any non-empty
        cell that is not a comment (and not already inside one) is a mistake —
        fail fast so a value typed into the wrong column is never silently lost.
        Comments are NOT processed for table rows (there '#' is a literal value).
        """
        in_comment = False
        for j in range(start_idx, len(row)):
            val = row[j]
            if val is None or in_comment:
                continue
            if str(val).lstrip().startswith('#'):
                in_comment = True
                continue
            raise PatternError(
                f"Unexpected content {val!r} in column {get_column_letter(j + 1)} "
                f"at pattern row {row_num} ({context}). Columns after the {context} "
                f"fields may only contain a comment starting with '#'."
            )

    def _apply_global_config(self, row, config: Config):
        key, val = row[1], row[2]
        if key == 'read.direction' and val:
            direction = str(val).strip().upper()
            if direction not in ('LR', 'TD'):
                raise PatternError(
                    f"Invalid read.direction '{val}'. "
                    f"Valid values: LR (left-to-right), TD (top-down)")
            config.read_direction = direction
        elif key == 'currency.sign' and val:
            config.currency_sign = str(val)
        elif key == 'empty.aliases' and val:
            config.empty_aliases.append(str(val))
        elif key == 'ignore.case' and val is not None:
            config.ignore_case = _truthy(val)
        elif key == 'lbl.match' and val is not None:
            mode = str(val).strip().lower()
            if mode not in LBL_MATCH_MODES:
                raise PatternError(
                    f"Invalid lbl.match value {val!r}. Valid values: "
                    f"{', '.join(sorted(LBL_MATCH_MODES))}."
                )
            config.lbl_match = mode
        elif key == 'pattern.version' and val is not None:
            config.pattern_version = _parse_pattern_version(val)
            config.pattern_version_explicit = True

    def _parse_field(self, row, role: str = 'var', row_num: int | None = None,
                     lbl_match_override: str | None = None) -> FieldDef:
        where = f' at pattern row {row_num}' if row_num is not None else ''
        name  = str(row[1]) if row[1] else ''
        if not name:
            raise PatternError(
                f"A {role}: definition{where} has no field name. "
                f"Expected:  {role}: | <name> | <type> | <regex>"
            )
        type_ = str(row[2]) if row[2] else 'string'
        if type_ not in _VALID_FIELD_TYPES:
            raise PatternError(
                f"Unknown field type {type_!r} for {role}: {name!r}{where}. "
                f"Valid types: {', '.join(sorted(_VALID_FIELD_TYPES))}."
            )
        regex = str(row[3]) if row[3] else '.*'
        # Skip regex safety check for lbl: fields in non-regexp modes — the
        # pattern is treated as a literal string or glob, not compiled as a regex.
        if role != 'lbl' or lbl_match_override not in ('literal', 'glob'):
            check_regex_safety(regex, field_name=name)
        return FieldDef(name=name, type=type_, regex=regex, role=role,
                        lbl_match=lbl_match_override)

    # backward-compat alias
    def _parse_def(self, row) -> FieldDef:
        return self._parse_field(row, role='var')
