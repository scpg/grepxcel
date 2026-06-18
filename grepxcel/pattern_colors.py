"""Colorize pattern .xlsx files for visual readability.

Applies row-level fills based on the row type (config:, lbl:, var:, etc.)
without changing any cell content. The color scheme is aligned with
``docs_generator.py`` so patterns and the reference sheet look consistent.

Usage::

    from grepxcel.pattern_colors import colorize_pattern_file
    colorize_pattern_file('pattern.xlsx')
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import PatternFill, Font

# ── Color scheme (aligned with docs_generator) ───────────────────────────────

ROW_FILLS = {
    'config':  PatternFill('solid', fgColor='FFE8D0'),  # orange
    'lbl':     PatternFill('solid', fgColor='D0E8FF'),  # blue
    'var':     PatternFill('solid', fgColor='D0FFD0'),  # green
    'doc':     PatternFill('solid', fgColor='FFFFD0'),  # yellow
    'comment': PatternFill('solid', fgColor='F0F0F0'),  # light grey
    'marker':  PatternFill('solid', fgColor='E0E0E0'),  # grey (START/END)
    'cell':    PatternFill('solid', fgColor='F0E0FF'),  # lavender
    'seek':    PatternFill('solid', fgColor='F0E0FF'),  # lavender
    'table':   PatternFill('solid', fgColor='E8D0FF'),  # purple
    'header':  PatternFill('solid', fgColor='D0E8FF'),  # blue (like lbl)
    'data':    PatternFill('solid', fgColor='D0FFD0'),  # green (like var)
    'footer':  PatternFill('solid', fgColor='FFE8D0'),  # orange
    'skip_if': PatternFill('solid', fgColor='FFE0C0'),  # light orange
    'splitter': PatternFill('solid', fgColor='F8F0FF'), # light lavender
}

_COMMENT_FONT = Font(italic=True, color='808080')

_TABLE_ROW_KEYWORDS = frozenset({
    'HEADER', 'DATA', 'FOOTER', 'SKIP_IF', 'SPLITTER',
})


def _classify_row(ws, row_num: int) -> str | None:
    """Return the fill key for a row, or None if unrecognized."""
    col_a = ws.cell(row=row_num, column=1).value
    col_b = ws.cell(row=row_num, column=2).value

    if col_a is not None:
        a = str(col_a).strip()

        if a.startswith('#'):
            return 'comment'

        a_upper = a.upper()
        if a_upper in ('START:', 'END:'):
            return 'marker'
        if a_upper.startswith('TABLE:'):
            return 'table'
        if a_upper.startswith('CONFIG:'):
            return 'config'

        a_key = a.rstrip(':').lower()
        if a_key in ('lbl', 'var', 'doc'):
            return a_key
        if a_key.startswith('cell:') or a_key == 'cell':
            return 'cell'
        if a_key.startswith('seek:') or a_key == 'seek':
            return 'seek'

    if col_a is None and col_b is not None:
        b_base = str(col_b).split(':')[0].strip().upper()
        if b_base in _TABLE_ROW_KEYWORDS:
            return b_base.lower()

    return None


def colorize_pattern_file(path: str) -> str:
    """Apply row-level color fills to a pattern .xlsx file in place.

    Returns the path for convenience (allows chaining).
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    # Track max content width per column for auto-sizing
    col_widths: dict[int, int] = {}

    for row_num in range(1, max_row + 1):
        fill_key = _classify_row(ws, row_num)
        fill = ROW_FILLS.get(fill_key) if fill_key else None

        for col_num in range(1, max_col + 1):
            cell = ws.cell(row=row_num, column=col_num)
            if fill:
                cell.fill = fill
            if fill_key == 'comment':
                cell.font = _COMMENT_FONT

            val = cell.value
            if val is not None:
                w = len(str(val)) + 2
                col_widths[col_num] = max(col_widths.get(col_num, 8), w)

    for col_num, width in col_widths.items():
        letter = openpyxl.utils.get_column_letter(col_num)
        ws.column_dimensions[letter].width = min(width, 60)

    wb.save(path)
    return path
