"""`grepxcel lint` — inspect an Excel file before extraction.

Reports, with a ✓ / ⚠ / ✗ / ℹ checklist, potential issues in a data file
that might cause extraction to fail or produce unexpected results. Checks:

* File access and format (extension, magic bytes, encryption/IRM)
* Sheet dimensions — declared vs real used extent
* Merged cells
* Formula cells (cached values may be stale)
* Empty sheets
* Multi-sheet inventory

Also includes advisory notes about known limitations not yet detected
(e.g. Microsoft Information Protection, password-protected sheets).
"""
from __future__ import annotations

import os
import sys
import zipfile

from .color import MARK_FAIL, MARK_INFO, MARK_OK, MARK_WARN, colorize_marks, should_color

OK, WARN, FAIL, INFO = 'ok', 'warn', 'fail', 'info'
_MARK = {OK: MARK_OK, WARN: MARK_WARN, FAIL: MARK_FAIL, INFO: MARK_INFO}

# OLE Compound Document magic (D0 CF 11 E0 A1 B1 1A E1)
_OLE_MAGIC = b'\xd0\xcf\x11\xe0'
_ZIP_MAGIC = b'PK\x03\x04'
_MACRO_EXTENSIONS = {'.xlsm', '.xlsb', '.xls'}

Result = tuple


def lint_file(path: str) -> list[Result]:
    """Run all lint checks on a single file. Returns a list of results."""
    results: list[Result] = []

    # ── 1. File existence and type ────────────────────────────────────────
    if not os.path.exists(path):
        results.append((FAIL, 'file access', f'File does not exist: {path}'))
        _add_advisory(results)
        return results

    if not os.path.isfile(path):
        results.append((FAIL, 'file access', f'Not a regular file: {path}'))
        _add_advisory(results)
        return results

    _, ext = os.path.splitext(path)
    ext = ext.lower()

    if ext in _MACRO_EXTENSIONS:
        results.append((FAIL, 'file format',
                         f'{ext!r} files are not accepted — macro-enabled and binary '
                         f'Excel formats are blocked. Save as .xlsx first.'))
        _add_advisory(results)
        return results

    if ext != '.xlsx':
        results.append((FAIL, 'file format',
                         f'Unsupported file type {ext!r}. Only .xlsx files are accepted.'))
        _add_advisory(results)
        return results

    # ── 2. Magic bytes — ZIP vs OLE vs corrupt ────────────────────────────
    try:
        with open(path, 'rb') as fh:
            header = fh.read(8)
    except OSError as exc:
        results.append((FAIL, 'file access', f'Cannot read file: {exc}'))
        _add_advisory(results)
        return results

    if header[:4] == _OLE_MAGIC:
        results.append((FAIL, 'file format',
                         'File is an OLE Compound Document, not a standard .xlsx (ZIP). '
                         'This usually means the file is encrypted, password-protected, '
                         'or restricted by Microsoft Information Protection (MIP/IRM). '
                         'To use it with grepxcel: (1) open the file in Excel, '
                         '(2) remove protection or save a decrypted copy as .xlsx, '
                         'then re-run lint.'))
        _add_advisory(results)
        return results

    if header[:4] != _ZIP_MAGIC:
        results.append((FAIL, 'file format',
                         f'File does not have a valid .xlsx (ZIP) signature. '
                         f'It may be corrupted or renamed from a different format.'))
        _add_advisory(results)
        return results

    # ── 3. ZIP integrity ──────────────────────────────────────────────────
    try:
        with zipfile.ZipFile(path) as zf:
            total_uncompressed = sum(e.file_size for e in zf.infolist())
            names = zf.namelist()
    except zipfile.BadZipFile:
        results.append((FAIL, 'file integrity',
                         'File has ZIP magic bytes but is not a valid ZIP archive. '
                         'It may be corrupted.'))
        _add_advisory(results)
        return results

    file_mb = os.path.getsize(path) / (1024 * 1024)
    results.append((OK, 'file access', f'{path} — {file_mb:.1f} MB'))

    # ── 4. Open with openpyxl ─────────────────────────────────────────────
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        wb_raw = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        results.append((FAIL, 'workbook load',
                         f'openpyxl cannot open this file: {exc}. '
                         f'The file may be corrupted, encrypted, or in an '
                         f'unsupported format variant.'))
        _add_advisory(results)
        return results

    # ── 5. Sheet inventory ────────────────────────────────────────────────
    sheet_names = wb.sheetnames
    n_sheets = len(sheet_names)
    if n_sheets == 1:
        results.append((OK, 'sheets', f'1 sheet: {sheet_names[0]!r}'))
    else:
        results.append((OK, 'sheets',
                         f'{n_sheets} sheets: {", ".join(repr(s) for s in sheet_names)}. '
                         f'Use --sheet to select (default: active sheet).'))

    # ── 6. Per-sheet checks ───────────────────────────────────────────────
    for ws, ws_raw in zip(wb.worksheets, wb_raw.worksheets):
        _check_sheet(ws, ws_raw, results)

    wb.close()
    wb_raw.close()

    # ── 7. Advisory notes ─────────────────────────────────────────────────
    _add_advisory(results)

    return results


def _check_sheet(ws, ws_raw, results: list[Result]) -> None:
    """Run per-sheet checks: dimensions, extent, merged cells, formulas.

    ws is loaded with data_only=True (cached values), ws_raw with
    data_only=False (preserves formula text).
    """
    title = ws.title
    declared_rows = ws.max_row or 0
    declared_cols = ws.max_column or 0

    # Empty sheet
    if declared_rows == 0 or declared_cols == 0:
        results.append((WARN, f'sheet {title!r}',
                         'Sheet is empty — nothing to extract.'))
        return

    # Real used extent (from data_only view)
    used_rows = 0
    used_cols = 0
    for row_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=declared_rows,
                         max_col=declared_cols, values_only=True), start=1):
        for col_idx, val in enumerate(row, start=1):
            if val is not None:
                used_rows = row_idx
                if col_idx > used_cols:
                    used_cols = col_idx

    # Formula cells (from raw view — preserves formula text)
    formula_cells = []
    raw_rows = ws_raw.max_row or 0
    raw_cols = ws_raw.max_column or 0
    for row in ws_raw.iter_rows(min_row=1, max_row=raw_rows,
                                max_col=raw_cols):
        for cell in row:
            if cell.data_type == 'f':
                formula_cells.append(cell.coordinate)

    if used_rows == 0:
        results.append((WARN, f'sheet {title!r}',
                         'Sheet has formatting but no data — empty for extraction.'))
        return

    # Dimension report
    if (declared_rows > used_rows * 2 or declared_cols > used_cols * 2) and (
            declared_rows > used_rows + 10 or declared_cols > used_cols + 10):
        results.append((WARN, f'sheet {title!r} dimensions',
                         f'Declared {declared_rows}×{declared_cols} but only '
                         f'{used_rows}×{used_cols} contains data — inflated by '
                         f'empty formatted cells. grepxcel uses real extent, '
                         f'so this is handled, but the file is larger than needed.'))
    else:
        results.append((OK, f'sheet {title!r} dimensions',
                         f'{used_rows} rows × {used_cols} columns'))

    # Merged cells
    merged = list(ws.merged_cells.ranges)
    if merged:
        examples = ', '.join(str(m) for m in merged[:5])
        suffix = f' (and {len(merged) - 5} more)' if len(merged) > 5 else ''
        results.append((WARN, f'sheet {title!r} merged cells',
                         f'{len(merged)} merged region(s): {examples}{suffix}. '
                         f'grepxcel reads the top-left cell value; other cells in '
                         f'the merge appear empty.'))
    else:
        results.append((OK, f'sheet {title!r} merged cells', 'none'))

    # Formula cells
    if formula_cells:
        examples = ', '.join(formula_cells[:5])
        suffix = f' (and {len(formula_cells) - 5} more)' if len(formula_cells) > 5 else ''
        results.append((WARN, f'sheet {title!r} formulas',
                         f'{len(formula_cells)} formula cell(s): {examples}{suffix}. '
                         f'grepxcel reads cached values (data_only=True). If the file '
                         f'was last saved by a tool that doesn\'t cache formula results, '
                         f'these cells may appear empty. Open and re-save in Excel to fix.'))
    else:
        results.append((OK, f'sheet {title!r} formulas', 'none'))


def _add_advisory(results: list[Result]) -> None:
    """Append advisory notes about conditions not yet auto-detected."""
    results.append((INFO, 'advisory',
                     'Microsoft Information Protection (MIP/IRM): if your organisation '
                     'applies sensitivity labels that encrypt or restrict Excel files, '
                     'grepxcel cannot read them directly. Open the file in Excel and '
                     'save an unprotected copy, or ask your IT admin for a decrypted '
                     'export. Lint detects the OLE signature of encrypted files but '
                     'cannot distinguish IRM labels from other encryption.'))
    results.append((INFO, 'advisory',
                     'Password-protected workbooks or sheets: grepxcel does not prompt '
                     'for passwords. A workbook-level password produces an encrypted '
                     'OLE file (detected above). A sheet-level password still allows '
                     'reading cell values — extraction works, but the file author may '
                     'not intend the data to be extracted.'))
    results.append((INFO, 'advisory',
                     'Conditional formatting, data validation, pivot tables, charts, '
                     'and VBA modules are ignored by grepxcel. They do not affect '
                     'extraction but may indicate the file is more complex than a '
                     'flat data sheet.'))


# ── runner (CLI entry point) ──────────────────────────────────────────────────

def run_lint(files: list[str], out=None) -> int:
    """Lint one or more files, print results, return exit code (0 or 1)."""
    out = out or sys.stderr
    color = should_color(out)
    any_fail = False

    for path in files:
        print(f'\ngrepxcel lint — {path}\n' + '─' * 62, file=out)
        results = lint_file(path)

        for status, name, detail in results:
            if status == FAIL:
                any_fail = True
            mark = _MARK[status]
            print(colorize_marks(
                f'  {mark}  {name:<32} {detail}', color), file=out)

        print('─' * 62, file=out)

    return 1 if any_fail else 0
