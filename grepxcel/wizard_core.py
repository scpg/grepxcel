"""Pure (non-TUI) utilities shared by the web wizard (wizard_api.py).

Extracted from the former wizard.py / wizard_tui.py when the TUI was removed.
This module has no Textual or terminal-UI dependency.
"""
from __future__ import annotations

import csv
import datetime
import io
import json
import os
import re
from dataclasses import dataclass, field
from typing import Any


# ── State ─────────────────────────────────────────────────────────────────────

@dataclass
class WizardState:
    direction: str = 'LR'
    ignore_case_labels: bool = False
    ignore_case_values: bool = False
    currency_sign: str = '€'
    sheet_name: str | None = None
    # Each lbl_def is a 4-tuple: (name, type, match, lbl_mode)
    # lbl_mode: '' = use global lbl.match; 'glob' or 'regexp' = per-field override
    lbl_defs: list[tuple] = field(default_factory=list)
    # Each var_def is a 4-tuple: (name, type, match, col_a_extra)
    # col_a_extra: '' = bare var:; otherwise colon-joined modifier tokens,
    # e.g. 'nullable', 'not-null:trim-whitespace', 'literal', 'glob:nullable'
    var_defs: list[tuple] = field(default_factory=list)
    body_rows: list[list[str]] = field(default_factory=list)
    template: bool = False
    trim_whitespace_labels: bool = False
    trim_whitespace_values: bool = False
    lbl_match: str = ''
    var_match: str = ''
    empty_aliases: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        import dataclasses
        return dataclasses.asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> 'WizardState':
        return cls(
            direction=d.get('direction', 'LR'),
            ignore_case_labels=d.get('ignore_case_labels', d.get('ignore_case', False)),
            ignore_case_values=d.get('ignore_case_values', d.get('ignore_case', False)),
            currency_sign=d.get('currency_sign', '€'),
            sheet_name=d.get('sheet_name'),
            lbl_defs=[tuple(t) for t in d.get('lbl_defs', [])],
            var_defs=[tuple(t) for t in d.get('var_defs', [])],
            body_rows=[list(r) for r in d.get('body_rows', [])],
            template=d.get('template', False),
            trim_whitespace_labels=d.get('trim_whitespace_labels', d.get('trim_whitespace', False)),
            trim_whitespace_values=d.get('trim_whitespace_values', d.get('trim_whitespace', False)),
            lbl_match=d.get('lbl_match', ''),
            var_match=d.get('var_match', ''),
            empty_aliases=list(d.get('empty_aliases', [])),
        )


# ── Cell-reference helpers ─────────────────────────────────────────────────────

def _slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r'[^a-z0-9]+', '_', text)
    return text.strip('_') or 'field'


def _col_label(col: int) -> str:
    """1-based column index → Excel letter (1→A, 26→Z, 27→AA)."""
    result = ''
    while col > 0:
        col, rem = divmod(col - 1, 26)
        result = chr(65 + rem) + result
    return result


def _parse_cell_ref(ref: str) -> tuple[int, int] | None:
    """Parse 'A5' → (row=5, col=1), or None if invalid."""
    m = re.match(r'^([A-Z]+)(\d+)$', ref.upper().strip())
    if not m:
        return None
    col_str, row_str = m.groups()
    col = 0
    for ch in col_str:
        col = col * 26 + (ord(ch) - 64)
    return int(row_str), col


def _cell_ref(row: int, col: int) -> str:
    return f'{_col_label(col)}{row}'


def _build_cell_order(ws, direction: str) -> list[tuple[int, int]]:
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    if direction == 'TD':
        return [(r, c) for c in range(1, max_col + 1) for r in range(1, max_row + 1)]
    return [(r, c) for r in range(1, max_row + 1) for c in range(1, max_col + 1)]


# ── Pattern serialisation ──────────────────────────────────────────────────────

def _pattern_rows(state: WizardState) -> list[list]:
    """Flat list of rows for the pattern — shared by CSV and xlsx writers."""
    rows: list[list] = []
    rows.append(['config:', 'read.direction', state.direction])
    if state.ignore_case_labels:
        rows.append(['config:', 'ignore.case.labels', 'yes'])
    if state.ignore_case_values:
        rows.append(['config:', 'ignore.case.values', 'yes'])
    if state.trim_whitespace_labels:
        rows.append(['config:', 'trim.whitespace.labels', 'yes'])
    if state.trim_whitespace_values:
        rows.append(['config:', 'trim.whitespace.values', 'yes'])
    rows.append(['config:', 'currency.sign', state.currency_sign])
    if state.lbl_match:
        rows.append(['config:', 'lbl.match', state.lbl_match])
    if state.var_match:
        rows.append(['config:', 'var.match', state.var_match])
    for alias in state.empty_aliases:
        rows.append(['config:', 'empty.aliases', alias])
    for t in state.lbl_defs:
        name, ltype, value = t[0], t[1], t[2]
        lbl_mode = t[3] if len(t) > 3 else ''
        col_a = f'lbl:{lbl_mode}' if lbl_mode and lbl_mode != '(default)' else 'lbl:'
        rows.append([col_a, name, ltype, value])
    for t in state.var_defs:
        name, vtype, match = t[0], t[1], t[2]
        col_a_extra = t[3] if len(t) > 3 else ''
        col_a = f'var:{col_a_extra}' if col_a_extra else 'var:'
        rows.append([col_a, name, vtype, match])
    rows.append(['START:'])
    for row in state.body_rows:
        rows.append(row)
    rows.append(['END:'])
    return rows


def _write_pattern(state: WizardState, output_path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    rows = _pattern_rows(state)
    if os.path.splitext(output_path)[1].lower() == '.xlsx':
        import openpyxl as _openpyxl
        from .pattern_colors import colorize_pattern_file
        wb = _openpyxl.Workbook()
        ws = wb.active
        ws.title = 'pattern'
        for row in rows:
            ws.append([str(c) if c is not None else '' for c in row])
        wb.save(output_path)
        colorize_pattern_file(output_path)
    else:
        with open(output_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            for row in rows:
                writer.writerow(row)


# ── Type inference ─────────────────────────────────────────────────────────────

def _infer_cell_type(cell) -> str:
    """Return the most likely grepxcel type for an openpyxl cell."""
    import datetime as _dt
    val = cell.value
    if val is None:
        return 'string'
    if isinstance(val, bool):
        return 'boolean'
    if isinstance(val, (_dt.datetime, _dt.date)):
        return 'date'
    if isinstance(val, (int, float)):
        fmt = (cell.number_format or '').lower()
        if any(p in fmt for p in ('yyyy', 'yy/', '/yy', 'dd', 'd-mmm',
                                   'd/m', 'm/d', 'mmm', 'mmmm')):
            return 'date'
        return 'number'
    return 'string'


def _propose_type(value: Any, ws=None, row: int = None, col: int = None) -> str:
    """Propose a wizard action for a single cell value."""
    if value is None:
        return 'skip'
    if isinstance(value, str):
        if not value.strip():
            return 'skip'
        s = value.strip()
        if s.endswith(':') or s.endswith('：'):
            return 'label'
        if '@' in s or '://' in s or s.lower().startswith('www.'):
            return 'var:string'
        words = s.split()
        if len(words) == 1:
            has_alpha = any(c.isalpha() for c in s)
            has_digit = any(c.isdigit() for c in s)
            if has_alpha and has_digit:
                return 'var:string'
            if has_alpha:
                if ws is not None and row is not None and col is not None and col > 1:
                    left = ws.cell(row=row, column=col - 1).value
                    if isinstance(left, str) and left.strip().endswith(':'):
                        return 'var:string'
                return 'label'
        if any(any(c.isdigit() for c in w) for w in words):
            return 'var:string'
        return 'var:string'
    if isinstance(value, bool):
        return 'var:string'
    if isinstance(value, datetime.datetime):
        return 'var:datetime'
    if isinstance(value, datetime.date):
        return 'var:date'
    if isinstance(value, (datetime.time, datetime.timedelta)):
        return 'var:string'
    if isinstance(value, float):
        return 'var:integer' if value == int(value) else 'var:currency'
    if isinstance(value, int):
        return 'var:integer'
    return 'var:string'


def _var_type_from_proposal(proposal: str) -> str:
    if proposal.startswith('var:'):
        return proposal[4:]
    return 'string'


# ── col_a_extra helpers ────────────────────────────────────────────────────────

def _col_a_extra_from_parts(var_mode_raw: str, modifiers_raw: str) -> str:
    """Build col_a_extra string from separate mode + modifier selections."""
    mode = '' if var_mode_raw == '(default)' else var_mode_raw
    mods = '' if modifiers_raw == 'none' else modifiers_raw
    parts = [p for p in ([mode] + mods.split(':')) if p]
    return ':'.join(parts)


def _col_a_extra_to_parts(col_a_extra: str) -> tuple[str, str]:
    """Split col_a_extra back into (var_mode_raw, modifiers_raw) for UI pre-fill."""
    _MODE_TOKENS = frozenset({'literal', 'glob', 'regexp', 're'})
    _MOD_TOKENS  = frozenset({'nullable', 'not-null', 'not-empty', 'trim-whitespace'})
    if not col_a_extra:
        return '(default)', 'none'
    tokens = col_a_extra.split(':')
    mode_parts = [t for t in tokens if t in _MODE_TOKENS]
    mod_parts  = [t for t in tokens if t in _MOD_TOKENS]
    var_mode_raw  = mode_parts[0] if mode_parts else '(default)'
    modifiers_raw = ':'.join(mod_parts) if mod_parts else 'none'
    return var_mode_raw, modifiers_raw


def _fd_to_col_a_extra(fd) -> str:
    """Reconstruct col_a_extra string from a FieldDef for pre-population."""
    parts: list[str] = []
    if fd.var_mode and fd.var_mode not in ('regexp', 're'):
        parts.append(fd.var_mode)
    if fd.required:
        parts.append('not-null')
    elif fd.nullable:
        parts.append('nullable')
    if fd.trim_whitespace:
        parts.append('trim-whitespace')
    return ':'.join(parts)


# ── Pattern builder ────────────────────────────────────────────────────────────

def _build_state_from_choices(
    ws,
    choices: dict[str, dict],
    cells: list[tuple[int, int]],
    direction: str,
    sheet_name: str,
    ignore_case_labels: bool = False,
    ignore_case_values: bool = False,
    currency_sign: str = '€',
    trim_whitespace_labels: bool = False,
    trim_whitespace_values: bool = False,
    lbl_match: str = '',
    var_match: str = '',
    empty_aliases: list | None = None,
    ignore_case: bool | None = None,
    trim_whitespace: bool | None = None,
) -> WizardState:
    """Build a WizardState from the choices dict, in cell-scan order."""
    if ignore_case is not None:
        ignore_case_labels = ignore_case_values = ignore_case
    if trim_whitespace is not None:
        trim_whitespace_values = trim_whitespace
    state = WizardState(
        direction=direction,
        sheet_name=sheet_name,
        ignore_case_labels=ignore_case_labels,
        ignore_case_values=ignore_case_values,
        currency_sign=currency_sign,
        trim_whitespace_labels=trim_whitespace_labels,
        trim_whitespace_values=trim_whitespace_values,
        lbl_match=lbl_match,
        var_match=var_match,
        empty_aliases=list(empty_aliases) if empty_aliases else [],
    )
    from openpyxl.utils import get_column_letter as _gcl

    def _t_struct_key(m: dict) -> tuple:
        if 'header_rows' in m:
            hdr = tuple(
                tuple(col.get('var_name') or col.get('orig_field') or ''
                      for col in row.get('cols', []))
                for row in m.get('header_rows', [])
            )
            dat = tuple(
                dv.get('var_name') or dv.get('orig_field') or ''
                for dv in m.get('data_vars', [])
            )
            ftr = tuple(
                tuple(col.get('var_name') or col.get('orig_field') or ''
                      for col in row.get('cols', []))
                for row in m.get('footer_rows', [])
            )
            return ('multi', hdr, dat, ftr)
        return ('legacy', tuple(m.get('columns', [])))

    _t_groups: dict[tuple, list[str]] = {}
    _t_ref_to_key: dict[str, tuple] = {}
    for _r2, _c2 in cells:
        _ref2 = _cell_ref(_r2, _c2)
        _m2   = choices.get(_ref2)
        if _m2 and _m2.get('choice') == 'T':
            _k = _t_struct_key(_m2)
            _t_ref_to_key[_ref2] = _k
            _t_groups.setdefault(_k, []).append(_ref2)

    seen_t_anchors: set[str] = set()
    prev_lbl = False
    prev_lbl_row: int | None = None
    prev_lbl_col: int | None = None
    for r, c in cells:
        ref  = _cell_ref(r, c)
        meta = choices.get(ref)
        if not meta:
            continue
        choice = meta.get('choice', '')
        value  = ws.cell(row=r, column=c).value
        name   = meta.get('name', '')
        if choice in ('L', 'C'):  # 'C' is a legacy alias (removed from UI)
            state.lbl_defs.append((name,
                                   meta.get('ltype', 'string'),
                                   meta.get('lmatch', str(value) if value is not None else ''),
                                   meta.get('lbl_mode', '')))
            state.body_rows.append([f'cell:{_gcl(c)}{r}', name])
            prev_lbl = True
            prev_lbl_row, prev_lbl_col = r, c
        elif choice == 'V':
            state.var_defs.append((name, meta.get('ftype', 'string'),
                                   meta.get('match', '.*'),
                                   meta.get('col_a_extra', '')))
            if prev_lbl and prev_lbl_row is not None:
                if state.direction == 'TD':
                    adjacent = (r == prev_lbl_row + 1 and c == prev_lbl_col)
                else:
                    adjacent = (r == prev_lbl_row and c == prev_lbl_col + 1)
                cell_instr = 'cell:1' if adjacent else 'cell:next'
            else:
                cell_instr = f'cell:{_gcl(c)}{r}'
            state.body_rows.append([cell_instr, name])
            prev_lbl = False
            prev_lbl_row = prev_lbl_col = None
        elif choice == 'T':
            t_key   = _t_ref_to_key.get(ref)
            t_group = _t_groups.get(t_key, [ref])
            if t_group[0] != ref:
                seen_t_anchors.add(ref)
                continue
            if ref in seen_t_anchors:
                continue
            seen_t_anchors.add(ref)
            orig_mult = meta.get('mult', '1')
            if orig_mult == '*':
                table_mult = '*'
            else:
                table_mult = str(len(t_group))
            state.body_rows.append([f'table:{table_mult}'])

            if 'header_rows' in meta:
                def _emit_header_footer_row(hf_row, row_list, row_label='HEADER:1'):
                    col_names = []
                    for col in hf_row['cols']:
                        role = col.get('role', 'label')
                        if role == 'ignore':
                            orig = col.get('orig_field', 'IGNORE')
                            col_names.append(orig if orig in ('EMPTY', 'IGNORE') else 'IGNORE')
                        elif role == 'var':
                            vn = col.get('var_name', 'IGNORE')
                            if (table_name and vn and vn != 'IGNORE'
                                    and not vn.startswith(table_name + '.')
                                    and '.' not in vn):
                                vn = f'{table_name}.{vn}'
                            col_names.append(vn)
                            if vn and vn != 'IGNORE':
                                state.var_defs.append((
                                    vn,
                                    col.get('var_type', 'string'),
                                    col.get('var_match', '.*'),
                                    col.get('col_a_extra', ''),
                                ))
                        else:
                            ln = col.get('lbl_name', 'IGNORE')
                            col_names.append(ln)
                            if ln and ln != 'IGNORE' and not col.get('no_global_lbl'):
                                state.lbl_defs.append((
                                    ln,
                                    col.get('lbl_type', 'string'),
                                    col.get('lbl_match', col.get('cell_value', '')),
                                    col.get('lbl_mode', ''),
                                ))
                    row_list.append(['', row_label] + col_names)

                _row_types  = meta.get('row_types', {})
                _h_nums     = sorted(r for r, t in _row_types.items() if t == 'H')
                _d_nums     = sorted(r for r, t in _row_types.items() if t == 'D')
                _f_nums     = sorted(r for r, t in _row_types.items() if t == 'F')
                _p_nums     = set(r for r, t in _row_types.items() if t == 'P')

                table_name = meta.get('name', '').strip()

                for h_row in meta['header_rows']:
                    _emit_header_footer_row(h_row, state.body_rows)

                if _h_nums and _d_nums and _p_nums:
                    if any(max(_h_nums) < p < min(_d_nums) for p in _p_nums):
                        state.body_rows.append(['', 'SPLITTER:1'])

                data_vars = meta.get('data_vars', [])
                var_names = []
                for item in data_vars:
                    if isinstance(item, dict):
                        role   = item.get('role', 'var')
                        if role == 'ignore' or item.get('var_name') == 'IGNORE':
                            var_names.append('IGNORE')
                        elif role == 'label':
                            ln = item.get('lbl_name', 'IGNORE')
                            var_names.append(ln)
                            if ln and ln != 'IGNORE':
                                state.lbl_defs.append((
                                    ln,
                                    item.get('lbl_type', 'string'),
                                    item.get('lbl_match', '.*'),
                                    item.get('lbl_mode', ''),
                                ))
                        else:
                            vn    = item.get('var_name', 'IGNORE')
                            vtype = item.get('var_type', 'string')
                            vmatch= item.get('var_match', '.*')
                            vcol_a = item.get('col_a_extra', '')
                            if (table_name and vn and vn != 'IGNORE'
                                    and not vn.startswith(table_name + '.')
                                    and '.' not in vn):
                                vn = f'{table_name}.{vn}'
                            var_names.append(vn)
                            if vn and vn != 'IGNORE':
                                state.var_defs.append((vn, vtype, vmatch, vcol_a))
                    else:
                        vn = item
                        vtype, vmatch = 'string', '.*'
                        if (table_name and vn and vn != 'IGNORE'
                                and not vn.startswith(table_name + '.')
                                and '.' not in vn):
                            vn = f'{table_name}.{vn}'
                        var_names.append(vn)
                        if vn and vn != 'IGNORE':
                            state.var_defs.append((vn, vtype, vmatch, ''))
                state.body_rows.append(['', 'DATA:*'] + var_names)

                for skip_cfg in meta.get('_web_skip_configs', []):
                    skip_col_names = []
                    for s_col in skip_cfg.get('cols', []):
                        cond = (s_col.get('condition') or 'IGNORE').upper()
                        if cond == 'IGNORE':
                            skip_col_names.append('IGNORE')
                        elif cond == 'EMPTY':
                            skip_col_names.append('EMPTY')
                        else:
                            ln     = s_col.get('lbl_name', 'IGNORE')
                            lmatch = s_col.get('lmatch', '')
                            skip_col_names.append(ln)
                            if ln and ln != 'IGNORE' and lmatch:
                                state.lbl_defs.append((ln, 'string', lmatch, ''))
                    state.body_rows.append(['', 'SKIP_IF'] + skip_col_names)

                if _d_nums and _f_nums and _p_nums:
                    if any(max(_d_nums) < p < min(_f_nums) for p in _p_nums):
                        state.body_rows.append(['', 'SPLITTER:1'])

                for f_row in meta.get('footer_rows', []):
                    _emit_header_footer_row(f_row, state.body_rows, 'FOOTER:1')
            else:
                cols      = meta.get('columns', [])
                lbl_names = []
                var_names = []
                for col in cols:
                    lbl_name  = col.get('lbl_name', 'IGNORE')
                    var_name  = col.get('var_name', 'IGNORE')
                    cell_val  = col.get('cell_value', '')
                    state.lbl_defs.append((lbl_name, 'string', cell_val, ''))
                    state.var_defs.append((var_name, col.get('var_type', 'string'),
                                           col.get('var_match', '.*'), ''))
                    lbl_names.append(lbl_name)
                    var_names.append(var_name)
                state.body_rows.append(['', 'HEADER:1'] + lbl_names)
                state.body_rows.append(['', 'DATA:*'] + var_names)
        elif choice in ('T-HEAD', 'T-DATA'):
            pass
        elif choice == 'I':
            state.body_rows.append(['cell:1', 'IGNORE'])
            prev_lbl = False
        if choice == 'T':
            prev_lbl = False
    return state


def _choices_to_csv(ws, choices, cells, direction, sheet_name,
                    ignore_case_labels: bool = False,
                    ignore_case_values: bool = False,
                    currency_sign: str = '€',
                    trim_whitespace_labels: bool = False,
                    trim_whitespace_values: bool = False,
                    lbl_match: str = '',
                    var_match: str = '', empty_aliases: list | None = None,
                    ignore_case: bool | None = None,
                    trim_whitespace: bool | None = None) -> str:
    if ignore_case is not None:
        ignore_case_labels = ignore_case_values = ignore_case
    if trim_whitespace is not None:
        trim_whitespace_values = trim_whitespace
    state = _build_state_from_choices(
        ws, choices, cells, direction, sheet_name,
        ignore_case_labels=ignore_case_labels,
        ignore_case_values=ignore_case_values,
        currency_sign=currency_sign,
        trim_whitespace_labels=trim_whitespace_labels,
        trim_whitespace_values=trim_whitespace_values,
        lbl_match=lbl_match,
        var_match=var_match, empty_aliases=empty_aliases,
    )
    buf = io.StringIO()
    w   = csv.writer(buf)
    for row in _pattern_rows(state):
        w.writerow(row)
    return buf.getvalue()


# ── Label-matching helper (mirrors engine._match_lbl, no circular import) ─────

def _lbl_cell_matches(cell_value, pattern: str, mode: str, ignore_case: bool) -> bool:
    import fnmatch as _fnmatch, re as _re2
    if not pattern:
        return True
    text = str(cell_value) if cell_value is not None else ''
    if mode == 'literal':
        return (text.lower() == pattern.lower()) if ignore_case else (text == pattern)
    if mode == 'glob':
        flags = _re2.DOTALL | (_re2.IGNORECASE if ignore_case else 0)
        return bool(_re2.match(_fnmatch.translate(pattern), text, flags))
    flags = _re2.IGNORECASE if ignore_case else 0
    try:
        return bool(_re2.search(pattern, text[:2000], flags))
    except Exception:
        return False


# ── Preload helpers ────────────────────────────────────────────────────────────

def _preload_table(
    ws,
    instr,
    defs: dict,
    global_config,
    choices: dict,
    claimed: set,
    max_row: int,
    max_col: int,
) -> list[str]:
    """Pre-populate *choices* with T / T-HEAD / T-DATA entries for one table."""
    from .models import TemplateRow  # noqa: F401 — imported for type context

    warnings: list[str] = []

    header_tmpl  = [r for r in instr.rows if r.row_type == 'HEADER']
    data_tmpl    = [r for r in instr.rows if r.row_type == 'DATA']
    footer_tmpl  = [r for r in instr.rows if r.row_type == 'FOOTER']

    all_fields = []
    for tr in instr.rows:
        for tc in tr.columns:
            fn = tc.field if hasattr(tc, 'field') else str(tc)
            if fn not in ('IGNORE', 'EMPTY', ''):
                all_fields.append(fn)

    if not header_tmpl:
        shown = ', '.join(all_fields[:5]) + (f' (+{len(all_fields)-5} more)' if len(all_fields) > 5 else '')
        warnings.append(
            f'TABLE has no HEADER rows — cannot auto-locate: {shown or "(no fields)"}'
        )
        return warnings

    first_h_tmpl = header_tmpl[0]
    num_cols = len(first_h_tmpl.columns)
    if num_cols == 0:
        warnings.append('TABLE HEADER row has no columns — cannot locate')
        return warnings

    scan_targets: list[tuple] = []
    for col_idx, tmpl_col in enumerate(first_h_tmpl.columns):
        field = tmpl_col.field if hasattr(tmpl_col, 'field') else str(tmpl_col)
        if field in ('IGNORE', 'EMPTY', ''):
            continue
        fd = defs.get(field)
        if fd and fd.role == 'lbl' and fd.regex:
            scan_targets.append((col_idx, fd))

    if not scan_targets:
        shown = ', '.join(all_fields[:5])
        warnings.append(
            f'TABLE header has no label columns with fixed text '
            f'(fields: {shown or "(none)"}) — cannot auto-locate; '
            f'press T on the header cell to classify manually'
        )
        return warnings

    first_offset, first_fd = scan_targets[0]
    ic      = global_config.ignore_case_labels
    mode    = first_fd.lbl_match or global_config.lbl_match

    header_row_num: int | None = None
    start_col:      int | None = None

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            if not _lbl_cell_matches(v, first_fd.regex, mode, ic):
                continue
            cand_start = c - first_offset
            if cand_start < 1:
                continue
            cand_anchor = _cell_ref(r, cand_start)
            if cand_anchor in claimed:
                continue
            ok = True
            for other_off, other_fd in scan_targets[1:]:
                tc = cand_start + other_off
                if tc < 1 or tc > max_col:
                    ok = False
                    break
                other_v = ws.cell(row=r, column=tc).value
                other_mode = other_fd.lbl_match or global_config.lbl_match
                if not _lbl_cell_matches(other_v, other_fd.regex, other_mode, ic):
                    ok = False
                    break
            if ok:
                header_row_num = r
                start_col      = cand_start
                break
        if header_row_num is not None:
            break

    if header_row_num is None:
        probe_texts = [fd.regex for _, fd in scan_targets[:3]]
        warnings.append(
            f'TABLE header row not found in sheet '
            f'(looking for: {", ".join(repr(t) for t in probe_texts)}) — '
            f'press T on the header cell to classify manually'
        )
        return warnings

    h_rows_sheet: list[dict] = []
    for hi, h_tmpl in enumerate(header_tmpl):
        sheet_r = header_row_num + hi
        if sheet_r > max_row:
            break
        hcols: list[dict] = []
        for ci, tmpl_col in enumerate(h_tmpl.columns):
            sheet_c = start_col + ci
            if sheet_c > max_col:
                break
            field = tmpl_col.field if hasattr(tmpl_col, 'field') else str(tmpl_col)
            fd    = defs.get(field) if field not in ('IGNORE', 'EMPTY', '') else None
            val   = ws.cell(row=sheet_r, column=sheet_c).value
            val_s = str(val) if val is not None else ''
            slug_v = _slugify(val_s) if val_s else _col_label(sheet_c).lower()

            if field in ('IGNORE', 'EMPTY', '') or fd is None:
                hcols.append({
                    'ref': _cell_ref(sheet_r, sheet_c), 'row': sheet_r, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'ignore',
                    'orig_field': field if field in ('IGNORE', 'EMPTY') else 'IGNORE',
                    'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': val_s,
                    'var_name': 'IGNORE', 'var_type': 'string', 'var_match': '.*', 'notes': '',
                })
            elif fd.role == 'lbl':
                hcols.append({
                    'ref': _cell_ref(sheet_r, sheet_c), 'row': sheet_r, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'label',
                    'lbl_name': field, 'lbl_type': fd.type, 'lbl_match': fd.regex,
                    'var_name': 'IGNORE', 'var_type': 'string', 'var_match': '.*',
                    'lbl_mode': fd.lbl_match or '', 'notes': '',
                })
            else:
                hcols.append({
                    'ref': _cell_ref(sheet_r, sheet_c), 'row': sheet_r, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'var',
                    'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': val_s,
                    'var_name': field, 'var_type': fd.type, 'var_match': fd.regex,
                    'notes': '', 'col_a_extra': _fd_to_col_a_extra(fd),
                })
        h_rows_sheet.append({'row': sheet_r, 'cols': hcols})

    last_h_row = header_row_num + len(h_rows_sheet) - 1

    _ser = [row for row in instr.rows if row.row_type == 'SKIP_EMPTY_ROW']
    _max_skip_empty = sum(int(row.multiplicity) for row in _ser)
    _consecutive_empty = 0
    d_rows_sheet: list[int] = []
    for r in range(last_h_row + 1, min(last_h_row + 201, max_row + 1)):
        has_data = any(
            ws.cell(row=r, column=start_col + ci).value is not None
            for ci in range(num_cols)
            if start_col + ci <= max_col
        )
        if has_data:
            _consecutive_empty = 0
            d_rows_sheet.append(r)
        elif _consecutive_empty < _max_skip_empty:
            _consecutive_empty += 1
        else:
            break

    data_vars: list[dict] = []
    d_tmpl_row = data_tmpl[0] if data_tmpl else None
    for ci in range(num_cols):
        if d_tmpl_row and ci < len(d_tmpl_row.columns):
            tmpl_col = d_tmpl_row.columns[ci]
            field = tmpl_col.field if hasattr(tmpl_col, 'field') else str(tmpl_col)
        else:
            field = 'IGNORE'
        fd = defs.get(field) if field not in ('IGNORE', 'EMPTY', '') else None

        if field in ('IGNORE', 'EMPTY', '') or fd is None:
            data_vars.append({'role': 'ignore', 'var_name': 'IGNORE',
                              'var_type': 'string', 'var_match': '.*',
                              'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': '.*',
                              'notes': '', 'col_a_extra': ''})
        elif fd.role == 'lbl':
            data_vars.append({'role': 'label', 'var_name': 'IGNORE',
                              'var_type': 'string', 'var_match': '.*',
                              'lbl_name': field, 'lbl_type': fd.type, 'lbl_match': fd.regex,
                              'notes': '', 'col_a_extra': ''})
        else:
            data_vars.append({'role': 'var', 'var_name': field,
                              'var_type': fd.type, 'var_match': fd.regex,
                              'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': '.*',
                              'notes': '', 'col_a_extra': _fd_to_col_a_extra(fd)})

    f_rows_sheet: list[dict] = []
    f_row_nums:   set[int]   = set()

    for fi, f_tmpl in enumerate(footer_tmpl):
        f_scan: list[tuple] = []
        for col_idx, tmpl_col in enumerate(f_tmpl.columns):
            field = tmpl_col.field if hasattr(tmpl_col, 'field') else str(tmpl_col)
            if field in ('IGNORE', 'EMPTY', ''):
                continue
            fd = defs.get(field)
            if fd and fd.role == 'lbl' and fd.regex:
                f_scan.append((col_idx, fd))

        matched_row: int | None = None

        if not f_scan:
            available = sorted(r for r in d_rows_sheet if r not in f_row_nums)
            if available:
                matched_row = available[-1]
            else:
                warnings.append(
                    f'TABLE footer row {fi + 1} has no label columns and no '
                    f'remaining rows; skipped (classify manually)'
                )
                continue
        else:
            for r in d_rows_sheet:
                if r in f_row_nums:
                    continue
                ok = True
                for col_off, fd in f_scan:
                    sheet_c = start_col + col_off
                    if sheet_c > max_col:
                        ok = False
                        break
                    v      = ws.cell(row=r, column=sheet_c).value
                    f_mode = fd.lbl_match or global_config.lbl_match
                    if not _lbl_cell_matches(v, fd.regex, f_mode, ic):
                        ok = False
                        break
                if ok:
                    matched_row = r
                    break

            if matched_row is None:
                probe = [fd.regex for _, fd in f_scan[:2]]
                warnings.append(
                    f'TABLE footer row {fi + 1}: label not matched in data rows '
                    f'(looking for: {", ".join(repr(t) for t in probe)}); '
                    f'classify manually'
                )
                continue

        f_row_nums.add(matched_row)
        fcols: list[dict] = []
        for ci, tmpl_col in enumerate(f_tmpl.columns):
            sheet_c = start_col + ci
            if sheet_c > max_col:
                break
            field = tmpl_col.field if hasattr(tmpl_col, 'field') else str(tmpl_col)
            fd    = defs.get(field) if field not in ('IGNORE', 'EMPTY', '') else None
            val   = ws.cell(row=matched_row, column=sheet_c).value
            val_s = str(val) if val is not None else ''

            if field in ('IGNORE', 'EMPTY', '') or fd is None:
                fcols.append({
                    'ref': _cell_ref(matched_row, sheet_c), 'row': matched_row, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'ignore',
                    'orig_field': field if field in ('IGNORE', 'EMPTY') else 'IGNORE',
                    'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': val_s,
                    'var_name': 'IGNORE', 'var_type': 'string', 'var_match': '.*', 'notes': '',
                })
            elif fd.role == 'lbl':
                fcols.append({
                    'ref': _cell_ref(matched_row, sheet_c), 'row': matched_row, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'label',
                    'lbl_name': field, 'lbl_type': fd.type, 'lbl_match': fd.regex,
                    'var_name': 'IGNORE', 'var_type': 'string', 'var_match': '.*',
                    'lbl_mode': fd.lbl_match or '', 'notes': '',
                })
            else:
                fcols.append({
                    'ref': _cell_ref(matched_row, sheet_c), 'row': matched_row, 'col': sheet_c,
                    'cell_value': val_s, 'role': 'var',
                    'lbl_name': 'IGNORE', 'lbl_type': 'string', 'lbl_match': val_s,
                    'var_name': field, 'var_type': fd.type, 'var_match': fd.regex,
                    'notes': '', 'col_a_extra': _fd_to_col_a_extra(fd),
                })
        f_rows_sheet.append({'row': matched_row, 'cols': fcols})

    d_rows_sheet = [r for r in d_rows_sheet if r not in f_row_nums]

    _skip_marker_checks: list[tuple[int, object]] = []
    for trow in instr.rows:
        if trow.row_type != 'SKIP_IF':
            continue
        for ci, tcol in enumerate(trow.columns):
            fld = tcol.field if hasattr(tcol, 'field') else str(tcol)
            if fld in ('EMPTY', 'IGNORE', ''):
                continue
            fd = defs.get(fld)
            if fd and fd.role == 'lbl' and fd.regex:
                _skip_marker_checks.append((ci, fd))

    skip_rows: list[int] = []
    if _skip_marker_checks:
        real_data: list[int] = []
        for r in d_rows_sheet:
            _is_skip = False
            for col_off, fd in _skip_marker_checks:
                sheet_c = start_col + col_off
                if sheet_c <= max_col:
                    v      = ws.cell(row=r, column=sheet_c).value
                    f_mode = fd.lbl_match or global_config.lbl_match
                    if _lbl_cell_matches(v, fd.regex, f_mode, ic):
                        _is_skip = True
                        break
            if _is_skip:
                skip_rows.append(r)
            else:
                real_data.append(r)
        d_rows_sheet = real_data

    row_types: dict[int, str] = {}
    for hd in h_rows_sheet:
        row_types[hd['row']] = 'H'
    for r in d_rows_sheet:
        row_types[r] = 'D'
    for fr in f_rows_sheet:
        row_types[fr['row']] = 'F'
    for r in skip_rows:
        row_types[r] = 'S'

    last_f_row = f_rows_sheet[-1]['row'] if f_rows_sheet else None
    end_row    = last_f_row or (d_rows_sheet[-1] if d_rows_sheet else last_h_row)
    end_col    = start_col + num_cols - 1
    anchor_ref = _cell_ref(header_row_num, start_col)
    range_str  = f'{anchor_ref}:{_cell_ref(header_row_num, end_col)}'

    dv_names = [
        dv['var_name'] for dv in data_vars
        if dv.get('var_name') and dv['var_name'] not in ('IGNORE', 'EMPTY', '')
    ]
    table_name = ''
    if dv_names:
        dot_prefixes = [n.split('.')[0] for n in dv_names if '.' in n]
        if dot_prefixes and len(set(dot_prefixes)) == 1:
            table_name = _slugify(dot_prefixes[0])
        if not table_name and len(dv_names) == 1:
            table_name = _slugify(dv_names[0].split('_')[0] or dv_names[0])
    if not table_name or table_name == 'field':
        raw_name   = first_fd.name
        table_name = _slugify(
            raw_name.removesuffix('_lbl').removesuffix('_label')
                    .removesuffix('_header').removesuffix('_col')
        ) or _slugify(raw_name)

    web_row_configs: list[dict] = []
    for hi, h_row in enumerate(h_rows_sheet):
        cols_cfg: list[dict] = []
        for col_dict in h_row['cols']:
            role = col_dict.get('role', 'ignore')
            if role == 'label':
                lbl_mode = col_dict.get('lbl_mode', '') or 'literal'
                cols_cfg.append({'role': 'L',
                                 'name': _slugify(col_dict.get('lbl_name', '') or col_dict.get('cell_value', '')),
                                 'lmatch': col_dict.get('lbl_match', col_dict.get('cell_value', '')),
                                 'lmatch_mode': lbl_mode})
            elif role == 'var':
                _, mods = _col_a_extra_to_parts(col_dict.get('col_a_extra', ''))
                cols_cfg.append({'role': 'V', 'name': col_dict.get('var_name', 'IGNORE'),
                                 'ftype': col_dict.get('var_type', 'string'),
                                 'match': col_dict.get('var_match', '.*'), 'modifiers': mods})
            else:
                cols_cfg.append({'role': 'I', 'orig_field': col_dict.get('orig_field', 'IGNORE')})
        web_row_configs.append({'sheet_row': h_row['row'], 'row_type': 'header',
                                'row_n': hi + 1, 'cols': cols_cfg})

    for di, r in enumerate(d_rows_sheet):
        if di == 0:
            cols_cfg = []
            for dv in data_vars:
                dv_role = dv.get('role', 'ignore')
                if dv_role == 'var':
                    _, mods = _col_a_extra_to_parts(dv.get('col_a_extra', ''))
                    cols_cfg.append({'role': 'V', 'name': dv.get('var_name', 'IGNORE'),
                                     'ftype': dv.get('var_type', 'string'),
                                     'match': dv.get('var_match', '.*'), 'modifiers': mods})
                elif dv_role == 'label':
                    lbl_mode = dv.get('lbl_mode', '') or 'literal'
                    cols_cfg.append({'role': 'L',
                                     'name': _slugify(dv.get('lbl_name', '') or ''),
                                     'lmatch': dv.get('lbl_match', ''), 'lmatch_mode': lbl_mode})
                else:
                    cols_cfg.append({'role': 'I', 'orig_field': dv.get('orig_field', 'IGNORE')})
            web_row_configs.append({'sheet_row': r, 'row_type': 'data', 'cols': cols_cfg})
        else:
            web_row_configs.append({'sheet_row': r, 'row_type': 'data_inherited'})

    for fi, f_row in enumerate(f_rows_sheet):
        cols_cfg = []
        for col_dict in f_row['cols']:
            role = col_dict.get('role', 'ignore')
            if role == 'label':
                lbl_mode = col_dict.get('lbl_mode', '') or 'literal'
                cols_cfg.append({'role': 'L',
                                 'name': _slugify(col_dict.get('lbl_name', '') or col_dict.get('cell_value', '')),
                                 'lmatch': col_dict.get('lbl_match', col_dict.get('cell_value', '')),
                                 'lmatch_mode': lbl_mode})
            elif role == 'var':
                _, mods = _col_a_extra_to_parts(col_dict.get('col_a_extra', ''))
                cols_cfg.append({'role': 'V', 'name': col_dict.get('var_name', 'IGNORE'),
                                 'ftype': col_dict.get('var_type', 'string'),
                                 'match': col_dict.get('var_match', '.*'), 'modifiers': mods})
            else:
                cols_cfg.append({'role': 'I', 'orig_field': col_dict.get('orig_field', 'IGNORE')})
        web_row_configs.append({'sheet_row': f_row['row'], 'row_type': 'footer',
                                'row_n': fi + 1, 'cols': cols_cfg})

    web_skip_configs: list[dict] = []
    for trow in instr.rows:
        if trow.row_type == 'SKIP_IF':
            cols = []
            for tcol in trow.columns:
                fld = tcol.field
                if fld == 'IGNORE':
                    cols.append({'condition': 'IGNORE'})
                elif fld == 'EMPTY':
                    cols.append({'condition': 'EMPTY'})
                else:
                    cols.append({'condition': 'LABEL', 'lbl_name': fld, 'lmatch': ''})
            web_skip_configs.append({'cols': cols})

    if anchor_ref in claimed:
        warnings.append(
            f'TABLE anchor {anchor_ref} already claimed by another field; '
            f'table not pre-loaded'
        )
        return warnings

    _anchor_col_role = (
        h_rows_sheet[0]['cols'][0].get('role', 'ignore')
        if h_rows_sheet and h_rows_sheet[0]['cols']
        else 'ignore'
    )
    _anchor_table_role = 'L' if _anchor_col_role == 'label' else ('V' if _anchor_col_role == 'var' else 'I')

    meta: dict = {
        'choice': 'T', 'name': table_name, 'mult': instr.multiplicity,
        'range': range_str, 'start_row': header_row_num, 'end_row': end_row,
        'start_col': start_col, 'end_col': end_col,
        'header_rows': h_rows_sheet, 'footer_rows': f_rows_sheet,
        'data_vars': data_vars, 'skip_rows': skip_rows, 'row_types': row_types,
        '_web_row_configs': web_row_configs, '_web_skip_configs': web_skip_configs,
        'table_role': _anchor_table_role, 'row_class': 'header',
    }
    choices[anchor_ref] = meta
    claimed.add(anchor_ref)

    for hd in h_rows_sheet:
        for col_dict in hd['cols']:
            ref = col_dict['ref']
            if ref != anchor_ref and ref not in claimed:
                role = col_dict.get('role', 'ignore')
                t_role = 'L' if role == 'label' else ('V' if role == 'var' else 'I')
                choices[ref] = {'choice': 'T-HEAD', 'anchor': anchor_ref,
                                'table_role': t_role, 'row_class': 'header'}
                claimed.add(ref)

    for fr in f_rows_sheet:
        for col_dict in fr['cols']:
            ref = col_dict['ref']
            if ref not in claimed:
                role = col_dict.get('role', 'ignore')
                t_role = 'L' if role == 'label' else ('V' if role == 'var' else 'I')
                choices[ref] = {'choice': 'T-HEAD', 'anchor': anchor_ref,
                                'table_role': t_role, 'row_class': 'footer'}
                claimed.add(ref)

    for r in d_rows_sheet:
        for ci in range(num_cols):
            ref = _cell_ref(r, start_col + ci)
            if ref not in claimed:
                dv = data_vars[ci] if ci < len(data_vars) else {'role': 'ignore'}
                dv_role = dv.get('role', 'ignore')
                t_role = 'V' if dv_role == 'var' else ('L' if dv_role == 'label' else 'I')
                choices[ref] = {'choice': 'T-DATA', 'anchor': anchor_ref,
                                'table_role': t_role, 'row_class': 'data'}
                claimed.add(ref)

    _end_ref = _cell_ref(end_row, end_col)
    if _end_ref in choices:
        choices[_end_ref]['is_table_end'] = True

    n_d = len(d_rows_sheet)
    if n_d == 200:
        warnings.append(
            f'TABLE {table_name!r}: data rows capped at 200 for pre-load; '
            f'adjust the range with T if needed'
        )

    return warnings


def _preload_from_pattern(ws, pattern_path: str) -> tuple[dict, dict, list[str]]:
    """Parse an existing pattern file and scan *ws* to locate matching cells.

    Returns (choices, preload_cfg, warnings).
    """
    from .pattern_parser import PatternParser
    from .models import (CellInstruction, TableInstruction,
                         SeekInstruction, DirectionInstruction)

    try:
        parser = PatternParser()
        global_config, defs, start_sequence = parser.parse(pattern_path)
    except Exception as exc:
        return {}, {}, [f'Could not parse pattern file: {exc}']

    preload_cfg: dict = {
        'direction':              global_config.read_direction,
        'ignore_case_labels':     global_config.ignore_case_labels,
        'ignore_case_values':     global_config.ignore_case_values,
        'trim_whitespace_labels': global_config.trim_whitespace_labels,
        'trim_whitespace_values': global_config.trim_whitespace_values,
        'currency_sign':          global_config.currency_sign,
        'lbl_match': global_config.lbl_match if global_config.lbl_match != 'literal' else '',
        'var_match': global_config.var_match if global_config.var_match != 'regexp' else '',
        'empty_aliases': list(global_config.empty_aliases),
    }

    from collections import defaultdict
    val_to_cells: dict[str, list[tuple]] = defaultdict(list)
    claimed: set[str] = set()
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                val_to_cells[str(v)].append((r, c))

    choices: dict[str, dict] = {}
    warnings: list[str] = []
    direction = global_config.read_direction
    last_lbl_pos: tuple | None = None

    if any(isinstance(instr, SeekInstruction) for instr in start_sequence):
        warnings.append(
            "This pattern uses seek: instructions, which the wizard cannot "
            "simulate during preload. Affected fields will not be pre-filled — "
            "classify them manually, or replace seek: with absolute cell: references."
        )

    def _find_lbl_pos(fd) -> tuple | None:
        mode = fd.lbl_match or global_config.lbl_match
        ic   = global_config.ignore_case_labels
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                ref = _cell_ref(r, c)
                if ref in claimed:
                    continue
                v = ws.cell(row=r, column=c).value
                if v is not None and _lbl_cell_matches(v, fd.regex, mode, ic):
                    return (r, c)
        return None

    def _find_var_by_content(fd) -> tuple | None:
        import re as _re
        import fnmatch as _fnmatch
        pat = fd.regex or ''
        if pat in ('.*', '.+', ''):
            return None
        mode = fd.var_mode or global_config.var_match or 'regexp'
        ic   = global_config.ignore_case_values
        try:
            if mode == 'literal':
                compiled = _re.compile(_re.escape(pat), _re.IGNORECASE if ic else 0)
            elif mode == 'glob':
                compiled = _re.compile(_fnmatch.translate(pat), _re.IGNORECASE if ic else 0)
            else:
                compiled = _re.compile(pat, _re.IGNORECASE if ic else 0)
        except _re.error:
            return None
        tw = global_config.trim_whitespace_values
        for rr in range(1, max_row + 1):
            for cc in range(1, max_col + 1):
                ref = _cell_ref(rr, cc)
                if ref in claimed:
                    continue
                v = ws.cell(row=rr, column=cc).value
                if v is not None:
                    sv = str(v).strip() if tw else str(v)
                    if compiled.fullmatch(sv):
                        return (rr, cc)
        return None

    def _adjacent(pos: tuple) -> tuple:
        r, c = pos
        return (r, c + 1) if direction == 'LR' else (r + 1, c)

    def _record_lbl(ref: str, fd) -> None:
        choices[ref] = {'choice': 'L', 'name': fd.name, 'ltype': fd.type,
                        'lmatch': fd.regex, 'lbl_mode': fd.lbl_match or ''}
        claimed.add(ref)

    def _record_var(ref: str, fd) -> None:
        choices[ref] = {'choice': 'V', 'name': fd.name, 'ftype': fd.type,
                        'match': fd.regex, 'col_a_extra': _fd_to_col_a_extra(fd)}
        claimed.add(ref)

    for instr in start_sequence:

        if isinstance(instr, DirectionInstruction):
            direction = instr.direction
            continue

        if isinstance(instr, SeekInstruction):
            last_lbl_pos = None
            warnings.append(
                f'seek:{instr.target} — var fields after a seek: instruction '
                f'cannot be auto-located; classify them manually'
            )
            continue

        if isinstance(instr, TableInstruction):
            mult_s = instr.multiplicity or '1'
            unbounded = (mult_s == '*')
            if unbounded:
                max_passes = 200
            elif mult_s.startswith('{') and ',' in mult_s:
                try:
                    max_passes = int(mult_s.strip('{}').split(',')[1])
                except (ValueError, IndexError):
                    max_passes = 1
            else:
                try:
                    max_passes = int(mult_s)
                except ValueError:
                    max_passes = 1

            found_count = 0
            for _ in range(max(1, max_passes)):
                tbl_warns = _preload_table(
                    ws, instr, defs, global_config,
                    choices, claimed, max_row, max_col,
                )
                miss = any(
                    'not found' in w or 'not pre-loaded' in w or 'cannot auto-locate' in w
                    for w in tbl_warns
                )
                if miss:
                    if not unbounded:
                        if found_count == 0:
                            warnings.extend(tbl_warns)
                        else:
                            warnings.append(
                                f'TABLE {mult_s}: only {found_count} instance(s) '
                                f'found in sheet (expected {max_passes}); '
                                f'classify remaining rows manually'
                            )
                    break
                warnings.extend(tbl_warns)
                found_count += 1

            last_lbl_pos = None
            continue

        if not isinstance(instr, CellInstruction):
            continue

        name = instr.field
        if name in ('IGNORE', 'EMPTY', ''):
            last_lbl_pos = None
            continue

        fd = defs.get(name)
        if fd is None:
            warnings.append(f'Field {name!r} referenced in body but not defined above START:')
            last_lbl_pos = None
            continue

        if instr.multiplicity == 'abs':
            try:
                from openpyxl.utils.cell import coordinate_to_tuple
                r, c = coordinate_to_tuple(instr.target)
            except Exception:
                warnings.append(f'Could not parse absolute ref {instr.target!r} for {name!r}')
                continue
            ref = _cell_ref(r, c)
            if fd.role == 'lbl':
                _record_lbl(ref, fd)
            else:
                _record_var(ref, fd)
            last_lbl_pos = (r, c)
            continue

        if fd.role == 'lbl':
            pos = _find_lbl_pos(fd)
            if pos is None:
                warnings.append(f'Label {name!r} (pattern: {fd.regex!r}) not found in sheet')
                last_lbl_pos = None
            else:
                ref = _cell_ref(*pos)
                _record_lbl(ref, fd)
                last_lbl_pos = pos
        else:
            if last_lbl_pos is None:
                found = _find_var_by_content(fd)
                if found:
                    ref = _cell_ref(*found)
                    _record_var(ref, fd)
                    last_lbl_pos = found
                else:
                    warnings.append(
                        f'Var {name!r}: no preceding label was matched — '
                        f'cannot determine cell position; classify manually'
                    )
            else:
                if instr.multiplicity == 'next':
                    ar, ac = _adjacent(last_lbl_pos)
                    while (1 <= ar <= max_row and 1 <= ac <= max_col
                           and (ws.cell(row=ar, column=ac).value is None
                                or _cell_ref(ar, ac) in claimed)):
                        ar, ac = _adjacent((ar, ac))
                else:
                    ar, ac = _adjacent(last_lbl_pos)
                if 1 <= ar <= max_row and 1 <= ac <= max_col:
                    ref = _cell_ref(ar, ac)
                    if ref in claimed:
                        warnings.append(
                            f'Var {name!r}: adjacent cell {ref} is already '
                            f'claimed by another field; classify manually'
                        )
                    else:
                        _record_var(ref, fd)
                        last_lbl_pos = (ar, ac)
                else:
                    warnings.append(
                        f'Var {name!r}: adjacent cell ({ar}, {ac}) is outside '
                        f'the sheet bounds; classify manually'
                    )

    return choices, preload_cfg, warnings
