"""Full-screen TUI wizard for grepxcel.

Requires textual: pip install 'grepxcel[wizard]'

Design principles
-----------------
* push_screen(modal, callback) throughout — no push_screen_wait, no workers needed.
* _choices dict is the single source of truth; WizardState is built on save.
* Panel has 4 fixed zones: CELL · CLASSIFY · NAVIGATE · LEGEND+STATS
* Color system: green=Label  bright_yellow=Value  blue=Header  magenta=Table  dim=Ignore
* Highlights: H key marks all unclassified non-empty cells in amber; auto-clears on
  any navigation or classification.
* Undo: Ctrl+Z / undo stack pops the last classification and reverts _choices entry.
* Web-service ready: WizardState is JSON-serialisable via dataclasses.asdict().
"""
from __future__ import annotations

import csv
import datetime
import hashlib
import io
import os
import sys
from typing import Any

try:
    from . import __version__ as _GX_VERSION
except Exception:
    _GX_VERSION = '?'

try:
    from textual.app import App, ComposeResult
    from textual.binding import Binding
    from textual.containers import Horizontal, Vertical
    from textual.screen import ModalScreen
    from textual.widgets import DataTable, Footer, Header, Input, Label, Select, Static
    from rich.text import Text as RichText
    _TEXTUAL_OK = True
except ImportError:
    _TEXTUAL_OK = False

from .wizard import (
    WizardState,
    _slugify,
    _propose_type,
    _write_pattern,
    _build_cell_order,
    _cell_ref,
    _col_label,
    _parse_cell_ref,
    _find_next_nonempty,
    _find_prev_nonempty,
    _save_history,
    _load_history,
    _detect_template,
)


# ── Style constants ────────────────────────────────────────────────────────────

_STYLE: dict[str, str] = {
    'L':      'bold green',
    'C':      'bold blue',
    'V':      'bold bright_yellow',
    'T':      'bold magenta',
    'T-HEAD': 'magenta',          # header / footer row cells
    'T-DATA': 'dim magenta',      # data row cells inside the table range
    'I':      'dim',
    'PENDING': 'bold black on dark_goldenrod',
}

_CHOICE_COLOR = {
    'L':      'green',
    'C':      'blue',
    'V':      'bright_yellow',
    'T':      'magenta',
    'T-HEAD': 'magenta',
    'T-DATA': 'magenta',
    'I':      'dim',
}

_CHOICE_NAME = {
    'L':      'Label',
    'C':      'Header',
    'V':      'Value',
    'T':      'Table anchor',
    'T-HEAD': 'Table column',
    'T-DATA': 'Table data row',
    'I':      'Ignore',
}

_SEP = '─' * 42   # visual divider for panel zones


def _trunc(value: Any, width: int = 20) -> str:
    """Truncate to width chars, adding … when text is cut."""
    if value is None:
        return ''
    s = str(value)
    return (s[:width - 1] + '…') if len(s) > width else s


def _infer_cell_type(cell) -> str:
    """Return the most likely grepxcel type for an openpyxl cell.

    Priority: Python value type > number_format string > 'string' fallback.
    Covers date, number (incl. currency / percentage), boolean, string.
    """
    import datetime as _dt
    val = cell.value
    if val is None:
        return 'string'
    if isinstance(val, bool):
        return 'boolean'
    if isinstance(val, (_dt.datetime, _dt.date)):
        return 'date'
    if isinstance(val, (int, float)):
        fmt = (cell.number_format or '').lower()
        # Date indicators in the format string (avoid matching 'mm' for minutes)
        if any(p in fmt for p in ('yyyy', 'yy/', '/yy', 'dd', 'd-mmm',
                                   'd/m', 'm/d', 'mmm', 'mmmm')):
            return 'date'
        return 'number'
    return 'string'


def _styled(value: Any, choice: str) -> 'RichText':
    if value is None:
        # Empty classified cell: show a colored dot so the user can see it was classified.
        # An empty RichText with a bold style renders as nothing — the dot is essential.
        marker = '○' if choice == 'I' else '●'
        return RichText(marker, style=_STYLE.get(choice, ''))
    if choice == 'I':
        return RichText(_trunc(value), style='dim strike')
    return RichText(_trunc(value), style=_STYLE.get(choice, ''))


# ── Pattern builder ────────────────────────────────────────────────────────────

def _build_state_from_choices(
    ws,
    choices: dict[str, dict],
    cells: list[tuple[int, int]],
    direction: str,
    sheet_name: str,
) -> WizardState:
    """Build a WizardState from the choices dict, in cell-scan order.

    This is the single aggregation point — keeping it here (not in action
    methods) means undo/reclassify only need to update _choices, and the
    pattern is always consistent on save.

    Web-service note: this is the function a REST handler would call after
    the client submits the final classification list.
    """
    state = WizardState(direction=direction, sheet_name=sheet_name)
    seen_t_anchors: set[str] = set()
    for r, c in cells:
        ref  = _cell_ref(r, c)
        meta = choices.get(ref)
        if not meta:
            continue
        choice = meta.get('choice', '')
        value  = ws.cell(row=r, column=c).value
        name   = meta.get('name', '')
        if choice == 'L':
            state.lbl_defs.append((name,
                                   meta.get('ltype', 'string'),
                                   meta.get('lmatch', str(value) if value is not None else '')))
            state.body_rows.append(['cell:1', name])
        elif choice == 'C':
            state.lbl_defs.append((name,
                                   meta.get('ltype', 'string'),
                                   meta.get('lmatch', str(value) if value is not None else '')))
            state.body_rows.append(['cell:1', name])
        elif choice == 'V':
            state.var_defs.append((name, meta.get('ftype', 'string'), meta.get('match', '.*')))
            state.body_rows.append(['cell:1', name])
        elif choice == 'T':
            # Each table is emitted once from its anchor cell; T-HEAD cells are skipped.
            if ref in seen_t_anchors:
                continue
            seen_t_anchors.add(ref)
            mult = meta.get('mult', '*')
            state.body_rows.append([f'table:{mult}'])

            if 'header_rows' in meta:
                # New multi-row model: header_rows + data_vars [+ footer_rows]
                def _emit_header_footer_row(hf_row, row_list, row_label='HEADER:1'):
                    """Emit one H/F row with role-aware column handling."""
                    col_names = []
                    for col in hf_row['cols']:
                        role = col.get('role', 'label')
                        if role == 'ignore':
                            col_names.append('IGNORE')
                        elif role == 'var':
                            vn = col.get('var_name', 'IGNORE')
                            if (table_name and vn and vn != 'IGNORE'
                                    and not vn.startswith(table_name + '.')):
                                vn = f'{table_name}.{vn}'
                            col_names.append(vn)
                            if vn and vn != 'IGNORE':
                                state.var_defs.append((
                                    vn,
                                    col.get('var_type', 'string'),
                                    col.get('var_match', '.*'),
                                ))
                        else:  # label (default for H/F)
                            ln = col.get('lbl_name', 'IGNORE')
                            col_names.append(ln)
                            if ln and ln != 'IGNORE':
                                state.lbl_defs.append((
                                    ln,
                                    col.get('lbl_type', 'string'),
                                    col.get('lbl_match', col.get('cell_value', '')),
                                ))
                    row_list.append(['', row_label] + col_names)

                # Detect SPLITTER rows (P) to emit SPLITTER:1 at the right position
                _row_types  = meta.get('row_types', {})
                _h_nums     = sorted(r for r, t in _row_types.items() if t == 'H')
                _d_nums     = sorted(r for r, t in _row_types.items() if t == 'D')
                _f_nums     = sorted(r for r, t in _row_types.items() if t == 'F')
                _p_nums     = set(r for r, t in _row_types.items() if t == 'P')

                table_name = meta.get('name', '').strip()

                for h_row in meta['header_rows']:
                    _emit_header_footer_row(h_row, state.body_rows)

                # SPLITTER between header and data
                if _h_nums and _d_nums and _p_nums:
                    if any(max(_h_nums) < p < min(_d_nums) for p in _p_nums):
                        state.body_rows.append(['', 'SPLITTER:1'])

                data_vars = meta.get('data_vars', [])
                var_names = []
                for item in data_vars:
                    if isinstance(item, dict):
                        role   = item.get('role', 'var')
                        if role == 'ignore' or item.get('var_name') == 'IGNORE':
                            var_names.append('IGNORE')
                        elif role == 'label':
                            ln = item.get('lbl_name', 'IGNORE')
                            var_names.append(ln)
                            if ln and ln != 'IGNORE':
                                state.lbl_defs.append((
                                    ln,
                                    item.get('lbl_type', 'string'),
                                    item.get('lbl_match', '.*'),
                                ))
                        else:  # var
                            vn    = item.get('var_name', 'IGNORE')
                            vtype = item.get('var_type', 'string')
                            vmatch= item.get('var_match', '.*')
                            # Auto-prefix with table name unless already namespaced
                            if (table_name and vn and vn != 'IGNORE'
                                    and not vn.startswith(table_name + '.')):
                                vn = f'{table_name}.{vn}'
                            var_names.append(vn)
                            if vn and vn != 'IGNORE':
                                state.var_defs.append((vn, vtype, vmatch))
                    else:
                        vn = item
                        vtype, vmatch = 'string', '.*'
                        if (table_name and vn and vn != 'IGNORE'
                                and not vn.startswith(table_name + '.')):
                            vn = f'{table_name}.{vn}'
                        var_names.append(vn)
                        if vn and vn != 'IGNORE':
                            state.var_defs.append((vn, vtype, vmatch))
                state.body_rows.append(['', f'DATA:{mult}'] + var_names)

                # SPLITTER between data and footer
                if _d_nums and _f_nums and _p_nums:
                    if any(max(_d_nums) < p < min(_f_nums) for p in _p_nums):
                        state.body_rows.append(['', 'SPLITTER:1'])

                for f_row in meta.get('footer_rows', []):
                    _emit_header_footer_row(f_row, state.body_rows, 'FOOTER:1')
            else:
                # Legacy single-header-row model (columns list)
                cols      = meta.get('columns', [])
                lbl_names = []
                var_names = []
                for col in cols:
                    lbl_name  = col.get('lbl_name', 'IGNORE')
                    var_name  = col.get('var_name', 'IGNORE')
                    cell_val  = col.get('cell_value', '')
                    state.lbl_defs.append((lbl_name, 'string', cell_val))
                    state.var_defs.append((var_name, col.get('var_type', 'string'),
                                           col.get('var_match', '.*')))
                    lbl_names.append(lbl_name)
                    var_names.append(var_name)
                state.body_rows.append(['', 'HEADER:1'] + lbl_names)
                state.body_rows.append(['', f'DATA:{mult}'] + var_names)
        elif choice in ('T-HEAD', 'T-DATA'):
            pass  # handled by the anchor cell above
        elif choice == 'I':
            state.body_rows.append(['cell:1', 'IGNORE'])
    return state


def _choices_to_csv(ws, choices, cells, direction, sheet_name) -> str:
    state = _build_state_from_choices(ws, choices, cells, direction, sheet_name)
    buf = io.StringIO()
    w   = csv.writer(buf)
    w.writerow(['config:', 'read.direction', state.direction])
    for name, typ, text in state.lbl_defs:
        w.writerow(['lbl:', name, typ, text])
    for name, typ, match in state.var_defs:
        w.writerow(['var:', name, typ, match])
    w.writerow(['START:'])
    for row in state.body_rows:
        w.writerow(row)
    w.writerow(['END:'])
    return buf.getvalue()


if _TEXTUAL_OK:

    # ── Modals ─────────────────────────────────────────────────────────────────

    _TYPE_OPTIONS = [
        'string', 'integer', 'number', 'currency', 'percentage',
        'boolean', 'date', 'datetime', 'time', 'duration',
    ]

    # (pattern, short_label) pairs — cycled with F4 in any Match/pattern Input
    _MATCH_PRESETS: list[tuple[str, str]] = [
        ('.*',                    'any'),
        (r'\d+',                  'integer'),
        (r'\d+\.?\d*',            'number'),
        (r'\d{4}-\d{2}-\d{2}',   'ISO date'),
        (r'\d{2}/\d{2}/\d{4}',   'US date'),
        (r'[\w.+]+@[\w.]+\.\w+', 'email'),
        (r'https?://\S+',         'URL'),
    ]

    class _FieldsModal(ModalScreen):
        """Multi-field input.
        Fields: list of (label, default) or (label, default, [opts]) for a Select dropdown.
        ENTER moves between Input fields; last ENTER submits → dismiss(list[str]).
        Tab moves between all fields (including Select). ESC → dismiss(None).
        """
        DEFAULT_CSS = """
        _FieldsModal              { align: center middle; }
        _FieldsModal > #dialog    { background: $surface; border: thick $primary;
                                    width: 68; height: auto; max-height: 82vh;
                                    padding: 1 2; overflow-y: auto; }
        _FieldsModal Label.title   { text-style: bold; margin-bottom: 1; }
        _FieldsModal Label.lbl     { color: $text-muted; margin-top: 1; }
        _FieldsModal Label.hint    { color: $text-muted; margin-top: 1; }
        _FieldsModal Label.presets { color: $text-muted; margin-top: 0; }
        _FieldsModal Select        { width: 100%; margin-top: 0; }
        """

        def __init__(self, title: str, fields: list[tuple]) -> None:
            super().__init__()
            self._title  = title
            self._fields = fields

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Label(self._title, classes='title')
                for i, field in enumerate(self._fields):
                    lbl, default = field[0], field[1]
                    opts    = field[2] if len(field) > 2 else None
                    presets = field[3] if len(field) > 3 else None
                    yield Label(lbl, classes='lbl')
                    if opts is not None:
                        sel_val = default if default in opts else opts[0]
                        yield Select(
                            [(o, o) for o in opts],
                            value=sel_val,
                            id=f'f{i}',
                            allow_blank=False,
                        )
                    else:
                        yield Input(value=default, id=f'f{i}')
                    if presets is not None:
                        labels = '  '.join(f'[dim]{lbl}[/dim]' for _, lbl in presets)
                        yield Label(f'F4 cycles: {labels}', classes='presets')
                yield Label('ENTER = confirm  •  Tab = next field  •  ESC = cancel', classes='hint')

        def on_mount(self) -> None:
            try:
                self.query_one('#f0').focus()
            except Exception:
                pass

        def on_input_submitted(self, event: Input.Submitted) -> None:
            for i in range(len(self._fields)):
                try:
                    w = self.query_one(f'#f{i}')
                    if w is event.input:
                        # Advance to the next INPUT field, skipping Selects.
                        # Textual's Select widget consumes ENTER internally
                        # (open/close overlay), so ENTER-based navigation must
                        # skip them. The user reaches Selects via Tab.
                        advanced = False
                        for next_i in range(i + 1, len(self._fields)):
                            try:
                                nw = self.query_one(f'#f{next_i}')
                                if isinstance(nw, Input):
                                    nw.focus()
                                    advanced = True
                                    break
                            except Exception:
                                pass
                        if not advanced:
                            self._submit()
                        return
                except Exception:
                    pass
            self._submit()

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(None)
            elif event.key == 'enter':
                # If Select dropdown is open, let Select handle ENTER (picks the item).
                # If closed, behave like Input: advance to next field, or submit if last.
                focused = self.focused
                if isinstance(focused, Select) and not focused.expanded:
                    for i in range(len(self._fields)):
                        try:
                            w = self.query_one(f'#f{i}')
                            if w is focused:
                                if i < len(self._fields) - 1:
                                    try:
                                        self.query_one(f'#f{i + 1}').focus()
                                    except Exception:
                                        pass
                                else:
                                    self._submit()
                                break
                        except Exception:
                            pass
            elif event.key == 'f4':
                # Cycle through presets for the focused Input field
                focused = self.focused
                if not isinstance(focused, Input):
                    return
                for i, field in enumerate(self._fields):
                    presets = field[3] if len(field) > 3 else None
                    if presets is None:
                        continue
                    try:
                        w = self.query_one(f'#f{i}')
                    except Exception:
                        continue
                    if w is focused:
                        vals = [p[0] for p in presets]
                        curr = focused.value
                        try:
                            nxt = vals[(vals.index(curr) + 1) % len(vals)]
                        except ValueError:
                            nxt = vals[0]
                        focused.value = nxt
                        focused.cursor_position = len(nxt)
                        event.stop()
                        break

        def _submit(self) -> None:
            results = []
            for i, field in enumerate(self._fields):
                default = field[1]
                try:
                    w = self.query_one(f'#f{i}')
                    if isinstance(w, Select):
                        val = w.value
                        results.append(
                            str(val) if val is not None and val is not Select.BLANK
                            else default
                        )
                    else:
                        results.append(w.value)
                except Exception:
                    results.append(default)
            self.dismiss(results)


    class _GotoModal(ModalScreen):
        DEFAULT_CSS = """
        _GotoModal              { align: center middle; }
        _GotoModal > #dialog    { background: $surface; border: thick $primary;
                                  width: 44; height: auto; padding: 1 2; }
        _GotoModal Label.title  { text-style: bold; margin-bottom: 1; }
        _GotoModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Label('Go to cell', classes='title')
                yield Input(placeholder='e.g. B5', id='ref')
                yield Label('ENTER = jump  •  ESC = cancel', classes='hint')

        def on_mount(self) -> None:
            self.query_one('#ref', Input).focus()

        def on_input_submitted(self, event: Input.Submitted) -> None:
            self.dismiss(event.value.strip().upper() or None)

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(None)


    class _ConfigModal(ModalScreen):
        DEFAULT_CSS = """
        _ConfigModal              { align: center middle; }
        _ConfigModal > #dialog    { background: $surface; border: thick $primary;
                                    width: 72; height: auto; padding: 1 3; }
        _ConfigModal Label.title  { text-style: bold; margin-bottom: 1; }
        _ConfigModal Label.sect   { text-style: bold; margin-top: 1; }
        _ConfigModal Label.desc   { color: $text-muted; margin-bottom: 1; }
        _ConfigModal Select       { margin-bottom: 1; }
        _ConfigModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def __init__(self, is_template: bool) -> None:
            super().__init__()
            self._is_template = is_template

        def compose(self) -> ComposeResult:
            note = '  [yellow]⚑ template detected[/yellow]' if self._is_template else ''
            with Vertical(id='dialog'):
                yield Label(f'Wizard configuration{note}', classes='title')
                yield Label('Scan direction', classes='sect')
                yield Label('In which direction does the data read?', classes='desc')
                yield Select(
                    options=[('Left → Right  (LR)', 'LR'), ('Top → Down  (TD)', 'TD')],
                    value='LR', id='dir',
                )
                yield Label('Template mode', classes='sect')
                yield Label('Empty cells after labels are treated as variable slots',
                            classes='desc')
                yield Select(
                    options=[('Yes — empty slots are variables', 'yes'),
                              ('No — classify each cell manually', 'no')],
                    value='yes' if self._is_template else 'no',
                    id='tpl',
                )
                yield Label('ENTER = start  •  ESC = cancel', classes='hint')

        def on_key(self, event) -> None:
            if event.key == 'enter':
                try:
                    direction = str(self.query_one('#dir', Select).value)
                    template  = str(self.query_one('#tpl', Select).value) == 'yes'
                except Exception:
                    direction, template = 'LR', self._is_template
                self.dismiss({'direction': direction, 'template': template})
            elif event.key == 'escape':
                self.dismiss(None)


    class _ZoomModal(ModalScreen):
        """Read-only full-content view of a cell (the table truncates to 20 chars)."""
        DEFAULT_CSS = """
        _ZoomModal              { align: center middle; }
        _ZoomModal > #dialog    { background: $surface; border: thick $primary;
                                  width: 72; height: auto; padding: 1 2; }
        _ZoomModal Label.title  { text-style: bold; color: $accent; margin-bottom: 1; }
        _ZoomModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def __init__(self, ref: str, value: Any, meta: dict | None) -> None:
            super().__init__()
            self._ref   = ref
            self._value = value
            self._meta  = meta

        def compose(self) -> ComposeResult:
            display = '' if self._value is None else str(self._value)
            status  = ''
            if self._meta:
                ch    = self._meta.get('choice', '')
                cname = _CHOICE_NAME.get(ch, ch)
                ccol  = _CHOICE_COLOR.get(ch, 'white')
                name  = self._meta.get('name', '')
                status = f'\n[bold {ccol}]✓ {cname}[/bold {ccol}]' + (f'  [dim]{name}[/dim]' if name else '')
            with Vertical(id='dialog'):
                yield Label(f'Cell {self._ref}', classes='title')
                yield Static(f'[white]{display}[/white]{status}')
                yield Label('ESC / ENTER to close', classes='hint')

        def on_key(self, event) -> None:
            if event.key in ('escape', 'enter', 'space'):
                self.dismiss(None)


    class _PreviewModal(ModalScreen):
        """Show the current pattern as it would be written to CSV."""
        DEFAULT_CSS = """
        _PreviewModal              { align: center middle; }
        _PreviewModal > #dialog    { background: $surface; border: thick $primary;
                                     width: 80; height: 28; padding: 1 2;
                                     overflow-y: auto; }
        _PreviewModal Label.title  { text-style: bold; color: $accent; margin-bottom: 1; }
        _PreviewModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def __init__(self, csv_text: str) -> None:
            super().__init__()
            self._csv = csv_text

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Label('Pattern preview (current state)', classes='title')
                yield Static(self._csv)
                yield Label('ESC to close', classes='hint')

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(None)


    class _HelpModal(ModalScreen):
        DEFAULT_CSS = """
        _HelpModal              { align: center middle; }
        _HelpModal > #dialog    { background: $surface; border: thick $primary;
                                  width: 84; height: auto; max-height: 44;
                                  padding: 1 2; overflow-y: auto; }
        _HelpModal Label.title  { text-style: bold; color: $accent; margin-bottom: 1; }
        _HelpModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        _HELP = """\
[bold cyan]CLASSIFY[/bold cyan]

  [bold green]L[/bold green]  [bold]Label[/bold]    — Cell whose text identifies a nearby value (e.g. "Invoice No:").
              In template mode, cursor jumps to the adjacent value cell next.

  [bold blue]C[/bold blue]  [bold]Header[/bold]   — Section title or structural marker; no value follows it.

  [bold bright_yellow]V[/bold bright_yellow]  [bold]Value[/bold]    — Extract this cell's content as a named variable.

  [bold magenta]T[/bold magenta]  [bold]Table[/bold]    — Define a repeating block (mini-table). A 3-step wizard opens:
              step 1 — set name, multiplicity, and row range;
              step 2 — classify each row as Header / Data / Footer / Skip;
              step 3 — name each column per row type.

  [dim]I[/dim]  [bold]Ignore[/bold]   — Skip this cell; it produces no pattern entry.

  [bold]R[/bold]  [bold]Remove[/bold]   — Clear the current cell's classification so you can redo it.

[bold cyan]NAVIGATE[/bold cyan]

  [bold]N[/bold]  Next non-empty cell     [bold]P[/bold]  Prev non-empty cell
  [bold]U[/bold]  Next unclassified cell  [bold]G[/bold]  Go to cell ref (e.g. D11)
  [bold]H[/bold]  Highlight all unclassified cells in amber (press again to clear)

[bold cyan]OTHER[/bold cyan]

  [bold]ENTER[/bold]      Auto-accept proposed classification (no modal, uses defaults)
  [bold]Space[/bold]      Zoom — view full untruncated cell content
  [bold]F1 / ?[/bold]     This help screen
  [bold]F3[/bold]         Preview current pattern
  [bold]F2[/bold]         Add internal note to current cell
  [bold]Ctrl+Z[/bold]     Undo last classification
  [bold]E[/bold]          End wizard and save pattern file
  [bold]Ctrl+Q[/bold]     Cancel without saving

[bold cyan]LEGEND[/bold cyan]

  [bold green]●[/bold green] Label (L)   [bold bright_yellow]●[/bold bright_yellow] Value (V)   [bold blue]●[/bold blue] Header (C)
  [bold magenta]●[/bold magenta] Table (T)   [dim]○[/dim] Ignore (I)  white = not yet classified\
"""

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Label('grepxcel wizard — Help  (F1)', classes='title')
                yield Static(self._HELP)
                yield Label('ESC to close', classes='hint')

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(None)


    class _ConfirmModal(ModalScreen):
        """Warning + Y/N confirmation before a destructive/questionable save."""
        DEFAULT_CSS = """
        _ConfirmModal              { align: center middle; }
        _ConfirmModal > #dialog    { background: $surface; border: thick $warning;
                                     width: 72; height: auto; padding: 1 2; }
        _ConfirmModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def __init__(self, message: str) -> None:
            super().__init__()
            self._message = message

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Static(self._message)
                yield Label('Y = save anyway  •  any other key = go back', classes='hint')

        def on_key(self, event) -> None:
            if event.key.lower() == 'y':
                self.dismiss(True)
            else:
                self.dismiss(False)


    class _LogModal(ModalScreen):
        """Session event log — hidden support tool, triggered by F12."""
        DEFAULT_CSS = """
        _LogModal              { align: center middle; }
        _LogModal > #dialog    { background: $surface; border: thick $warning;
                                 width: 90; height: 34; padding: 1 2;
                                 overflow-y: auto; }
        _LogModal Label.title  { text-style: bold; color: $warning; margin-bottom: 1; }
        _LogModal Label.hint   { color: $text-muted; margin-top: 1; }
        """

        def __init__(self, lines: list[str], data_file: str) -> None:
            super().__init__()
            self._lines     = lines
            self._data_file = data_file
            self._wrote: str | None = None

        def compose(self) -> ComposeResult:
            fname   = os.path.basename(self._data_file)
            content = '\n'.join(self._lines) if self._lines else '[dim](no events yet)[/dim]'
            with Vertical(id='dialog'):
                yield Label(f'Session event log — {fname}', classes='title')
                yield Static(content)
                yield Label('ESC to close  •  W to write log to file', classes='hint')

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(self._wrote)
            elif event.key == 'w':
                ts       = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
                stem     = os.path.splitext(os.path.basename(self._data_file))[0]
                log_path = os.path.join(
                    os.path.dirname(os.path.abspath(self._data_file)),
                    f'grepxcel-wizard-log-{stem}-{ts}.txt',
                )
                with open(log_path, 'w', encoding='utf-8') as fh:
                    fh.write('\n'.join(self._lines))
                self._wrote = log_path
                # Update hint to confirm
                try:
                    self.query_one(Label.hint if False else 'Label.hint', Label).update(
                        f'[bold green]Written:[/bold green] {log_path}  •  ESC to close'
                    )
                except Exception:
                    pass


    # ── Main application ───────────────────────────────────────────────────────

    class _TableSetupModal(ModalScreen):
        """Table definition — step 1: name, full table range, multiplicity.

        The range covers ALL rows of the table (headers + data + footers).
        Dismissed with {'name', 'range_str', 'mult', 'start', 'end'} or None.
        """
        DEFAULT_CSS = """
        _TableSetupModal              { align: center middle; }
        _TableSetupModal > #dialog    { background: $surface; border: thick $primary;
                                        width: 76; height: auto; padding: 1 3; }
        _TableSetupModal Label.title  { text-style: bold; color: $accent; margin-bottom: 1; }
        _TableSetupModal Label.sect   { text-style: bold; margin-top: 1; }
        _TableSetupModal Label.desc   { color: $text-muted; margin-bottom: 1; }
        _TableSetupModal Label.err    { color: $error; margin-top: 1; }
        _TableSetupModal Label.hint   { color: $text-muted; margin-top: 1; }
        _TableSetupModal Input        { margin-bottom: 1; }
        """

        def __init__(self, default_name: str, default_range: str,
                     default_mult: str = '*') -> None:
            super().__init__()
            self._default_name  = default_name
            self._default_range = default_range
            self._default_mult  = default_mult

        def compose(self) -> ComposeResult:
            with Vertical(id='dialog'):
                yield Label('[bold magenta]Table[/bold magenta] — repeating row block',
                            classes='title')
                yield Label('Table name', classes='sect')
                yield Label('Short identifier used in the output JSON key', classes='desc')
                yield Input(value=self._default_name, id='name')
                yield Label('Full table range', classes='sect')
                yield Label(
                    'Include ALL rows: headers, data rows, footers  (e.g. B11:F29)',
                    classes='desc')
                yield Input(value=self._default_range, id='range',
                            placeholder='e.g. B11:F29')
                yield Label('Multiplicity', classes='sect')
                yield Label(
                    '*=any instances  ·  1=exactly one  ·  {n,m}=bounded range',
                    classes='desc')
                yield Input(value=self._default_mult, id='mult', placeholder='*')
                yield Label(id='errmsg', classes='err')
                yield Label('ENTER = next field  •  ESC = cancel', classes='hint')

        def on_mount(self) -> None:
            self.query_one('#name', Input).focus()

        def on_input_submitted(self, event: Input.Submitted) -> None:
            inputs = list(self.query(Input))
            try:
                idx = inputs.index(event.input)
            except ValueError:
                idx = len(inputs) - 1
            if idx < len(inputs) - 1:
                inputs[idx + 1].focus()
            else:
                self._validate_and_submit()

        def _validate_and_submit(self) -> None:
            name       = self.query_one('#name',  Input).value.strip()
            range_str  = self.query_one('#range', Input).value.strip().upper()
            mult_raw   = self.query_one('#mult',  Input).value.strip() or '*'
            errlbl     = self.query_one('#errmsg', Label)
            if not name:
                errlbl.update('Table name cannot be empty.')
                self.query_one('#name', Input).focus()
                return
            parts = range_str.split(':')
            if len(parts) != 2:
                errlbl.update('Range must be like B11:F11 (two cell refs separated by :)')
                self.query_one('#range', Input).focus()
                return
            start = _parse_cell_ref(parts[0].strip())
            end   = _parse_cell_ref(parts[1].strip())
            if not start or not end:
                errlbl.update('Could not parse the range — use the format B11:F11.')
                self.query_one('#range', Input).focus()
                return
            if start[0] > end[0]:
                errlbl.update('Start row must be ≤ end row (first:last).')
                self.query_one('#range', Input).focus()
                return
            if start[1] > end[1]:
                errlbl.update('Start column must be ≤ end column.')
                self.query_one('#range', Input).focus()
                return
            # Normalise multiplicity
            if mult_raw == '1':
                mult = '1'
            elif mult_raw == '*':
                mult = '*'
            elif mult_raw.startswith('{') and mult_raw.endswith('}'):
                mult = mult_raw
            elif '..' in mult_raw:
                a, b = mult_raw.split('..', 1)
                mult = f'{{{a.strip()},{b.strip()}}}'
            elif mult_raw.isdigit():
                mult = mult_raw
            else:
                mult = '*'
            errlbl.update('')
            self.dismiss({'name': name, 'range_str': range_str,
                          'mult': mult, 'start': start, 'end': end})

        def on_key(self, event) -> None:
            if event.key == 'escape':
                self.dismiss(None)


    class _TableRangeModal(ModalScreen):
        """Step 2 of table definition: show the full range; classify each row.

        User presses H/D/F/S on the highlighted row to mark it as:
          H = HEADER (fixed label row — column values become label text)
          D = DATA   (repeating data rows — variables extracted here)
          F = FOOTER (fixed summary row — column values become label text)
          S = SKIP   (ignore this row entirely)
        ENTER confirms; ESC cancels.
        Dismissed with {row_num: 'H'|'D'|'F'|'S'} dict, or None.
        """
        DEFAULT_CSS = """
        _TableRangeModal            { align: center middle; }
        _TableRangeModal > #dialog  { background: $surface; border: thick $primary;
                                      width: 94; height: auto; max-height: 85vh;
                                      padding: 1 2; }
        _TableRangeModal Label.t    { text-style: bold; color: $accent; margin-bottom: 1; }
        _TableRangeModal Label.h    { color: $text-muted; margin-bottom: 1; }
        _TableRangeModal DataTable  { height: auto; max-height: 56vh; }
        """

        BINDINGS = [
            Binding('h', 'set_h', 'Header',   show=True),
            Binding('d', 'set_d', 'Data',     show=True),
            Binding('f', 'set_f', 'Footer',   show=True),
            Binding('p', 'set_p', 'Splitter', show=True),
            Binding('s', 'set_s', 'Skip',     show=True),
        ]

        _TYPE_LABEL = {
            'H': '[bold green]HEADER  [/bold green]',
            'D': '[bold blue]DATA    [/bold blue]',
            'F': '[bold cyan]FOOTER  [/bold cyan]',
            'P': '[magenta]SPLITTER[/magenta]',
            'S': '[dim]SKIP    [/dim]',
        }

        def __init__(self, ws, start_row: int, end_row: int,
                     start_col: int, end_col: int,
                     name: str, mult: str,
                     existing_row_types: 'dict | None' = None) -> None:
            super().__init__()
            self._ws        = ws
            self._start_row = start_row
            self._end_row   = min(end_row, start_row + 35)  # cap display at 36 rows
            self._start_col = start_col
            self._end_col   = end_col
            self._name      = name
            self._mult      = mult

            if existing_row_types:
                # Re-editing: pre-fill from saved classification
                self._row_types = {r: existing_row_types.get(r, 'S')
                                   for r in range(self._start_row, self._end_row + 1)}
            else:
                # Auto-detect row types
                self._row_types: dict[int, str] = {}
                first_header_set = False
                for r in range(self._start_row, self._end_row + 1):
                    vals = [ws.cell(row=r, column=c).value
                            for c in range(start_col, end_col + 1)]
                    has_text  = any(isinstance(v, str) and v.strip() for v in vals)
                    has_value = any(v is not None for v in vals)
                    if not first_header_set and has_text:
                        self._row_types[r] = 'H'
                        first_header_set = True
                    elif has_value:
                        self._row_types[r] = 'D'
                    else:
                        self._row_types[r] = 'S'

        def _col_letter(self, col: int) -> str:
            from openpyxl.utils import get_column_letter
            return get_column_letter(col)

        def compose(self) -> ComposeResult:
            start_ref = _cell_ref(self._start_row, self._start_col)
            end_ref   = _cell_ref(self._end_row, self._end_col)
            with Vertical(id='dialog'):
                yield Label(
                    f'[bold magenta]Table "{self._name}"[/bold magenta]'
                    f'  range {start_ref}:{end_ref}  mult={self._mult}',
                    classes='t',
                )
                yield Label(
                    '[bold green]H[/bold green]=Header  '
                    '[bold blue]D[/bold blue]=Data  '
                    '[bold cyan]F[/bold cyan]=Footer  '
                    '[magenta]P[/magenta]=Splitter (blank row)  '
                    '[dim]S[/dim]=Skip  '
                    '↑↓ navigate  ENTER confirm  ESC cancel',
                    classes='h',
                )
                yield DataTable(id='rdt', cursor_type='row')

        def on_mount(self) -> None:
            dt = self.query_one('#rdt', DataTable)
            col_keys  = ['_row', '_type'] + [
                self._col_letter(c)
                for c in range(self._start_col, self._end_col + 1)
            ]
            col_labels = ['Row', 'Type'] + [
                self._col_letter(c)
                for c in range(self._start_col, self._end_col + 1)
            ]
            for key, label in zip(col_keys, col_labels):
                dt.add_column(label, key=key)
            for r in range(self._start_row, self._end_row + 1):
                vals = [
                    _trunc(self._ws.cell(row=r, column=c).value, 11)
                    for c in range(self._start_col, self._end_col + 1)
                ]
                rtype = self._row_types.get(r, 'D')
                dt.add_row(
                    str(r),
                    self._TYPE_LABEL[rtype],
                    *vals,
                    key=str(r),
                )
            dt.focus()

        def _set_current_type(self, new_type: str) -> None:
            dt      = self.query_one('#rdt', DataTable)
            row_idx = dt.cursor_row
            ws_row  = self._start_row + row_idx
            self._row_types[ws_row] = new_type
            dt.update_cell(str(ws_row), '_type', self._TYPE_LABEL[new_type])

        def action_set_h(self) -> None:
            self._set_current_type('H')

        def action_set_d(self) -> None:
            self._set_current_type('D')

        def action_set_f(self) -> None:
            self._set_current_type('F')

        def action_set_p(self) -> None:
            self._set_current_type('P')

        def action_set_s(self) -> None:
            self._set_current_type('S')

        def on_key(self, event) -> None:
            if event.key == 'enter':
                self.dismiss(dict(self._row_types))
            elif event.key == 'escape':
                self.dismiss(None)


    class _Panel(Static):
        """Scrollable right-hand classification panel."""


    class WizardTUIApp(App):
        """Full-screen interactive pattern wizard."""

        CSS = """
        Screen  { layers: default modal; }
        #main   { height: 1fr; }

        DataTable {
            width: 1fr;
            height: 1fr;
        }

        _Panel {
            width: 46;
            height: 1fr;
            border-left: solid $primary-darken-1;
            padding: 0 1;
            overflow-y: auto;
        }

        Footer { height: 1; }
        """

        BINDINGS = [
            # Classify (shown in footer)
            Binding('l', 'act_L',  'Label',       show=True),
            Binding('c', 'act_C',  'Header',       show=True),
            Binding('v', 'act_V',  'Value',        show=True),
            Binding('t', 'act_T',  'Table',        show=True),
            Binding('i', 'act_I',  'Ignore',       show=True),
            # Navigate (shown in footer)
            Binding('n', 'nav_next',  'Next',      show=True),
            Binding('p', 'nav_prev',  'Prev',      show=True),
            Binding('g', 'nav_goto',  'Goto',      show=True),
            Binding('e', 'end_save',  'End & Save', show=True),
            # Classify extras (hidden, discoverable via ^P palette)
            Binding('r',      'act_R',             'Remove classif.', show=False),
            # Navigate extras (hidden)
            Binding('u',      'nav_unclassified',  'Next unclassified', show=False),
            Binding('h',      'highlight_pending', 'Highlight pending', show=False),
            # Other (hidden)
            Binding('enter',  'accept',             'Auto-accept',     show=False),
            Binding('space',  'zoom',               'Zoom cell',       show=False),
            Binding('f1',             'show_help', 'Help (F1/?)', show=True),
            Binding('question_mark',  'show_help', 'Help',       show=False),
            Binding('f3',     'preview',            'Pattern preview', show=False),
            Binding('ctrl+z', 'undo',               'Undo',            show=False),
            Binding('ctrl+q', 'cancel',             'Cancel',          show=False),
            Binding('q',      'cancel',             'Quit',            show=False),
            # Debug / support hotkeys (hidden)
            Binding('f2',     'add_note',           'Cell note',       show=False),
            Binding('semicolon', 'add_comment',     'Comment',         show=False),
            Binding('f11',    'screenshot',         'Screenshot',      show=False),
            Binding('f12',    'show_log',           'Session log',     show=False),
        ]

        def __init__(self, ws, state: WizardState, data_file: str) -> None:
            super().__init__()
            self._ws        = ws
            self._state     = state   # used for direction / sheet_name only; pattern built at save
            self._data_file = data_file

            self._max_row = ws.max_row or 1
            self._max_col = ws.max_column or 1

            self._history  = _load_history(data_file)
            # Primary state: all classifications live here
            self._choices: dict[str, dict] = {}
            self._last_label_base: str | None = None
            self._current_prefix: str = ''  # dot-notation prefix for variable names

            self._ws_row = 1
            self._ws_col = 1
            self._is_template     = False
            self._cells: list[tuple[int, int]] = []
            self._initialized     = False   # NOTE: must not be named _ready (conflicts with App._ready)
            self._total_nonempty  = 0

            # Undo / highlight state
            self._undo_stack: list[dict] = []
            self._highlighted: set[str]  = set()

            # Session event log (F12 debug tool — tracks all classify/navigate events)
            self._session_start = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            self._event_log: list[str] = [
                f'grepxcel wizard session — {os.path.basename(data_file)}',
                f'Started: {self._session_start}',
                f'Data file: {os.path.abspath(data_file)}',
            ]

            # Per-cell debug notes (internal_notes, F2) — separate from _choices
            self._notes: dict[str, str] = {}

            # ── Persistent log file ─────────────────────────────────────────
            # Layout: logs/wizard/<stem>/<YYYY-MM-DD-HHmmss>_<sha8>/session.log
            # Screenshots land in the same session directory.
            self._log_file = None
            self._log_file_path = None
            self._log_dir: str | None = None
            self._screenshot_count = 0

            ts_file  = datetime.datetime.now().strftime('%Y-%m-%d-%H%M%S')
            stem     = os.path.splitext(os.path.basename(data_file))[0]
            abs_path = os.path.abspath(data_file)

            # Compute SHA-256 of the data file (streaming to handle large files)
            sha256 = hashlib.sha256()
            file_size = 0
            try:
                with open(data_file, 'rb') as _fh:
                    for _chunk in iter(lambda: _fh.read(65536), b''):
                        sha256.update(_chunk)
                        file_size += len(_chunk)
                sha256_hex = sha256.hexdigest()
                sha8 = sha256_hex[:8]
            except OSError:
                sha256_hex = 'unavailable'
                sha8 = 'unknown'
                file_size = 0

            session_dir = os.path.join(
                os.getcwd(), 'logs', 'wizard', stem, f'{ts_file}_{sha8}',
            )
            try:
                os.makedirs(session_dir, exist_ok=True)
                log_path = os.path.join(session_dir, 'session.log')
                self._log_file = open(log_path, 'w', encoding='utf-8')  # noqa: WPS515
                self._log_file_path = log_path
                self._log_dir = session_dir
                sep = '═' * 51
                self._log_file.write(f'grepxcel wizard session\n')
                self._log_file.write(f'{sep}\n')
                self._log_file.write(f'Data file : {abs_path}\n')
                self._log_file.write(f'SHA-256   : {sha256_hex}\n')
                self._log_file.write(f'File size : {file_size} bytes\n')
                self._log_file.write(f'Sheet     : {state.sheet_name or "(active)"}\n')
                self._log_file.write(f'Version   : grepxcel {_GX_VERSION}\n')
                self._log_file.write(f'Started   : {self._session_start}\n')
                self._log_file.write(f'{sep}\n')
                self._log_file.write(f'# Columns: timestamp   EVENT_TYPE    detail\n')
                self._log_file.write(f'# COMMENT lines = user notes typed with ;\n')
                self._log_file.write(f'# NOTE lines    = per-cell notes typed with F2\n')
                self._log_file.write(f'{sep}\n\n')
                self._log_file.flush()
            except OSError:
                pass

        def _log(self, event_type: str, detail: str) -> None:
            _now = datetime.datetime.now()
            ts   = _now.strftime('%H:%M:%S.') + f'{_now.microsecond // 1000:03d}'
            line = f'{ts}  {event_type:<12}  {detail}'
            self._event_log.append(line)
            if self._log_file:
                try:
                    self._log_file.write(line + '\n')
                    self._log_file.flush()
                except OSError:
                    pass

        def _close_log_file(self) -> None:
            """Write summary footer and close the persistent log file."""
            if not self._log_file:
                return
            try:
                ts = datetime.datetime.now().strftime('%H:%M:%S')
                self._log_file.write(f'#\n')
                self._log_file.write(f'# Session ended: {ts}\n')
                counts = {}
                for m in self._choices.values():
                    counts[m.get('choice', '?')] = counts.get(m.get('choice', '?'), 0) + 1
                self._log_file.write(f'# Classifications: {counts}\n')
                if self._notes:
                    self._log_file.write(f'#\n# Internal notes:\n')
                    for ref, note in sorted(self._notes.items()):
                        self._log_file.write(f'#   {ref}: {note}\n')
                self._log_file.close()
                self._log_file = None
            except OSError:
                pass

        def compose(self) -> ComposeResult:
            yield Header(show_clock=False)
            with Horizontal(id='main'):
                yield DataTable(id='sheet', cursor_type='cell', zebra_stripes=True)
                yield _Panel(id='panel')
            yield Footer()

        async def on_mount(self) -> None:
            fname = os.path.basename(self._data_file)
            self.title = f'grepxcel wizard — {fname} ({self._state.sheet_name})'
            is_tpl = _detect_template(self._ws, self._data_file)
            self.push_screen(_ConfigModal(is_tpl), self._on_config_done)

        def _on_config_done(self, cfg: dict | None) -> None:
            if cfg is None:
                self.exit(result=None)
                return
            self._state.direction = cfg['direction']
            self._is_template     = cfg['template']
            self._cells           = _build_cell_order(self._ws, self._state.direction)
            self._total_nonempty  = sum(
                1 for r, c in self._cells
                if self._ws.cell(row=r, column=c).value is not None
            )
            self._log('CONFIG',
                      f'direction={cfg["direction"]}  template={cfg["template"]}'
                      f'  sheet={self._state.sheet_name}'
                      f'  cells={self._total_nonempty} non-empty')
            self._populate_table()
            first = _find_next_nonempty(self._cells, self._ws, 0)
            if first is not None:
                r, c = self._cells[first]
                self._ws_row, self._ws_col = r, c
                self._move_cursor(r, c)
            self._initialized = True
            self._refresh_panel()

        # ── DataTable helpers ──────────────────────────────────────────────────

        def _populate_table(self) -> None:
            table = self.query_one('#sheet', DataTable)
            table.clear(columns=True)
            table.add_column('#', width=4, key='__rn__')
            for col in range(1, self._max_col + 1):
                table.add_column(_col_label(col), key=str(col))
            for row in range(1, self._max_row + 1):
                cells: list[Any] = [RichText(str(row), style='dim')]
                for col in range(1, self._max_col + 1):
                    v    = self._ws.cell(row=row, column=col).value
                    ref  = _cell_ref(row, col)
                    meta = self._choices.get(ref)
                    cells.append(
                        _styled(v, meta['choice']) if meta
                        else _trunc(v)
                    )
                table.add_row(*cells, key=str(row))
            table.fixed_columns = 1

        def _move_cursor(self, ws_row: int, ws_col: int) -> None:
            table = self.query_one('#sheet', DataTable)
            # DataTable col 0 = row-pin; col N = worksheet col N
            table.move_cursor(row=ws_row - 1, column=ws_col, animate=False)

        def _restyle_cell(self, ws_row: int, ws_col: int) -> None:
            table = self.query_one('#sheet', DataTable)
            ref   = _cell_ref(ws_row, ws_col)
            value = self._ws.cell(row=ws_row, column=ws_col).value
            meta  = self._choices.get(ref)
            new   = _styled(value, meta['choice']) if meta else _trunc(value)
            try:
                table.update_cell(str(ws_row), str(ws_col), new, update_width=False)
            except Exception:
                pass

        def on_data_table_cell_highlighted(
            self, event: DataTable.CellHighlighted
        ) -> None:
            """Update panel whenever the DataTable cursor moves to a new cell."""
            if not self._initialized:
                return
            dt_col = event.coordinate.column
            dt_row = event.coordinate.row
            if dt_col == 0:     # row-pin column — ignore
                return
            ws_row = dt_row + 1
            ws_col = dt_col     # DataTable col N == worksheet col N
            if ws_row == self._ws_row and ws_col == self._ws_col:
                return
            self._clear_highlights()
            self._ws_row = ws_row
            self._ws_col = ws_col
            self._refresh_panel()

        # ── Panel (4 zones) ────────────────────────────────────────────────────

        def _proposal(self) -> str:
            v = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            if v is None:
                return 'var:string' if (self._is_template and self._last_label_base) else 'skip'
            return _propose_type(v, ws=self._ws, row=self._ws_row, col=self._ws_col)

        def _refresh_panel(self) -> None:
            ref      = _cell_ref(self._ws_row, self._ws_col)
            raw      = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            proposal = self._proposal()
            meta     = self._choices.get(ref)
            prior    = self._history.get(ref)

            # ── Zone 1 — CURRENT CELL ─────────────────────────────────────────
            if raw is None:
                if self._is_template and self._last_label_base:
                    val_line = '[dim](empty — template slot)[/dim]'
                else:
                    val_line = '[dim](empty)[/dim]'
            else:
                s = str(raw)
                val_line = f'[white]"{s[:36]}{"…" if len(s) > 36 else ""}[/white]"'

            if proposal == 'label':
                prop_line = '[bold green]label[/bold green]'
            elif proposal.startswith('var:'):
                prop_line = f'[bold bright_yellow]{proposal}[/bold bright_yellow]'
            else:
                prop_line = f'[dim]{proposal}[/dim]'

            # Build status line — special handling for table cells
            if meta:
                ch   = meta['choice']
                ccol = _CHOICE_COLOR.get(ch, 'white')
                if ch == 'T':
                    tname = meta.get('name', '')
                    trng  = meta.get('range', ref)
                    mult  = meta.get('mult', '*')
                    if 'header_rows' in meta:
                        n_h = len(meta.get('header_rows', []))
                        n_d = len(meta.get('data_vars', []))
                        n_f = len(meta.get('footer_rows', []))
                        n_s = len(meta.get('skip_rows', []))
                        detail = (f'H={n_h} D={n_d} F={n_f}'
                                  + (f' S={n_s}' if n_s else '')
                                  + f'  mult={mult}')
                    else:
                        ncols  = len(meta.get('columns', []))
                        detail = f'{ncols} col{"s" if ncols != 1 else ""}  mult={mult}'
                    status = (f'[bold magenta]✓ Table anchor[/bold magenta]'
                              f'  [dim]{tname}[/dim]')
                    extra  = [
                        f'  Range:  [magenta]{trng}[/magenta]  {detail}',
                        '  [dim]R = remove whole table[/dim]',
                    ]
                elif ch == 'T-HEAD':
                    anchor_ref  = meta.get('anchor', '')
                    anchor_meta = self._choices.get(anchor_ref, {})
                    tname  = anchor_meta.get('name', '')
                    trng   = anchor_meta.get('range', anchor_ref)
                    # Find role of this cell: header/footer, column index
                    col_name = ''
                    role = 'header col'
                    if 'header_rows' in anchor_meta:
                        for h_row in anchor_meta.get('header_rows', []):
                            for col in h_row['cols']:
                                if col['ref'] == ref:
                                    col_name = col['lbl_name']
                                    role = 'header col'
                        for f_row in anchor_meta.get('footer_rows', []):
                            for col in f_row['cols']:
                                if col['ref'] == ref:
                                    col_name = col['lbl_name']
                                    role = 'footer col'
                    else:
                        col_idx = next(
                            (i for i, c in enumerate(anchor_meta.get('columns', []))
                             if c.get('ref') == ref), -1
                        )
                        if col_idx >= 0:
                            col_name = anchor_meta['columns'][col_idx].get('var_name', '')
                    status = (f'[magenta]✓ Table {role}[/magenta]'
                              f'  [dim]{tname}: {col_name}[/dim]')
                    extra  = [
                        f'  Table:  [magenta]{tname}[/magenta]  range {trng}',
                        '  [dim]T = edit table  ·  R = remove table[/dim]',
                    ]
                elif ch == 'T-DATA':
                    anchor_ref  = meta.get('anchor', '')
                    anchor_meta = self._choices.get(anchor_ref, {})
                    tname = anchor_meta.get('name', '')
                    trng  = anchor_meta.get('range', anchor_ref)
                    status = (f'[dim magenta]✓ Table data row[/dim magenta]'
                              f'  [dim]{tname}[/dim]')
                    extra  = [
                        f'  Table:  [magenta]{tname}[/magenta]  range {trng}',
                        '  [dim]T = edit table  ·  R = remove table[/dim]',
                    ]
                else:
                    cname  = _CHOICE_NAME.get(ch, ch)
                    status = f'[bold {ccol}]✓ {cname}[/bold {ccol}]'
                    if 'name' in meta:
                        status += f'  [dim]{meta["name"]}[/dim]'
                    extra = []
            elif prior:
                pname  = _CHOICE_NAME.get(prior, prior)
                status = f'[magenta]↺ prev: {pname}[/magenta]'
                extra  = []
            else:
                status = '[dim]not classified[/dim]'
                extra  = []

            note = self._notes.get(ref, '')
            lines: list[str] = [
                f'[bold cyan]─ {ref} {"─" * (40 - len(ref))}[/bold cyan]',
                f'  {val_line}',
                f'  Proposal:  {prop_line}',
                f'  Status:    {status}',
            ] + extra
            if note:
                lines.append(f'  [bold yellow]✎ Note:[/bold yellow] {_trunc(note, 36)}')
            if self._last_label_base and not meta:
                lines += [
                    f'  [dim magenta]← label: {self._last_label_base}_label[/dim magenta]',
                    f'  [dim cyan]  hint: variable → {self._last_label_base}[/dim cyan]',
                ]
            lines.append('')

            # ── Zone 2 — CLASSIFY ─────────────────────────────────────────────
            lines += [
                f'[bold cyan]─ Classify {"─" * 32}[/bold cyan]',
                '  [dim]ENTER[/dim]  auto-accept proposal',
                '  [bold green]L[/bold green]  Label  — text anchors a value',
                '  [bold blue]C[/bold blue]  Header — section title only',
                '  [bold bright_yellow]V[/bold bright_yellow]  Value  — extract this cell',
                '  [bold magenta]T[/bold magenta]  Table  — define range + column names',
                '  [dim]I[/dim]  Ignore — skip',
                '  [bold]R[/bold]  Remove classification',
                '',
            ]

            # ── Zone 3 — NAVIGATE ─────────────────────────────────────────────
            lines += [
                f'[bold cyan]─ Navigate {"─" * 32}[/bold cyan]',
                '  [bold]N[/bold]  Next non-empty',
                '  [bold]P[/bold]  Prev non-empty',
                '  [bold]U[/bold]  Next unclassified',
                '  [bold]G[/bold]  Go to cell (e.g. D11)',
                '  [bold red]E[/bold red]  End & save pattern',
                '',
                '  [dim]Space[/dim] Zoom  [dim]F3[/dim] Preview  [dim]F1[/dim] Help',
                '  [dim]H[/dim] Highlight pending  [dim]^Z[/dim] Undo',
                '  [dim]^D[/dim] Dark/light  [dim]^P[/dim] Palette',
                '  [dim]F2[/dim] Cell note  [dim];[/dim] Comment  [dim]F11[/dim] Screenshot  [dim]F12[/dim] View log',
                '',
            ]

            # ── Zone 4 — LEGEND + STATS ───────────────────────────────────────
            counts = {}
            for m in self._choices.values():
                c = m.get('choice', '')
                counts[c] = counts.get(c, 0) + 1

            done    = len(self._choices)
            total   = self._total_nonempty
            pct     = int(done / total * 100) if total else 0

            lines += [
                f'[bold cyan]─ Legend {"─" * 34}[/bold cyan]',
                '  [bold green]●[/bold green] green   = Label (L)',
                '  [bold bright_yellow]●[/bold bright_yellow] bright_yellow = Value (V)',
                '  [bold blue]●[/bold blue] blue    = Header (C)',
                '  [bold magenta]●[/bold magenta] magenta = Table (T)',
                '  [dim]○[/dim] dim     = Ignore (I)',
                '',
                f'[bold cyan]─ Stats {"─" * 35}[/bold cyan]',
                f'  Sheet:  {self._max_row} rows × {self._max_col} cols',
                f'  Done:   [bold]{done}[/bold] / {total} non-empty  ({pct}%)',
                f'  Labels: {counts.get("L", 0)}   Values: {counts.get("V", 0)}'
                f'   Hdrs: {counts.get("C", 0)}',
                f'  Tables: {counts.get("T", 0)}   Ignored: {counts.get("I", 0)}'
                f'   Dir: {self._state.direction}',
                f'  Template: {"[yellow]ON[/yellow]" if self._is_template else "[dim]off[/dim]"}',
            ]
            if self._undo_stack:
                lines.append(f'  [dim]Undo depth: {len(self._undo_stack)}[/dim]')

            self.query_one('#panel', _Panel).update('\n'.join(lines))

        # ── Table helpers ─────────────────────────────────────────────────────

        def _detect_table_end_col(self, ws_row: int, ws_col: int) -> int:
            """Return the last col in ws_row that has a value, starting from ws_col."""
            end = ws_col
            for c in range(ws_col, self._max_col + 1):
                if self._ws.cell(row=ws_row, column=c).value is not None:
                    end = c
                else:
                    break
            return end

        def _extract_table_columns(self,
                                   start_row: int, start_col: int,
                                   end_col: int) -> list[dict]:
            """Return column spec list from the header row range."""
            cols = []
            for c in range(start_col, end_col + 1):
                val  = self._ws.cell(row=start_row, column=c).value
                ref  = _cell_ref(start_row, c)
                slug = _slugify(str(val)) if val is not None else f'col{c}'
                cols.append({
                    'ref':        ref,
                    'row':        start_row,
                    'col':        c,
                    'cell_value': str(val) if val is not None else '',
                    'lbl_name':   f'col_{slug}_label',
                    'var_name':   slug,
                    'var_type':   'string',
                    'var_match':  '.*',
                })
            return cols

        def _commit_table(self, anchor_ref: str, name: str, mult: str,
                          range_str: str, cols: list[dict]) -> None:
            """Write T + T-HEAD entries (legacy single-header format)."""
            anchor_meta = {
                'choice':  'T',
                'name':    name,
                'mult':    mult,
                'range':   range_str,
                'columns': cols,
            }
            self._choices[anchor_ref] = anchor_meta
            parsed = _parse_cell_ref(anchor_ref)
            if parsed:
                self._restyle_cell(parsed[0], parsed[1])
            for col in cols[1:]:
                self._choices[col['ref']] = {'choice': 'T-HEAD', 'anchor': anchor_ref}
                self._restyle_cell(col['row'], col['col'])

        def _all_thead_refs(self, anchor_meta: dict) -> list[str]:
            """Return all T-HEAD cell refs for a table (both old and new model)."""
            refs = []
            if 'header_rows' in anchor_meta:
                for h_row in anchor_meta.get('header_rows', []):
                    for col in h_row['cols']:
                        refs.append(col['ref'])
                for f_row in anchor_meta.get('footer_rows', []):
                    for col in f_row['cols']:
                        refs.append(col['ref'])
            else:
                for col in anchor_meta.get('columns', []):
                    refs.append(col['ref'])
            return refs

        def _remove_table(self, anchor_ref: str, meta: dict) -> None:
            """Remove anchor + all T-HEAD and T-DATA cells for a table."""
            # Collect via _choices scan (handles T-HEAD, T-DATA, any future types)
            members = {anchor_ref}
            for ref, m in list(self._choices.items()):
                if m.get('anchor') == anchor_ref:
                    members.add(ref)
            # Also include legacy header refs derived from meta structure
            for r in self._all_thead_refs(meta):
                members.add(r)
            for ref in members:
                self._choices.pop(ref, None)
                p = _parse_cell_ref(ref)
                if p:
                    self._restyle_cell(p[0], p[1])

        def _col_modal_seq(
            self,
            col_infos: list[dict],
            idx: int,
            existing_cols: 'list | None',
            mode: str,
            ref_row: int,
            results: list,
            on_done,
            h_rows: 'list[int] | None' = None,
            d_rows: 'list[int] | None' = None,
            col_titles: 'list[str] | None' = None,
        ) -> None:
            """Push one _FieldsModal per column in sequence (recursive).

            col_infos   — list of {col, letter} dicts (one per table column)
            idx         — current column index
            existing_cols — pre-fill data (list of dicts) or None
            mode        — 'DATA' | 'HEADER' | 'FOOTER'
            ref_row     — worksheet row used to read the actual cell value
            results     — mutable list[dict|None], filled in as modals complete
            on_done     — callback(results) when all done, callback(None) if cancelled
            h_rows      — (DATA mode) header row numbers; used to infer variable name
            d_rows      — (DATA mode) all data row numbers; used for type inference
            col_titles  — short legend title per column (from the legend step); shown in title
            """
            from openpyxl.utils import get_column_letter as _gcl  # local import is ok
            if idx >= len(col_infos):
                on_done(results)
                return

            ci      = col_infos[idx]
            col_c   = ci['col']
            ltr     = ci['letter']
            val     = self._ws.cell(row=ref_row, column=col_c).value
            val_str = str(val) if val is not None else ''
            val_disp = f'"{_trunc(val, 22)}"' if val is not None else f'(empty {ltr})'
            n_total  = len(col_infos)
            existing = (existing_cols[idx]
                        if existing_cols and idx < len(existing_cols)
                        else None)

            if mode == 'DATA':
                # Infer variable name from the matching header cell (may be empty in data row)
                header_val = None
                if h_rows:
                    for hr in h_rows:
                        hv = self._ws.cell(row=hr, column=col_c).value
                        if hv is not None:
                            header_val = hv
                            break
                slug_src = header_val if header_val is not None else (val or ltr)
                slug_v   = _slugify(str(slug_src))
                n_def    = (existing.get('var_name', slug_v) if existing else slug_v)

                # Show header label (or data cell text) in the title
                if header_val is not None:
                    val_disp = f'"{_trunc(val, 15)}"  ← {_trunc(header_val, 22)}' if val is not None else f'← {_trunc(header_val, 28)}'
                # else val_disp already set above

                # Type inference: scan d_rows for first non-empty cell
                _inf_type = 'string'
                scan_rows = d_rows if d_rows else [ref_row]
                for _dr in scan_rows:
                    _cell_obj = self._ws.cell(row=_dr, column=col_c)
                    if _cell_obj.value is not None:
                        _inf_type = _infer_cell_type(_cell_obj)
                        break
                t_def     = (existing.get('var_type', _inf_type)
                             if existing else _inf_type)
                m_def     = (existing.get('var_match', '.*')
                             if existing else '.*')
                note_def = (existing.get('notes', '')
                            if existing else
                            self._notes.get(_cell_ref(ref_row, col_c), ''))
                # Pre-fill name: restore role prefix for display when editing
                ex_role = existing.get('role', 'var') if existing else 'var'
                if existing:
                    if ex_role == 'label':
                        n_def = f"lbl:{existing.get('lbl_name', n_def)}"
                    elif ex_role == 'ignore':
                        n_def = 'IGNORE'
                    # else var: n_def already set
                fields = [
                    ('Name  (plain = variable · lbl:name = label · empty/IGNORE = skip)', n_def),
                    ('Type', t_def, _TYPE_OPTIONS),
                    ('Match pattern  (F4 cycles presets)', m_def, None, _MATCH_PRESETS),
                    ('Notes  (written to session log — optional)', note_def),
                ]
                row_label = 'DATA'
            else:
                pfx = 'col' if mode == 'HEADER' else 'foot'
                slug_v = _slugify(val_str) if val_str else ltr.lower()
                ex_role = existing.get('role', 'label') if existing else 'label'
                if existing:
                    if ex_role == 'var':
                        n_def = f"var:{existing.get('var_name', slug_v)}"
                    elif ex_role == 'ignore':
                        n_def = 'IGNORE'
                    else:
                        n_def = existing.get('lbl_name', f'{pfx}_{slug_v}_label')
                else:
                    # Default to IGNORE for empty cells — natural action (ENTER)
                    # skips the column instead of creating a dummy label name.
                    n_def = f'{pfx}_{slug_v}_label' if val_str else 'IGNORE'
                t_def = (existing.get('lbl_type', existing.get('var_type', 'string'))
                         if existing else 'string')
                m_def = (existing.get('lbl_match', existing.get('var_match', val_str))
                         if existing else val_str)
                note_def = (existing.get('notes', '') if existing else '')
                fields = [
                    ('Name  (plain = label · var:name = variable · empty/IGNORE = skip)', n_def),
                    ('Type', t_def, _TYPE_OPTIONS),
                    ('Match  (label: exact text · var: regexp · F4 cycles presets)', m_def, None, _MATCH_PRESETS),
                    ('Notes  (written to session log — optional)', note_def),
                ]
                row_label = mode.title()

            legend_tag = ''
            if col_titles and idx < len(col_titles) and col_titles[idx]:
                legend_tag = f'  [bold dim]{_trunc(col_titles[idx], 20)}[/bold dim]'
            title = (
                f'[bold magenta]{row_label}  col {idx + 1}/{n_total}'
                f'[/bold magenta]  [dim]{ltr}{legend_tag}: {val_disp}[/dim]'
            )

            def _on_col(values, _idx=idx, _val_str=val_str, _ltr=ltr):
                if values is None:
                    on_done(None)
                    return
                raw_name   = values[0].strip()
                type_val   = values[1].strip() or 'string'
                match_val  = values[2].strip()
                notes_val  = values[3].strip()

                if mode == 'DATA':
                    if not raw_name or raw_name.upper() == 'IGNORE':
                        results[_idx] = {
                            'role': 'ignore', 'var_name': 'IGNORE',
                            'var_type': type_val, 'var_match': match_val or '.*',
                            'notes': notes_val,
                        }
                    elif raw_name.lower().startswith('lbl:'):
                        ln = raw_name[4:].strip() or f'col_{_ltr.lower()}_label'
                        results[_idx] = {
                            'role': 'label', 'lbl_name': ln,
                            'lbl_type': type_val,
                            'lbl_match': match_val or _val_str,
                            'notes': notes_val,
                        }
                    else:
                        results[_idx] = {
                            'role': 'var', 'var_name': raw_name,
                            'var_type': type_val,
                            'var_match': match_val or '.*',
                            'notes': notes_val,
                        }
                else:  # HEADER / FOOTER
                    pfx2 = 'col' if mode == 'HEADER' else 'foot'
                    slug_fb = _slugify(_val_str) if _val_str else _ltr.lower()
                    if not raw_name or raw_name.upper() == 'IGNORE':
                        results[_idx] = {'role': 'ignore'}
                    elif raw_name.lower().startswith('var:'):
                        vn = raw_name[4:].strip() or f'{_ltr.lower()}'
                        # Auto-correct match if user left cell text (wrong for variables)
                        if match_val == _val_str:
                            match_val = '.*'
                        results[_idx] = {
                            'role': 'var', 'var_name': vn,
                            'var_type': type_val,
                            'var_match': match_val or '.*',
                            'notes': notes_val,
                        }
                    else:
                        results[_idx] = {
                            'role': 'label',
                            'lbl_name':  raw_name or f'{pfx2}_{slug_fb}_label',
                            'lbl_type':  type_val,
                            'lbl_match': match_val or _val_str,
                            'notes':     notes_val,
                        }
                self._col_modal_seq(col_infos, _idx + 1, existing_cols,
                                    mode, ref_row, results, on_done,
                                    h_rows=h_rows, d_rows=d_rows,
                                    col_titles=col_titles)

            self.push_screen(_FieldsModal(title, fields), _on_col)

        # ── Highlights ────────────────────────────────────────────────────────

        def _clear_highlights(self) -> None:
            if not self._highlighted:
                return
            table = self.query_one('#sheet', DataTable)
            for ref in self._highlighted:
                parsed = _parse_cell_ref(ref)
                if not parsed:
                    continue
                r, c = parsed
                val  = self._ws.cell(row=r, column=c).value
                meta = self._choices.get(ref)
                disp = _styled(val, meta['choice']) if meta else ('' if val is None else str(val)[:20])
                try:
                    table.update_cell(str(r), str(c), disp, update_width=False)
                except Exception:
                    pass
            self._highlighted.clear()

        async def action_highlight_pending(self) -> None:
            if self._highlighted:           # toggle off
                self._log('HIGHLIGHT', f'cleared ({len(self._highlighted)} cells)')
                self._clear_highlights()
                return
            table = self.query_one('#sheet', DataTable)
            count = 0
            for r, c in self._cells:
                ref = _cell_ref(r, c)
                if ref in self._choices:    # already classified
                    continue
                val = self._ws.cell(row=r, column=c).value
                if val is None:             # empty cell
                    continue
                disp = RichText(str(val)[:20], style=_STYLE['PENDING'])
                try:
                    table.update_cell(str(r), str(c), disp, update_width=False)
                    self._highlighted.add(ref)
                    count += 1
                except Exception:
                    pass
            if count:
                self._log('HIGHLIGHT', f'{count} unclassified cells highlighted')
                self.notify(
                    f'{count} unclassified cells highlighted.  '
                    'Navigate or press H to clear.',
                    timeout=4,
                )
            else:
                self._log('HIGHLIGHT', 'all cells classified — nothing to highlight')
                self.notify('All non-empty cells are classified!', timeout=2)

        # ── Navigation ────────────────────────────────────────────────────────

        def _scan_idx(self) -> int:
            try:
                return self._cells.index((self._ws_row, self._ws_col))
            except ValueError:
                return -1

        def _advance(self) -> None:
            self._clear_highlights()
            nxt = _find_next_nonempty(self._cells, self._ws, self._scan_idx() + 1)
            if nxt is not None:
                r, c = self._cells[nxt]
                self._ws_row, self._ws_col = r, c
                self._move_cursor(r, c)
            else:
                self.notify('No more non-empty cells.', timeout=2)

        def _advance_adjacent(self) -> None:
            """Move exactly one cell in the scan direction (LR: right, TD: down).
            Used after L in template mode so the cursor lands on the adjacent value
            cell even when it is empty. Falls back to _advance() at sheet boundaries.
            """
            self._clear_highlights()
            r, c = self._ws_row, self._ws_col
            if self._state.direction == 'LR':
                nc = c + 1
                if nc <= self._ws.max_column:
                    self._ws_row, self._ws_col = r, nc
                    self._move_cursor(r, nc)
                    return
            else:
                nr = r + 1
                if nr <= self._ws.max_row:
                    self._ws_row, self._ws_col = nr, c
                    self._move_cursor(nr, c)
                    return
            self._advance()

        async def action_nav_next(self) -> None:
            self._clear_highlights()
            self._log('NAV-NEXT', f'from {_cell_ref(self._ws_row, self._ws_col)}')
            self._advance()

        async def action_nav_prev(self) -> None:
            self._clear_highlights()
            idx = self._scan_idx()
            prv = _find_prev_nonempty(self._cells, self._ws, idx - 1)
            if prv is not None:
                r, c = self._cells[prv]
                self._ws_row, self._ws_col = r, c
                self._move_cursor(r, c)
                self._log('NAV-PREV', f'→ {_cell_ref(r, c)}')
            else:
                self.notify('No previous non-empty cell.', timeout=2)
                self._log('NAV-PREV', 'no previous non-empty cell')

        async def action_nav_unclassified(self) -> None:
            """Jump to the next unclassified non-empty cell, wrapping around."""
            self._clear_highlights()
            start = self._scan_idx() + 1
            # Forward pass
            for i in range(start, len(self._cells)):
                r, c = self._cells[i]
                if self._ws.cell(row=r, column=c).value is not None:
                    if _cell_ref(r, c) not in self._choices:
                        self._ws_row, self._ws_col = r, c
                        self._move_cursor(r, c)
                        self._log('NAV-UNCLASSIFIED', f'→ {_cell_ref(r, c)}')
                        return
            # Wrap to beginning
            for i in range(0, start):
                r, c = self._cells[i]
                if self._ws.cell(row=r, column=c).value is not None:
                    if _cell_ref(r, c) not in self._choices:
                        self._ws_row, self._ws_col = r, c
                        self._move_cursor(r, c)
                        self.notify('Wrapped to first unclassified cell.', timeout=2)
                        self._log('NAV-UNCLASSIFIED', f'→ {_cell_ref(r, c)}  (wrapped)')
                        return
            self.notify('All non-empty cells are classified!', timeout=2)
            self._log('NAV-UNCLASSIFIED', 'all cells classified')

        async def action_nav_goto(self) -> None:
            def _on_ref(ref: str | None) -> None:
                self._clear_highlights()
                if not ref:
                    return
                parsed = _parse_cell_ref(ref)
                if parsed:
                    wr, wc = parsed
                    if 1 <= wr <= self._max_row and 1 <= wc <= self._max_col:
                        self._ws_row, self._ws_col = wr, wc
                        self._move_cursor(wr, wc)
                        self._log('GOTO', f'→ {ref}')
                    else:
                        self.notify(f'Cell {ref} is out of range.', timeout=2)
                        self._log('GOTO-X', f'{ref}  out of range')
                else:
                    self.notify(f'Invalid reference: {ref}', timeout=2)
                    self._log('GOTO-X', f'"{ref}"  invalid reference')
            self.push_screen(_GotoModal(), _on_ref)

        # ── Classification helpers ─────────────────────────────────────────────

        def _update_prefix(self, name: str) -> None:
            """Extract dot-notation prefix from a variable name and store it."""
            if '.' in name:
                prefix_part, _ = name.rsplit('.', 1)
                self._current_prefix = prefix_part + '.'
            else:
                self._current_prefix = ''

        def _push_undo(self, ref: str) -> None:
            self._undo_stack.append({
                'ref':             ref,
                'prev_choice':     self._choices.get(ref),  # dict | None
                'prev_label_base': self._last_label_base,
            })

        def _commit(self, ref: str, meta: dict) -> None:
            self._choices[ref] = meta
            parsed = _parse_cell_ref(ref)
            if parsed:
                self._restyle_cell(parsed[0], parsed[1])

        # ── Classification actions ─────────────────────────────────────────────

        async def action_accept(self) -> None:
            """Silently commit the proposed classification with defaults — no modal."""
            p     = self._proposal()
            value = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            ref   = _cell_ref(self._ws_row, self._ws_col)

            if p == 'label':
                slug  = _slugify(str(value)) if value is not None else 'label'
                name  = f'{slug}_label'
                match = str(value) if value is not None else ''
                ltype = _infer_cell_type(self._ws.cell(row=self._ws_row, column=self._ws_col))
                old   = self._choices.get(ref, {}).get('choice')
                self._push_undo(ref)
                base  = name[:-6] if name.endswith('_label') else name
                self._last_label_base = base
                self._commit(ref, {'choice': 'L', 'name': name,
                                   'ltype': ltype, 'lmatch': match})
                rawval = '' if value is None else f'  "{str(value)[:30]}"'
                reclassify = f'  (was {old})' if old else ''
                self._log('LABEL',
                          f'{ref}{rawval}  →  {name}  [{ltype}, {match}]{reclassify}  [auto]')
                self._refresh_panel()
                if self._is_template:
                    self._advance_adjacent()
                else:
                    self._advance()
            elif p.startswith('var:') or (p == 'skip' and self._is_template
                                          and self._last_label_base):
                slug  = _slugify(str(value)) if value is not None else 'field'
                base  = self._last_label_base or slug
                name  = (self._current_prefix + base) if self._current_prefix else base
                ftype = (p[4:] if p.startswith('var:')
                         else _infer_cell_type(
                             self._ws.cell(row=self._ws_row, column=self._ws_col)))
                old   = self._choices.get(ref, {}).get('choice')
                self._push_undo(ref)
                self._last_label_base = None
                self._update_prefix(name)
                self._commit(ref, {'choice': 'V', 'name': name,
                                   'ftype': ftype, 'match': '.*'})
                rawval = '(empty)' if value is None else f'"{str(value)[:30]}"'
                reclassify = f'  (was {old})' if old else ''
                self._log('VALUE',
                          f'{ref}  {rawval}  →  {name}  [{ftype}, .*]{reclassify}  [auto]')
                self._refresh_panel()
                self._advance()
            else:
                self._advance()

        async def action_act_L(self) -> None:
            value = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            ref   = _cell_ref(self._ws_row, self._ws_col)
            slug  = _slugify(str(value)) if value is not None else 'label'
            default_name  = f'{slug}_label'
            default_match = str(value) if value is not None else ''
            existing_meta = self._choices.get(ref, {})
            existing_note = self._notes.get(ref, '')

            def _done(result: list[str] | None) -> None:
                self._clear_highlights()
                if result is None:
                    self._log('LABEL-X', f'{ref}  cancelled')
                    return
                name, ltype, lmatch, notes = (result[0], result[1],
                                              result[2], result[3].strip())
                old  = self._choices.get(ref, {}).get('choice')
                self._push_undo(ref)
                base = name[:-6] if name.endswith('_label') else name
                self._last_label_base = base
                self._commit(ref, {'choice': 'L', 'name': name,
                                   'ltype': ltype, 'lmatch': lmatch})
                if notes:
                    self._notes[ref] = notes
                    self._log('NOTE', f'{ref}: {notes}')
                elif ref in self._notes:
                    del self._notes[ref]
                reclassify = f'  (was {old})' if old else ''
                rawval = '' if value is None else f'  "{str(value)[:30]}"'
                self._log('LABEL',
                          f'{ref}{rawval}  →  {name}  [{ltype}, {lmatch}]{reclassify}')
                self._refresh_panel()
                if self._is_template:
                    self._advance_adjacent()
                else:
                    self._advance()

            self.push_screen(
                _FieldsModal(
                    '[bold green]Label[/bold green] — text that identifies a nearby value',
                    [('Label anchor name', existing_meta.get('name', default_name)),
                     ('Type', existing_meta.get('ltype', _infer_cell_type(self._ws.cell(row=self._ws_row, column=self._ws_col))), _TYPE_OPTIONS),
                     ('Match  (exact text · lbl:regexp for regex · F4 cycles presets)', existing_meta.get('lmatch', default_match), None, _MATCH_PRESETS),
                     ('Notes  (written to session log — optional)', existing_note)],
                ),
                _done,
            )

        async def action_act_C(self) -> None:
            value = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            ref   = _cell_ref(self._ws_row, self._ws_col)
            slug  = _slugify(str(value)) if value is not None else 'header'

            def _done(result: list[str] | None) -> None:
                self._clear_highlights()
                if result is None:
                    self._log('HEADER-X', f'{ref}  cancelled')
                    return
                name = result[0]
                old  = self._choices.get(ref, {}).get('choice')
                self._push_undo(ref)
                self._last_label_base = None
                self._commit(ref, {'choice': 'C', 'name': name})
                reclassify = f'  (was {old})' if old else ''
                rawval = '' if value is None else f'  "{str(value)[:30]}"'
                self._log('HEADER', f'{ref}{rawval}  →  {name}{reclassify}')
                self._refresh_panel()
                self._advance()

            self.push_screen(
                _FieldsModal(
                    '[bold blue]Header[/bold blue] — section title, no value follows',
                    [('Header name', slug)],
                ),
                _done,
            )

        async def action_act_V(self) -> None:
            value        = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            ref          = _cell_ref(self._ws_row, self._ws_col)
            slug         = _slugify(str(value)) if value is not None else 'field'
            base         = self._last_label_base or slug
            default_name = (self._current_prefix + base) if self._current_prefix else base
            p            = self._proposal()
            default_type = (p[4:] if p.startswith('var:') else
                            _infer_cell_type(self._ws.cell(row=self._ws_row, column=self._ws_col)))

            existing_meta = self._choices.get(ref, {})
            existing_note = self._notes.get(ref, '')

            def _done(result: list[str] | None) -> None:
                self._clear_highlights()
                if result is None:
                    self._log('VALUE-X', f'{ref}  cancelled')
                    return
                name, ftype, match, notes = result[0], result[1], result[2], result[3].strip()
                old  = self._choices.get(ref, {}).get('choice')
                self._push_undo(ref)
                self._last_label_base = None
                self._update_prefix(name)
                self._commit(ref, {'choice': 'V', 'name': name,
                                   'ftype': ftype, 'match': match})
                if notes:
                    self._notes[ref] = notes
                    self._log('NOTE', f'{ref}: {notes}')
                elif ref in self._notes:
                    del self._notes[ref]
                reclassify = f'  (was {old})' if old else ''
                rawval = '(empty)' if value is None else f'"{str(value)[:30]}"'
                self._log('VALUE',
                          f'{ref}  {rawval}  →  {name}  [{ftype}, {match}]{reclassify}')
                self._refresh_panel()
                self._advance()

            self.push_screen(
                _FieldsModal(
                    '[bold bright_yellow]Value[/bold bright_yellow] — extract this cell\'s content',
                    [('Field name', existing_meta.get('name', default_name)),
                     ('Type',       existing_meta.get('ftype', default_type), _TYPE_OPTIONS),
                     ('Match', existing_meta.get('match', '.*'), None, _MATCH_PRESETS),
                     ('Notes  (written to session log — optional)', existing_note)],
                ),
                _done,
            )

        async def action_act_T(self) -> None:
            """Table definition or edit — multi-step flow:

            Step 1  _TableSetupModal   name / full range / mult
            Step 2  _TableRangeModal   classify each row H / D / F / S
            Step 3  per-column modals  one _FieldsModal per column:
                      • For each HEADER row: label name + type + match + notes
                      • For DATA rows (once): var name + type + match + notes
                      • For each FOOTER row: label name + type + match + notes
            Commit  write T + T-HEAD into _choices

            Pressing T on an already-classified T or T-HEAD cell re-opens
            the flow with all fields pre-filled from the existing definition.
            """
            from openpyxl.utils import get_column_letter as _gcl

            cur_ref  = _cell_ref(self._ws_row, self._ws_col)
            cur_meta = self._choices.get(cur_ref, {})
            cur_ch   = cur_meta.get('choice', '')

            # ── Detect edit vs. new ───────────────────────────────────────────
            existing_anchor: str | None = None
            existing_meta:   dict | None = None
            if cur_ch == 'T':
                existing_anchor = cur_ref
                existing_meta   = cur_meta
            elif cur_ch in ('T-HEAD', 'T-DATA'):
                existing_anchor = cur_meta.get('anchor', cur_ref)
                existing_meta   = self._choices.get(existing_anchor)

            cur_val  = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            slug     = _slugify(str(cur_val)) if cur_val is not None else 'table'
            end_col  = self._detect_table_end_col(self._ws_row, self._ws_col)
            auto_rng = (f'{cur_ref}:{_cell_ref(self._ws_row, end_col)}')

            setup_name = existing_meta['name']  if existing_meta else slug
            setup_rng  = existing_meta['range'] if existing_meta else auto_rng
            setup_mult = existing_meta['mult']  if existing_meta else '*'

            def _on_setup(setup: dict | None) -> None:
                self._clear_highlights()
                if setup is None:
                    self._log('TABLE-X', f'{cur_ref}  cancelled at setup')
                    return
                start_row, start_col = setup['start']
                end_row,   end_col2  = setup['end']
                name = setup['name']
                mult = setup['mult']
                rng  = setup['range_str']
                col_count = end_col2 - start_col + 1
                col_infos = [
                    {'col': start_col + i, 'letter': _gcl(start_col + i)}
                    for i in range(col_count)
                ]

                def _on_rows(row_types: dict | None) -> None:
                    self._clear_highlights()
                    if row_types is None:
                        self._log('TABLE-X', f'{cur_ref}  cancelled at row classification')
                        return

                    h_rows    = sorted(r for r, t in row_types.items() if t == 'H')
                    d_rows    = sorted(r for r, t in row_types.items() if t == 'D')
                    f_rows    = sorted(r for r, t in row_types.items() if t == 'F')
                    skip_rows = sorted(r for r, t in row_types.items() if t == 'S')

                    # Column legend titles — set by the legend step before col modals
                    col_titles: list[str] = [''] * col_count

                    # State collectors (filled progressively by column modals)
                    h_results: list[list | None] = [None] * len(h_rows)
                    d_results: list[dict | None] = [None] * col_count
                    f_results: list[list | None] = [None] * len(f_rows)

                    # ── Helpers to extract pre-fill data from existing_meta ────
                    def _existing_h_cols(row_idx: int) -> 'list | None':
                        if not existing_meta:
                            return None
                        hrows = existing_meta.get('header_rows', [])
                        return hrows[row_idx]['cols'] if row_idx < len(hrows) else None

                    def _existing_f_cols(row_idx: int) -> 'list | None':
                        if not existing_meta:
                            return None
                        frows = existing_meta.get('footer_rows', [])
                        return frows[row_idx]['cols'] if row_idx < len(frows) else None

                    def _existing_d_cols() -> 'list | None':
                        if not existing_meta:
                            return None
                        dvs = existing_meta.get('data_vars', [])
                        if not dvs:
                            return None
                        result = []
                        for item in dvs:
                            if isinstance(item, dict):
                                result.append(item)
                            else:
                                result.append({
                                    'var_name': item, 'var_type': 'string',
                                    'var_match': '.*', 'notes': '',
                                })
                        return result

                    # ── Commit ────────────────────────────────────────────────
                    def _commit_all() -> None:
                        header_rows = []
                        for ri, r in enumerate(h_rows):
                            hcols = []
                            row_def = h_results[ri] or []
                            for ci, cinfo in enumerate(col_infos):
                                val = self._ws.cell(row=r, column=cinfo['col']).value
                                val_s = str(val) if val is not None else ''
                                slug_v = _slugify(val_s) if val_s else cinfo['letter'].lower()
                                cd = row_def[ci] if ci < len(row_def) else {}
                                hcols.append({
                                    'ref':       _cell_ref(r, cinfo['col']),
                                    'row':       r, 'col': cinfo['col'],
                                    'cell_value': val_s,
                                    'role':      cd.get('role', 'label'),
                                    'lbl_name':  cd.get('lbl_name', f'col_{slug_v}_label'),
                                    'lbl_type':  cd.get('lbl_type', 'string'),
                                    'lbl_match': cd.get('lbl_match', val_s),
                                    'var_name':  cd.get('var_name', 'IGNORE'),
                                    'var_type':  cd.get('var_type', 'string'),
                                    'var_match': cd.get('var_match', '.*'),
                                    'notes':     cd.get('notes', ''),
                                })
                            header_rows.append({'row': r, 'cols': hcols})

                        footer_rows = []
                        for ri, r in enumerate(f_rows):
                            fcols = []
                            row_def = f_results[ri] or []
                            for ci, cinfo in enumerate(col_infos):
                                val = self._ws.cell(row=r, column=cinfo['col']).value
                                val_s = str(val) if val is not None else ''
                                slug_v = _slugify(val_s) if val_s else cinfo['letter'].lower()
                                cd = row_def[ci] if ci < len(row_def) else {}
                                fcols.append({
                                    'ref':       _cell_ref(r, cinfo['col']),
                                    'row':       r, 'col': cinfo['col'],
                                    'cell_value': val_s,
                                    'role':      cd.get('role', 'label'),
                                    'lbl_name':  cd.get('lbl_name', f'foot_{slug_v}_label'),
                                    'lbl_type':  cd.get('lbl_type', 'string'),
                                    'lbl_match': cd.get('lbl_match', val_s),
                                    'var_name':  cd.get('var_name', 'IGNORE'),
                                    'var_type':  cd.get('var_type', 'string'),
                                    'var_match': cd.get('var_match', '.*'),
                                    'notes':     cd.get('notes', ''),
                                })
                            footer_rows.append({'row': r, 'cols': fcols})

                        data_vars = []
                        for ci, cinfo in enumerate(col_infos):
                            cd = d_results[ci] or {}
                            vn = cd.get('var_name', 'IGNORE') or 'IGNORE'
                            data_vars.append({
                                'role':      cd.get('role', 'var'),
                                'var_name':  vn,
                                'var_type':  cd.get('var_type',  'string') or 'string',
                                'var_match': cd.get('var_match', '.*')    or '.*',
                                'lbl_name':  cd.get('lbl_name',  'IGNORE'),
                                'lbl_type':  cd.get('lbl_type',  'string'),
                                'lbl_match': cd.get('lbl_match', '.*'),
                                'notes':     cd.get('notes',     ''),
                            })
                            note = cd.get('notes', '').strip()
                            if note:
                                d_ref = _cell_ref(
                                    d_rows[0] if d_rows else start_row, cinfo['col']
                                )
                                self._notes[d_ref] = note
                                self._log('NOTE', f'{d_ref}: {note}  (DATA col {cinfo["letter"]})')

                        # Remove old table if editing
                        if existing_meta and existing_anchor:
                            self._remove_table(existing_anchor, existing_meta)

                        actual_anchor = _cell_ref(start_row, start_col)
                        all_head_cells = (
                            [col for hr in header_rows for col in hr['cols']] +
                            [col for fr in footer_rows for col in fr['cols']]
                        )
                        # DATA row cells (T-DATA) — all cells in d_rows
                        data_cells = [
                            {'ref': _cell_ref(r, cinfo['col']),
                             'row': r, 'col': cinfo['col']}
                            for r in d_rows for cinfo in col_infos
                        ]

                        # Undo snapshot: anchor + T-HEAD + T-DATA
                        affected = [{'ref': actual_anchor,
                                     'prev': self._choices.get(actual_anchor)}]
                        for hc in all_head_cells:
                            if hc['ref'] != actual_anchor:
                                affected.append({'ref':  hc['ref'],
                                                 'prev': self._choices.get(hc['ref'])})
                        for dc in data_cells:
                            if dc['ref'] != actual_anchor:
                                affected.append({'ref':  dc['ref'],
                                                 'prev': self._choices.get(dc['ref'])})
                        self._undo_stack.append({
                            'type':  'table', 'cells': affected,
                            'prev_label_base': self._last_label_base,
                        })
                        self._last_label_base = None

                        meta = {
                            'choice':      'T',
                            'name':        name,
                            'mult':        mult,
                            'range':       rng,
                            'start_row':   start_row, 'end_row':   end_row,
                            'start_col':   start_col, 'end_col':   end_col2,
                            'header_rows': header_rows,
                            'footer_rows': footer_rows,
                            'data_vars':   data_vars,
                            'skip_rows':   skip_rows,
                            'row_types':   row_types,
                        }
                        self._choices[actual_anchor] = meta
                        self._restyle_cell(start_row, start_col)
                        # Mark header/footer cells T-HEAD
                        for hc in all_head_cells:
                            if hc['ref'] != actual_anchor:
                                self._choices[hc['ref']] = {
                                    'choice': 'T-HEAD', 'anchor': actual_anchor,
                                }
                                self._restyle_cell(hc['row'], hc['col'])
                        # Mark data row cells T-DATA
                        for dc in data_cells:
                            if dc['ref'] != actual_anchor:
                                self._choices[dc['ref']] = {
                                    'choice': 'T-DATA', 'anchor': actual_anchor,
                                }
                                self._restyle_cell(dc['row'], dc['col'])

                        n_h = len(h_rows); n_f = len(f_rows)
                        n_d = len(d_rows); n_s = len(skip_rows)
                        n_p = sum(1 for t in row_types.values() if t == 'P')
                        var_names = [dv['var_name'] for dv in data_vars]
                        self._log(
                            'TABLE',
                            f'{rng}  name={name}  mult={mult}  '
                            f'H={n_h} D={n_d} F={n_f} S={n_s} P={n_p}  '
                            f'vars=[{", ".join(var_names)}]',
                        )
                        self._refresh_panel()
                        self._advance()

                    # ── Footer rows sequence ───────────────────────────────────
                    def _run_f_row(row_idx: int) -> None:
                        if row_idx >= len(f_rows):
                            _commit_all()
                            return
                        r = f_rows[row_idx]
                        ex_f = _existing_f_cols(row_idx)
                        res_f: list[dict | None] = [None] * col_count

                        def _on_f_done(results, _ri=row_idx) -> None:
                            if results is None:
                                self._log('TABLE-X',
                                          f'{cur_ref}  cancelled at FOOTER col def row {_ri+1}')
                                return
                            f_results[_ri] = results
                            _run_f_row(_ri + 1)

                        self._col_modal_seq(col_infos, 0, ex_f, 'FOOTER', r, res_f, _on_f_done,
                                            col_titles=col_titles)

                    # ── DATA column sequence ───────────────────────────────────
                    def _run_d() -> None:
                        if not d_rows:
                            _run_f_row(0)
                            return
                        ref_r = d_rows[0]
                        ex_d  = _existing_d_cols()
                        res_d: list[dict | None] = [None] * col_count

                        def _on_d_done(results) -> None:
                            if results is None:
                                self._log('TABLE-X', f'{cur_ref}  cancelled at DATA col def')
                                return
                            for i in range(col_count):
                                d_results[i] = results[i]
                            _run_f_row(0)

                        self._col_modal_seq(col_infos, 0, ex_d, 'DATA', ref_r, res_d, _on_d_done,
                                            h_rows=h_rows, d_rows=d_rows, col_titles=col_titles)

                    # ── Header rows sequence ───────────────────────────────────
                    def _run_h_row(row_idx: int) -> None:
                        if row_idx >= len(h_rows):
                            _run_d()
                            return
                        r    = h_rows[row_idx]
                        ex_h = _existing_h_cols(row_idx)
                        res_h: list[dict | None] = [None] * col_count

                        def _on_h_done(results, _ri=row_idx) -> None:
                            if results is None:
                                self._log('TABLE-X',
                                          f'{cur_ref}  cancelled at HEADER col def row {_ri+1}')
                                return
                            h_results[_ri] = results
                            _run_h_row(_ri + 1)

                        self._col_modal_seq(col_infos, 0, ex_h, 'HEADER', r, res_h, _on_h_done,
                                            col_titles=col_titles)

                    # ── Column legend (Step 2.5) — names each column once ──────
                    def _run_col_legend() -> None:
                        # Pre-fill from first header row values (best source of col names)
                        pre: list[str] = [''] * col_count
                        src_row = h_rows[0] if h_rows else (d_rows[0] if d_rows else None)
                        if src_row:
                            for ci, cinfo in enumerate(col_infos):
                                v = self._ws.cell(row=src_row, column=cinfo['col']).value
                                pre[ci] = str(v) if v is not None else ''

                        fields = [
                            (f'Column {cinfo["letter"]}', pre[ci])
                            for ci, cinfo in enumerate(col_infos)
                        ]

                        def _on_legend(result: 'list[str] | None') -> None:
                            nonlocal col_titles
                            if result is None:
                                self._log('TABLE-X', f'{cur_ref}  cancelled at column legend')
                                return
                            col_titles = result
                            _run_h_row(0)

                        self.push_screen(
                            _FieldsModal(
                                f'[bold magenta]Column legend[/bold magenta] — '
                                f'give each column a short title (used as hints in next steps)',
                                fields,
                            ),
                            _on_legend,
                        )

                    # ── Start the column-definition chain ─────────────────────
                    _run_col_legend()

                self.push_screen(
                    _TableRangeModal(
                        self._ws,
                        start_row, end_row,
                        start_col, end_col2,
                        name, mult,
                        existing_row_types=existing_meta.get('row_types') if existing_meta else None,
                    ),
                    _on_rows,
                )

            self.push_screen(
                _TableSetupModal(setup_name, setup_rng, setup_mult),
                _on_setup,
            )

        async def action_act_I(self) -> None:
            ref   = _cell_ref(self._ws_row, self._ws_col)
            old   = self._choices.get(ref, {}).get('choice')
            val   = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            rawval = '(empty)' if val is None else f'"{str(val)[:30]}"'
            self._push_undo(ref)
            self._last_label_base = None
            self._clear_highlights()
            self._commit(ref, {'choice': 'I'})
            reclassify = f'  (was {old})' if old else ''
            self._log('IGNORE', f'{ref}  {rawval}{reclassify}')
            self._refresh_panel()
            self._advance()

        async def action_act_R(self) -> None:
            """Remove classification.  For T/T-HEAD cells removes the whole table."""
            ref  = _cell_ref(self._ws_row, self._ws_col)
            meta = self._choices.get(ref)
            if not meta:
                self.notify('Cell is not classified yet.', timeout=2)
                return
            choice = meta.get('choice', '?')

            if choice in ('T', 'T-HEAD', 'T-DATA'):
                # Resolve the anchor
                if choice in ('T-HEAD', 'T-DATA'):
                    anchor_ref  = meta.get('anchor', ref)
                    anchor_meta = self._choices.get(anchor_ref, {})
                else:
                    anchor_ref  = ref
                    anchor_meta = meta
                tname = anchor_meta.get('name', '')
                trng  = anchor_meta.get('range', anchor_ref)
                # Snapshot all cells for undo (T-HEAD + T-DATA + legacy)
                all_members = {anchor_ref}
                for cref, cm in self._choices.items():
                    if cm.get('anchor') == anchor_ref:
                        all_members.add(cref)
                for cref in self._all_thead_refs(anchor_meta):
                    all_members.add(cref)
                affected = [{'ref': ar, 'prev': self._choices.get(ar)}
                            for ar in all_members]
                self._undo_stack.append({
                    'type':            'table',
                    'cells':           affected,
                    'prev_label_base': self._last_label_base,
                })
                self._remove_table(anchor_ref, anchor_meta)
                self._log('REMOVE',
                          f'table "{tname}" ({trng}) — all {len(affected)} cells cleared')
                self._refresh_panel()
                self.notify(f'Table "{tname}" removed ({trng}).', timeout=3)
            else:
                old_name = meta.get('name', '')
                self._push_undo(ref)
                del self._choices[ref]
                self._restyle_cell(self._ws_row, self._ws_col)
                self._log('REMOVE',
                          f'{ref}  cleared  was: {_CHOICE_NAME.get(choice, choice)}'
                          + (f' "{old_name}"' if old_name else ''))
                self._refresh_panel()
                self.notify(f'{ref} cleared — choose a new type.', timeout=2)

        # ── Undo ──────────────────────────────────────────────────────────────

        async def action_undo(self) -> None:
            if not self._undo_stack:
                self.notify('Nothing to undo.', timeout=2)
                return
            entry = self._undo_stack.pop()
            self._last_label_base = entry.get('prev_label_base')

            if entry.get('type') == 'table':
                # Multi-cell table undo
                for cell_entry in entry['cells']:
                    cref = cell_entry['ref']
                    prev = cell_entry['prev']
                    if prev is None:
                        self._choices.pop(cref, None)
                    else:
                        self._choices[cref] = prev
                    p = _parse_cell_ref(cref)
                    if p:
                        self._restyle_cell(p[0], p[1])
                first_ref = entry['cells'][0]['ref'] if entry['cells'] else None
                if first_ref:
                    p = _parse_cell_ref(first_ref)
                    if p:
                        self._ws_row, self._ws_col = p[0], p[1]
                        self._move_cursor(p[0], p[1])
                self._log('UNDO', f'table  {len(entry["cells"])} cells restored')
                self.notify(f'Table undo: {len(entry["cells"])} cells restored', timeout=2)
            else:
                ref     = entry['ref']
                prev    = entry['prev_choice']
                curr_ch = self._choices.get(ref, {}).get('choice', 'unclassified')
                if prev is None:
                    self._choices.pop(ref, None)
                    restored = 'unclassified'
                else:
                    self._choices[ref] = prev
                    restored = _CHOICE_NAME.get(prev.get('choice', ''), prev.get('choice', '?'))
                self._log('UNDO',
                          f'{ref}  {_CHOICE_NAME.get(curr_ch, curr_ch)}  →  {restored}')
                p = _parse_cell_ref(ref)
                if p:
                    self._ws_row, self._ws_col = p[0], p[1]
                    self._move_cursor(p[0], p[1])
                    self._restyle_cell(p[0], p[1])
                self.notify(f'Undone: {ref}', timeout=2)
            self._refresh_panel()

        # ── Other actions ──────────────────────────────────────────────────────

        async def action_zoom(self) -> None:
            ref   = _cell_ref(self._ws_row, self._ws_col)
            value = self._ws.cell(row=self._ws_row, column=self._ws_col).value
            meta  = self._choices.get(ref)
            self._log('ZOOM', f'{ref}  value={value!r}')
            self.push_screen(_ZoomModal(ref, value, meta), lambda _: None)

        async def action_preview(self) -> None:
            self._log('PREVIEW', f'{len(self._choices)} cells classified so far')
            csv_text = _choices_to_csv(
                self._ws, self._choices, self._cells,
                self._state.direction, self._state.sheet_name,
            )
            self.push_screen(_PreviewModal(csv_text), lambda _: None)

        async def action_show_help(self) -> None:
            self._log('HELP', 'opened help modal')
            self.push_screen(_HelpModal(), lambda _: None)

        async def action_show_log(self) -> None:
            """F12: hidden session event log for debugging and support reports."""
            self._log('LOG-OPEN', f'{len(self._event_log)} events so far')
            def _on_close(log_path: str | None) -> None:
                if log_path:
                    self.notify(f'Log written: {log_path}', timeout=6)
            self.push_screen(_LogModal(list(self._event_log), self._data_file), _on_close)

        # ── Exit ──────────────────────────────────────────────────────────────

        async def action_end_save(self) -> None:
            done  = len(self._choices)
            total = self._total_nonempty
            state = _build_state_from_choices(
                self._ws, self._choices, self._cells,
                self._state.direction, self._state.sheet_name,
            )

            # Detect duplicate lbl: names
            seen: set[str] = set()
            duplicates: list[str] = []
            for name, _, _ in state.lbl_defs:
                if name in seen and name not in duplicates:
                    duplicates.append(name)
                seen.add(name)

            def _do_save() -> None:
                self._log('END-SAVE',
                          f'{done}/{total} classified  →  saving pattern')
                if self._notes:
                    for ref, note in sorted(self._notes.items()):
                        self._log('NOTE-SAVED', f'{ref}: {note}')
                _save_history(self._data_file,
                              {r: m['choice'] for r, m in self._choices.items()})
                self._close_log_file()
                self.exit(result=state)

            if duplicates:
                dup_str = ', '.join(f'"{n}"' for n in duplicates[:4])
                if len(duplicates) > 4:
                    dup_str += f' and {len(duplicates) - 4} more'
                msg = (
                    f'[bold yellow]⚠  Duplicate label names[/bold yellow]\n\n'
                    f'{dup_str}\n\n'
                    '[dim]Labels with the same name may match the wrong cell '
                    'during extraction.\nRename the duplicates to fix this.[/dim]'
                )
                self.push_screen(_ConfirmModal(msg), lambda ok: _do_save() if ok else None)
            else:
                _do_save()

        async def action_cancel(self) -> None:
            self._log('CANCEL', 'user cancelled — no pattern saved')
            _save_history(self._data_file,
                          {r: m['choice'] for r, m in self._choices.items()})
            self._close_log_file()
            self.exit(result=None)

        async def action_add_note(self) -> None:
            """F2: add/edit an internal debug note for the current cell."""
            ref     = _cell_ref(self._ws_row, self._ws_col)
            current = self._notes.get(ref, '')

            def _on_note(result: list[str] | None) -> None:
                if result is None:
                    return
                note = result[0].strip()
                if note:
                    self._notes[ref] = note
                    self._log('NOTE', f'{ref}: {note}')
                elif ref in self._notes:
                    del self._notes[ref]
                    self._log('NOTE-DEL', f'{ref}: removed')
                self._refresh_panel()

            self.push_screen(
                _FieldsModal(
                    f'[bold yellow]Internal note[/bold yellow]  {ref}'
                    f'{"  [dim](existing)[/dim]" if current else ""}',
                    [('Note (empty = delete)', current)],
                ),
                _on_note,
            )

        async def action_add_comment(self) -> None:
            """;  Add a free-text comment into the session log."""
            ref = _cell_ref(self._ws_row, self._ws_col)

            def _on_result(result: list[str] | None) -> None:
                if not result:
                    return
                text = result[0].strip()
                if text:
                    self._log('COMMENT', f'[{ref}] {text}')
                    self.notify('Comment saved to log.', timeout=2)

            self.push_screen(
                _FieldsModal(
                    '[bold cyan]Session comment[/bold cyan]'
                    f'  [dim]{ref}[/dim]',
                    [('Comment (saved to session log)', '')],
                ),
                _on_result,
            )

        async def action_screenshot(self) -> None:
            """F11: save an SVG screenshot of the current TUI state."""
            if not self._log_dir:
                self.notify('No session log directory — screenshot unavailable.',
                            severity='warning', timeout=4)
                return
            self._screenshot_count += 1
            n        = self._screenshot_count
            filename = f'screenshot_{n:03d}.svg'
            try:
                saved = self.save_screenshot(filename=filename, path=self._log_dir)
                ref   = _cell_ref(self._ws_row, self._ws_col)
                self._log('SCREENSHOT', f'{filename}  at {ref}  ({len(self._choices)} classified)')
                self.notify(f'Screenshot saved: {filename}', timeout=4)
            except Exception as exc:
                self._screenshot_count -= 1
                self.notify(f'Screenshot failed: {exc}', severity='warning', timeout=5)
                self._log('SCREENSHOT-X', str(exc))


# ── Public entry point ─────────────────────────────────────────────────────────

def run_wizard_tui(
    data_file: str,
    sheet: str | None = None,
    output: str | None = None,
    save_state: str | None = None,
) -> int:
    """Run the TUI wizard.

    Returns
    -------
    0   pattern saved successfully
    1   user cancelled
    2   textual not installed (caller falls back to sequential wizard)

    ``save_state``
        If given, write the final WizardState to this JSON path after saving
        the pattern.  Enables replay, scripted testing, and ``--load-state``.
    """
    if not _TEXTUAL_OK:
        return 2

    try:
        import openpyxl
        wb = openpyxl.load_workbook(data_file, data_only=True)
    except Exception as exc:
        print(f'Error loading {data_file}: {exc}', file=sys.stderr)
        return 1

    if sheet:
        if sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            try:
                ws = wb.worksheets[int(sheet)]
            except (ValueError, IndexError):
                print(f'Sheet {sheet!r} not found. Available: {wb.sheetnames}',
                      file=sys.stderr)
                return 1
    elif len(wb.sheetnames) > 1:
        print(f'Multiple sheets: {wb.sheetnames}')
        print('Use  --sheet NAME  to choose; defaulting to first sheet.')
        ws = wb.worksheets[0]
    else:
        ws = wb.active

    if not output:
        stem   = os.path.splitext(os.path.basename(data_file))[0]
        output = os.path.join(
            os.path.dirname(os.path.abspath(data_file)),
            f'pattern-{stem}.csv',
        )

    state  = WizardState(sheet_name=ws.title)
    app    = WizardTUIApp(ws, state, data_file)
    result = app.run()

    if result is None:
        print('Wizard cancelled.', file=sys.stderr)
        return 1

    _write_pattern(result, output)
    print(f'\n✓  Pattern written: {output}')
    if save_state:
        from .wizard import _save_state_json
        _save_state_json(result, save_state)
    print(f'   Try:  grepxcel extract -p {output} {data_file}')
    print(f'         grepxcel validate-pattern {output}')
    return 0
