import openpyxl

wb = openpyxl.load_workbook('pattern-example-1.xlsx')
print('Sheets:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== Sheet: {sheet_name} ===')
    print(f'Max row: {ws.max_row}, Max col: {ws.max_column}')
    print()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                print(f'  [{cell.coordinate}] ({cell.data_type}) = {repr(cell.value)}')
