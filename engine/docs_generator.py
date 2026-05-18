"""
DocsGenerator — writes a self-documenting pattern-reference.xlsx.

The output file is both human-readable documentation and a valid grepxcel
pattern that the engine can parse.  Every keyword is shown in context with
colour coding:

  doc:    — yellow   (inline comment, ignored by engine)
  lbl:    — blue     (anchor label, never in output JSON)
  var:    — green    (variable, extracted to output JSON)
  config: — orange   (global settings)
  cell:   — lavender (cell extraction instruction)
  table:  — lavender (table extraction instruction)
  START:/END: — grey (section markers)
  HEADER/DATA/FOOTER — light lavender (table template rows)
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment


# Row colours (ARGB hex, no leading '#')
_FILL = {
    'doc':     PatternFill('solid', fgColor='FFFFD0'),  # yellow
    'lbl':     PatternFill('solid', fgColor='D0E8FF'),  # blue
    'var':     PatternFill('solid', fgColor='D0FFD0'),  # green
    'config':  PatternFill('solid', fgColor='FFE8D0'),  # orange
    'cell':    PatternFill('solid', fgColor='F0E0FF'),  # lavender
    'table':   PatternFill('solid', fgColor='F0E0FF'),  # lavender
    'marker':  PatternFill('solid', fgColor='E0E0E0'),  # grey  (START/END)
    'tmpl':    PatternFill('solid', fgColor='F8F0FF'),  # light lavender
}

_BOLD = Font(bold=True)


def _row(ws, row_num: int, cells: list, fill_key: str) -> None:
    fill = _FILL.get(fill_key)
    for col, val in enumerate(cells, start=1):
        c = ws.cell(row=row_num, column=col, value=val)
        if fill:
            c.fill = fill


class DocsGenerator:
    def write(self, output_path: str) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'pattern-reference'

        # Column widths (A, B, C, D, E)
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 28
        ws.column_dimensions['E'].width = 50

        r = 1

        def row(cells, fill_key=''):
            nonlocal r
            _row(ws, r, cells, fill_key)
            r += 1

        def blank():
            nonlocal r
            r += 1

        # ── Title ──────────────────────────────────────────────────────────────
        ws.cell(r, 1, 'grepxcel pattern reference').font = _BOLD
        ws.cell(r, 5, 'Column colour key:')
        ws.cell(r + 1, 5, 'doc:  — yellow  (comment, ignored)')
        ws.cell(r + 2, 5, 'lbl:  — blue    (anchor label, never in output)')
        ws.cell(r + 3, 5, 'var:  — green   (extracted variable)')
        ws.cell(r + 4, 5, 'config: — orange  (global setting)')
        ws.cell(r + 5, 5, 'cell:/table: — lavender  (instructions)')
        r += 2
        blank()

        # ── CONFIG section ─────────────────────────────────────────────────────
        row(['doc:', '', '', '', 'Config rows set global options. Must appear before lbl:/var: rows.'], 'doc')
        row(['config:', 'read.direction', 'LR', '', 'LR = left-to-right scan (default). TD = top-to-bottom.'], 'config')
        row(['config:', 'currency.sign', '€', '', 'Symbol used when parsing currency values.'], 'config')
        blank()

        # ── LBL section ────────────────────────────────────────────────────────
        row(['doc:', '', '', '', 'lbl: defines an anchor label. Matched for position only — NEVER written to output JSON.'], 'doc')
        row(['doc:', '', '', '', 'Use lbl: for literal text like "Invoice No:" or table column headers like "Product".'], 'doc')
        row(['lbl:', 'po_label',     'string',   'PO Number:',  'Matches the literal text "PO Number:" in the sheet.'], 'lbl')
        row(['lbl:', 'date_label',   'string',   'Date:',       'Matches "Date:".'], 'lbl')
        row(['lbl:', 'vendor_label', 'string',   'Vendor:',     'Matches "Vendor:".'], 'lbl')
        row(['lbl:', 'col_item',     'string',   'Item',        'Table column header anchor.'], 'lbl')
        row(['lbl:', 'col_qty',      'string',   'Qty',         'Table column header anchor.'], 'lbl')
        row(['lbl:', 'col_price',    'string',   'Unit Price',  'Table column header anchor.'], 'lbl')
        row(['lbl:', 'col_total',    'string',   'Total',       'Table column header anchor.'], 'lbl')
        blank()

        # ── VAR section ────────────────────────────────────────────────────────
        row(['doc:', '', '', '', 'var: defines a data field. Extracted and written to output JSON.'], 'doc')
        row(['doc:', '', '', '', 'Dot notation creates nested JSON: po.number → {"po": {"number": ...}}'], 'doc')
        row(['doc:', '', '', '', 'All var: fields in one table DATA row must share the same group prefix.'], 'doc')
        row(['doc:', '', '', '', 'Types: string  integer  currency  date  datetime'], 'doc')
        row(['doc:', '', '', '', 'Regex: Python re.fullmatch pattern. Leave blank for date/datetime. .* matches anything.'], 'doc')
        row(['var:', 'po.number',    'string',   r'PO-[0-9]{4}',    'Matches e.g. "PO-2026". Output key: {"po": {"number": ...}}'], 'var')
        row(['var:', 'po.date',      'date',     '',                 'Any date cell. Leave regex empty for date/datetime.'], 'var')
        row(['var:', 'vendor.name',  'string',   '.+',              'Any non-empty string.'], 'var')
        row(['var:', 'line.item',    'string',   '.+',              'Table DATA field. Group prefix "line" → {"line": [...]}'], 'var')
        row(['var:', 'line.qty',     'integer',  r'[1-9][0-9]*',    'Positive integer.'], 'var')
        row(['var:', 'line.price',   'currency', '.*',              'Any currency value.'], 'var')
        row(['var:', 'line.total',   'currency', '.*',              'Any currency value.'], 'var')
        row(['var:', 'footer.label', 'string',   'Grand Total',     'FOOTER field — matches the literal "Grand Total".'], 'var')
        row(['var:', 'footer.value', 'currency', '.*',              'FOOTER value field.'], 'var')
        blank()

        # ── START section ──────────────────────────────────────────────────────
        row(['doc:', '', '', '', 'Everything between START: and END: defines the extraction order.'], 'doc')
        row(['doc:', '', '', '', 'cell:1  reads the next non-empty cell into a field (or skips it with IGNORE).'], 'doc')
        row(['doc:', '', '', '', 'table:* finds all instances of a repeating table block.'], 'doc')
        row(['START:'], 'marker')

        # cell instructions
        row(['cell:1', 'po_label',     '', '', 'Consume "PO Number:" as an anchor (lbl: field — not in output).'], 'cell')
        row(['cell:1', 'po.number',    '', '', 'Extract the PO number.'], 'cell')
        row(['cell:1', 'date_label',   '', '', 'Consume "Date:" anchor.'], 'cell')
        row(['cell:1', 'po.date',      '', '', 'Extract the PO date.'], 'cell')
        row(['cell:1', 'vendor_label', '', '', 'Consume "Vendor:" anchor.'], 'cell')
        row(['cell:1', 'vendor.name',  '', '', 'Extract the vendor name.'], 'cell')
        row(['doc:', '', '', '', 'Use IGNORE to skip a non-empty cell without defining a field for it.'], 'doc')
        row(['cell:1', 'IGNORE',       '', '', 'Skip one non-empty cell without capturing it.'], 'cell')
        blank()

        # table instruction + template rows
        row(['doc:', '', '', '', 'table:* matches 0-or-more mini-table instances in greedy order.'], 'doc')
        row(['doc:', '', '', '', 'Use table:1 when exactly one instance is expected.'], 'doc')
        row(['table:*'], 'table')
        row([None, 'HEADER:1', 'col_item', 'col_qty', 'col_price', 'col_total'], 'tmpl')
        row([None, 'DATA:*',   'line.item', 'line.qty', 'line.price', 'line.total'], 'tmpl')
        row([None, 'FOOTER:1', 'footer.label', 'IGNORE', 'IGNORE', 'footer.value'], 'tmpl')
        row(['doc:', '', '', '', 'HEADER/FOOTER are strict: wrong or missing value = no match.'], 'doc')
        row(['doc:', '', '', '', 'DATA rows are lenient: wrong value emits a warning but extraction continues.'], 'doc')
        row(['doc:', '', '', '', 'EMPTY column keyword: cell must be blank. IGNORE: consume without matching.'], 'doc')
        blank()

        row(['END:'], 'marker')
        blank()

        # ── Quick reference table ──────────────────────────────────────────────
        row(['doc:', '', '', '', '── QUICK REFERENCE ──────────────────────────────────────────────'], 'doc')
        ref = [
            ('doc:',     'doc: | free text',          'Inline comment. Ignored by engine.'),
            ('config:',  'config: | key | value',     'Global setting. Keys: read.direction, currency.sign, empty.aliases'),
            ('lbl:',     'lbl: | name | type | regex', 'Anchor label. Matched but never in output.'),
            ('var:',     'var: | name | type | regex', 'Extracted variable. Use dot notation for nesting.'),
            ('cell:1',   'cell:1 | FieldName',         'Read next non-empty cell into FieldName.'),
            ('cell:1',   'cell:1 | IGNORE',            'Skip next non-empty cell.'),
            ('table:*',  'table:*  (or table:1)',       'Begin a repeating table block.'),
            ('HEADER:N', ' | HEADER:1 | F1 | F2',     'Strict header row template (col A must be blank).'),
            ('DATA:*',   ' | DATA:* | F1 | F2',        'Data row template (lenient validation).'),
            ('FOOTER:N', ' | FOOTER:1 | F1 | F2',      'Strict footer row template.'),
            ('SPLITTER', ' | SPLITTER:1',              'All columns must be blank (separator row).'),
        ]
        for keyword, syntax, description in ref:
            fill_key = keyword.rstrip(':*1').lower()
            if fill_key.startswith('header') or fill_key.startswith('data') or fill_key.startswith('footer') or fill_key.startswith('splitter'):
                fill_key = 'tmpl'
            row([keyword, syntax, '', description], fill_key)

        wb.save(output_path)
