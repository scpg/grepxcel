import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from .models import Config, CellInstruction, TableInstruction, TemplateRow
from .utils import is_empty, validate_type, _MAX_REGEX_INPUT_LEN
from .pattern_parser import PatternParser, PatternError
from .logger import Logger, LogRecord, EngineError, cell_ref
from .security import validate_file, SecurityError, DEFAULT_MAX_UNCOMPRESSED_MB


# ── Output helpers ─────────────────────────────────────────────────────────────

def _range_ref(r1: int, c1: int, r2: int, c2: int) -> str:
    """Row/col bounds → Excel A1-notation range string, e.g. 'B3:E10'."""
    return f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}'


def _set_nested(d: dict, dotted_key: str, value) -> None:
    """d['a']['b']['c'] = value  for  dotted_key='a.b.c'."""
    parts = dotted_key.split('.')
    for part in parts[:-1]:
        d = d.setdefault(part, {})
    d[parts[-1]] = value


def _field_local(field_name: str) -> str:
    """Everything after the first dot, or the full name if there is no dot."""
    _, _, rest = field_name.partition('.')
    return rest if rest else field_name


def _field_group(field_name: str) -> str:
    """Everything before the first dot, or the full name if there is no dot."""
    return field_name.split('.', 1)[0]


def _table_group(raw_tables_entry: dict, defs: dict) -> str:
    """
    Derive the top-level output key for a table match from its DATA var: fields.
    Falls back to 'table_<index>' when no var: field has a dot.
    """
    counts: dict[str, int] = {}
    for raw_row in raw_tables_entry.get('data', []):
        for field in raw_row:
            fd = defs.get(field)
            if fd and fd.role == 'var' and '.' in field:
                g = _field_group(field)
                counts[g] = counts.get(g, 0) + 1
    if counts:
        return max(counts, key=lambda k: counts[k])
    return f"table_{raw_tables_entry.get('table_index', 0)}"


def _build_row_obj(raw_row: dict, defs: dict) -> dict:
    """Convert a flat field→value dict into a nested object, skipping lbl: fields."""
    obj: dict = {}
    for field, value in raw_row.items():
        fd = defs.get(field)
        if fd and fd.role == 'lbl':
            continue
        _set_nested(obj, _field_local(field), value)
    return obj


def _build_nested_output(raw: dict, defs: dict) -> dict:
    """
    Convert the internal raw result (cells/tables) into the public nested JSON:
      - lbl: fields are stripped
      - var: cell fields: dot-notation → nested dicts
      - var: table fields: per-instance {data, header, footer} objects in an array
    """
    out: dict = {}

    # ── scalar cells ──────────────────────────────────────────────────────────
    for field, value in raw.get('cells', {}).items():
        fd = defs.get(field)
        if fd and fd.role == 'lbl':
            continue
        _set_nested(out, field, value)

    # ── table groups ──────────────────────────────────────────────────────────
    for match in raw.get('tables', []):
        group = _table_group(match, defs)
        instance: dict = {}

        if '_source' in match:
            instance['_source'] = match['_source']

        if match.get('headers'):
            header_obj: dict = {}
            for raw_row in match['headers']:
                header_obj.update(_build_row_obj(raw_row, defs))
            if header_obj:
                instance['header'] = header_obj

        if match.get('data'):
            data_rows = [_build_row_obj(r, defs) for r in match['data']]
            data_rows = [r for r in data_rows if r]  # drop fully-empty rows
            if data_rows:
                instance['data'] = data_rows

        if match.get('footers'):
            footer_obj: dict = {}
            for raw_row in match['footers']:
                footer_obj.update(_build_row_obj(raw_row, defs))
            if footer_obj:
                instance['footer'] = footer_obj

        out.setdefault(group, []).append(instance)

    return out


# ── Sheet preparation helpers ──────────────────────────────────────────────────

def _expand_merged_cells(ws) -> None:
    """
    Fill every cell in each merged range with the top-left value so the
    scanner sees the value in every visually merged cell.

    openpyxl makes non-top-left merged cells into read-only MergedCell proxy
    objects — writing to them raises AttributeError. The fix is to snapshot
    each range's bounds and top-left value, unmerge everything (which removes
    the proxies and makes all cells writable again), then fill each cell.
    The original file on disk is never modified — this operates on the
    in-memory workbook object only.
    """
    snapshots = []
    for merged_range in list(ws.merged_cells.ranges):
        snapshots.append((
            merged_range.min_row, merged_range.min_col,
            merged_range.max_row, merged_range.max_col,
            ws.cell(merged_range.min_row, merged_range.min_col).value,
        ))

    for min_row, min_col, max_row, max_col, _ in snapshots:
        ws.unmerge_cells(
            start_row=min_row, start_column=min_col,
            end_row=max_row, end_column=max_col,
        )

    for min_row, min_col, max_row, max_col, top_val in snapshots:
        for row_num in range(min_row, max_row + 1):
            for col_num in range(min_col, max_col + 1):
                ws.cell(row=row_num, column=col_num).value = top_val


def _warn_uncached_formulas(ws, logger: Logger) -> None:
    """
    Warn once if any cell still has data_type 'f' (formula) after loading
    with data_only=True, which means the cached value was never written.
    """
    for row in ws.iter_rows():
        for cell in row:
            if getattr(cell, 'data_type', None) == 'f':
                logger.warn_uncached_formulas()
                return


class SheetScanner:
    """
    Scans an openpyxl worksheet in a given direction.
    Tracks consumed cells. The main cursor only advances during cell:1 processing.
    Table scanning uses a separate local cursor so it never moves the main cursor.
    """

    def __init__(self, ws, config: Config):
        self.ws = ws
        self.config = config
        self.consumed: set = set()
        self.scan_order = self._build_scan_order()
        self.scan_order_index: dict[tuple, int] = {pos: i for i, pos in enumerate(self.scan_order)}
        self.cursor = 0

    def _build_scan_order(self) -> list:
        cells = []
        if self.config.read_direction == 'LR':
            for r in range(1, self.ws.max_row + 1):
                for c in range(1, self.ws.max_column + 1):
                    cells.append((r, c))
        elif self.config.read_direction == 'TD':
            for c in range(1, self.ws.max_column + 1):
                for r in range(1, self.ws.max_row + 1):
                    cells.append((r, c))
        return cells

    def cell_value(self, row: int, col: int):
        return self.ws.cell(row=row, column=col).value

    def cell_empty(self, row: int, col: int) -> bool:
        return is_empty(self.cell_value(row, col), self.config.empty_aliases)

    def is_consumed(self, row: int, col: int) -> bool:
        return (row, col) in self.consumed

    def consume(self, row: int, col: int):
        self.consumed.add((row, col))

    def advance_to_next(self) -> tuple:
        """Advance main cursor to the next non-empty, non-consumed cell and return it."""
        while self.cursor < len(self.scan_order):
            r, c = self.scan_order[self.cursor]
            if not self.is_consumed(r, c) and not self.cell_empty(r, c):
                return r, c
            self.cursor += 1
        return None, None

    def advance_one(self):
        self.cursor += 1

    def find_next_from(self, search_cursor: int) -> tuple:
        """
        Find the next non-empty, non-consumed cell starting from search_cursor.
        Does NOT modify self.cursor. Returns (row, col, new_search_cursor).
        """
        i = search_cursor
        while i < len(self.scan_order):
            r, c = self.scan_order[i]
            if not self.is_consumed(r, c) and not self.cell_empty(r, c):
                return r, c, i
            i += 1
        return None, None, i


class Engine:
    def process(self, pattern_file: str, data_file: str,
                logger: Logger = None,
                max_file_mb: float = 5,
                max_uncompressed_mb: float = DEFAULT_MAX_UNCOMPRESSED_MB,
                max_cell_len: int = _MAX_REGEX_INPUT_LEN,
                sheet: str | int | None = None,
                output_format: str = 'nested') -> dict:
        if logger is None:
            logger = Logger()

        self._max_cell_len = max_cell_len

        defs = {}
        _raw = {'cells': {}, 'tables': []}  # internal flat format (stats + legacy output)

        try:
            # Security: validate both files before openpyxl touches them
            for path in (pattern_file, data_file):
                try:
                    validate_file(path,
                                  max_file_mb=max_file_mb,
                                  max_uncompressed_mb=max_uncompressed_mb)
                except SecurityError as exc:
                    logger.fatal(str(exc), found=path)

            try:
                global_config, defs, start_sequence = PatternParser().parse(pattern_file)
            except (SecurityError, PatternError) as exc:
                logger.fatal(str(exc), found=pattern_file)

            try:
                wb = openpyxl.load_workbook(data_file, data_only=True)
            except Exception as exc:
                logger.fatal(
                    f'Failed to open the data file: {exc}',
                    found=data_file,
                    expected='a valid, uncorrupted .xlsx workbook',
                )

            if sheet is None:
                ws = wb.active
            elif isinstance(sheet, int):
                if sheet < 0 or sheet >= len(wb.worksheets):
                    logger.fatal(
                        f'Sheet index {sheet} is out of range '
                        f'(workbook has {len(wb.worksheets)} sheet(s))',
                        found=str(sheet),
                        expected=f'an index between 0 and {len(wb.worksheets) - 1}',
                    )
                ws = wb.worksheets[sheet]
            else:
                if sheet not in wb.sheetnames:
                    logger.fatal(
                        f'Sheet {sheet!r} not found in workbook',
                        found=sheet,
                        expected=f'one of: {", ".join(wb.sheetnames)}',
                    )
                ws = wb[sheet]
            logger.sheet_name = ws.title

            _expand_merged_cells(ws)
            _warn_uncached_formulas(ws, logger)

            logger.engine_start(pattern_file, data_file)
            logger.sheet_info(ws.title, ws.max_row, ws.max_column, global_config.read_direction)

            scanner = SheetScanner(ws, global_config)
            table_index = 0

            try:
                for instruction in start_sequence:
                    if isinstance(instruction, CellInstruction):
                        self._process_cell(instruction, scanner, defs, global_config, _raw, logger)
                    elif isinstance(instruction, TableInstruction):
                        self._process_table(instruction, scanner, defs, _raw, table_index, logger)
                        table_index += 1
            except EngineError:
                pass  # already logged; return partial result

        except EngineError:
            pass  # setup-phase fatal; already logged, return partial result

        logger.summary(_raw)

        if output_format == 'legacy':
            return _raw
        return _build_nested_output(_raw, defs)

    # -------------------------------------------------------------------------
    # cell:1 processing
    # -------------------------------------------------------------------------

    def _resolve_abs_ref(self, instr: CellInstruction, scanner: SheetScanner,
                         defs: dict, config: Config, logger: Logger) -> tuple:
        """
        Resolve an absolute cell reference (e.g. 'B5') for cell:B5 instructions.
        Validates the cursor hasn't already passed the target.
        Returns (row, col) on success; calls logger.fatal (raises EngineError) on failure.
        Empty-cell handling is role-aware: lbl: → fatal, var: → allow None.
        """
        row, col = coordinate_to_tuple(instr.target)

        idx = scanner.scan_order_index.get((row, col))
        if idx is not None and idx < scanner.cursor:
            logger.fatal(
                f"cell:{instr.target} is unreachable — the scanner has already advanced past it",
                location=cell_ref(row, col, logger.sheet_name),
                expected='a cell that has not yet been scanned',
                found=f'cursor is at scan-order position {scanner.cursor}; '
                      f'{instr.target} is at position {idx}',
            )

        value = scanner.cell_value(row, col)
        if is_empty(value, config.empty_aliases) and instr.field != 'IGNORE':
            fd = defs.get(instr.field)
            fd_role = fd.role if fd else 'var'
            if fd_role == 'lbl':
                logger.fatal(
                    f"Expected label '{instr.field}' at {instr.target} but cell is empty",
                    location=cell_ref(row, col, logger.sheet_name),
                    expected=(f"string matching {fd.regex!r}" if fd else 'a label string'),
                    found='empty cell',
                )

        return row, col

    def _process_cell(self, instr: CellInstruction, scanner: SheetScanner,
                      defs: dict, config: Config, result: dict, logger: Logger):
        if instr.multiplicity == 'abs':
            row, col = self._resolve_abs_ref(instr, scanner, defs, config, logger)
            value = scanner.cell_value(row, col)
            scanner.consume(row, col)
            idx = scanner.scan_order_index.get((row, col))
            if idx is not None:
                scanner.cursor = idx + 1
        else:
            row, col = scanner.advance_to_next()
            if row is None:
                logger.fatal(
                    f'Expected cell:{instr.multiplicity} ({instr.field!r}) but sheet is exhausted',
                    expected=f'a cell containing field {instr.field!r}',
                    found='no more non-empty cells on the sheet',
                )
            value = scanner.cell_value(row, col)
            scanner.consume(row, col)
            scanner.advance_one()

        if instr.field == 'IGNORE':
            logger.cell_ignored(row, col, value)
            return

        fd = defs.get(instr.field)
        if fd is None:
            logger.fatal(
                f'Field {instr.field!r} referenced in START: but not found in def: section',
                location=cell_ref(row, col, logger.sheet_name),
                expected=f'a def: entry named {instr.field!r}',
                found='no matching def: row in pattern file',
            )

        logger.cell_processed(row, col, instr.field, value)

        if value is not None:
            ok, _ = validate_type(value, fd.type, fd.regex, config.currency_sign, self._max_cell_len)
            if not ok:
                rec = logger.warn_validation(row, col, instr.field, fd.type, fd.regex, value)
                logger.commit_warnings([rec])

        result['cells'][instr.field] = value

    # -------------------------------------------------------------------------
    # table:* processing
    # -------------------------------------------------------------------------

    def _process_table(self, instr: TableInstruction, scanner: SheetScanner,
                       defs: dict, result: dict, table_index: int, logger: Logger):
        """
        Greedy search for all mini-table instances.
        Uses a local search cursor — never modifies scanner.cursor.
        """
        logger.table_group_start(table_index, instr.config.read_direction)
        search_cursor = scanner.cursor
        instance_index = 0

        while True:
            match, search_cursor = self._try_match_mini_table(
                instr, scanner, defs, search_cursor, logger
            )
            if match is None:
                break

            anchor_row, anchor_col = match['anchor']
            end_row = match.pop('_end_row', anchor_row)
            end_col = match.pop('_end_col', anchor_col)

            match['_source'] = {
                'sheet': logger.sheet_name,
                'ref': _range_ref(anchor_row, anchor_col, end_row, end_col),
            }
            match['table_index'] = table_index
            match['instance_index'] = instance_index
            result['tables'].append(match)

            logger.mini_table_matched(
                table_index, instance_index,
                anchor_row, anchor_col, end_row, end_col,
            )
            instance_index += 1

        logger.table_group_done(table_index, instance_index)

    def _try_match_mini_table(self, instr: TableInstruction, scanner: SheetScanner,
                               defs: dict, search_cursor: int,
                               logger: Logger) -> tuple:
        """
        Search from search_cursor for the next valid mini-table.
        Returns (match_dict, new_search_cursor) on success or (None, cursor) when exhausted.
        Warnings are only committed for successful matches — failed attempts are silent.
        """
        config = instr.config

        while True:
            anchor_row, anchor_col, search_cursor = scanner.find_next_from(search_cursor)
            if anchor_row is None:
                return None, search_cursor

            logger.anchor_probe(anchor_row, anchor_col, scanner.cell_value(anchor_row, anchor_col))

            local_warnings: list[LogRecord] = []
            match = self._attempt_match(
                anchor_row, anchor_col, instr, scanner, defs, config, local_warnings, logger
            )

            if match is not None:
                logger.commit_warnings(local_warnings)
                return match, search_cursor + 1

            logger.anchor_rejected(anchor_row, anchor_col, 'mini-table pattern did not match')
            search_cursor += 1

    def _attempt_match(self, anchor_row: int, anchor_col: int,
                       instr: TableInstruction, scanner: SheetScanner,
                       defs: dict, config: Config,
                       local_warnings: list, logger: Logger) -> dict | None:
        """
        Tentatively match the full mini-table at (anchor_row, anchor_col).
        On success: consumes all cells, returns extracted data.
        On failure: returns None, no cells consumed.
        Warnings go into local_warnings (caller decides whether to keep them).
        """
        tentative_consumed: set = set()
        extracted = {
            'headers': [], 'data': [], 'footers': [],
            'anchor': (anchor_row, anchor_col),
        }

        header_rows  = [r for r in instr.rows if r.row_type == 'HEADER']
        splitter_rows = [r for r in instr.rows if r.row_type == 'SPLITTER']
        data_rows    = [r for r in instr.rows if r.row_type == 'DATA']
        footer_rows  = [r for r in instr.rows if r.row_type == 'FOOTER']

        current_row = anchor_row

        # --- HEADER rows (strict: empty non-EMPTY field = fail) ---
        for tmpl_row in header_rows:
            row_data, ok = self._match_row(
                current_row, anchor_col, tmpl_row, defs, config,
                scanner, tentative_consumed, local_warnings, logger, strict=True,
            )
            if not ok:
                return None
            if row_data:
                extracted['headers'].append(row_data)
            current_row += 1

        # --- SPLITTER rows (all columns must be empty) ---
        num_cols = len(instr.rows[0].columns)
        for _ in splitter_rows:
            for c_offset in range(num_cols):
                col = anchor_col + c_offset
                val = scanner.ws.cell(row=current_row, column=col).value
                if not is_empty(val, config.empty_aliases):
                    return None
                tentative_consumed.add((current_row, col))
            current_row += 1

        # --- DATA rows ---
        data_tmpl = data_rows[0] if data_rows else None
        footer_tmpl_first = footer_rows[0] if footer_rows else None

        if data_tmpl:
            while True:
                if self._row_is_end_of_data(current_row, anchor_col, data_tmpl, config, scanner):
                    logger.footer_detected(current_row, anchor_col,
                                           scanner.cell_value(current_row, anchor_col))
                    current_row += 1
                    break

                if footer_tmpl_first and self._row_matches_footer(
                    current_row, anchor_col, footer_tmpl_first, defs, config, scanner
                ):
                    logger.footer_detected(current_row, anchor_col,
                                           scanner.cell_value(current_row, anchor_col))
                    break

                row_data, ok = self._match_row(
                    current_row, anchor_col, data_tmpl, defs, config,
                    scanner, tentative_consumed, local_warnings, logger, strict=False,
                )
                if not ok:
                    return None

                logger.data_row(current_row, len(data_tmpl.columns))
                extracted['data'].append(row_data)
                current_row += 1

                if data_tmpl.multiplicity == '1':
                    break

        # --- FOOTER rows (strict: empty non-EMPTY field = fail) ---
        for tmpl_row in footer_rows:
            row_data, ok = self._match_row(
                current_row, anchor_col, tmpl_row, defs, config,
                scanner, tentative_consumed, local_warnings, logger, strict=True,
            )
            if not ok:
                return None
            if row_data:
                extracted['footers'].append(row_data)
            current_row += 1

        # Full match — commit consumed cells and record span for caller
        scanner.consumed |= tentative_consumed
        extracted['_end_row'] = current_row - 1
        extracted['_end_col'] = anchor_col + num_cols - 1
        return extracted

    # -------------------------------------------------------------------------
    # Row-level matching helpers
    # -------------------------------------------------------------------------

    def _match_row(self, sheet_row: int, anchor_col: int, tmpl_row: TemplateRow,
                   defs: dict, config: Config, scanner: SheetScanner,
                   tentative_consumed: set, local_warnings: list, logger: Logger,
                   strict: bool = False) -> tuple:
        """
        Match a single template row against a sheet row.
        strict=True: empty value in a non-EMPTY field causes immediate failure (HEADER/FOOTER).
        strict=False: empty value is warned but allowed (DATA).
        Returns (row_data: dict, success: bool).
        """
        row_data = {}

        for c_offset, tmpl_col in enumerate(tmpl_row.columns):
            col = anchor_col + c_offset
            val = scanner.ws.cell(row=sheet_row, column=col).value

            if tmpl_col.field == 'EMPTY':
                if not is_empty(val, config.empty_aliases):
                    return {}, False
                tentative_consumed.add((sheet_row, col))
                continue

            if tmpl_col.field == 'IGNORE':
                tentative_consumed.add((sheet_row, col))
                continue

            fd = defs.get(tmpl_col.field)

            if is_empty(val, config.empty_aliases):
                if strict:
                    return {}, False  # HEADER/FOOTER: missing field = no match
                fd_type = fd.type if fd else 'unknown'
                local_warnings.append(
                    logger.warn_empty_field(sheet_row, col, tmpl_col.field, fd_type)
                )
                row_data[tmpl_col.field] = None
            else:
                if fd is None:
                    local_warnings.append(
                        logger.warn_undefined_field(sheet_row, col, tmpl_col.field)
                    )
                else:
                    ok, _ = validate_type(val, fd.type, fd.regex, config.currency_sign, self._max_cell_len)
                    if not ok:
                        if strict:
                            return {}, False  # HEADER/FOOTER: wrong value = no match
                        local_warnings.append(
                            logger.warn_validation(sheet_row, col, tmpl_col.field,
                                                   fd.type, fd.regex, val)
                        )
                row_data[tmpl_col.field] = val

            tentative_consumed.add((sheet_row, col))

        return row_data, True

    def _row_is_end_of_data(self, sheet_row: int, anchor_col: int,
                             data_tmpl: TemplateRow, config: Config,
                             scanner: SheetScanner) -> bool:
        for c_offset, tmpl_col in enumerate(data_tmpl.columns):
            if tmpl_col.field == 'IGNORE':
                continue
            col = anchor_col + c_offset
            val = scanner.ws.cell(row=sheet_row, column=col).value
            if not is_empty(val, config.empty_aliases):
                return False
        return True

    def _row_matches_footer(self, sheet_row: int, anchor_col: int,
                             footer_tmpl: TemplateRow, defs: dict,
                             config: Config, scanner: SheetScanner) -> bool:
        """Check if a row matches the first FOOTER template row (no cell consumption)."""
        for c_offset, tmpl_col in enumerate(footer_tmpl.columns):
            col = anchor_col + c_offset
            val = scanner.ws.cell(row=sheet_row, column=col).value

            if tmpl_col.field == 'EMPTY':
                if not is_empty(val, config.empty_aliases):
                    return False
                continue

            if tmpl_col.field == 'IGNORE':
                continue

            fd = defs.get(tmpl_col.field)
            if fd is None:
                continue
            if is_empty(val, config.empty_aliases):
                return False
            ok, _ = validate_type(val, fd.type, fd.regex, config.currency_sign, self._max_cell_len)
            if not ok:
                return False

        return True
