"""
Pattern drafting: analyse a target Excel file and use a locally-running
GGUF model (via llama-cpp-python) to produce a starter grepxcel pattern file.

No data leaves the process during inference — the model runs entirely in-process.
The only network activity is the one-time model download (and daily update checks)
handled by ModelManager.
"""

import os
import platform
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

from .model_manager import MODEL_CHAT_FORMAT, ModelManager
from .pattern_parser import PatternError, PatternParser
from .security import SecurityError, validate_file
from .utils import infer_cell_type, is_empty


# ── Prompts ───────────────────────────────────────────────────────────────────

_SYSTEM_PROMPT = """\
You are an expert at generating grepxcel pattern files.

A grepxcel pattern file is an Excel workbook that describes how to extract structured data
from another Excel file. Each row uses up to 4 columns (A, B, C, D).
In your output, separate columns with ' | ' (space-pipe-space).

─── SECTIONS (in order) ───────────────────────────────────────────────────────

1. Config rows (optional):
   config: | read.direction | LR          (or TD for top-down column scanning)
   config: | currency.sign  | €

2. Label definitions — anchor cells, NEVER written to output JSON:
   lbl: | FieldName | type | regex

   Use lbl: for literal text that marks where a value lives (e.g. "Invoice No:",
   column headers like "Product", "Qty").  These are matched for position only.

3. Variable definitions — extracted to the output JSON:
   var: | field.name | type | regex

   Use var: for every value you want to capture.
   Dot notation creates nested JSON: po.number → {"po": {"number": ...}}
   All var: field names in one table DATA row must share the same group prefix.

   Types:  string  integer  currency  percentage  date  datetime
   Regex:  Python re.fullmatch pattern.  Use .* to match anything.
           Leave blank for date/datetime (type check only, no regex needed).

   Examples:
     var: | po.number    | string     | PO-\\d{4,8}
     var: | inv.total    | currency   | \\d+(\\.\\d{1,2})?
     var: | line.qty     | integer    | \\d+
     var: | line.margin  | percentage |
     var: | inv.date     | date       |

4. Extraction sequence between START: and END:

   For scattered key-value cells — two addressing modes:

     cell:next | FieldName     (next non-empty cell in scan order → FieldName)
     cell:next | IGNORE        (skip the next non-empty cell — use for lbl: anchors)
     cell:B5   | FieldName     (jump directly to cell B5 → FieldName)
     cell:B5   | IGNORE        (jump to B5 and discard it)

   Prefer cell:next for sequential layouts.
   Use cell:B5 (A1-notation) when you know the exact address — makes the pattern
   self-documenting and independent of scan order.
   Absolute references must appear in forward reading order (LR: row then column).

   For repeating tables:
     table:*                   (bare keyword, no pipe, starts a table block)
       | HEADER:1 | ColA | ColB | ColC
       | DATA:*   | ColA | ColB | ColC
       | FOOTER:1 | ColA | ColB | ColC
     (Table template rows have a blank column A — start the line with ' | ')

   Row type suffixes:  :1 (exactly one)  :* (one or more)  :N (exactly N)
   Column keywords:  FieldName  IGNORE  EMPTY

─── KEY DESIGN RULES ──────────────────────────────────────────────────────────

- Use lbl: for label cells ("Invoice No:", "Total:", column headers).
- Use var: for the values that follow those labels.
- In HEADER rows, use lbl: field names.  In DATA rows, use var: field names.
- Give var: fields a dot-notation name: group.field (e.g. po.number, line.qty).
- Use percentage for cells that hold a percentage (e.g. 0.625 representing 62.5%).

─── OUTPUT RULES ──────────────────────────────────────────────────────────────

- Output ONLY the pattern rows. No explanations, no markdown fences, no comments.
- One row per line, columns separated by ' | '.
- Table template rows must start with ' | ' (blank column A).
- START: and END: are bare keywords with no pipe separator.

─── EXAMPLE OUTPUT ────────────────────────────────────────────────────────────

config: | read.direction | LR
lbl: | inv_label | string | Invoice No:
lbl: | total_label | string | Total
lbl: | col_product | string | Product
lbl: | col_qty | string | Qty
var: | inv.number | string | INV-\\d+
var: | inv.total | currency | \\d+(\\.\\d{2})?
var: | inv.date | date |
var: | line.product | string | .*
var: | line.qty | integer | \\d+
START:
cell:next | inv_label
cell:next | inv.number
cell:next | total_label
cell:next | inv.total
cell:next | inv.date
table:*
 | HEADER:1 | col_product | col_qty
 | DATA:* | line.product | line.qty
END:
"""

_USER_PROMPT_TEMPLATE = """\
Analyse the Excel structure below and produce a grepxcel pattern file that would extract its key data.

{analysis}

Generate the pattern file now:
"""


# ── Number-format type hint ───────────────────────────────────────────────────

def _type_from_number_format(fmt: str) -> str | None:
    """Derive a grepxcel type from an Excel number format string.

    Conservative — returns None for ambiguous formats so value-based inference
    can serve as the fallback.
    """
    if not fmt or fmt in ('General', '@', '0', '#,##0'):
        return None
    if '%' in fmt:
        return 'percentage'
    fmt_lower = fmt.lower()
    if 'y' in fmt_lower:                          # year token → date family
        return 'datetime' if ('h' in fmt_lower and ':' in fmt) else 'date'
    if any(c in fmt for c in ('$', '€', '£', '¥', '₹')) or '[$' in fmt:
        return 'currency'
    if '#,##0.00' in fmt or ('0.00' in fmt and '#' in fmt):
        return 'currency'
    return None


# ── Excel analyser ────────────────────────────────────────────────────────────

class ExcelAnalyzer:
    """Reads an xlsx file and returns a plain-text structural description for the LLM.

    Two-pass load strategy:
      Pass 1 (data_only=False): formula detection + number-format strings.
      Pass 2 (data_only=True):  cached cell values for type inference and samples.
    """

    _MAX_SAMPLE_ROWS = 100   # rows used for type inference
    _MAX_SAMPLE_COLS = 20    # columns described per section
    _DISPLAY_SAMPLES = 3     # value examples shown per column in the output text
    _ROW_TOLERANCE   = 1.2   # if actual_rows <= MAX × tolerance, include all rows

    def __init__(
        self,
        path: str,
        sheet=None,
        max_file_mb: float = 5.0,
        max_uncompressed_mb: float = 50.0,
    ):
        self.path                = path
        self.sheet               = sheet
        self.max_file_mb         = max_file_mb
        self.max_uncompressed_mb = max_uncompressed_mb

    def analyse(self) -> str:
        validate_file(self.path, self.max_file_mb, self.max_uncompressed_mb)
        wb_raw = openpyxl.load_workbook(self.path, data_only=False)
        wb_val = openpyxl.load_workbook(self.path, data_only=True)
        try:
            preamble = self._workbook_preamble(wb_val)
            ws_raw   = self._select_sheet(wb_raw)
            ws_val   = self._select_sheet(wb_val)
            body     = self._analyse_sheet(ws_raw, ws_val)
        finally:
            wb_raw.close()
            wb_val.close()
        return (preamble + body) if preamble else body

    def _select_sheet(self, wb):
        if self.sheet is None:
            return wb.active
        if isinstance(self.sheet, int):
            return wb.worksheets[self.sheet]
        return wb[self.sheet]

    def _workbook_preamble(self, wb) -> str:
        """List all sheets with dimensions. Empty for single-sheet workbooks."""
        sheets = wb.worksheets
        if len(sheets) <= 1:
            return ''
        target_title = self._select_sheet(wb).title
        lines = [f'Workbook: {len(sheets)} sheets']
        for i, ws in enumerate(sheets, 1):
            r = ws.max_row or 0
            c = ws.max_column or 0
            marker = ' ← target' if ws.title == target_title else ''
            lines.append(f"  Sheet {i}: '{ws.title}' — {r} rows × {c} columns{marker}")
        return '\n'.join(lines) + '\n\n'

    def _analyse_sheet(self, ws_raw, ws_val) -> str:
        max_row = ws_val.max_row or 0
        max_col = ws_val.max_column or 0

        header = [
            f"Sheet: '{ws_val.title}'",
            f'Dimensions: {max_row} rows × {max_col} columns',
        ]

        if max_row == 0:
            return '\n'.join(header) + '\n\nEmpty sheet — no data found.'

        # Apply tolerance: include all rows when count is only slightly above the cap.
        load_rows = (
            max_row if max_row <= int(self._MAX_SAMPLE_ROWS * self._ROW_TOLERANCE)
            else self._MAX_SAMPLE_ROWS
        )
        load_cols = min(max_col, self._MAX_SAMPLE_COLS)

        formula_cells, number_formats = self._cell_metadata(ws_raw, load_rows, load_cols)

        rows = list(ws_val.iter_rows(
            min_row=1, max_row=load_rows, max_col=load_cols,
            values_only=True,
        ))

        if not rows or all(all(v is None for v in r) for r in rows):
            return '\n'.join(header) + '\n\nEmpty sheet — no data found.'

        notes = []
        if max_row > load_rows:
            notes.append(f'(analysis covers first {load_rows} of {max_row} rows)')
        if max_col > self._MAX_SAMPLE_COLS:
            notes.append(f'(first {self._MAX_SAMPLE_COLS} of {max_col} columns shown)')

        sections = self._split_into_sections(rows)
        if not sections:
            return '\n'.join(header) + '\n\nEmpty sheet — no data found.'

        lines   = header + notes
        is_multi = len(sections) > 1

        for sec_idx, (sec_start, sec_rows) in enumerate(sections, 1):
            sec_end = sec_start + len(sec_rows)
            label   = (
                f'Section {sec_idx} (rows {sec_start + 1}–{sec_end})'
                if is_multi else None
            )
            non_empty_first = [v for v in sec_rows[0] if not is_empty(v)]
            # TABLE requires >= 2 rows (header + at least one data row)
            # and >= 2 non-empty values in the first row.
            if len(non_empty_first) >= 2 and len(sec_rows) >= 2:
                lines += self._describe_table_section(
                    sec_rows, sec_start, formula_cells, number_formats, label)
            else:
                lines += self._describe_kv_section(
                    sec_rows, sec_start, formula_cells, number_formats, label)

        return '\n'.join(lines)

    def _cell_metadata(self, ws, max_row: int, max_col: int) -> tuple[set, dict]:
        """Return (formula_cells, number_formats) using 1-based (row, col) keys."""
        formula_cells  = set()
        number_formats = {}
        for cell_row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in cell_row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    formula_cells.add((cell.row, cell.column))
                fmt = getattr(cell, 'number_format', None)
                if fmt and fmt not in ('General', '@'):
                    number_formats[(cell.row, cell.column)] = fmt
        return formula_cells, number_formats

    def _split_into_sections(self, rows: list) -> list[tuple[int, list]]:
        """Split rows at runs of fully-empty rows. Returns [(0based_start, rows), ...]."""
        sections: list[tuple[int, list]] = []
        current:  list                   = []
        start = 0
        for i, row in enumerate(rows):
            if all(v is None or is_empty(v) for v in row):
                if current:
                    sections.append((start, current))
                    current = []
            else:
                if not current:
                    start = i
                current.append(row)
        if current:
            sections.append((start, current))
        return sections

    def _col_type(
        self, col_idx: int, sec_start: int,
        data_rows: list, number_formats: dict,
    ) -> tuple[str, str | None]:
        """Return (type_str, format_hint_or_None) for one column."""
        vals = [r[col_idx] for r in data_rows if col_idx < len(r)]
        # First non-None format for this column in any data row (1-based indexing)
        fmt = next(
            (number_formats[(sec_start + 2 + ri, col_idx + 1)]
             for ri in range(len(data_rows))
             if (sec_start + 2 + ri, col_idx + 1) in number_formats),
            number_formats.get((sec_start + 1, col_idx + 1)),  # fall back to header row
        )
        return (_type_from_number_format(fmt) or infer_cell_type(vals)), fmt

    def _describe_table_section(
        self, rows, sec_start, formula_cells, number_formats, label,
    ) -> list:
        headers     = list(rows[0])
        data_rows   = rows[1:]
        non_empty_h = [(c, h) for c, h in enumerate(headers) if not is_empty(h)
                       ][:self._MAX_SAMPLE_COLS]

        prefix = (f'{label}: TABLE layout' if label
                  else 'Layout: TABLE (first row appears to be column headers)')
        lines  = ['', prefix, f'Columns ({len(non_empty_h)}):']

        for col_idx, header in non_empty_h:
            col_type, fmt = self._col_type(col_idx, sec_start, data_rows, number_formats)
            is_formula    = any(
                (sec_start + 2 + ri, col_idx + 1) in formula_cells
                for ri in range(len(data_rows))
            )
            vals    = [r[col_idx] for r in data_rows
                       if col_idx < len(r) and not is_empty(r[col_idx])]
            samples = [repr(v) for v in vals[:self._DISPLAY_SAMPLES]]
            line    = f"  - '{header}' [{col_type}]  samples: {', '.join(samples) or '(none)'}"
            if fmt:
                line += f'  (format: {fmt})'
            if is_formula:
                line += '  [formula — consider var:]'
            lines.append(line)

        if data_rows:
            lines.append(f'\nData rows: {len(data_rows)}')
        return lines

    def _describe_kv_section(
        self, rows, sec_start, formula_cells, number_formats, label,
    ) -> list:
        prefix = (f'{label}: KEY-VALUE layout' if label
                  else 'Layout: KEY-VALUE (scattered cells, not a standard table)')
        lines  = ['', prefix, 'Cell pairs found:']
        count  = 0

        for ri, row in enumerate(rows):
            global_row = sec_start + ri + 1          # 1-based sheet row
            non_empty  = [(ci, v) for ci, v in enumerate(row) if not is_empty(v)]
            if not non_empty:
                continue
            for ci, val in non_empty:
                next_val = row[ci + 1] if ci + 1 < len(row) else None
                fmt      = number_formats.get((global_row, ci + 2)) if next_val is not None else None
                val_type = (_type_from_number_format(fmt)
                            or (infer_cell_type([next_val]) if next_val is not None else 'string'))
                is_fml   = (global_row, ci + 1) in formula_cells
                fml_note = '  [formula]' if is_fml else ''
                if next_val is not None:
                    lines.append(f"  - '{val}': {repr(next_val)} [{val_type}]{fml_note}")
                else:
                    lines.append(f"  - '{val}': (standalone){fml_note}")
                count += 1
                if count >= 50:
                    lines.append('  ... (additional pairs omitted)')
                    return lines
        return lines


# ── Hardware detection ────────────────────────────────────────────────────────

_CUDA_INSTALL  = "CMAKE_ARGS=\"-DGGML_CUDA=on\"  pip install llama-cpp-python --force-reinstall"
_METAL_INSTALL = "CMAKE_ARGS=\"-DGGML_METAL=on\" pip install llama-cpp-python --force-reinstall"
_VULKAN_INSTALL= "CMAKE_ARGS=\"-DGGML_VULKAN=on\" pip install llama-cpp-python --force-reinstall"


def _detect_gpu() -> tuple[int, str | None, str | None]:
    """
    Probe the local machine for GPU/NPU acceleration and return
    (n_gpu_layers, backend_label, install_hint).

    Priority:
      1. GREPXCEL_GPU_LAYERS env var — user override, no messages printed.
      2. NVIDIA GPU via nvidia-smi (CUDA).
      3. Apple Silicon via platform check (Metal).
      4. Any Vulkan-capable GPU via vulkaninfo (Windows/Linux fallback).
      5. CPU-only fallback (n_gpu_layers=0).

    NPU note: Intel NPU, Qualcomm Hexagon, and Apple ANE are not yet supported
    by llama.cpp in mainstream builds.  Apple Silicon users get equivalent
    acceleration through Metal (step 3 above).
    """
    override = os.environ.get("GREPXCEL_GPU_LAYERS")
    if override is not None:
        try:
            return int(override), None, None
        except ValueError:
            pass

    # NVIDIA CUDA (works on Linux, Windows, and WSL2 with NVIDIA WSL2 driver)
    try:
        proc = subprocess.run(
            ["nvidia-smi", "--query-gpu=name", "--format=csv,noheader"],
            capture_output=True, text=True, timeout=5,
        )
        if proc.returncode == 0:
            gpu_name = proc.stdout.strip().splitlines()[0]
            return -1, f"NVIDIA GPU ({gpu_name}) — CUDA", _CUDA_INSTALL
    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        pass

    # Apple Silicon Metal (macOS only — ANE is bypassed; Metal is the fast path)
    if platform.system() == "Darwin":
        import struct
        if struct.calcsize("P") == 8 and platform.machine() == "arm64":
            return -1, "Apple Silicon — Metal", _METAL_INSTALL
        # Intel Mac: Metal exists but GPU offload yields little benefit
        return 0, None, None

    # Vulkan (any GPU on Windows/Linux — AMD, Intel, NVIDIA without CUDA driver)
    try:
        proc = subprocess.run(
            ["vulkaninfo", "--summary"],
            capture_output=True, text=True, timeout=5,
        )
        if proc.returncode == 0 and "GPU" in proc.stdout:
            for line in proc.stdout.splitlines():
                if "deviceName" in line:
                    gpu_name = line.split("=", 1)[-1].strip()
                    return -1, f"Vulkan GPU ({gpu_name})", _VULKAN_INSTALL
            return -1, "Vulkan GPU", _VULKAN_INSTALL
    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        pass

    return 0, None, None


# ── Local LLM client ──────────────────────────────────────────────────────────

class LlamaCppClient:
    """
    Runs inference entirely in-process using llama-cpp-python.
    The model is loaded lazily on the first chat() call and stays loaded
    for the lifetime of this instance — making it safe to reuse across
    multiple requests in a web service.
    """

    def __init__(self, model_path: str):
        self.model_path = model_path
        self._llm       = None

    def _load(self) -> None:
        if self._llm is not None:
            return
        try:
            from llama_cpp import Llama
        except ImportError:
            print(
                "Error: 'llama-cpp-python' is not installed.\n"
                "Fix:   pip install llama-cpp-python",
                file=sys.stderr,
            )
            sys.exit(1)

        n_gpu_layers, backend, install_hint = _detect_gpu()

        if backend:
            print(f"  Hardware: {backend}", file=sys.stderr)
            layers_label = "all layers" if n_gpu_layers == -1 else f"{n_gpu_layers} layers"
            print(f"  GPU offload: {layers_label}", file=sys.stderr)
            print(
                f"  Note: GPU acceleration requires a GPU-compiled build of llama-cpp-python.\n"
                f"  If inference seems slow, reinstall with:\n"
                f"    {install_hint}",
                file=sys.stderr,
            )
        else:
            print("  Hardware: CPU (no GPU detected — set GREPXCEL_GPU_LAYERS to override)",
                  file=sys.stderr)

        kwargs: dict = dict(
            model_path   = self.model_path,
            n_ctx        = 4096,
            n_threads    = os.cpu_count() or 4,
            n_gpu_layers = n_gpu_layers,
            verbose      = False,
        )
        if MODEL_CHAT_FORMAT is not None:
            kwargs['chat_format'] = MODEL_CHAT_FORMAT
        self._llm = Llama(**kwargs)

    def chat(self, system: str, user: str) -> str:
        self._load()
        response = self._llm.create_chat_completion(
            messages=[
                {"role": "system", "content": system},
                {"role": "user",   "content": user},
            ],
            temperature = 0.1,
            max_tokens  = 2048,
        )
        return response["choices"][0]["message"]["content"]


# ── Pattern writer ────────────────────────────────────────────────────────────

_TABLE_ROW_PREFIXES = ('HEADER:', 'DATA:', 'FOOTER:', 'SPLITTER:')


class PatternWriter:
    """Parses LLM pipe-delimited text and writes a pattern xlsx."""

    def write(self, llm_text: str, output_path: str) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title  = 'Pattern'
        row_num   = 1

        for line in llm_text.splitlines():
            stripped = line.strip()
            if not stripped or stripped.startswith('```'):
                continue

            if stripped.upper() in ('START:', 'END:'):
                ws.cell(row=row_num, column=1, value=stripped.upper())
                row_num += 1
                continue

            is_table_row = stripped.startswith('|') or any(
                stripped.upper().startswith(p) for p in _TABLE_ROW_PREFIXES
            )
            if is_table_row:
                # Strip leading '|'; remaining parts map to columns B, C, D, …
                body  = stripped.lstrip('|')
                parts = [p.strip() for p in body.split('|')]
                for rel_i, val in enumerate(parts):
                    if val:
                        ws.cell(row=row_num, column=rel_i + 2, value=val)
                row_num += 1
                continue

            if '|' in stripped:
                parts = [p.strip() for p in stripped.split('|')]
                for col_num, val in enumerate(parts, start=1):
                    if val:
                        ws.cell(row=row_num, column=col_num, value=val)
                row_num += 1
                continue

            # Bare keyword (e.g. 'table:*')
            parts = stripped.split(None, 1)
            ws.cell(row=row_num, column=1, value=parts[0])
            if len(parts) > 1:
                ws.cell(row=row_num, column=2, value=parts[1])
            row_num += 1

        wb.save(output_path)


# ── Validation helpers ────────────────────────────────────────────────────────

def _failed_path(output_path: str) -> str:
    """Derive a _FAILED.txt path alongside the intended output file."""
    p = Path(output_path)
    return str(p.parent / f'{p.stem}_FAILED.txt')


def _write_failed_draft(llm_text: str, error: str, output_path: str) -> None:
    path = _failed_path(output_path)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(f'# Draft validation error\n# {error}\n\n')
        f.write(llm_text)


def _validate_draft(xlsx_path: str) -> None:
    """Parse the draft xlsx with the engine's own PatternParser.

    Raises PatternError or SecurityError if the LLM output is structurally invalid.
    """
    PatternParser().parse(xlsx_path)


# ── Orchestrator ──────────────────────────────────────────────────────────────

class PatternDrafter:
    """Analyse → ensure model → run LLM → write pattern xlsx."""

    def __init__(
        self,
        input_path: str,
        output_path: str,
        sheet=None,
        max_file_mb: float = 5.0,
        max_uncompressed_mb: float = 50.0,
        verbose: bool = False,
        dry_run: bool = False,
    ):
        self.input_path          = input_path
        self.output_path         = output_path
        self.sheet               = sheet
        self.max_file_mb         = max_file_mb
        self.max_uncompressed_mb = max_uncompressed_mb
        self.verbose             = verbose
        self.dry_run             = dry_run

    def run(self) -> int:
        """Run the full pipeline. Returns exit code (0 = success, 1 = error)."""
        # 1. Analyse Excel structure
        print('Analysing Excel structure...', file=sys.stderr)
        try:
            analysis = ExcelAnalyzer(
                self.input_path, self.sheet,
                self.max_file_mb, self.max_uncompressed_mb,
            ).analyse()
        except SecurityError as exc:
            print(f'Security error: {exc}', file=sys.stderr)
            return 1

        if self.verbose or self.dry_run:
            print('\n── Excel analysis ──────────────────────────────', file=sys.stderr)
            print(analysis, file=sys.stderr)
            print('────────────────────────────────────────────────\n', file=sys.stderr)

        if self.dry_run:
            print('[dry-run] Model inference skipped.', file=sys.stderr)
            return 0

        # 2. Ensure the model is present and up to date
        model_path = ModelManager().ensure_ready(verbose=self.verbose)

        # 3. Run inference (entirely in-process — no network)
        print('Running local model inference...', file=sys.stderr)
        user_prompt = _USER_PROMPT_TEMPLATE.format(analysis=analysis)
        llm_text    = LlamaCppClient(str(model_path)).chat(_SYSTEM_PROMPT, user_prompt)

        # 4. Text preview → stdout (pipe-friendly)
        print(llm_text)

        # 5. Write to a temp xlsx, validate, then move to final path
        output_dir = os.path.dirname(os.path.abspath(self.output_path))
        os.makedirs(output_dir, exist_ok=True)
        fd, tmp_path = tempfile.mkstemp(suffix='.xlsx', dir=output_dir)
        os.close(fd)
        try:
            PatternWriter().write(llm_text, tmp_path)
            _validate_draft(tmp_path)
        except (PatternError, SecurityError) as exc:
            os.unlink(tmp_path)
            _write_failed_draft(llm_text, str(exc), self.output_path)
            print(
                f'\n[!] Draft validation failed: {exc}\n'
                f'    Raw LLM output saved to: {_failed_path(self.output_path)}',
                file=sys.stderr,
            )
            return 1

        os.replace(tmp_path, self.output_path)
        print(f'\nDraft pattern written to: {self.output_path}', file=sys.stderr)
        print(
            f'Open it in Excel/LibreOffice, refine the regexes, then run:\n'
            f'  grepxcel extract -p {self.output_path} your_data.xlsx',
            file=sys.stderr,
        )
        return 0
