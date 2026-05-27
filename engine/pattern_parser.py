import re

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from .models import Config, FieldDef, TemplateColumn, TemplateRow, CellInstruction, TableInstruction
from .security import check_regex_safety, SecurityError

_MAX_PATTERN_CELL_LEN = 1_000  # max characters in any pattern file cell value

_A1_RE = re.compile(r'^[A-Z]{1,3}[1-9][0-9]*$', re.IGNORECASE)


class PatternError(Exception):
    """Raised when a pattern file contains a structural or ordering error."""


def _coord_after(prev: tuple, new: tuple, direction: str) -> bool:
    """Return True if `new` comes strictly after `prev` in `direction` reading order."""
    pr, pc = prev
    nr, nc = new
    if direction == 'LR':
        return (nr, nc) > (pr, pc)
    return (nc, nr) > (pc, pr)  # TD: column primary


def _coord_str(pos: tuple) -> str:
    return f'{get_column_letter(pos[1])}{pos[0]}'


class PatternParser:
    def parse(self, filepath: str) -> tuple:
        """
        Parse a pattern Excel file.
        Returns (global_config, defs, start_sequence).
        """
        wb = openpyxl.load_workbook(filepath)  # data_only=False: we want to see formulas
        ws = wb.active

        rows = self._read_and_validate(ws)

        global_config = Config()
        defs = {}
        start_sequence = []

        i = 0
        in_start = False
        last_abs_pos: tuple | None = None  # (row, col) of last cell:XY seen

        while i < len(rows):
            row = self._pad(rows[i])
            col_a = row[0]

            if col_a == 'END:':
                break

            if col_a == 'START:':
                in_start = True
                i += 1
                continue

            if not in_start:
                if col_a == 'config:':
                    self._apply_global_config(row, global_config)
                elif col_a in ('def:', 'var:'):
                    fd = self._parse_field(row, role='var')
                    defs[fd.name] = fd
                elif col_a == 'lbl:':
                    fd = self._parse_field(row, role='lbl')
                    defs[fd.name] = fd
                elif col_a in ('doc:', 'info:'):
                    pass  # inline documentation — ignored by engine
                i += 1
                continue

            # Inside START: section
            if col_a and col_a.startswith('cell:'):
                raw = col_a.split(':', 1)[1]
                field = str(row[1] or 'IGNORE')
                if raw in ('1', 'next'):
                    start_sequence.append(CellInstruction(multiplicity=raw, field=field))
                elif _A1_RE.match(raw):
                    ref = raw.upper()
                    new_pos = coordinate_to_tuple(ref)
                    if last_abs_pos is not None and not _coord_after(last_abs_pos, new_pos, global_config.read_direction):
                        raise PatternError(
                            f"cell:{ref} at pattern row {i + 1} is before or equal to the "
                            f"previous absolute reference {_coord_str(last_abs_pos)} in "
                            f"{global_config.read_direction} reading order — unreachable"
                        )
                    last_abs_pos = new_pos
                    start_sequence.append(CellInstruction(multiplicity='abs', field=field, target=ref))
                else:
                    raise PatternError(
                        f"Unknown cell instruction 'cell:{raw}' at pattern row {i + 1}. "
                        f"Use 'cell:next', 'cell:1', or a cell coordinate like 'cell:B5'."
                    )
                i += 1

            elif col_a and col_a.startswith('table:'):
                mult = col_a.split(':', 1)[1]
                table_config = Config(
                    read_direction=global_config.read_direction,
                    currency_sign=global_config.currency_sign,
                    empty_aliases=list(global_config.empty_aliases),
                )
                template_rows = []
                i += 1

                while i < len(rows):
                    sub = self._pad(rows[i])
                    sub_a = sub[0]
                    sub_b = sub[1]

                    # Non-empty col A means we're back at the top level
                    if sub_a is not None:
                        break

                    # Blank row inside table block
                    if sub_b is None:
                        i += 1
                        continue

                    if sub_b == 'config:':
                        key, val = sub[2], sub[3]
                        if key == 'read.direction' and val:
                            table_config.read_direction = str(val)
                        i += 1
                        continue

                    # HEADER / SPLITTER / DATA / FOOTER row
                    if sub_b and ':' in str(sub_b):
                        row_type, row_mult = str(sub_b).rsplit(':', 1)
                        cols_raw = list(sub[2:])
                        while cols_raw and cols_raw[-1] is None:
                            cols_raw.pop()
                        columns = [
                            TemplateColumn(field=str(v) if v is not None else 'EMPTY')
                            for v in cols_raw
                        ]
                        template_rows.append(TemplateRow(
                            row_type=row_type,
                            multiplicity=row_mult,
                            columns=columns,
                        ))

                    i += 1

                start_sequence.append(TableInstruction(
                    multiplicity=mult,
                    config=table_config,
                    rows=template_rows,
                ))

            else:
                i += 1

        return global_config, defs, start_sequence

    def _read_and_validate(self, ws) -> list:
        """
        Read all rows from the pattern worksheet, enforcing that every cell is
        either empty or a plain string.  Formulas, numbers, dates, and booleans
        are all rejected — the pattern file is a configuration document, not a
        spreadsheet.
        """
        rows = []
        for row in ws.iter_rows():
            row_values = []
            for cell in row:
                val = cell.value
                if val is None:
                    row_values.append(None)
                    continue

                # Formulas are never allowed in pattern files
                if cell.data_type == 'f' or (isinstance(val, str) and val.startswith('=')):
                    raise SecurityError(
                        f'Formulas are not allowed in pattern files. '
                        f'Cell {cell.coordinate} contains: {val!r}  '
                        f'Replace it with a plain text value.'
                    )

                # Only plain strings are accepted
                if not isinstance(val, str):
                    raise SecurityError(
                        f'Pattern file cells must contain plain text only. '
                        f'Cell {cell.coordinate} contains a {type(val).__name__} value: {val!r}  '
                        f'All values in a pattern file must be strings.'
                    )

                # Guard against excessively long values
                if len(val) > _MAX_PATTERN_CELL_LEN:
                    raise SecurityError(
                        f'Pattern file cell {cell.coordinate} value is too long '
                        f'({len(val)} chars, limit is {_MAX_PATTERN_CELL_LEN}). '
                        f'Pattern values should be short identifiers or regex patterns.'
                    )

                row_values.append(val)
            rows.append(row_values)
        return rows

    def _pad(self, row, length=10) -> list:
        """Ensure a row list has at least `length` elements."""
        row = list(row)
        while len(row) < length:
            row.append(None)
        return row

    def _apply_global_config(self, row, config: Config):
        key, val = row[1], row[2]
        if key == 'read.direction' and val:
            config.read_direction = str(val)
        elif key == 'currency.sign' and val:
            config.currency_sign = str(val)
        elif key == 'empty.aliases' and val:
            config.empty_aliases.append(str(val))

    def _parse_field(self, row, role: str = 'var') -> FieldDef:
        name  = str(row[1]) if row[1] else ''
        type_ = str(row[2]) if row[2] else 'string'
        regex = str(row[3]) if row[3] else '.*'
        check_regex_safety(regex, field_name=name)
        return FieldDef(name=name, type=type_, regex=regex, role=role)

    # backward-compat alias
    def _parse_def(self, row) -> FieldDef:
        return self._parse_field(row, role='var')
